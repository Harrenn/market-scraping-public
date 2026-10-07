#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import re
import sys
import warnings
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from html import unescape
from pathlib import Path
from typing import Iterable
from urllib.parse import parse_qs, urlencode, urljoin, urlparse, urlunparse
from xml.etree import ElementTree as ET
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup, MarkupResemblesLocatorWarning, Tag
from openpyxl import Workbook


warnings.filterwarnings("ignore", category=MarkupResemblesLocatorWarning)

USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36"
)
TIMEZONE = ZoneInfo("Asia/Manila")
REST_PAGE_SIZE = 100
EXCEL_CELL_LIMIT = 32767
SITEMAP_NAMESPACE = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
MOTORTRADE_DOMAIN = "https://motortrade.com.ph"
FORMS_DOMAIN = "https://forms.motortrade.com.ph"
DEALER_ID = "motortrade"
DEALER_NAME = "Motortrade"

TABLE_SCHEMAS: dict[str, list[str]] = {
    "dealer": [
        "dealer_id",
        "dealer_name",
        "primary_domain",
        "forms_domain",
        "crawl_timezone",
        "crawl_date",
        "notes",
    ],
    "crawl_run": [
        "run_id",
        "dealer_id",
        "run_started_at",
        "run_finished_at",
        "source_count",
        "notes",
    ],
    "source_page": [
        "source_page_id",
        "run_id",
        "dealer_id",
        "source_type",
        "wp_post_type",
        "wp_id",
        "url",
        "slug",
        "title",
        "status",
        "published_at",
        "modified_at",
        "author_id",
        "excerpt_text",
        "content_html",
        "content_text",
        "featured_media_id",
        "featured_image_url",
        "category_ids",
        "tag_ids",
        "job_type_ids",
        "seo_title",
        "seo_description",
        "canonical_url",
    ],
    "taxonomy_term": [
        "term_id",
        "taxonomy",
        "name",
        "slug",
        "parent_term_id",
        "count",
        "link",
        "used_by_post_types",
    ],
    "taxonomy_assignment": [
        "assignment_id",
        "dealer_id",
        "entity_type",
        "entity_id",
        "taxonomy",
        "term_slug",
        "term_label",
        "source_url",
        "confidence",
        "notes",
    ],
    "filter_option": [
        "filter_name",
        "option_value",
        "option_label",
        "source_url",
        "notes",
    ],
    "motorcycle": [
        "motorcycle_id",
        "dealer_id",
        "source_page_id",
        "detail_url",
        "listing_url",
        "display_name",
        "brand_name",
        "model_name",
        "vehicle_type",
        "vehicle_type_source",
        "raw_taxonomies",
        "classification_notes",
        "spec_make",
        "spec_model",
        "spec_transmission",
        "spec_engine_type",
        "spec_seat_height",
        "spec_brake_system",
        "spec_displacement_cc",
        "spec_front_tire",
        "spec_rear_tire",
        "spec_wheels_type",
        "spec_starting_system",
        "spec_overall_dimensions",
        "spec_ground_clearance",
        "spec_fuel_system",
        "is_blue_tag",
        "is_featured",
        "is_top_pick",
        "has_vehicle_promo_flag",
        "primary_image_url",
        "excerpt_text",
        "apply_loan_url",
        "inquiry_url",
        "service_url",
        "compare_enabled",
        "published_at",
        "modified_at",
    ],
    "motorcycle_spec": [
        "motorcycle_id",
        "spec_group",
        "spec_label",
        "spec_value_raw",
        "spec_value_numeric",
        "spec_unit",
        "display_order",
        "source_url",
    ],
    "pricing": [
        "pricing_id",
        "motorcycle_id",
        "dealer_id",
        "source_url",
        "source_context",
        "currency",
        "srp_text",
        "srp_amount",
        "downpayment_text",
        "downpayment_amount",
        "monthly_text",
        "monthly_amount",
        "price_note",
        "is_indicative",
        "captured_at",
    ],
    "financing": [
        "financing_id",
        "motorcycle_id",
        "dealer_id",
        "source_url",
        "financing_type",
        "program_name",
        "monthly_amount",
        "downpayment_amount",
        "loan_amount",
        "loan_term_months",
        "estimated_income_amount",
        "form_url",
        "cta_url",
        "note_text",
        "captured_at",
    ],
    "branch": [
        "branch_id",
        "dealer_id",
        "source_url",
        "locator_id",
        "branch_name",
        "slug",
        "street",
        "city",
        "state_or_province",
        "postal_code",
        "full_address",
        "latitude",
        "longitude",
        "phone",
        "fax",
        "email",
        "website",
        "open_hours_raw",
        "days_str",
        "brand_raw",
        "product_raw",
        "categories_raw",
        "inquire_url",
    ],
    "branch_product_coverage": [
        "coverage_id",
        "dealer_id",
        "branch_id",
        "coverage_type",
        "coverage_value_raw",
        "coverage_value_normalized",
        "source_url",
        "source_field",
        "confidence",
        "observed_at",
        "notes",
    ],
    "branch_model_availability": [
        "availability_id",
        "dealer_id",
        "branch_id",
        "motorcycle_id",
        "availability_status",
        "source_url",
        "source_field",
        "evidence_text",
        "raw_payload",
        "confidence",
        "observed_at",
        "notes",
    ],
    "relationship_probe": [
        "probe_id",
        "dealer_id",
        "relationship_name",
        "source_url",
        "source_type",
        "status",
        "evidence_found",
        "rows_created",
        "notes",
        "checked_at",
    ],
    "promo": [
        "promo_id",
        "dealer_id",
        "source_page_id",
        "source_url",
        "title",
        "promo_type",
        "description_text",
        "image_url",
        "cta_url",
        "target_listing_url",
        "applies_to_brand",
        "applies_to_vehicle_type",
        "applies_to_tag",
        "has_vehicle_promo_flag",
        "published_at",
        "modified_at",
    ],
    "motorcycle_promo_link": [
        "motorcycle_id",
        "promo_id",
        "link_type",
        "source_url",
        "notes",
    ],
    "service_offering": [
        "service_id",
        "dealer_id",
        "source_page_id",
        "service_category",
        "title",
        "description_text",
        "partner_name",
        "cta_url",
        "external_form_url",
        "notes",
    ],
    "form_definition": [
        "form_id",
        "dealer_id",
        "form_name",
        "form_category",
        "source_url",
        "final_url",
        "page_title",
        "form_action",
        "form_method",
        "requires_file_upload",
        "requires_otp",
        "multi_step",
        "has_privacy_checkbox",
        "notes",
    ],
    "form_field": [
        "form_id",
        "field_order",
        "field_name",
        "field_label",
        "field_type",
        "is_required",
        "default_value",
        "placeholder",
        "option_value",
        "option_label",
        "section_name",
        "source_url",
    ],
    "job_listing": [
        "job_id",
        "dealer_id",
        "source_page_id",
        "title",
        "slug",
        "job_location",
        "application_target",
        "company_name",
        "company_website",
        "company_tagline",
        "job_salary",
        "salary_currency",
        "salary_unit",
        "is_featured",
        "is_filled",
        "is_remote",
        "job_type_ids",
        "description_text",
        "published_at",
        "modified_at",
    ],
    "parts_catalog": [
        "part_id",
        "dealer_id",
        "source_page_id",
        "title",
        "slug",
        "detail_url",
        "image_url",
        "description_text",
        "published_at",
        "modified_at",
    ],
    "preowned_motorcycle": [
        "preowned_id",
        "dealer_id",
        "source_page_id",
        "title",
        "slug",
        "detail_url",
        "image_url",
        "description_text",
        "published_at",
        "modified_at",
    ],
    "external_cta": [
        "cta_id",
        "dealer_id",
        "source_page_id",
        "source_url",
        "target_url",
        "target_domain",
        "cta_type",
        "anchor_text",
        "notes",
    ],
    "data_quality_issue": [
        "issue_id",
        "dealer_id",
        "entity_type",
        "entity_id",
        "source_url",
        "issue_type",
        "issue_detail",
        "severity",
        "captured_at",
    ],
}

REST_COLLECTIONS = {
    "pages": ("page", "pages"),
    "posts": ("post", "posts"),
    "featured_products": ("featured_products", "featured_products"),
    "parts": ("parts", "parts"),
    "job-listings": ("job_listing", "job-listings"),
    "preowned": ("preowned", "preowned"),
    "categories": ("category", "categories"),
    "tags": ("post_tag", "tags"),
    "job-types": ("job_listing_type", "job-types"),
}

TRACKED_FORM_URLS = [
    "https://forms.motortrade.com.ph/loan?src=MTgxMg",
    "https://forms.motortrade.com.ph/inquiry?src=MTgxMg",
    "https://forms.motortrade.com.ph/service",
    "https://forms.motortrade.com.ph/bigbike",
    "https://forms.motortrade.com.ph/customercare",
    "https://forms.motortrade.com.ph/ltrs",
    "https://forms.motortrade.com.ph/careers",
]

TRACKED_HTML_PAGES = [
    ("promo_surface", "https://motortrade.com.ph/promos/"),
    ("promo_surface", "https://motortrade.com.ph/promotion/"),
    ("service_page", "https://motortrade.com.ph/insurance/"),
    ("service_page", "https://motortrade.com.ph/maintenance-repair/"),
    ("service_page", "https://motortrade.com.ph/ltrs/"),
    ("service_page", "https://motortrade.com.ph/forms/"),
    ("service_page", "https://motortrade.com.ph/customer-care/"),
    ("service_page", "https://motortrade.com.ph/service/"),
    ("service_page", "https://motortrade.com.ph/big-bikes/"),
    ("service_page", "https://motortrade.com.ph/loan-calculator/"),
    ("service_page", "https://motortrade.com.ph/loan-application/"),
    ("service_page", "https://motortrade.com.ph/pre-owned-motorcycle/"),
]

PROMO_QUERY_LABELS = {
    "/motorcycles/?motorcycle-category=automatic&vehicle_promo=yes": "Automatic promo listing",
    "/motorcycles/?motorcycle-category=pang-araw-araw&vehicle_promo=yes": "Pang-araw-araw promo listing",
    "/motorcycles/?motorcycle-category=pang-negosyo&vehicle_promo=yes": "Pang-negosyo promo listing",
    "/motorcycles/?motorcycle-category=pang-sports&vehicle_promo=yes": "Pang-sports promo listing",
}

PUBLIC_FILTER_OPTIONS = {
    "mot-type": [
        ("regular-bike", "Regular Bike"),
        ("big-bike", "Big Bike"),
        ("tri-wheeler", "Tri-Wheeler"),
        ("electric", "Electric"),
    ],
    "brand-name": [
        ("bajaj", "Bajaj"),
        ("honda", "Honda"),
        ("kawasaki", "Kawasaki"),
        ("ktm", "KTM"),
        ("kymco", "Kymco"),
        ("suzuki", "Suzuki"),
        ("yamaha", "Yamaha"),
    ],
    "price-min": [
        ("lowtohigh", "Price: Low to High"),
        ("hightolow", "Price: High to Low"),
    ],
    "prod-tags": [
        ("featured", "Featured"),
        ("toppicks", "Top Picks"),
    ],
    "motorcycle-category": [
        ("automatic", "Automatic"),
        ("pang-araw-araw", "Pang-araw-araw"),
        ("pang-negosyo", "Pang-negosyo"),
        ("pang-sports", "Pang-sports"),
    ],
}

SECTION_SHEETS = [
    "motorcycle_catalog",
    "motorcycle_specs",
    "motorcycle_pricing",
    "branches",
    "branch_coverage_notes",
    "promos",
    "financing",
    "forms",
    "aftersales_services",
    "jobs",
    "parts",
    "preowned",
    "external_links",
    "taxonomy_filters",
    "source_pages",
    "data_quality_notes",
]

SECTION_SCHEMAS: dict[str, list[str]] = {
    "motorcycle_catalog": [
        "display_name",
        "brand_name",
        "model_name",
        "vehicle_type",
        "detail_url",
        "listing_url",
        "srp_amount",
        "downpayment_amount",
        "monthly_amount",
        "price_contexts",
        "promo_titles",
        "is_blue_tag",
        "is_featured",
        "is_top_pick",
        "has_vehicle_promo_flag",
        "spec_make",
        "spec_model",
        "spec_transmission",
        "spec_engine_type",
        "spec_seat_height",
        "spec_brake_system",
        "spec_displacement_cc",
        "spec_front_tire",
        "spec_rear_tire",
        "spec_wheels_type",
        "spec_starting_system",
        "spec_overall_dimensions",
        "spec_ground_clearance",
        "spec_fuel_system",
        "primary_image_url",
        "excerpt_text",
        "apply_loan_url",
        "inquiry_url",
        "service_url",
        "compare_enabled",
        "vehicle_type_source",
        "classification_notes",
        "raw_taxonomies",
        "source_page_url",
        "motorcycle_id",
    ],
    "motorcycle_specs": [
        "display_name",
        "brand_name",
        "model_name",
        "spec_group",
        "spec_label",
        "spec_value_raw",
        "spec_value_numeric",
        "spec_unit",
        "display_order",
        "source_url",
        "motorcycle_id",
    ],
    "motorcycle_pricing": [
        "display_name",
        "brand_name",
        "model_name",
        "source_context",
        "source_url",
        "currency",
        "srp_text",
        "srp_amount",
        "downpayment_text",
        "downpayment_amount",
        "monthly_text",
        "monthly_amount",
        "price_note",
        "is_indicative",
        "captured_at",
        "motorcycle_id",
        "pricing_id",
    ],
    "branches": TABLE_SCHEMAS["branch"],
    "branch_coverage_notes": [
        "note_type",
        "branch_name",
        "coverage_type",
        "coverage_value_raw",
        "coverage_value_normalized",
        "source_field",
        "confidence",
        "source_url",
        "status",
        "evidence_found",
        "rows_created",
        "notes",
        "branch_id",
        "probe_id",
    ],
    "promos": [
        "title",
        "promo_type",
        "description_text",
        "source_url",
        "cta_url",
        "target_listing_url",
        "applies_to_brand",
        "applies_to_vehicle_type",
        "applies_to_tag",
        "has_vehicle_promo_flag",
        "linked_motorcycle_count",
        "linked_motorcycles",
        "link_types",
        "link_source_urls",
        "image_url",
        "published_at",
        "modified_at",
        "source_page_url",
        "promo_id",
    ],
    "financing": [
        "program_name",
        "financing_type",
        "display_name",
        "brand_name",
        "model_name",
        "monthly_amount",
        "downpayment_amount",
        "loan_amount",
        "loan_term_months",
        "estimated_income_amount",
        "source_url",
        "form_url",
        "cta_url",
        "note_text",
        "captured_at",
        "motorcycle_id",
        "financing_id",
    ],
    "forms": [
        "form_name",
        "form_category",
        "source_url",
        "final_url",
        "page_title",
        "form_action",
        "form_method",
        "requires_file_upload",
        "requires_otp",
        "multi_step",
        "has_privacy_checkbox",
        "field_order",
        "field_name",
        "field_label",
        "field_type",
        "is_required",
        "default_value",
        "placeholder",
        "option_value",
        "option_label",
        "section_name",
        "notes",
        "form_id",
    ],
    "aftersales_services": TABLE_SCHEMAS["service_offering"],
    "jobs": TABLE_SCHEMAS["job_listing"],
    "parts": TABLE_SCHEMAS["parts_catalog"],
    "preowned": TABLE_SCHEMAS["preowned_motorcycle"],
    "external_links": TABLE_SCHEMAS["external_cta"],
    "taxonomy_filters": [
        "row_type",
        "taxonomy_or_filter",
        "term_or_option_value",
        "label",
        "parent_term_id",
        "count",
        "link_or_source_url",
        "used_by_post_types",
        "notes",
    ],
    "source_pages": TABLE_SCHEMAS["source_page"],
    "data_quality_notes": [
        "note_type",
        "entity_type",
        "entity_id",
        "source_url",
        "issue_type",
        "issue_detail",
        "severity",
        "relationship_name",
        "status",
        "evidence_found",
        "rows_created",
        "notes",
        "captured_at",
    ],
}

DATABASE_SHEETS = [
    "dealer",
    "crawl_run",
    "source_page",
    "motorcycle",
    "motorcycle_spec",
    "pricing",
    "financing",
    "branch",
    "branch_product_coverage",
    "branch_model_availability",
    "promo",
    "motorcycle_promo_link",
    "aftersales_service",
    "form_definition",
    "form_field",
    "job_listing",
    "parts_catalog",
    "preowned_motorcycle",
    "external_cta",
    "website_filter_reference",
    "relationship_probe",
    "data_quality_issue",
]

DATABASE_SCHEMAS: dict[str, list[str]] = {
    "dealer": TABLE_SCHEMAS["dealer"],
    "crawl_run": TABLE_SCHEMAS["crawl_run"],
    "source_page": TABLE_SCHEMAS["source_page"],
    "motorcycle": [
        "motorcycle_id",
        "dealer_id",
        "display_name",
        "brand_name",
        "model_name",
        "vehicle_type",
        "detail_url",
        "listing_url",
        "srp_amount",
        "downpayment_amount",
        "monthly_amount",
        "price_contexts",
        "promo_titles",
        "is_blue_tag",
        "is_featured",
        "is_top_pick",
        "has_vehicle_promo_flag",
        "spec_make",
        "spec_model",
        "spec_transmission",
        "spec_engine_type",
        "spec_seat_height",
        "spec_brake_system",
        "spec_displacement_cc",
        "spec_front_tire",
        "spec_rear_tire",
        "spec_wheels_type",
        "spec_starting_system",
        "spec_overall_dimensions",
        "spec_ground_clearance",
        "spec_fuel_system",
        "primary_image_url",
        "excerpt_text",
        "apply_loan_url",
        "inquiry_url",
        "service_url",
        "compare_enabled",
        "published_at",
        "modified_at",
        "vehicle_type_source",
        "classification_notes",
        "raw_taxonomies",
        "source_page_id",
        "source_page_url",
    ],
    "motorcycle_spec": [
        "motorcycle_id",
        "dealer_id",
        "display_name",
        "brand_name",
        "model_name",
        "spec_group",
        "spec_label",
        "spec_value_raw",
        "spec_value_numeric",
        "spec_unit",
        "display_order",
        "source_url",
    ],
    "pricing": [
        "pricing_id",
        "motorcycle_id",
        "dealer_id",
        "display_name",
        "brand_name",
        "model_name",
        "source_url",
        "source_context",
        "currency",
        "srp_text",
        "srp_amount",
        "downpayment_text",
        "downpayment_amount",
        "monthly_text",
        "monthly_amount",
        "price_note",
        "is_indicative",
        "captured_at",
    ],
    "financing": [
        "financing_id",
        "motorcycle_id",
        "dealer_id",
        "display_name",
        "brand_name",
        "model_name",
        "source_url",
        "financing_type",
        "program_name",
        "monthly_amount",
        "downpayment_amount",
        "loan_amount",
        "loan_term_months",
        "estimated_income_amount",
        "form_url",
        "cta_url",
        "note_text",
        "captured_at",
    ],
    "branch": TABLE_SCHEMAS["branch"],
    "branch_product_coverage": TABLE_SCHEMAS["branch_product_coverage"],
    "branch_model_availability": TABLE_SCHEMAS["branch_model_availability"],
    "promo": [
        "promo_id",
        "dealer_id",
        "source_page_id",
        "source_page_url",
        "source_url",
        "title",
        "promo_type",
        "description_text",
        "image_url",
        "cta_url",
        "target_listing_url",
        "applies_to_brand",
        "applies_to_vehicle_type",
        "applies_to_tag",
        "has_vehicle_promo_flag",
        "linked_motorcycle_count",
        "linked_motorcycles",
        "link_types",
        "link_source_urls",
        "published_at",
        "modified_at",
    ],
    "motorcycle_promo_link": [
        "motorcycle_id",
        "promo_id",
        "dealer_id",
        "display_name",
        "promo_title",
        "link_type",
        "source_url",
        "notes",
    ],
    "aftersales_service": TABLE_SCHEMAS["service_offering"],
    "form_definition": TABLE_SCHEMAS["form_definition"],
    "form_field": [
        "form_id",
        "dealer_id",
        "form_name",
        "form_category",
        "field_order",
        "field_name",
        "field_label",
        "field_type",
        "is_required",
        "default_value",
        "placeholder",
        "option_value",
        "option_label",
        "section_name",
        "source_url",
    ],
    "job_listing": TABLE_SCHEMAS["job_listing"],
    "parts_catalog": TABLE_SCHEMAS["parts_catalog"],
    "preowned_motorcycle": TABLE_SCHEMAS["preowned_motorcycle"],
    "external_cta": TABLE_SCHEMAS["external_cta"],
    "website_filter_reference": [
        "dealer_id",
        "reference_type",
        "filter_name",
        "option_value",
        "option_label",
        "source_url",
        "notes",
    ],
    "relationship_probe": TABLE_SCHEMAS["relationship_probe"],
    "data_quality_issue": TABLE_SCHEMAS["data_quality_issue"],
}


@dataclass
class HtmlPage:
    url: str
    final_url: str
    status_code: int
    title: str
    html: str
    soup: BeautifulSoup


class ScrapeError(Exception):
    pass


class TableStore:
    def __init__(self, schemas: dict[str, list[str]]):
        self.schemas = schemas
        self.rows = {name: [] for name in schemas}

    def append(self, table: str, row: dict[str, object]) -> None:
        schema = self.schemas[table]
        clean = {column: sanitize_excel_value(row.get(column)) for column in schema}
        self.rows[table].append(clean)

    def row_count(self, table: str) -> int:
        return len(self.rows[table])


class MotortradeAnalyticsScraper:
    def __init__(self, output_path: Path, output_style: str = "section"):
        self.output_path = output_path
        self.output_style = output_style
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": USER_AGENT})
        self.tables = TableStore(TABLE_SCHEMAS)
        self.media_cache: dict[int, str] = {}
        self.source_page_by_url: dict[str, str] = {}
        self.source_page_publish_meta: dict[str, tuple[str, str]] = {}
        self.taxonomy_usage: dict[str, set[str]] = defaultdict(set)
        self.term_rows_seen: set[tuple[str, str]] = set()
        self.motorcycle_rows: dict[str, dict[str, object]] = {}
        self.motorcycle_pricing_seen: set[tuple[str, str]] = set()
        self.promo_rows_by_url: dict[str, str] = {}
        self.taxonomy_assignment_seen: set[tuple[str, str, str, str]] = set()
        self.external_cta_seen: set[tuple[str, str]] = set()
        self.branch_product_coverage_seen: set[tuple[str, str, str]] = set()
        self.source_counter = 0

        now = datetime.now(TIMEZONE)
        self.crawl_date = now.date().isoformat()
        self.run_started_at = now.isoformat()
        self.run_id = f"{DEALER_ID}-run-{now.strftime('%Y%m%d%H%M%S')}"

    def build_motorcycle_row(
        self,
        detail_url: str,
        *,
        source_page_id: str = "",
        listing_url: str = "",
        display_name: str = "",
        brand_name: str = "",
        model_name: str = "",
        primary_image_url: str = "",
        excerpt_text: str = "",
        compare_enabled: bool = False,
    ) -> dict[str, object]:
        return {
            "motorcycle_id": f"motorcycle-{slug_from_url(detail_url)}",
            "dealer_id": DEALER_ID,
            "source_page_id": source_page_id,
            "detail_url": detail_url,
            "listing_url": listing_url,
            "display_name": display_name,
            "brand_name": brand_name,
            "model_name": model_name,
            "vehicle_type": "",
            "vehicle_type_source": "",
            "raw_taxonomies": json.dumps([]),
            "classification_notes": "",
            "spec_make": "",
            "spec_model": "",
            "spec_transmission": "",
            "spec_engine_type": "",
            "spec_seat_height": "",
            "spec_brake_system": "",
            "spec_displacement_cc": "",
            "spec_front_tire": "",
            "spec_rear_tire": "",
            "spec_wheels_type": "",
            "spec_starting_system": "",
            "spec_overall_dimensions": "",
            "spec_ground_clearance": "",
            "spec_fuel_system": "",
            "is_blue_tag": False,
            "is_featured": False,
            "is_top_pick": False,
            "has_vehicle_promo_flag": False,
            "primary_image_url": primary_image_url,
            "excerpt_text": excerpt_text,
            "apply_loan_url": "",
            "inquiry_url": "",
            "service_url": "",
            "compare_enabled": compare_enabled,
            "published_at": "",
            "modified_at": "",
        }

    def record_data_quality_issue(
        self,
        *,
        entity_type: str,
        entity_id: str,
        source_url: str,
        issue_type: str,
        issue_detail: str,
        severity: str = "warning",
    ) -> None:
        self.tables.append(
            "data_quality_issue",
            {
                "issue_id": f"dq-{self.tables.row_count('data_quality_issue') + 1:05d}",
                "dealer_id": DEALER_ID,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "source_url": normalize_url(source_url) if source_url else "",
                "issue_type": issue_type,
                "issue_detail": issue_detail,
                "severity": severity,
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def record_taxonomy_assignment(
        self,
        *,
        entity_type: str,
        entity_id: str,
        taxonomy: str,
        term_slug: str,
        term_label: str,
        source_url: str,
        confidence: str,
        notes: str,
    ) -> None:
        key = (entity_type, entity_id, taxonomy, term_slug)
        if key in self.taxonomy_assignment_seen:
            return
        self.taxonomy_assignment_seen.add(key)
        self.tables.append(
            "taxonomy_assignment",
            {
                "assignment_id": f"taxassign-{len(self.taxonomy_assignment_seen):05d}",
                "dealer_id": DEALER_ID,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "taxonomy": taxonomy,
                "term_slug": term_slug,
                "term_label": term_label,
                "source_url": normalize_url(source_url),
                "confidence": confidence,
                "notes": notes,
            },
        )

    def assign_vehicle_type(
        self,
        row: dict[str, object],
        *,
        vehicle_type: str,
        source_url: str,
        source_kind: str,
        confidence: str = "high",
    ) -> None:
        vehicle_type = vehicle_type.strip()
        if not vehicle_type:
            return
        motorcycle_id = str(row["motorcycle_id"])
        self.record_taxonomy_assignment(
            entity_type="motorcycle",
            entity_id=motorcycle_id,
            taxonomy="motorcycle-type",
            term_slug=vehicle_type,
            term_label=vehicle_type.replace("-", " ").title(),
            source_url=source_url,
            confidence=confidence,
            notes=f"Detected from {source_kind}.",
        )
        raw_taxonomies = json_loads_list(row.get("raw_taxonomies"))
        raw_taxonomies.append(
            {
                "taxonomy": "motorcycle-type",
                "term_slug": vehicle_type,
                "source": source_kind,
                "source_url": normalize_url(source_url),
            }
        )
        row["raw_taxonomies"] = json.dumps(dedupe_json_objects(raw_taxonomies))
        current = str(row.get("vehicle_type") or "")
        if not current:
            row["vehicle_type"] = vehicle_type
            row["vehicle_type_source"] = source_kind
            return
        if current != vehicle_type:
            row["classification_notes"] = append_note(
                row.get("classification_notes"),
                f"Conflicting vehicle type `{vehicle_type}` from {source_kind}; kept `{current}`.",
            )
            self.record_data_quality_issue(
                entity_type="motorcycle",
                entity_id=motorcycle_id,
                source_url=source_url,
                issue_type="vehicle_type_conflict",
                issue_detail=f"Existing vehicle_type `{current}` conflicts with `{vehicle_type}` from {source_kind}.",
            )

    def record_external_ctas(self, page: HtmlPage, source_page_id: str, cta_type: str) -> None:
        allowed_domains = {"motortrade.com.ph", "www.motortrade.com.ph"}
        for anchor in page.soup.select("a[href]"):
            href = anchor.get("href", "").strip()
            if not href or href.startswith(("#", "javascript:", "mailto:", "tel:")):
                continue
            target_url = normalize_url(urljoin(page.final_url, href))
            domain = urlparse(target_url).netloc.lower()
            if domain in allowed_domains:
                continue
            key = (source_page_id, target_url)
            if key in self.external_cta_seen:
                continue
            self.external_cta_seen.add(key)
            self.tables.append(
                "external_cta",
                {
                    "cta_id": f"cta-{len(self.external_cta_seen):05d}",
                    "dealer_id": DEALER_ID,
                    "source_page_id": source_page_id,
                    "source_url": page.final_url,
                    "target_url": target_url,
                    "target_domain": domain,
                    "cta_type": cta_type,
                    "anchor_text": anchor.get_text(" ", strip=True),
                    "notes": "Public link captured as CTA; target was not scraped as a data source unless separately in scope.",
                },
            )

    def record_branch_product_coverage(
        self,
        *,
        branch_id: str,
        coverage_type: str,
        coverage_value_raw: object,
        source_url: str,
        source_field: str,
        confidence: str,
        notes: str,
    ) -> None:
        normalized_values = normalize_coverage_values(coverage_value_raw)
        if not normalized_values:
            return
        raw_text = json.dumps(coverage_value_raw)
        for normalized in normalized_values:
            key = (branch_id, coverage_type, normalized)
            if key in self.branch_product_coverage_seen:
                continue
            self.branch_product_coverage_seen.add(key)
            self.tables.append(
                "branch_product_coverage",
                {
                    "coverage_id": f"coverage-{len(self.branch_product_coverage_seen):05d}",
                    "dealer_id": DEALER_ID,
                    "branch_id": branch_id,
                    "coverage_type": coverage_type,
                    "coverage_value_raw": raw_text,
                    "coverage_value_normalized": normalized,
                    "source_url": normalize_url(source_url),
                    "source_field": source_field,
                    "confidence": confidence,
                    "observed_at": datetime.now(TIMEZONE).isoformat(),
                    "notes": notes,
                },
            )

    def record_relationship_probe(
        self,
        *,
        relationship_name: str,
        source_url: str,
        source_type: str,
        status: str,
        evidence_found: bool,
        rows_created: int,
        notes: str,
    ) -> None:
        self.tables.append(
            "relationship_probe",
            {
                "probe_id": f"probe-{self.tables.row_count('relationship_probe') + 1:05d}",
                "dealer_id": DEALER_ID,
                "relationship_name": relationship_name,
                "source_url": normalize_url(source_url) if source_url else "",
                "source_type": source_type,
                "status": status,
                "evidence_found": evidence_found,
                "rows_created": rows_created,
                "notes": notes,
                "checked_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def run(self) -> None:
        self.tables.append(
            "dealer",
            {
                "dealer_id": DEALER_ID,
                "dealer_name": DEALER_NAME,
                "primary_domain": MOTORTRADE_DOMAIN,
                "forms_domain": FORMS_DOMAIN,
                "crawl_timezone": "Asia/Manila",
                "crawl_date": self.crawl_date,
                "notes": "Public Motortrade scrape across WordPress, locator feed, motorcycle detail pages, and forms.",
            },
        )

        types_payload = self.fetch_json("https://motortrade.com.ph/wp-json/wp/v2/types")
        taxonomies_payload = self.fetch_json("https://motortrade.com.ph/wp-json/wp/v2/taxonomies")
        self.collect_taxonomy_terms(types_payload, taxonomies_payload)
        rest_items = self.collect_rest_objects()
        self.collect_filter_options_and_listing_pricing()
        self.collect_locator_branches()
        self.collect_motorcycles_from_vehicle_sitemap()
        self.collect_service_and_promo_pages(rest_items)
        self.collect_forms()
        self.finalize_crawl_run()
        self.write_workbook()

    def finalize_crawl_run(self) -> None:
        finished = datetime.now(TIMEZONE).isoformat()
        self.tables.append(
            "crawl_run",
            {
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "run_started_at": self.run_started_at,
                "run_finished_at": finished,
                "source_count": self.source_counter,
                "notes": "Current-state public snapshot with raw and normalized fields preserved.",
            },
        )

    def write_workbook(self) -> None:
        if self.output_style == "analytics":
            self.write_analytics_workbook()
            return
        if self.output_style == "database":
            self.write_database_workbook()
            return
        self.write_section_workbook()

    def write_analytics_workbook(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        for table_name, schema in TABLE_SCHEMAS.items():
            sheet = workbook.create_sheet(title=table_name[:31])
            sheet.append(schema)
            for row in self.tables.rows[table_name]:
                sheet.append([row.get(column, "") for column in schema])

        workbook.save(self.output_path)

    def write_section_workbook(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        section_rows = build_section_rows(self.tables.rows)

        for sheet_name in SECTION_SHEETS:
            schema = SECTION_SCHEMAS[sheet_name]
            sheet = workbook.create_sheet(title=sheet_name[:31])
            sheet.append(schema)
            for row in section_rows.get(sheet_name, []):
                sheet.append([sanitize_excel_value(row.get(column)) for column in schema])

        workbook.save(self.output_path)

    def write_database_workbook(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        database_rows = build_database_rows(self.tables.rows)

        for sheet_name in DATABASE_SHEETS:
            schema = DATABASE_SCHEMAS[sheet_name]
            sheet = workbook.create_sheet(title=sheet_name[:31])
            sheet.append(schema)
            for row in database_rows.get(sheet_name, []):
                sheet.append([sanitize_excel_value(row.get(column)) for column in schema])

        workbook.save(self.output_path)

    def fetch(self, url: str, *, params: dict[str, object] | None = None, timeout: int = 30) -> requests.Response:
        response = self.session.get(url, params=params, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
        return response

    def fetch_json(self, url: str, *, params: dict[str, object] | None = None) -> object:
        return self.fetch(url, params=params).json()

    def fetch_html(self, url: str) -> HtmlPage:
        response = self.fetch(url)
        soup = BeautifulSoup(response.text, "lxml")
        self.source_counter += 1
        return HtmlPage(
            url=url,
            final_url=normalize_url(response.url),
            status_code=response.status_code,
            title=(soup.title.string.strip() if soup.title and soup.title.string else ""),
            html=response.text,
            soup=soup,
        )

    def collect_rest_objects(self) -> dict[str, list[dict[str, object]]]:
        objects: dict[str, list[dict[str, object]]] = {}
        for rest_base, (post_type_slug, _) in REST_COLLECTIONS.items():
            endpoint = f"{MOTORTRADE_DOMAIN}/wp-json/wp/v2/{rest_base}"
            rows = self.fetch_all_rest_rows(endpoint)
            objects[rest_base] = rows
            if rest_base in {"categories", "tags", "job-types"}:
                self.collect_rest_terms(rows, rest_base, post_type_slug)
                continue
            for item in rows:
                self.record_rest_source_page(item, post_type_slug)
                if rest_base == "job-listings":
                    self.record_job_listing(item)
                elif rest_base == "parts":
                    self.record_parts_catalog(item)
                elif rest_base == "preowned":
                    self.record_preowned(item)
        return objects

    def fetch_all_rest_rows(self, endpoint: str) -> list[dict[str, object]]:
        page = 1
        rows: list[dict[str, object]] = []
        while True:
            response = self.fetch(endpoint, params={"per_page": REST_PAGE_SIZE, "page": page})
            payload = response.json()
            if not isinstance(payload, list):
                break
            rows.extend(payload)
            total_pages = int(response.headers.get("X-WP-TotalPages", "1"))
            if page >= total_pages:
                break
            page += 1
        self.source_counter += 1
        return rows

    def collect_rest_terms(self, rows: list[dict[str, object]], rest_base: str, taxonomy_slug: str) -> None:
        taxonomy_name = {
            "categories": "category",
            "tags": "post_tag",
            "job-types": "job_listing_type",
        }[rest_base]
        used_by = {
            "category": ["page", "post", "featured_products"],
            "post_tag": ["page", "post", "featured_products"],
            "job_listing_type": ["job_listing"],
        }[taxonomy_name]
        for term in rows:
            term_id = str(term.get("id"))
            self.term_rows_seen.add((taxonomy_name, term_id))
            self.tables.append(
                "taxonomy_term",
                {
                    "term_id": term_id,
                    "taxonomy": taxonomy_name,
                    "name": html_to_text(term.get("name", "")),
                    "slug": term.get("slug", ""),
                    "parent_term_id": null_if_zero(term.get("parent")),
                    "count": term.get("count"),
                    "link": term.get("link", ""),
                    "used_by_post_types": json.dumps(used_by),
                },
            )

    def collect_taxonomy_terms(self, types_payload: object, taxonomies_payload: object) -> None:
        if isinstance(types_payload, dict):
            for type_data in types_payload.values():
                if not isinstance(type_data, dict):
                    continue
                post_type_slug = str(type_data.get("slug", ""))
                for taxonomy in type_data.get("taxonomies", []) or []:
                    self.taxonomy_usage[str(taxonomy)].add(post_type_slug)

        if isinstance(taxonomies_payload, dict):
            for taxonomy_slug, info in taxonomies_payload.items():
                rest_base = info.get("rest_base")
                if not rest_base:
                    continue
                if rest_base in {"categories", "tags", "job-types"}:
                    continue
                endpoint = f"{MOTORTRADE_DOMAIN}/wp-json/wp/v2/{rest_base}"
                try:
                    rows = self.fetch_all_rest_rows(endpoint)
                except requests.RequestException:
                    continue
                for term in rows:
                    term_id = str(term.get("id"))
                    self.term_rows_seen.add((taxonomy_slug, term_id))
                    self.tables.append(
                        "taxonomy_term",
                        {
                            "term_id": term_id,
                            "taxonomy": taxonomy_slug,
                            "name": html_to_text(term.get("name", "")),
                            "slug": term.get("slug", ""),
                            "parent_term_id": null_if_zero(term.get("parent")),
                            "count": term.get("count"),
                            "link": term.get("link", ""),
                            "used_by_post_types": json.dumps(sorted(self.taxonomy_usage.get(taxonomy_slug, []))),
                        },
                    )

        for sitemap_name, taxonomy_slug in (
            ("motorcycle-category-sitemap.xml", "motorcycle-category"),
            ("motorcycle-type-sitemap.xml", "motorcycle-type"),
            ("vehicle_promo-sitemap.xml", "vehicle_promo"),
            ("make-sitemap.xml", "make"),
            ("model-sitemap.xml", "model"),
            ("agency-sitemap.xml", "agency"),
            ("agency-category-sitemap.xml", "agency-category"),
        ):
            for url in self.fetch_sitemap_urls(f"{MOTORTRADE_DOMAIN}/{sitemap_name}"):
                parsed = urlparse(url)
                slug = parsed.path.rstrip("/").split("/")[-1]
                name = slug.replace("-", " ").replace("_", " ").title()
                term_id = slug or url
                if (taxonomy_slug, term_id) in self.term_rows_seen:
                    continue
                self.term_rows_seen.add((taxonomy_slug, term_id))
                self.tables.append(
                    "taxonomy_term",
                    {
                        "term_id": term_id,
                        "taxonomy": taxonomy_slug,
                        "name": name,
                        "slug": slug,
                        "parent_term_id": "",
                        "count": "",
                        "link": url,
                        "used_by_post_types": json.dumps([]),
                    },
                )

    def record_rest_source_page(self, item: dict[str, object], wp_post_type: str) -> str:
        url = normalize_url(str(item.get("link", "")))
        source_page_id = self.source_page_by_url.get(url)
        if source_page_id:
            return source_page_id
        source_page_id = f"sp-{len(self.source_page_by_url) + 1:05d}"
        self.source_page_by_url[url] = source_page_id
        published_at = str(item.get("date", ""))
        modified_at = str(item.get("modified", ""))
        self.source_page_publish_meta[url] = (published_at, modified_at)

        category_ids = item.get("categories", []) or []
        tag_ids = item.get("tags", []) or []
        job_type_ids = item.get("job-types", []) or item.get("job_types", []) or []
        featured_media_id = int(item.get("featured_media", 0) or 0)
        featured_image_url = self.resolve_featured_image(item, featured_media_id)
        yoast_json = item.get("yoast_head_json", {}) or {}

        self.tables.append(
            "source_page",
            {
                "source_page_id": source_page_id,
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "source_type": "wordpress_rest",
                "wp_post_type": wp_post_type,
                "wp_id": item.get("id", ""),
                "url": url,
                "slug": item.get("slug", ""),
                "title": html_to_text(get_rendered(item.get("title"))),
                "status": item.get("status", ""),
                "published_at": published_at,
                "modified_at": modified_at,
                "author_id": item.get("author", ""),
                "excerpt_text": html_to_text(get_rendered(item.get("excerpt"))),
                "content_html": get_rendered(item.get("content")),
                "content_text": html_to_text(get_rendered(item.get("content"))),
                "featured_media_id": featured_media_id or "",
                "featured_image_url": featured_image_url,
                "category_ids": json.dumps(category_ids),
                "tag_ids": json.dumps(tag_ids),
                "job_type_ids": json.dumps(job_type_ids),
                "seo_title": yoast_json.get("title", ""),
                "seo_description": yoast_json.get("description", "") or yoast_json.get("og_description", ""),
                "canonical_url": yoast_json.get("canonical", url),
            },
        )
        return source_page_id

    def record_html_source_page(
        self,
        page: HtmlPage,
        *,
        source_type: str,
        wp_post_type: str = "",
        title_override: str = "",
        content_html: str = "",
        content_text: str = "",
        published_at: str = "",
        modified_at: str = "",
        excerpt_text: str = "",
        featured_image_url: str = "",
        seo_title: str = "",
        seo_description: str = "",
        canonical_url: str = "",
    ) -> str:
        url = normalize_url(page.final_url)
        source_page_id = self.source_page_by_url.get(url)
        if source_page_id:
            return source_page_id
        source_page_id = f"sp-{len(self.source_page_by_url) + 1:05d}"
        self.source_page_by_url[url] = source_page_id
        self.source_page_publish_meta[url] = (published_at, modified_at)

        soup = page.soup
        title = title_override or page.title or text_or_blank(soup.find("h1"))
        if not content_html:
            main = soup.select_one(".vehicle-content") or soup.select_one(".page-content") or soup.select_one(".container")
            content_html = str(main) if main else page.html
        if not content_text:
            content_text = html_to_text(content_html)
        if not excerpt_text:
            excerpt_text = content_text[:500]

        self.tables.append(
            "source_page",
            {
                "source_page_id": source_page_id,
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "source_type": source_type,
                "wp_post_type": wp_post_type,
                "wp_id": "",
                "url": url,
                "slug": slug_from_url(url),
                "title": title,
                "status": "publish" if page.status_code == 200 else page.status_code,
                "published_at": published_at,
                "modified_at": modified_at,
                "author_id": "",
                "excerpt_text": excerpt_text,
                "content_html": content_html,
                "content_text": content_text,
                "featured_media_id": "",
                "featured_image_url": featured_image_url,
                "category_ids": json.dumps([]),
                "tag_ids": json.dumps([]),
                "job_type_ids": json.dumps([]),
                "seo_title": seo_title or meta_content(soup, "title"),
                "seo_description": seo_description or meta_content(soup, "description"),
                "canonical_url": canonical_url or canonical_href(soup) or url,
            },
        )
        return source_page_id

    def resolve_featured_image(self, item: dict[str, object], featured_media_id: int) -> str:
        yoast_json = item.get("yoast_head_json", {}) or {}
        og_image = yoast_json.get("og_image")
        if isinstance(og_image, list) and og_image:
            first = og_image[0]
            if isinstance(first, dict):
                return first.get("url", "")
        if featured_media_id <= 0:
            return ""
        if featured_media_id in self.media_cache:
            return self.media_cache[featured_media_id]
        try:
            media = self.fetch_json(f"{MOTORTRADE_DOMAIN}/wp-json/wp/v2/media/{featured_media_id}")
        except requests.RequestException:
            self.media_cache[featured_media_id] = ""
            return ""
        image_url = ""
        if isinstance(media, dict):
            image_url = str(media.get("source_url", ""))
        self.media_cache[featured_media_id] = image_url
        return image_url

    def record_job_listing(self, item: dict[str, object]) -> None:
        source_page_id = self.record_rest_source_page(item, "job_listing")
        meta = item.get("meta", {}) or {}
        self.tables.append(
            "job_listing",
            {
                "job_id": f"job-{item.get('id')}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "title": html_to_text(get_rendered(item.get("title"))),
                "slug": item.get("slug", ""),
                "job_location": meta.get("_job_location", ""),
                "application_target": meta.get("_application", ""),
                "company_name": meta.get("_company_name", ""),
                "company_website": meta.get("_company_website", ""),
                "company_tagline": meta.get("_company_tagline", ""),
                "job_salary": meta.get("_job_salary", ""),
                "salary_currency": meta.get("_job_salary_currency", ""),
                "salary_unit": meta.get("_job_salary_unit", ""),
                "is_featured": truthy(meta.get("_featured")),
                "is_filled": truthy(meta.get("_filled")),
                "is_remote": truthy(meta.get("_remote_position")),
                "job_type_ids": json.dumps(item.get("job-types", []) or []),
                "description_text": html_to_text(get_rendered(item.get("content"))),
                "published_at": item.get("date", ""),
                "modified_at": item.get("modified", ""),
            },
        )

    def record_parts_catalog(self, item: dict[str, object]) -> None:
        source_page_id = self.record_rest_source_page(item, "parts")
        self.tables.append(
            "parts_catalog",
            {
                "part_id": f"part-{item.get('id')}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "title": html_to_text(get_rendered(item.get("title"))),
                "slug": item.get("slug", ""),
                "detail_url": item.get("link", ""),
                "image_url": self.resolve_featured_image(item, int(item.get("featured_media", 0) or 0)),
                "description_text": html_to_text(get_rendered(item.get("content"))),
                "published_at": item.get("date", ""),
                "modified_at": item.get("modified", ""),
            },
        )

    def record_preowned(self, item: dict[str, object]) -> None:
        source_page_id = self.record_rest_source_page(item, "preowned")
        self.tables.append(
            "preowned_motorcycle",
            {
                "preowned_id": f"preowned-{item.get('id')}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "title": html_to_text(get_rendered(item.get("title"))),
                "slug": item.get("slug", ""),
                "detail_url": item.get("link", ""),
                "image_url": self.resolve_featured_image(item, int(item.get("featured_media", 0) or 0)),
                "description_text": html_to_text(get_rendered(item.get("content"))),
                "published_at": item.get("date", ""),
                "modified_at": item.get("modified", ""),
            },
        )

    def collect_filter_options_and_listing_pricing(self) -> None:
        archive = self.fetch_html(f"{MOTORTRADE_DOMAIN}/motorcycles/")
        archive_source_id = self.record_html_source_page(
            archive,
            source_type="motorcycle_archive",
            title_override="Motorcycles Archive",
            content_html=str(archive.soup.select_one(".widget") or archive.soup.select_one("form") or archive.soup),
            content_text=html_to_text(str(archive.soup.select_one(".widget") or archive.soup.select_one("form") or archive.soup)),
        )
        filter_rows_before = self.tables.row_count("filter_option")
        for select in archive.soup.select("form select[name]"):
            filter_name = select.get("name", "").strip()
            if filter_name not in {"mot-type", "brand-name", "price-min", "prod-tags"}:
                continue
            for option in select.select("option"):
                value = option.get("value", "").strip()
                label = option.get_text(" ", strip=True)
                self.tables.append(
                    "filter_option",
                    {
                        "filter_name": filter_name,
                        "option_value": value,
                        "option_label": label,
                        "source_url": archive.final_url,
                        "notes": "Motorcycles archive filter option.",
                    },
                )
        if self.tables.row_count("filter_option") == filter_rows_before:
            for filter_name, options in PUBLIC_FILTER_OPTIONS.items():
                for value, label in options:
                    self.tables.append(
                        "filter_option",
                        {
                            "filter_name": filter_name,
                            "option_value": value,
                            "option_label": label,
                            "source_url": archive.final_url,
                            "notes": "Public query/taxonomy option; current rendered archive does not expose these as select controls.",
                        },
                    )

        cards = archive.soup.select(".item")
        for index, card in enumerate(cards, start=1):
            detail_link = absolute_href(archive.final_url, card.select_one("a[href*='/motorcycles/']"))
            if not detail_link or "/motorcycles/" not in detail_link:
                continue
            detail_link = normalize_url(detail_link)
            title_link = card.select_one(".Vehicle-Title a")
            display_name = text_or_blank(title_link) or text_or_blank(card.select_one(".Vehicle-Title"))
            excerpt = text_or_blank(card.select_one(".vehicle-feature-des"))
            image_url = image_src(card.select_one("img"))
            compare_enabled = bool(card.select_one(".btn_detail_compare"))
            srp_text = text_or_blank(card.select_one(".pcd-price"))
            downpayment_text = text_or_blank(card.select_one(".pcd-downpayment"))
            monthly_text = text_or_blank(card.select_one(".pcd-price-msrp"))

            if detail_link not in self.motorcycle_rows:
                self.motorcycle_rows[detail_link] = self.build_motorcycle_row(
                    detail_link,
                    source_page_id=archive_source_id,
                    listing_url=archive.final_url,
                    display_name=display_name,
                    primary_image_url=image_url,
                    excerpt_text=excerpt,
                    compare_enabled=compare_enabled,
                )

            self.record_pricing(
                motorcycle_id=self.motorcycle_rows[detail_link]["motorcycle_id"],
                source_url=archive.final_url,
                source_context="motorcycle_archive_card",
                srp_text=srp_text,
                downpayment_text=downpayment_text,
                monthly_text=monthly_text,
                note="Pricing captured from the public motorcycles archive card.",
            )

        self.collect_listing_flags_and_promos()

    def collect_listing_flags_and_promos(self) -> None:
        query_configs = [
            ("featured", {"prod-tags": "featured"}, "Featured listing filter"),
            ("top_picks", {"prod-tags": "toppicks"}, "Top picks listing filter"),
            ("vehicle_promo", {"motorcycle-category": "automatic", "vehicle_promo": "yes"}, "Automatic promo listing"),
            ("vehicle_promo", {"motorcycle-category": "pang-araw-araw", "vehicle_promo": "yes"}, "Pang-araw-araw promo listing"),
            ("vehicle_promo", {"motorcycle-category": "pang-negosyo", "vehicle_promo": "yes"}, "Pang-negosyo promo listing"),
            ("vehicle_promo", {"motorcycle-category": "pang-sports", "vehicle_promo": "yes"}, "Pang-sports promo listing"),
        ]
        for vehicle_type in ("regular-bike", "big-bike", "tri-wheeler", "electric"):
            query_configs.append(
                (
                    "vehicle_type",
                    {"mot-type": vehicle_type},
                    f"{vehicle_type.replace('-', ' ').title()} listing filter",
                )
            )
        for flag_type, query_params, label in query_configs:
            url = normalize_url(f"{MOTORTRADE_DOMAIN}/motorcycles/?{urlencode(query_params)}")
            try:
                page = self.fetch_html(url)
            except requests.RequestException:
                continue
            promo_id = ""
            source_page_id = ""
            if flag_type == "vehicle_promo":
                source_page_id = self.record_html_source_page(
                    page,
                    source_type="promo_listing",
                    title_override=label,
                    content_html=str(page.soup.select_one(".vehicle-section-content") or page.soup),
                    content_text=html_to_text(str(page.soup.select_one(".vehicle-section-content") or page.soup)),
                )
                promo_id = self.record_promo_surface(
                    source_url=url,
                    source_page_id=source_page_id,
                    title=label,
                    promo_type="vehicle_promo_query",
                    description_text=f"Query-driven promo listing for {query_params.get('motorcycle-category', '')}.",
                    image_url="",
                    cta_url="",
                    target_listing_url=url,
                    applies_to_brand="",
                    applies_to_vehicle_type=query_params.get("motorcycle-category", ""),
                    applies_to_tag="",
                    has_vehicle_promo_flag=True,
                    published_at="",
                    modified_at="",
                )

            for detail_url in extract_motorcycle_links(page.soup, page.final_url):
                detail_url = normalize_url(detail_url)
                row = self.motorcycle_rows.setdefault(
                    detail_url,
                    self.build_motorcycle_row(detail_url, listing_url=page.final_url),
                )
                if flag_type == "featured":
                    row["is_featured"] = True
                elif flag_type == "top_picks":
                    row["is_top_pick"] = True
                elif flag_type == "vehicle_type":
                    self.assign_vehicle_type(
                        row,
                        vehicle_type=str(query_params.get("mot-type", "")),
                        source_url=page.final_url,
                        source_kind="mot-type_filter",
                    )
                elif flag_type == "vehicle_promo":
                    row["has_vehicle_promo_flag"] = True
                    if promo_id:
                        self.tables.append(
                            "motorcycle_promo_link",
                            {
                                "motorcycle_id": row["motorcycle_id"],
                                "promo_id": promo_id,
                                "link_type": "query_listing",
                                "source_url": page.final_url,
                                "notes": "Motorcycle appeared in a public vehicle promo query listing.",
                            },
                        )

        for type_url in self.fetch_sitemap_urls(f"{MOTORTRADE_DOMAIN}/motorcycle-type-sitemap.xml"):
            try:
                page = self.fetch_html(type_url)
            except requests.RequestException:
                continue
            vehicle_type = slug_from_url(type_url)
            for detail_url in extract_motorcycle_links(page.soup, page.final_url):
                detail_url = normalize_url(detail_url)
                row = self.motorcycle_rows.setdefault(
                    detail_url,
                    self.build_motorcycle_row(detail_url, listing_url=page.final_url),
                )
                self.assign_vehicle_type(
                    row,
                    vehicle_type=vehicle_type,
                    source_url=page.final_url,
                    source_kind="motorcycle-type_taxonomy_page",
                )

        for tax_url in self.fetch_sitemap_urls(f"{MOTORTRADE_DOMAIN}/vehicle_promo-sitemap.xml"):
            try:
                page = self.fetch_html(tax_url)
            except requests.RequestException:
                continue
            slug = slug_from_url(tax_url)
            source_page_id = self.record_html_source_page(
                page,
                source_type="promo_taxonomy",
                title_override=page.title or slug.replace("-", " ").title(),
                content_html=str(page.soup.select_one(".container") or page.soup),
                content_text=html_to_text(str(page.soup.select_one(".container") or page.soup)),
            )
            promo_id = self.record_promo_surface(
                source_url=tax_url,
                source_page_id=source_page_id,
                title=page.title or slug.replace("-", " ").title(),
                promo_type="vehicle_promo_taxonomy",
                description_text=html_to_text(str(page.soup.select_one(".container") or page.soup))[:2000],
                image_url=image_src(page.soup.select_one("img")),
                cta_url="",
                target_listing_url="",
                applies_to_brand="",
                applies_to_vehicle_type="",
                applies_to_tag=slug,
                has_vehicle_promo_flag=(slug in {"blue-tag", "yes", "new-promo", "top-pick"}),
                published_at="",
                modified_at="",
            )
            for detail_url in extract_motorcycle_links(page.soup, page.final_url):
                detail_url = normalize_url(detail_url)
                row = self.motorcycle_rows.get(detail_url)
                if row:
                    if slug == "blue-tag":
                        row["is_blue_tag"] = True
                    row["has_vehicle_promo_flag"] = True
                    self.tables.append(
                        "motorcycle_promo_link",
                        {
                            "motorcycle_id": row["motorcycle_id"],
                            "promo_id": promo_id,
                            "link_type": "taxonomy_page",
                            "source_url": tax_url,
                            "notes": f"Linked from vehicle promo taxonomy `{slug}`.",
                        },
                    )

    def collect_locator_branches(self) -> None:
        page = self.fetch(f"{MOTORTRADE_DOMAIN}/find-a-store/")
        match = re.search(r'var ASL_REMOTE = \{"ajax_url":"[^"]+","nonce":"([^"]+)"\}', page.text)
        if not match:
            raise ScrapeError("Motortrade locator nonce was not found.")

        payload = self.fetch_json(
            f"{MOTORTRADE_DOMAIN}/wp-admin/admin-ajax.php",
            params={
                "action": "asl_load_stores",
                "nonce": match.group(1),
                "load_all": "1",
                "layout": "1",
                "lat": "12.8797",
                "lng": "121.7740",
                "nw[]": ["21.5", "116.0"],
                "se[]": ["4.5", "127.5"],
            },
        )
        if not isinstance(payload, list):
            raise ScrapeError("Motortrade locator feed did not return a list.")

        locator_source = normalize_url(page.url)
        coverage_rows_before = self.tables.row_count("branch_product_coverage")
        for row in payload:
            locator_id = row.get("id", "")
            branch_id = f"branch-{locator_id}"
            self.tables.append(
                "branch",
                {
                    "branch_id": branch_id,
                    "dealer_id": DEALER_ID,
                    "source_url": locator_source,
                    "locator_id": locator_id,
                    "branch_name": html_to_text(row.get("title", "")),
                    "slug": row.get("slug", ""),
                    "street": row.get("street", ""),
                    "city": row.get("city", ""),
                    "state_or_province": row.get("state", ""),
                    "postal_code": row.get("postal_code", ""),
                    "full_address": " ".join(part for part in [row.get("street", ""), row.get("city", ""), row.get("state", ""), row.get("postal_code", "")] if part).strip(),
                    "latitude": parse_float(row.get("lat")),
                    "longitude": parse_float(row.get("lng")),
                    "phone": row.get("phone", ""),
                    "fax": row.get("fax", ""),
                    "email": row.get("email", ""),
                    "website": row.get("website", ""),
                    "open_hours_raw": row.get("open_hours", ""),
                    "days_str": row.get("days_str", ""),
                    "brand_raw": json.dumps(row.get("brand", "")),
                    "product_raw": json.dumps(row.get("product", "")),
                    "categories_raw": json.dumps(row.get("categories", "")),
                    "inquire_url": row.get("inquire_url", ""),
                },
            )
            self.record_branch_product_coverage(
                branch_id=branch_id,
                coverage_type="brand",
                coverage_value_raw=row.get("brand", ""),
                source_url=locator_source,
                source_field="brand",
                confidence="medium",
                notes="Coverage-level locator field. This is not exact branch-to-model availability.",
            )
            self.record_branch_product_coverage(
                branch_id=branch_id,
                coverage_type="product",
                coverage_value_raw=row.get("product", ""),
                source_url=locator_source,
                source_field="product",
                confidence="medium",
                notes="Coverage-level locator field. This is not exact branch-to-model availability.",
            )
            self.record_branch_product_coverage(
                branch_id=branch_id,
                coverage_type="category",
                coverage_value_raw=row.get("categories", ""),
                source_url=locator_source,
                source_field="categories",
                confidence="low",
                notes="Sparse locator category field. It should not be expanded into exact motorcycle models without source evidence.",
            )

        coverage_rows_created = self.tables.row_count("branch_product_coverage") - coverage_rows_before
        branch_model_evidence = find_branch_model_evidence(payload)
        self.record_relationship_probe(
            relationship_name="branch_to_motorcycle_model_availability",
            source_url=locator_source,
            source_type="locator_ajax",
            status="exact_relationship_found" if branch_model_evidence else ("coverage_only" if coverage_rows_created else "no_relationship_exposed"),
            evidence_found=branch_model_evidence,
            rows_created=self.tables.row_count("branch_model_availability"),
            notes=(
                "Locator AJAX was checked for records containing both a branch identifier and exact model/motorcycle identifier. "
                "Only coverage-level branch fields were present."
                if not branch_model_evidence
                else "Locator AJAX exposed branch/model evidence; availability rows should be reviewed."
            ),
        )

    def collect_motorcycles_from_vehicle_sitemap(self) -> None:
        vehicle_urls = [
            url
            for url in self.fetch_sitemap_urls(f"{MOTORTRADE_DOMAIN}/vehicle-sitemap.xml")
            if "/motorcycles/" in url and not url.rstrip("/").endswith("/motorcycles")
        ]
        model_term_urls = set(self.fetch_sitemap_urls(f"{MOTORTRADE_DOMAIN}/model-sitemap.xml"))

        for detail_url in vehicle_urls:
            try:
                page = self.fetch_html(detail_url)
            except requests.RequestException:
                continue
            detail_url = page.final_url
            detail_soup = page.soup
            detail_slug = slug_from_url(detail_url)
            spec_blocks = detail_soup.select(".pcd-specs div")
            specs: list[tuple[str, str]] = []
            for block in spec_blocks:
                label = text_or_blank(block.find("label"))
                value = text_or_blank(block.find("span"))
                if label and value:
                    specs.append((label, value))

            pricing = extract_pricing_from_soup(detail_soup)
            apply_loan_url = find_motortrade_form_link(detail_soup, "loan")
            inquiry_url = find_motortrade_form_link(detail_soup, "inquiry")
            service_url = find_motortrade_form_link(detail_soup, "service")
            excerpt = text_or_blank(detail_soup.select_one(".vehicle-feature-des")) or text_or_blank(detail_soup.select_one(".vehicle-content p"))
            display_name = extract_motorcycle_display_name(detail_soup, page.title)
            source_page_id = self.record_html_source_page(
                page,
                source_type="motorcycle_detail",
                title_override=display_name,
                content_html=str(detail_soup.select_one(".vehicle-content") or detail_soup.select_one(".col-md-3.autoshowroom-sidebar") or detail_soup),
                content_text=html_to_text(str(detail_soup.select_one(".vehicle-content") or detail_soup)),
                excerpt_text=excerpt,
                featured_image_url=image_src(detail_soup.select_one(".Vehicle-Feature-Image img")) or image_src(detail_soup.select_one(".vehicle-box img")),
            )

            spec_model_name = find_spec_value(specs, "Model")
            promo_heading = display_name if is_promo_title(display_name) else ""
            if promo_heading and spec_model_name:
                display_name = spec_model_name
            model_name = display_name
            brand_name = normalize_brand(find_spec_value(specs, "Make"))
            motorcycle_id = f"motorcycle-{detail_slug}"
            row = self.motorcycle_rows.setdefault(
                detail_url,
                self.build_motorcycle_row(
                    detail_url,
                    source_page_id=source_page_id,
                    listing_url=f"{MOTORTRADE_DOMAIN}/motorcycles/",
                    display_name=display_name,
                    brand_name=brand_name,
                    model_name=model_name,
                    primary_image_url=image_src(detail_soup.select_one(".Vehicle-Feature-Image img")) or "",
                    excerpt_text=excerpt,
                    compare_enabled=bool(detail_soup.select_one(".btn_detail_compare")),
                ),
            )
            row["source_page_id"] = source_page_id
            row["display_name"] = display_name or row["display_name"]
            row["brand_name"] = brand_name or row["brand_name"]
            row["model_name"] = model_name or row["model_name"]
            apply_specs_to_motorcycle_row(row, specs)
            row["excerpt_text"] = excerpt or row["excerpt_text"]
            row["apply_loan_url"] = apply_loan_url or row["apply_loan_url"]
            row["inquiry_url"] = inquiry_url or row["inquiry_url"]
            row["service_url"] = service_url or row["service_url"]
            row["primary_image_url"] = row["primary_image_url"] or image_src(detail_soup.select_one(".Vehicle-Feature-Image img")) or ""
            publish_meta = self.source_page_publish_meta.get(detail_url)
            if publish_meta:
                row["published_at"], row["modified_at"] = publish_meta
            else:
                model_variant = model_terms_urls_by_slug(model_term_urls).get(detail_slug) or ""
                if model_variant:
                    row["model_name"] = row["model_name"] or model_variant

            if spec_model_name and not model_names_compatible(display_name, spec_model_name):
                self.record_data_quality_issue(
                    entity_type="motorcycle",
                    entity_id=motorcycle_id,
                    source_url=detail_url,
                    issue_type="detail_title_spec_model_mismatch",
                    issue_detail=(
                        f"Page title '{display_name}' conflicts with visible specification Model '{spec_model_name}'. "
                        "motorcycle.model_name uses the page title; the raw specification remains in motorcycle.spec_model."
                    ),
                    severity="warning",
                )
            if promo_heading:
                self.record_data_quality_issue(
                    entity_type="motorcycle",
                    entity_id=motorcycle_id,
                    source_url=detail_url,
                    issue_type="promotional_heading_rejected_as_model_name",
                    issue_detail=(
                        f"Promotional heading '{promo_heading}' was not used as motorcycle.model_name. "
                        f"The visible specification Model '{spec_model_name}' was used instead."
                    ),
                    severity="info",
                )

            vehicle_type = infer_scoped_vehicle_type(detail_soup, detail_url)
            if vehicle_type:
                self.assign_vehicle_type(
                    row,
                    vehicle_type=vehicle_type,
                    source_url=detail_url,
                    source_kind="scoped_detail_taxonomy",
                    confidence="medium",
                )
            elif not row.get("vehicle_type"):
                row["classification_notes"] = append_note(
                    row.get("classification_notes"),
                    "No reliable vehicle type source found during detail-page scrape.",
                )
                self.record_data_quality_issue(
                    entity_type="motorcycle",
                    entity_id=motorcycle_id,
                    source_url=detail_url,
                    issue_type="missing_vehicle_type",
                    issue_detail="No mot-type filter, motorcycle-type taxonomy page, or scoped detail taxonomy identified this motorcycle.",
                    severity="info",
                )

            self.record_pricing(
                motorcycle_id=motorcycle_id,
                source_url=detail_url,
                source_context="motorcycle_detail_page",
                srp_text=pricing["srp_text"],
                downpayment_text=pricing["downpayment_text"],
                monthly_text=pricing["monthly_text"],
                note=pricing["note_text"],
            )

            if (
                pricing["srp_amount"] not in ("", None)
                and pricing["monthly_amount"] == pricing["srp_amount"]
            ):
                self.record_data_quality_issue(
                    entity_type="motorcycle",
                    entity_id=motorcycle_id,
                    source_url=detail_url,
                    issue_type="published_monthly_equals_srp",
                    issue_detail=(
                        "The public detail page labels MONTHLY with the same amount as SRP. "
                        "The scraper preserves both source-labelled values without correction."
                    ),
                    severity="warning",
                )

            if any(pricing[key] for key in ("monthly_amount", "downpayment_amount", "srp_amount")):
                self.tables.append(
                    "financing",
                    {
                        "financing_id": f"financing-{motorcycle_id}",
                        "motorcycle_id": motorcycle_id,
                        "dealer_id": DEALER_ID,
                        "source_url": detail_url,
                        "financing_type": "motorcycle_detail_offer",
                        "program_name": "Public motorcycle detail offer",
                        "monthly_amount": pricing["monthly_amount"],
                        "downpayment_amount": pricing["downpayment_amount"],
                        "loan_amount": "",
                        "loan_term_months": "",
                        "estimated_income_amount": "",
                        "form_url": apply_loan_url,
                        "cta_url": f"{MOTORTRADE_DOMAIN}/loan-calculator/",
                        "note_text": pricing["note_text"],
                        "captured_at": datetime.now(TIMEZONE).isoformat(),
                    },
                )

            for display_order, (label, value) in enumerate(specs, start=1):
                numeric_value, unit = split_numeric_unit(value)
                self.tables.append(
                    "motorcycle_spec",
                    {
                        "motorcycle_id": motorcycle_id,
                        "spec_group": "Specifications",
                        "spec_label": label,
                        "spec_value_raw": value,
                        "spec_value_numeric": numeric_value,
                        "spec_unit": unit,
                        "display_order": display_order,
                        "source_url": detail_url,
                    },
                )

            related_links = [link for link in extract_motorcycle_links(detail_soup, detail_url) if normalize_url(link) != detail_url]
            for related in related_links:
                related = normalize_url(related)
                self.motorcycle_rows.setdefault(
                    related,
                    self.build_motorcycle_row(related, listing_url=detail_url),
                )

        for row in sorted(self.motorcycle_rows.values(), key=lambda item: str(item["detail_url"])):
            if not row.get("display_name"):
                self.record_data_quality_issue(
                    entity_type="motorcycle",
                    entity_id=str(row["motorcycle_id"]),
                    source_url=str(row["detail_url"]),
                    issue_type="missing_display_name",
                    issue_detail="Motorcycle was discovered as a related/listing link but no detail title was parsed.",
                    severity="info",
                )
            self.tables.append("motorcycle", row)

    def collect_service_and_promo_pages(self, rest_items: dict[str, list[dict[str, object]]]) -> None:
        pages_by_slug = {str(item.get("slug")): item for item in rest_items.get("pages", [])}
        for source_type, url in TRACKED_HTML_PAGES:
            try:
                page = self.fetch_html(url)
            except requests.RequestException:
                continue
            slug = slug_from_url(page.final_url)
            page_item = pages_by_slug.get(slug)
            source_page_id = self.record_html_source_page(
                page,
                source_type=source_type,
                wp_post_type="page" if page_item else "",
                title_override=html_to_text(get_rendered(page_item.get("title"))) if page_item else page.title,
                content_html=get_rendered(page_item.get("content")) if page_item else str(page.soup.select_one(".container") or page.soup),
                content_text=html_to_text(get_rendered(page_item.get("content"))) if page_item else html_to_text(str(page.soup.select_one(".container") or page.soup)),
                published_at=str(page_item.get("date", "")) if page_item else "",
                modified_at=str(page_item.get("modified", "")) if page_item else "",
                excerpt_text=html_to_text(get_rendered(page_item.get("excerpt"))) if page_item else "",
                featured_image_url=self.resolve_featured_image(page_item, int(page_item.get("featured_media", 0) or 0)) if page_item else image_src(page.soup.select_one("img")),
                seo_title=(page_item.get("yoast_head_json", {}) or {}).get("title", "") if page_item else "",
                seo_description=(page_item.get("yoast_head_json", {}) or {}).get("description", "") if page_item else "",
                canonical_url=(page_item.get("yoast_head_json", {}) or {}).get("canonical", "") if page_item else "",
            )
            self.record_external_ctas(page, source_page_id, source_type)

            text = html_to_text(str(page.soup.select_one(".container") or page.soup))
            form_links = [normalize_url(link) for link in extract_links(page.soup, page.final_url) if "forms.motortrade.com.ph" in link]
            external_ctas = [normalize_url(link) for link in extract_links(page.soup, page.final_url) if any(token in link for token in ("cportal.motortrade.com.ph", "docs.google.com/forms", "vehicle_promo=yes"))]
            if source_type == "service_page":
                self.tables.append(
                    "service_offering",
                    {
                        "service_id": f"service-{slug}",
                        "dealer_id": DEALER_ID,
                        "source_page_id": source_page_id,
                        "service_category": slug,
                        "title": page.title or slug.replace("-", " ").title(),
                        "description_text": text[:4000],
                        "partner_name": detect_partner_name(text),
                        "cta_url": form_links[0] if form_links else (external_ctas[0] if external_ctas else ""),
                        "external_form_url": form_links[0] if form_links else "",
                        "notes": json.dumps({"all_form_links": form_links, "external_ctas": external_ctas}),
                    },
                )
                if slug in {"loan-calculator", "loan-application", "service", "insurance", "maintenance-repair", "ltrs", "forms"}:
                    self.tables.append(
                        "financing",
                        {
                            "financing_id": f"financing-{slug}",
                            "motorcycle_id": "",
                            "dealer_id": DEALER_ID,
                            "source_url": page.final_url,
                            "financing_type": slug,
                            "program_name": page.title or slug.replace("-", " ").title(),
                            "monthly_amount": parse_currency_amount(text) if "₱" in text else "",
                            "downpayment_amount": "",
                            "loan_amount": "",
                            "loan_term_months": "",
                            "estimated_income_amount": "",
                            "form_url": form_links[0] if form_links else "",
                            "cta_url": external_ctas[0] if external_ctas else (form_links[0] if form_links else ""),
                            "note_text": text[:2000],
                            "captured_at": datetime.now(TIMEZONE).isoformat(),
                        },
                    )
            elif source_type == "promo_surface":
                self.record_promo_surface(
                    source_url=page.final_url,
                    source_page_id=source_page_id,
                    title=page.title or slug.replace("-", " ").title(),
                    promo_type="promo_page",
                    description_text=text[:3000],
                    image_url=image_src(page.soup.select_one("img")),
                    cta_url=form_links[0] if form_links else (external_ctas[0] if external_ctas else ""),
                    target_listing_url=first_matching_link(page.soup, page.final_url, "vehicle_promo=yes"),
                    applies_to_brand="",
                    applies_to_vehicle_type="",
                    applies_to_tag="",
                    has_vehicle_promo_flag=("vehicle_promo=yes" in page.html),
                    published_at=str(page_item.get("date", "")) if page_item else "",
                    modified_at=str(page_item.get("modified", "")) if page_item else "",
                )

                for link in extract_links(page.soup, page.final_url):
                    if "vehicle_promo=yes" not in link:
                        continue
                    query = parse_qs(urlparse(link).query)
                    promo_id = self.record_promo_surface(
                        source_url=normalize_url(link),
                        source_page_id=source_page_id,
                        title=PROMO_QUERY_LABELS.get(urlparse(link).path + "?" + urlparse(link).query, page.title or "Promo query"),
                        promo_type="promo_query_link",
                        description_text=f"Promo query discovered on {page.final_url}",
                        image_url="",
                        cta_url="",
                        target_listing_url=normalize_url(link),
                        applies_to_brand=query.get("maker", [""])[0],
                        applies_to_vehicle_type=query.get("motorcycle-category", [""])[0],
                        applies_to_tag=query.get("prod-tags", [""])[0],
                        has_vehicle_promo_flag=True,
                        published_at=str(page_item.get("date", "")) if page_item else "",
                        modified_at=str(page_item.get("modified", "")) if page_item else "",
                    )
                    try:
                        listing_page = self.fetch_html(link)
                    except requests.RequestException:
                        continue
                    for detail_url in extract_motorcycle_links(listing_page.soup, listing_page.final_url):
                        detail_url = normalize_url(detail_url)
                        row = self.motorcycle_rows.get(detail_url)
                        if row:
                            row["has_vehicle_promo_flag"] = True
                            self.tables.append(
                                "motorcycle_promo_link",
                                {
                                    "motorcycle_id": row["motorcycle_id"],
                                    "promo_id": promo_id,
                                    "link_type": "promo_query",
                                    "source_url": normalize_url(link),
                                    "notes": f"Linked from promo surface {page.final_url}.",
                                },
                            )

    def record_promo_surface(
        self,
        *,
        source_url: str,
        source_page_id: str,
        title: str,
        promo_type: str,
        description_text: str,
        image_url: str,
        cta_url: str,
        target_listing_url: str,
        applies_to_brand: str,
        applies_to_vehicle_type: str,
        applies_to_tag: str,
        has_vehicle_promo_flag: bool,
        published_at: str,
        modified_at: str,
    ) -> str:
        source_url = normalize_url(source_url)
        promo_id = self.promo_rows_by_url.get(source_url)
        if promo_id:
            return promo_id
        promo_id = f"promo-{len(self.promo_rows_by_url) + 1:04d}"
        self.promo_rows_by_url[source_url] = promo_id
        self.tables.append(
            "promo",
            {
                "promo_id": promo_id,
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "source_url": source_url,
                "title": title,
                "promo_type": promo_type,
                "description_text": description_text,
                "image_url": image_url,
                "cta_url": cta_url,
                "target_listing_url": target_listing_url,
                "applies_to_brand": applies_to_brand,
                "applies_to_vehicle_type": applies_to_vehicle_type,
                "applies_to_tag": applies_to_tag,
                "has_vehicle_promo_flag": has_vehicle_promo_flag,
                "published_at": published_at,
                "modified_at": modified_at,
            },
        )
        return promo_id

    def collect_forms(self) -> None:
        for form_url in TRACKED_FORM_URLS:
            try:
                page = self.fetch_html(form_url)
            except requests.RequestException:
                continue
            forms = page.soup.find_all("form")
            if not forms:
                continue
            for index, form in enumerate(forms, start=1):
                form_category = slug_from_url(form_url).split("?")[0].split("/")[-1] or "form"
                form_id = f"form-{form_category}-{index}"
                form_action = normalize_url(urljoin(page.final_url, form.get("action", "").strip() or page.final_url))
                field_rows = extract_form_fields(form, page.final_url)
                requires_file = any(row["field_type"] == "file" for row in field_rows)
                requires_otp = any("otp" in (row["field_name"] or "").lower() for row in field_rows)
                multi_step = "tabs" in form_action.lower() or bool(page.soup.select(".tab-pane, .step, .steps, .wizard"))
                has_privacy_checkbox = any(
                    row["field_type"] in {"checkbox", "radio"}
                    and any(token in (row["field_name"] or "").lower() or token in (row["field_label"] or "").lower() for token in ["privacy", "agreement", "consent"])
                    for row in field_rows
                )

                self.tables.append(
                    "form_definition",
                    {
                        "form_id": form_id,
                        "dealer_id": DEALER_ID,
                        "form_name": page.title or form_category.replace("-", " ").title(),
                        "form_category": form_category,
                        "source_url": form_url,
                        "final_url": page.final_url,
                        "page_title": page.title,
                        "form_action": form_action,
                        "form_method": (form.get("method") or "GET").upper(),
                        "requires_file_upload": requires_file,
                        "requires_otp": requires_otp,
                        "multi_step": multi_step,
                        "has_privacy_checkbox": has_privacy_checkbox,
                        "notes": json.dumps(
                            {
                                "form_index": index,
                                "field_count": len(field_rows),
                                "source_links": [link for link in extract_links(page.soup, page.final_url) if "forms.motortrade.com.ph" in link][:20],
                            }
                        ),
                    },
                )

                for field_order, row in enumerate(field_rows, start=1):
                    option_rows = row.pop("options")
                    if option_rows:
                        for option in option_rows:
                            self.tables.append(
                                "form_field",
                                {
                                    "form_id": form_id,
                                    "field_order": field_order,
                                    "field_name": row["field_name"],
                                    "field_label": row["field_label"],
                                    "field_type": row["field_type"],
                                    "is_required": row["is_required"],
                                    "default_value": row["default_value"],
                                    "placeholder": row["placeholder"],
                                    "option_value": option["option_value"],
                                    "option_label": option["option_label"],
                                    "section_name": row["section_name"],
                                    "source_url": page.final_url,
                                },
                            )
                    else:
                        self.tables.append(
                            "form_field",
                            {
                                "form_id": form_id,
                                "field_order": field_order,
                                "field_name": row["field_name"],
                                "field_label": row["field_label"],
                                "field_type": row["field_type"],
                                "is_required": row["is_required"],
                                "default_value": row["default_value"],
                                "placeholder": row["placeholder"],
                                "option_value": "",
                                "option_label": "",
                                "section_name": row["section_name"],
                                "source_url": page.final_url,
                            },
                        )

                if form_category in {"loan", "service", "ltrs", "bigbike", "inquiry"}:
                    loan_terms = find_numeric_hint(field_rows, ("term", "months"))
                    estimated_income = find_numeric_hint(field_rows, ("income",))
                    self.tables.append(
                        "financing",
                        {
                            "financing_id": f"financing-form-{form_category}",
                            "motorcycle_id": "",
                            "dealer_id": DEALER_ID,
                            "source_url": page.final_url,
                            "financing_type": f"{form_category}_form",
                            "program_name": page.title or form_category,
                            "monthly_amount": "",
                            "downpayment_amount": "",
                            "loan_amount": "",
                            "loan_term_months": loan_terms,
                            "estimated_income_amount": estimated_income,
                            "form_url": page.final_url,
                            "cta_url": form_action,
                            "note_text": f"Public form schema with {len(field_rows)} fields.",
                            "captured_at": datetime.now(TIMEZONE).isoformat(),
                        },
                    )

    def fetch_sitemap_urls(self, sitemap_url: str) -> list[str]:
        response = self.fetch(sitemap_url)
        root = ET.fromstring(response.content)
        urls = []
        for node in root.findall(".//sm:url/sm:loc", SITEMAP_NAMESPACE):
            if node.text:
                urls.append(node.text.strip())
        self.source_counter += 1
        return urls

    def record_pricing(
        self,
        *,
        motorcycle_id: str,
        source_url: str,
        source_context: str,
        srp_text: str,
        downpayment_text: str,
        monthly_text: str,
        note: str,
    ) -> None:
        key = (motorcycle_id, source_context)
        if key in self.motorcycle_pricing_seen:
            return
        self.motorcycle_pricing_seen.add(key)
        self.tables.append(
            "pricing",
            {
                "pricing_id": f"pricing-{len(self.motorcycle_pricing_seen):05d}",
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "source_url": source_url,
                "source_context": source_context,
                "currency": "PHP",
                "srp_text": srp_text,
                "srp_amount": parse_currency_amount(srp_text),
                "downpayment_text": downpayment_text,
                "downpayment_amount": parse_currency_amount(downpayment_text),
                "monthly_text": monthly_text,
                "monthly_amount": parse_currency_amount(monthly_text),
                "price_note": note,
                "is_indicative": True if note else "",
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )


def sanitize_excel_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        return value
    if isinstance(value, (int,)):
        return value
    text = str(value)
    if len(text) > EXCEL_CELL_LIMIT:
        return text[: EXCEL_CELL_LIMIT - 22] + "...[TRUNCATED_FOR_EXCEL]"
    return text


def build_section_rows(tables: dict[str, list[dict[str, object]]]) -> dict[str, list[dict[str, object]]]:
    motorcycles = {str(row.get("motorcycle_id", "")): row for row in tables.get("motorcycle", [])}
    source_pages = {str(row.get("source_page_id", "")): row for row in tables.get("source_page", [])}
    promos = {str(row.get("promo_id", "")): row for row in tables.get("promo", [])}
    forms = {str(row.get("form_id", "")): row for row in tables.get("form_definition", [])}
    branches = {str(row.get("branch_id", "")): row for row in tables.get("branch", [])}

    pricing_by_motorcycle: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in tables.get("pricing", []):
        pricing_by_motorcycle[str(row.get("motorcycle_id", ""))].append(row)

    promo_links_by_promo: dict[str, list[dict[str, object]]] = defaultdict(list)
    promo_links_by_motorcycle: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in tables.get("motorcycle_promo_link", []):
        promo_links_by_promo[str(row.get("promo_id", ""))].append(row)
        promo_links_by_motorcycle[str(row.get("motorcycle_id", ""))].append(row)

    section_rows: dict[str, list[dict[str, object]]] = {name: [] for name in SECTION_SHEETS}

    for motorcycle_id, row in motorcycles.items():
        prices = pricing_by_motorcycle.get(motorcycle_id, [])
        best_price = choose_best_pricing(prices)
        promo_titles = []
        for link in promo_links_by_motorcycle.get(motorcycle_id, []):
            promo = promos.get(str(link.get("promo_id", "")))
            if promo and promo.get("title"):
                promo_titles.append(str(promo.get("title")))
        section_rows["motorcycle_catalog"].append(
            {
                **row,
                "srp_amount": best_price.get("srp_amount", ""),
                "downpayment_amount": best_price.get("downpayment_amount", ""),
                "monthly_amount": best_price.get("monthly_amount", ""),
                "price_contexts": join_unique(price.get("source_context", "") for price in prices),
                "promo_titles": join_unique(promo_titles),
                "source_page_url": source_pages.get(str(row.get("source_page_id", "")), {}).get("url", ""),
            }
        )

    for row in tables.get("motorcycle_spec", []):
        motorcycle = motorcycles.get(str(row.get("motorcycle_id", "")), {})
        section_rows["motorcycle_specs"].append(
            {
                **row,
                "display_name": motorcycle.get("display_name", ""),
                "brand_name": motorcycle.get("brand_name", ""),
                "model_name": motorcycle.get("model_name", ""),
            }
        )

    for row in tables.get("pricing", []):
        motorcycle = motorcycles.get(str(row.get("motorcycle_id", "")), {})
        section_rows["motorcycle_pricing"].append(
            {
                **row,
                "display_name": motorcycle.get("display_name", ""),
                "brand_name": motorcycle.get("brand_name", ""),
                "model_name": motorcycle.get("model_name", ""),
            }
        )

    section_rows["branches"] = list(tables.get("branch", []))

    for row in tables.get("branch_product_coverage", []):
        branch = branches.get(str(row.get("branch_id", "")), {})
        section_rows["branch_coverage_notes"].append(
            {
                **row,
                "note_type": "coverage_only",
                "branch_name": branch.get("branch_name", ""),
                "status": "",
                "evidence_found": "",
                "rows_created": "",
                "probe_id": "",
            }
        )
    for row in tables.get("relationship_probe", []):
        section_rows["branch_coverage_notes"].append(
            {
                "note_type": "relationship_probe",
                "source_url": row.get("source_url", ""),
                "status": row.get("status", ""),
                "evidence_found": row.get("evidence_found", ""),
                "rows_created": row.get("rows_created", ""),
                "notes": row.get("notes", ""),
                "probe_id": row.get("probe_id", ""),
            }
        )

    for promo_id, row in promos.items():
        links = promo_links_by_promo.get(promo_id, [])
        linked_names = []
        linked_ids = []
        for link in links:
            motorcycle_id = str(link.get("motorcycle_id", ""))
            linked_ids.append(motorcycle_id)
            motorcycle = motorcycles.get(motorcycle_id, {})
            linked_names.append(str(motorcycle.get("display_name") or motorcycle_id))
        section_rows["promos"].append(
            {
                **row,
                "linked_motorcycle_count": len(links),
                "linked_motorcycles": join_unique(linked_names),
                "link_types": join_unique(link.get("link_type", "") for link in links),
                "link_source_urls": join_unique(link.get("source_url", "") for link in links),
                "source_page_url": source_pages.get(str(row.get("source_page_id", "")), {}).get("url", ""),
            }
        )

    for row in tables.get("financing", []):
        motorcycle = motorcycles.get(str(row.get("motorcycle_id", "")), {})
        section_rows["financing"].append(
            {
                **row,
                "display_name": motorcycle.get("display_name", ""),
                "brand_name": motorcycle.get("brand_name", ""),
                "model_name": motorcycle.get("model_name", ""),
            }
        )

    for row in tables.get("form_field", []):
        form = forms.get(str(row.get("form_id", "")), {})
        section_rows["forms"].append({**form, **row})
    form_ids_with_fields = {str(row.get("form_id", "")) for row in tables.get("form_field", [])}
    for form_id, form in forms.items():
        if form_id not in form_ids_with_fields:
            section_rows["forms"].append(form)

    copy_mappings = {
        "aftersales_services": "service_offering",
        "jobs": "job_listing",
        "parts": "parts_catalog",
        "preowned": "preowned_motorcycle",
        "external_links": "external_cta",
        "source_pages": "source_page",
    }
    for section, table in copy_mappings.items():
        section_rows[section] = list(tables.get(table, []))

    for row in tables.get("taxonomy_term", []):
        section_rows["taxonomy_filters"].append(
            {
                "row_type": "taxonomy_term",
                "taxonomy_or_filter": row.get("taxonomy", ""),
                "term_or_option_value": row.get("slug", ""),
                "label": row.get("name", ""),
                "parent_term_id": row.get("parent_term_id", ""),
                "count": row.get("count", ""),
                "link_or_source_url": row.get("link", ""),
                "used_by_post_types": row.get("used_by_post_types", ""),
                "notes": "",
            }
        )
    for row in tables.get("filter_option", []):
        section_rows["taxonomy_filters"].append(
            {
                "row_type": "filter_option",
                "taxonomy_or_filter": row.get("filter_name", ""),
                "term_or_option_value": row.get("option_value", ""),
                "label": row.get("option_label", ""),
                "link_or_source_url": row.get("source_url", ""),
                "notes": row.get("notes", ""),
            }
        )

    for row in tables.get("data_quality_issue", []):
        section_rows["data_quality_notes"].append({**row, "note_type": "data_quality_issue", "captured_at": row.get("captured_at", "")})
    for row in tables.get("relationship_probe", []):
        section_rows["data_quality_notes"].append(
            {
                "note_type": "relationship_probe",
                "entity_type": "relationship",
                "source_url": row.get("source_url", ""),
                "relationship_name": row.get("relationship_name", ""),
                "status": row.get("status", ""),
                "evidence_found": row.get("evidence_found", ""),
                "rows_created": row.get("rows_created", ""),
                "notes": row.get("notes", ""),
                "captured_at": row.get("checked_at", ""),
            }
        )

    return section_rows


def build_database_rows(tables: dict[str, list[dict[str, object]]]) -> dict[str, list[dict[str, object]]]:
    motorcycles = {str(row.get("motorcycle_id", "")): row for row in tables.get("motorcycle", [])}
    promos = {str(row.get("promo_id", "")): row for row in tables.get("promo", [])}
    forms = {str(row.get("form_id", "")): row for row in tables.get("form_definition", [])}
    section_rows = build_section_rows(tables)
    database_rows: dict[str, list[dict[str, object]]] = {name: [] for name in DATABASE_SHEETS}

    direct_tables = {
        "dealer": "dealer",
        "crawl_run": "crawl_run",
        "source_page": "source_page",
        "branch": "branch",
        "branch_product_coverage": "branch_product_coverage",
        "branch_model_availability": "branch_model_availability",
        "aftersales_service": "service_offering",
        "form_definition": "form_definition",
        "job_listing": "job_listing",
        "parts_catalog": "parts_catalog",
        "preowned_motorcycle": "preowned_motorcycle",
        "external_cta": "external_cta",
        "relationship_probe": "relationship_probe",
        "data_quality_issue": "data_quality_issue",
    }
    for database_table, source_table in direct_tables.items():
        database_rows[database_table] = list(tables.get(source_table, []))

    database_rows["motorcycle"] = list(section_rows.get("motorcycle_catalog", []))
    database_rows["promo"] = list(section_rows.get("promos", []))

    for row in section_rows.get("motorcycle_specs", []):
        motorcycle = motorcycles.get(str(row.get("motorcycle_id", "")), {})
        database_rows["motorcycle_spec"].append({**row, "dealer_id": motorcycle.get("dealer_id", "")})

    for row in section_rows.get("motorcycle_pricing", []):
        motorcycle = motorcycles.get(str(row.get("motorcycle_id", "")), {})
        database_rows["pricing"].append({**row, "dealer_id": row.get("dealer_id") or motorcycle.get("dealer_id", "")})

    for row in section_rows.get("financing", []):
        motorcycle = motorcycles.get(str(row.get("motorcycle_id", "")), {})
        database_rows["financing"].append({**row, "dealer_id": row.get("dealer_id") or motorcycle.get("dealer_id", DEALER_ID)})

    for row in tables.get("motorcycle_promo_link", []):
        motorcycle = motorcycles.get(str(row.get("motorcycle_id", "")), {})
        promo = promos.get(str(row.get("promo_id", "")), {})
        database_rows["motorcycle_promo_link"].append(
            {
                **row,
                "dealer_id": motorcycle.get("dealer_id") or promo.get("dealer_id", ""),
                "display_name": motorcycle.get("display_name", ""),
                "promo_title": promo.get("title", ""),
            }
        )

    for row in tables.get("form_field", []):
        form = forms.get(str(row.get("form_id", "")), {})
        database_rows["form_field"].append(
            {
                **row,
                "dealer_id": form.get("dealer_id", ""),
                "form_name": form.get("form_name", ""),
                "form_category": form.get("form_category", ""),
            }
        )

    for row in tables.get("filter_option", []):
        database_rows["website_filter_reference"].append(
            {
                "dealer_id": DEALER_ID,
                "reference_type": "public_listing_filter",
                "filter_name": row.get("filter_name", ""),
                "option_value": row.get("option_value", ""),
                "option_label": row.get("option_label", ""),
                "source_url": row.get("source_url", ""),
                "notes": "Website filter/reference metadata only. Do not join to motorcycles unless a separate evidence table maps this exact option to exact rows.",
            }
        )

    return database_rows


def choose_best_pricing(rows: list[dict[str, object]]) -> dict[str, object]:
    for context in ("motorcycle_detail_page", "motorcycle_archive_card"):
        for row in rows:
            if row.get("source_context") == context:
                return row
    return rows[0] if rows else {}


def join_unique(values: Iterable[object]) -> str:
    cleaned = [str(value).strip() for value in values if str(value or "").strip()]
    return " | ".join(dedupe(cleaned))


def normalize_url(url: str) -> str:
    parsed = urlparse(url.strip())
    scheme = parsed.scheme or "https"
    netloc = parsed.netloc.lower()
    path = re.sub(r"/{2,}", "/", parsed.path or "/")
    if path != "/" and path.endswith("/"):
        path = path.rstrip("/")
    query = parsed.query.rstrip("&")
    return urlunparse((scheme, netloc, path, "", query, ""))


def slug_from_url(url: str) -> str:
    path = urlparse(url).path.rstrip("/")
    return path.split("/")[-1] if path else "index"


def get_rendered(value: object) -> str:
    if isinstance(value, dict):
        return str(value.get("rendered", ""))
    return str(value or "")


def html_to_text(html: object) -> str:
    return " ".join(BeautifulSoup(str(html or ""), "lxml").get_text(" ", strip=True).split())


def text_or_blank(tag: Tag | None) -> str:
    return tag.get_text(" ", strip=True) if tag else ""


def meta_content(soup: BeautifulSoup, name: str) -> str:
    tag = soup.find("meta", attrs={"name": name}) or soup.find("meta", attrs={"property": name})
    return tag.get("content", "").strip() if tag else ""


def canonical_href(soup: BeautifulSoup) -> str:
    tag = soup.find("link", rel=lambda value: value and "canonical" in value)
    return tag.get("href", "").strip() if tag else ""


def image_src(tag: Tag | None) -> str:
    if not tag:
        return ""
    return tag.get("src", "").strip() or tag.get("data-src", "").strip()


def absolute_href(base_url: str, tag: Tag | None) -> str:
    if not tag:
        return ""
    href = tag.get("href", "").strip()
    return normalize_url(urljoin(base_url, href)) if href else ""


def extract_links(soup: BeautifulSoup, base_url: str) -> list[str]:
    links = []
    for anchor in soup.select("a[href]"):
        href = anchor.get("href", "").strip()
        if not href or href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        links.append(normalize_url(urljoin(base_url, href)))
    return dedupe(links)


def extract_motorcycle_links(soup: BeautifulSoup, base_url: str) -> list[str]:
    links = []
    for anchor in soup.select("a[href]"):
        href = anchor.get("href", "").strip()
        if not href:
            continue
        absolute = normalize_url(urljoin(base_url, href))
        if is_motorcycle_detail_url(absolute):
            links.append(absolute)
    return dedupe(links)


def is_motorcycle_detail_url(url: str) -> bool:
    parsed = urlparse(normalize_url(url))
    path = parsed.path.rstrip("/")
    if not path.startswith("/motorcycles/"):
        return False
    if path == "/motorcycles":
        return False
    if re.fullmatch(r"/motorcycles/page/\d+", path):
        return False
    return True


def dedupe(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        ordered.append(value)
    return ordered


def parse_currency_amount(text: object) -> float | str:
    cleaned = str(text or "")
    match = re.search(r"₱\s*([\d,]+(?:\.\d+)?)", cleaned)
    if not match:
        return ""
    return float(match.group(1).replace(",", ""))


def parse_float(value: object) -> float | str:
    try:
        return float(str(value))
    except (TypeError, ValueError):
        return ""


def parse_numeric_text(text: str) -> float | str:
    match = re.search(r"(-?\d[\d,]*(?:\.\d+)?)", text)
    if not match:
        return ""
    return float(match.group(1).replace(",", ""))


def split_numeric_unit(raw: str) -> tuple[float | str, str]:
    number = parse_numeric_text(raw)
    if number == "":
        return "", ""
    match = re.search(r"(-?\d[\d,]*(?:\.\d+)?)\s*(.*)", raw.strip())
    if not match:
        return number, ""
    unit = match.group(2).strip()
    return number, unit


def find_spec_value(specs: list[tuple[str, str]], label: str) -> str:
    for spec_label, spec_value in specs:
        if spec_label.strip().lower() == label.strip().lower():
            return spec_value
    return ""


def model_names_compatible(page_name: str, spec_name: str) -> bool:
    def normalize(value: str) -> str:
        value = value.lower()
        value = re.sub(r"\b(honda|yamaha|suzuki|kawasaki|kymco|bajaj|ktm)\b", " ", value)
        return re.sub(r"[^a-z0-9]+", "", value)

    page_value = normalize(page_name)
    spec_value = normalize(spec_name)
    return bool(page_value and spec_value and (page_value in spec_value or spec_value in page_value))


def normalize_brand(raw: str) -> str:
    if not raw:
        return ""
    return re.split(r"\s+[-–—]\s+", raw, maxsplit=1)[0].strip()


def truthy(value: object) -> bool:
    return str(value).strip() not in {"", "0", "False", "false", "None"}


def null_if_zero(value: object) -> str:
    if value in {0, "0", None, ""}:
        return ""
    return str(value)


def normalize_coverage_values(value: object) -> list[str]:
    if value in (None, "", [], {}):
        return []
    if isinstance(value, (list, tuple, set)):
        raw_values = list(value)
    elif isinstance(value, dict):
        raw_values = list(value.values())
    else:
        text = str(value).strip()
        if not text:
            return []
        try:
            loaded = json.loads(text)
        except json.JSONDecodeError:
            loaded = text
        if isinstance(loaded, (list, tuple, set)):
            raw_values = list(loaded)
        elif isinstance(loaded, dict):
            raw_values = list(loaded.values())
        else:
            raw_values = re.split(r"[,|;/]+", str(loaded))

    normalized = []
    for raw in raw_values:
        text = html_to_text(raw).strip().lower()
        text = re.sub(r"\s+", " ", text)
        if text:
            normalized.append(text)
    return dedupe(normalized)


def find_branch_model_evidence(rows: list[object]) -> bool:
    branch_keys = {"id", "locator_id", "branch_id", "branch", "store_id", "store"}
    model_keys = {"motorcycle_id", "model_id", "vehicle_id", "model", "motorcycle", "vehicle"}
    for row in rows:
        if not isinstance(row, dict):
            continue
        populated_keys = {str(key).lower() for key, value in row.items() if value not in (None, "", [], {})}
        if populated_keys & branch_keys and populated_keys & model_keys:
            return True
    return False


def json_loads_list(value: object) -> list[object]:
    if isinstance(value, list):
        return value
    try:
        loaded = json.loads(str(value or "[]"))
    except json.JSONDecodeError:
        return []
    return loaded if isinstance(loaded, list) else []


def dedupe_json_objects(values: Iterable[object]) -> list[object]:
    seen: set[str] = set()
    deduped: list[object] = []
    for value in values:
        key = json.dumps(value, sort_keys=True)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(value)
    return deduped


def append_note(existing: object, note: str) -> str:
    existing_text = str(existing or "").strip()
    if not existing_text:
        return note
    if note in existing_text:
        return existing_text
    return f"{existing_text} {note}"


def extract_pricing_from_soup(soup: BeautifulSoup) -> dict[str, object]:
    block = soup.select_one(".pcd-pricing")
    note_text = text_or_blank(soup.select_one(".psnote"))
    srp_text = text_or_blank(block.select_one(".pcd-price")) if block else ""
    downpayment_text = text_or_blank(block.select_one(".pcd-downpayment")) if block else ""
    monthly_text = text_or_blank(block.select_one(".pcd-price-msrp")) if block else ""
    return {
        "srp_text": srp_text,
        "downpayment_text": downpayment_text,
        "monthly_text": monthly_text,
        "srp_amount": parse_currency_amount(srp_text),
        "downpayment_amount": parse_currency_amount(downpayment_text),
        "monthly_amount": parse_currency_amount(monthly_text),
        "note_text": note_text,
    }


def find_motortrade_form_link(soup: BeautifulSoup, category: str) -> str:
    for link in extract_links(soup, MOTORTRADE_DOMAIN):
        if "forms.motortrade.com.ph" in link and category in link.lower():
            return link
    return ""


def extract_motorcycle_display_name(soup: BeautifulSoup, page_title: str = "") -> str:
    # Primary page headings are stronger identity evidence than promo copy in
    # nested content blocks.
    for selector in (
        ".vehicle-content h1",
        ".autoshowroom-detail h1",
        ".single-vehicle h1",
        "article h1",
        "main h1",
        "body h1",
    ):
        for candidate in soup.select(selector):
            if candidate.find_parent(class_=re.compile("item|related|carousel|owl|slick", re.I)):
                continue
            text = candidate.get_text(" ", strip=True)
            if text and not is_promo_title(text):
                return text

    og_title = meta_content(soup, "og:title")
    fallback = og_title or page_title
    fallback = re.sub(r"\s*[-|]\s*Motortrade.*$", "", fallback, flags=re.I).strip()
    title_parts = [part.strip() for part in fallback.split("|") if part.strip()]
    if len(title_parts) > 1:
        fallback = title_parts[-1] if title_parts[0].lower() == "motortrade" else title_parts[0]
    if fallback and not is_promo_title(fallback):
        return fallback

    for selector in (
        ".vehicle-content h2",
        ".autoshowroom-detail h2",
        ".single-vehicle h2",
        "article h2",
        "main h2",
        ".vehicle-title",
        ".Vehicle-Title",
    ):
        for candidate in soup.select(selector):
            if candidate.find_parent(class_=re.compile("item|related|carousel|owl|slick", re.I)):
                continue
            text = candidate.get_text(" ", strip=True)
            if text and not is_promo_title(text):
                return text

    return ""


def is_promo_title(title: str) -> bool:
    return bool(re.search(r"\brebate\s*:", title or "", flags=re.I))


def infer_scoped_vehicle_type(soup: BeautifulSoup, detail_url: str) -> str:
    scoped_containers = []
    for selector in (
        ".breadcrumb",
        ".breadcrumbs",
        ".rank-math-breadcrumb",
        "[typeof='BreadcrumbList']",
        ".vehicle-content",
        ".autoshowroom-detail",
        ".single-vehicle",
        "article",
        "main",
    ):
        scoped_containers.extend(soup.select(selector))
    links = []
    for container in scoped_containers:
        if container.find_parent(["header", "footer", "nav"]):
            continue
        links.extend(extract_links(BeautifulSoup(str(container), "lxml"), detail_url))
    links = [link.lower() for link in links]
    for token, label in (
        ("motorcycle-type/big-bike", "big-bike"),
        ("motorcycle-type/electric", "electric"),
        ("motorcycle-type/tri-wheeler", "tri-wheeler"),
        ("motorcycle-type/regular-bike", "regular-bike"),
    ):
        if any(token in link for link in links):
            return label
    return ""


SPEC_WIDE_COLUMNS = {
    "make": "spec_make",
    "model": "spec_model",
    "transmission": "spec_transmission",
    "engine type": "spec_engine_type",
    "seat height": "spec_seat_height",
    "brake system": "spec_brake_system",
    "brake system (front / rear)": "spec_brake_system",
    "displacement (cc)": "spec_displacement_cc",
    "displacement": "spec_displacement_cc",
    "front tire": "spec_front_tire",
    "rear tire": "spec_rear_tire",
    "wheels type": "spec_wheels_type",
    "starting system": "spec_starting_system",
    "overall dimensions": "spec_overall_dimensions",
    "overall dimensions (lenght x width x height)": "spec_overall_dimensions",
    "overall dimensions (length x width x height)": "spec_overall_dimensions",
    "minimum ground clearance": "spec_ground_clearance",
    "fuel system": "spec_fuel_system",
}


def apply_specs_to_motorcycle_row(row: dict[str, object], specs: list[tuple[str, str]]) -> None:
    for label, value in specs:
        normalized_label = re.sub(r"\s+", " ", label.strip().lower())
        column = SPEC_WIDE_COLUMNS.get(normalized_label)
        if column and value and not row.get(column):
            row[column] = value


def first_matching_link(soup: BeautifulSoup, base_url: str, needle: str) -> str:
    for link in extract_links(soup, base_url):
        if needle in link:
            return link
    return ""


def detect_partner_name(text: str) -> str:
    lowered = text.lower()
    if "pioneer motorprotect" in lowered:
        return "Pioneer MotorProtect"
    return ""


def model_terms_urls_by_slug(urls: Iterable[str]) -> dict[str, str]:
    mapping = {}
    for url in urls:
        mapping[slug_from_url(url)] = slug_from_url(url).replace("-", " ").replace("_", " ").title()
    return mapping


def extract_form_fields(form: Tag, base_url: str) -> list[dict[str, object]]:
    fields: list[dict[str, object]] = []
    for field in form.select("input, select, textarea"):
        if field.name == "input" and field.get("type", "").lower() in {"submit", "button", "reset", "image"}:
            continue
        field_name = field.get("name", "").strip()
        field_id = field.get("id", "").strip()
        field_label = find_field_label(field, form)
        field_type = field.get("type", "").lower() if field.name == "input" else field.name.lower()
        placeholder = field.get("placeholder", "").strip()
        default_value = field.get("value", "").strip()
        options: list[dict[str, str]] = []
        if field.name == "select":
            for option in field.select("option"):
                options.append(
                    {
                        "option_value": option.get("value", "").strip(),
                        "option_label": option.get_text(" ", strip=True),
                    }
                )
        elif field_type in {"radio", "checkbox"}:
            option_label = field_label or field.get("value", "").strip()
            options.append(
                {
                    "option_value": field.get("value", "").strip(),
                    "option_label": option_label,
                }
            )
        fields.append(
            {
                "field_name": field_name or field_id,
                "field_label": field_label,
                "field_type": field_type,
                "is_required": field.has_attr("required"),
                "default_value": default_value,
                "placeholder": placeholder,
                "section_name": nearest_section_heading(field),
                "options": options,
            }
        )
    return fields


def find_field_label(field: Tag, form: Tag) -> str:
    field_id = field.get("id", "").strip()
    if field_id:
        label = form.find("label", attrs={"for": field_id})
        if label:
            return label.get_text(" ", strip=True)
    parent_label = field.find_parent("label")
    if parent_label:
        return parent_label.get_text(" ", strip=True)
    container = field.find_parent(class_=re.compile("form-group|field|input", re.I))
    if container:
        label = container.find("label")
        if label:
            return label.get_text(" ", strip=True)
    return field.get("aria-label", "").strip() or field.get("placeholder", "").strip()


def nearest_section_heading(field: Tag) -> str:
    parent = field.parent
    while parent:
        for selector in ("legend", "h1", "h2", "h3", "h4", ".title", ".step-title"):
            heading = parent.select_one(selector)
            if heading:
                text = heading.get_text(" ", strip=True)
                if text:
                    return text
        parent = parent.parent if isinstance(parent.parent, Tag) else None
    return ""


def find_numeric_hint(fields: list[dict[str, object]], keywords: tuple[str, ...]) -> float | str:
    for field in fields:
        haystack = f"{field['field_name']} {field['field_label']} {field['placeholder']}".lower()
        if all(keyword in haystack for keyword in keywords):
            parsed = parse_numeric_text(str(field["default_value"]))
            if parsed != "":
                return parsed
    return ""


def text_of_interest(soup: BeautifulSoup) -> str:
    main = soup.select_one(".container") or soup
    return html_to_text(str(main))


def main() -> int:
    parser = argparse.ArgumentParser(description="Scrape public Motortrade analytics data into a structured workbook.")
    default_output = Path(__file__).resolve().with_name("motortrade_data.xlsx")
    parser.add_argument("--output", default=str(default_output), help="Workbook output path.")
    parser.add_argument(
        "--output-style",
        choices=("section", "analytics", "database"),
        default="section",
        help=(
            "Workbook layout. `section` is the human-readable default; "
            "`database` is the analyst-facing SQL-friendly layout; "
            "`analytics` keeps the legacy normalized table layout."
        ),
    )
    args = parser.parse_args()

    scraper = MotortradeAnalyticsScraper(Path(args.output), output_style=args.output_style)
    scraper.run()
    print(f"[OK] Wrote {Path(args.output).resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
