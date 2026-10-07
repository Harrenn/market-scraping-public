#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from collections import Counter, deque
from dataclasses import dataclass, field
from html import unescape
from pathlib import Path
from typing import Iterable
from urllib.parse import parse_qs, urljoin, urlparse, urlunparse
from xml.etree import ElementTree as ET

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

try:
    from playwright.sync_api import Error as PlaywrightError
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
    from playwright.sync_api import sync_playwright
except ImportError:  # pragma: no cover
    PlaywrightError = Exception
    PlaywrightTimeoutError = Exception
    sync_playwright = None


USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36"
)
REQUEST_TIMEOUT = 20
REQUEST_DELAY_SECONDS = 0.15
MAX_PLAYWRIGHT_PAGES_PER_SITE = 6
DEFAULT_FAMILY_FETCH_LIMIT = 5
MAX_ENDPOINT_SAMPLES_PER_SITE = 12
SITEMAP_NAMESPACE = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
ASSET_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".webp",
    ".svg",
    ".ico",
    ".css",
    ".js",
    ".mjs",
    ".pdf",
    ".zip",
    ".mp4",
    ".webm",
    ".mov",
    ".mp3",
    ".wav",
    ".xml",
    ".rss",
    ".atom",
    ".txt",
    ".woff",
    ".woff2",
    ".ttf",
    ".otf",
}
KEYWORD_PRIORITY = (
    "branch",
    "store",
    "locator",
    "product",
    "motorcycle",
    "vehicle",
    "promo",
    "contact",
    "loan",
    "career",
    "job",
    "news",
    "blog",
    "shop",
    "about",
    "financing",
    "application",
)


@dataclass
class DealerRow:
    dealer: str
    website: str


@dataclass
class FormRecord:
    action: str
    method: str
    visible_fields: list[str] = field(default_factory=list)
    hidden_fields: list[str] = field(default_factory=list)
    submit_labels: list[str] = field(default_factory=list)


@dataclass
class EndpointRecord:
    source_url: str
    endpoint: str
    kind: str
    auth_required: bool
    blocked: bool
    notes: str
    sample_fields: list[str] = field(default_factory=list)


@dataclass
class ObjectRecord:
    kind: str
    source_url: str
    fields: list[str]
    extraction_hint: str
    confidence: str


@dataclass
class PageRecord:
    url: str
    final_url: str
    status: int | None
    content_type: str
    title: str
    meta_description: str
    canonical_url: str
    headings: list[str] = field(default_factory=list)
    internal_links: list[str] = field(default_factory=list)
    scripts: list[str] = field(default_factory=list)
    forms: list[FormRecord] = field(default_factory=list)
    json_ld_types: list[str] = field(default_factory=list)
    tech_flags: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    embeds: list[str] = field(default_factory=list)
    text_sample: str = ""


@dataclass
class CrawlSummary:
    homepage_status: int | None
    final_homepage_url: str
    robots_status: int | None
    sitemap_urls: list[str] = field(default_factory=list)
    discovered_urls: int = 0
    fetched_pages: int = 0
    html_pages: int = 0
    blocked_pages: int = 0
    errors: int = 0


@dataclass
class ScrapeRecommendation:
    primary_method: str
    fallback_method: str
    tooling: str
    blockers: str
    confidence: str


@dataclass
class ReportBundle:
    dealer: str
    website: str
    crawl_summary: CrawlSummary
    url_families: list[tuple[str, int, list[str], str]]
    objects: list[ObjectRecord]
    endpoints: list[EndpointRecord]
    forms: list[tuple[str, FormRecord]]
    tech_flags: list[str]
    constraints: list[str]
    overview_notes: list[str]
    open_questions: list[str]
    scrape_recommendation: ScrapeRecommendation


class BrowserInspector:
    def __init__(self, headed: bool = False):
        self.headed = headed
        self._playwright = None
        self.browser = None

    def __enter__(self) -> "BrowserInspector":
        if sync_playwright is None:
            raise RuntimeError("Playwright is not available in the current environment.")
        self._playwright = sync_playwright().start()
        self.browser = self._playwright.chromium.launch(headless=not self.headed)
        return self

    def inspect(self, url: str) -> tuple[list[str], list[str], str, bool]:
        page = self.browser.new_page(viewport={"width": 1440, "height": 900})
        requests_seen: list[str] = []

        def on_request(request) -> None:
            if request.resource_type in {"xhr", "fetch", "document"}:
                requests_seen.append(request.url)

        page.on("request", on_request)
        challenge = False
        notes: list[str] = []
        text_sample = ""

        try:
            page.goto(url, wait_until="domcontentloaded", timeout=30000)
            page.wait_for_timeout(2500)
            text_sample = (page.locator("body").inner_text() or "").strip()
            body_lower = text_sample.lower()
            if "just a moment" in body_lower or "security verification" in body_lower:
                challenge = True
                notes.append("Browser session encountered an anti-bot or verification page.")
            selector_counts = []
            for label, selector in (
                ("store cards", ".store, .store-card, .sl-item, .branch, .branch-card"),
                ("product cards", ".product, .product-card, .vehicle-item"),
                ("promo cards", ".promo, .promo-card, .promotion-item"),
            ):
                try:
                    count = page.locator(selector).count()
                    if count:
                        selector_counts.append(f"{label}: {count}")
                except PlaywrightError:
                    continue
            if selector_counts:
                notes.append("Rendered DOM counts: " + "; ".join(selector_counts))
        except PlaywrightTimeoutError:
            notes.append("Browser inspection timed out before the page fully stabilized.")
        except PlaywrightError as exc:
            notes.append(f"Browser inspection failed: {exc}")
        finally:
            page.close()

        return dedupe_preserve_order(requests_seen), notes, text_sample[:1500], challenge

    def __exit__(self, exc_type, exc, tb) -> None:
        if self.browser:
            self.browser.close()
        if self._playwright:
            self._playwright.stop()


def slugify(value: str) -> str:
    value = unescape(value).lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    value = re.sub(r"-{2,}", "-", value).strip("-")
    return value or "dealer"


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    return session


def load_dealers(workbook_path: Path) -> list[DealerRow]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook.active
    dealers: list[DealerRow] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        dealer, website = (row + (None, None))[:2]
        if not dealer or not website:
            continue
        dealers.append(DealerRow(str(dealer).strip(), str(website).strip()))
    return dealers


def registrable_domain(hostname: str) -> str:
    parts = hostname.lower().split(".")
    if len(parts) <= 2:
        return hostname.lower()
    if len(parts[-1]) == 2 and len(parts[-2]) <= 3 and len(parts) >= 3:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def normalize_url(url: str) -> str:
    parsed = urlparse(url.strip())
    scheme = "https" if parsed.scheme in {"http", "https"} else parsed.scheme
    netloc = parsed.netloc.lower()
    path = re.sub(r"/{2,}", "/", parsed.path or "/")
    if path != "/" and path.endswith("/"):
        path = path.rstrip("/")
    query = parsed.query.rstrip("&")
    return urlunparse((scheme, netloc, path or "/", "", query, ""))


def is_same_site(candidate: str, root_url: str) -> bool:
    candidate_host = urlparse(candidate).hostname or ""
    root_host = urlparse(root_url).hostname or ""
    if not candidate_host or not root_host:
        return False
    return registrable_domain(candidate_host) == registrable_domain(root_host)


def is_html_response(content_type: str, url: str) -> bool:
    lowered = content_type.lower()
    if "text/html" in lowered or "application/xhtml+xml" in lowered:
        return True
    suffix = Path(urlparse(url).path).suffix.lower()
    return suffix == "" and "json" not in lowered and "xml" not in lowered


def is_asset_url(url: str) -> bool:
    path = Path(urlparse(url).path)
    return path.suffix.lower() in ASSET_EXTENSIONS


def is_noise_endpoint(url: str) -> bool:
    lowered = url.lower()
    host = urlparse(lowered).netloc
    if is_asset_url(url):
        return True
    if any(token in lowered for token in ("/cdn-cgi/rum", "challenge-platform", "turnstile", "/pat/")):
        return True
    if any(
        domain in host
        for domain in (
            "analytics.google.com",
            "www.google-analytics.com",
            "googleads.g.doubleclick.net",
            "www.googleadservices.com",
            "connect.facebook.net",
            "platform.twitter.com",
        )
    ):
        return True
    if "facebook.com/sharer" in lowered:
        return True
    return False


def dedupe_preserve_order(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


def parse_robots_txt(text: str) -> tuple[list[str], list[str]]:
    disallows: list[str] = []
    sitemaps: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        key, _, value = stripped.partition(":")
        key = key.strip().lower()
        value = value.strip()
        if key == "disallow":
            disallows.append(value)
        elif key == "sitemap":
            sitemaps.append(value)
    return disallows, sitemaps


def parse_sitemap_xml(content: bytes) -> tuple[str, list[str]]:
    root = ET.fromstring(content)
    tag = root.tag.split("}")[-1]

    def items(name: str):
        found = root.findall(f"sm:{name}", SITEMAP_NAMESPACE)
        return found if found else root.findall(name)

    def child_loc(node: ET.Element) -> str | None:
        loc = node.find("sm:loc", SITEMAP_NAMESPACE)
        if loc is None:
            loc = node.find("loc")
        return loc.text.strip() if loc is not None and loc.text else None

    urls: list[str] = []
    if tag == "urlset":
        for node in items("url"):
            loc = child_loc(node)
            if loc:
                urls.append(loc)
    elif tag == "sitemapindex":
        for node in items("sitemap"):
            loc = child_loc(node)
            if loc:
                urls.append(loc)
    else:
        raise ValueError(f"Unsupported sitemap tag: {tag}")
    return tag, urls


def fetch(session: requests.Session, url: str, timeout: int = REQUEST_TIMEOUT) -> requests.Response | None:
    try:
        response = session.get(url, timeout=timeout, allow_redirects=True)
        time.sleep(REQUEST_DELAY_SECONDS)
        return response
    except requests.RequestException:
        return None


def discover_sitemaps(session: requests.Session, base_url: str) -> tuple[int | None, list[str], list[str]]:
    robots_url = urljoin(base_url if base_url.endswith("/") else base_url + "/", "robots.txt")
    robots_response = fetch(session, robots_url)
    disallows: list[str] = []
    sitemap_urls: list[str] = []
    robots_status = robots_response.status_code if robots_response is not None else None
    if robots_response is not None and robots_response.ok:
        disallows, sitemap_urls = parse_robots_txt(robots_response.text)

    fallback_urls = [
        urljoin(base_url, "sitemap.xml"),
        urljoin(base_url, "sitemap_index.xml"),
        urljoin(base_url, "wp-sitemap.xml"),
    ]
    combined = dedupe_preserve_order([*sitemap_urls, *fallback_urls])
    valid_sitemaps: list[str] = []
    for sitemap_url in combined:
        response = fetch(session, sitemap_url)
        if response is None or response.status_code != 200:
            continue
        if "xml" in response.headers.get("content-type", "").lower() or response.text.lstrip().startswith("<?xml"):
            valid_sitemaps.append(sitemap_url)
    return robots_status, disallows, dedupe_preserve_order(valid_sitemaps)


def gather_sitemap_urls(
    session: requests.Session,
    sitemap_urls: list[str],
    base_url: str,
    max_depth: int = 2,
) -> list[str]:
    discovered: list[str] = []
    seen_sitemaps: set[str] = set()

    def walk(url: str, depth: int) -> None:
        if url in seen_sitemaps or depth > max_depth:
            return
        seen_sitemaps.add(url)
        response = fetch(session, url)
        if response is None or response.status_code != 200:
            return
        try:
            kind, urls = parse_sitemap_xml(response.content)
        except Exception:
            return
        if kind == "urlset":
            for entry in urls:
                normalized = normalize_url(entry)
                if is_same_site(normalized, base_url) and not is_asset_url(normalized):
                    discovered.append(normalized)
        else:
            for child in urls:
                walk(child, depth + 1)

    for sitemap_url in sitemap_urls:
        walk(sitemap_url, 0)
    return dedupe_preserve_order(discovered)


def score_link(url: str) -> tuple[int, str]:
    lowered = url.lower()
    for index, keyword in enumerate(KEYWORD_PRIORITY):
        if keyword in lowered:
            return (index, lowered)
    return (len(KEYWORD_PRIORITY) + 1, lowered)


def extract_internal_links(soup: BeautifulSoup, base_url: str, current_url: str) -> list[str]:
    links: list[str] = []
    for tag in soup.select("a[href]"):
        href = tag.get("href", "").strip()
        if not href or href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        absolute = normalize_url(urljoin(current_url, href))
        if is_same_site(absolute, base_url) and not is_asset_url(absolute):
            links.append(absolute)
    links = dedupe_preserve_order(links)
    return sorted(links, key=score_link)


def bfs_discover_urls(session: requests.Session, base_url: str, homepage_response: requests.Response | None) -> list[str]:
    if homepage_response is None or not homepage_response.ok or not is_html_response(homepage_response.headers.get("content-type", ""), homepage_response.url):
        return [normalize_url(base_url)]
    soup = BeautifulSoup(homepage_response.text, "lxml")
    queue = deque(extract_internal_links(soup, base_url, homepage_response.url))
    seen = {normalize_url(homepage_response.url), normalize_url(base_url)}
    discovered = [normalize_url(homepage_response.url)]

    while queue:
        url = queue.popleft()
        if url in seen:
            continue
        seen.add(url)
        discovered.append(url)
        response = fetch(session, url)
        if response is None or not response.ok:
            continue
        if not is_html_response(response.headers.get("content-type", ""), response.url):
            continue
        page_soup = BeautifulSoup(response.text, "lxml")
        for link in extract_internal_links(page_soup, base_url, response.url):
            if link not in seen:
                queue.append(link)
    return dedupe_preserve_order(discovered)


def meta_content(soup: BeautifulSoup, name: str, *, attr: str = "name") -> str:
    tag = soup.find("meta", attrs={attr: name})
    return tag.get("content", "").strip() if tag else ""


def detect_tech_flags(
    html: str,
    final_url: str,
    headers: dict[str, str] | requests.structures.CaseInsensitiveDict[str],
    scripts: list[str],
) -> list[str]:
    lowered = html.lower()
    flags: list[str] = []
    markers = {
        "wordpress": "wp-content" in lowered or "/wp-json/" in lowered,
        "woocommerce": "woocommerce" in lowered,
        "elementor": "elementor" in lowered,
        "nextjs": "__next_data__" in lowered or "/_next/" in lowered or headers.get("x-powered-by", "").lower() == "next.js",
        "wix": "parastorage.com" in lowered or "static.wixstatic.com" in lowered or headers.get("server", "").lower() == "pepyaka",
        "cloudflare": "cloudflare" in headers.get("server", "").lower() or "cf-ray" in {k.lower() for k in headers.keys()} or "just a moment" in lowered,
        "google-my-maps": "google.com/maps/d" in lowered,
        "elfsight": "elfsight" in lowered,
        "wp-json": "/wp-json/" in lowered,
        "admin-ajax": "admin-ajax.php" in lowered,
        "agile-store-locator": "agile-store-locator" in lowered or "asl_load_stores" in lowered,
    }
    for key, enabled in markers.items():
        if enabled:
            flags.append(key)
    if any("/_next/" in script for script in scripts) and "nextjs" not in flags:
        flags.append("nextjs")
    return sorted(flags)


def extract_json_ld_types(soup: BeautifulSoup) -> list[str]:
    types: list[str] = []
    for script in soup.select("script[type='application/ld+json']"):
        text = script.string or script.get_text()
        if not text.strip():
            continue
        try:
            payload = json.loads(text)
        except Exception:
            continue
        stack = [payload]
        while stack:
            current = stack.pop()
            if isinstance(current, dict):
                if "@type" in current:
                    value = current["@type"]
                    if isinstance(value, list):
                        for item in value:
                            types.append(str(item))
                    else:
                        types.append(str(value))
                stack.extend(current.values())
            elif isinstance(current, list):
                stack.extend(current)
    return sorted(set(types))


def summarize_text(soup: BeautifulSoup) -> str:
    text = " ".join(chunk.strip() for chunk in soup.get_text(" ", strip=True).split())
    return text[:1200]


def extract_forms(soup: BeautifulSoup, current_url: str) -> list[FormRecord]:
    forms: list[FormRecord] = []
    for form in soup.find_all("form"):
        action = normalize_url(urljoin(current_url, form.get("action", "").strip() or current_url))
        method = (form.get("method") or "GET").upper()
        visible_fields: list[str] = []
        hidden_fields: list[str] = []
        submit_labels: list[str] = []
        for field in form.find_all(["input", "select", "textarea", "button"]):
            name = field.get("name") or field.get("id") or field.get("aria-label") or field.get("placeholder")
            if not name:
                continue
            name = str(name).strip()
            field_type = field.get("type", "").lower()
            if field.name == "button" or field_type in {"submit", "button"}:
                label = field.get_text(" ", strip=True) or field.get("value", "").strip()
                if label:
                    submit_labels.append(label)
                continue
            if field_type == "hidden":
                hidden_fields.append(name)
            else:
                visible_fields.append(name)
        forms.append(
            FormRecord(
                action=action,
                method=method,
                visible_fields=dedupe_preserve_order(visible_fields),
                hidden_fields=dedupe_preserve_order(hidden_fields),
                submit_labels=dedupe_preserve_order(submit_labels),
            )
        )
    return forms


def extract_embeds(soup: BeautifulSoup, current_url: str) -> list[str]:
    embeds: list[str] = []
    for tag in soup.find_all(["iframe", "script"], src=True):
        src = normalize_url(urljoin(current_url, tag.get("src", "").strip()))
        if src:
            embeds.append(src)
    return dedupe_preserve_order(embeds)


def classify_url_family(url: str, *, title: str = "") -> str:
    path = urlparse(url).path.lower().strip("/")
    title_lower = title.lower()
    segments = path.split("/") if path else []
    first = segments[0] if segments else ""

    if path == "":
        return "homepage"
    if any(keyword in path for keyword in ("branch", "store", "locator", "find-a-store", "contact")):
        return "branches / store locator"
    if any(keyword in path for keyword in ("product", "motorcycle", "vehicle", "model", "shop", "parts")):
        return "products / vehicles"
    if any(keyword in path for keyword in ("promo", "promotion", "sale")):
        return "promos / campaigns"
    if any(keyword in path for keyword in ("news", "blog", "article")):
        return "blog / news"
    if any(keyword in path for keyword in ("career", "job")):
        return "careers"
    if any(keyword in path for keyword in ("loan", "apply", "application", "calculator", "credit")):
        return "forms / applications / calculators"
    if any(keyword in path for keyword in ("about", "company", "history")):
        return "about / company"
    if any(keyword in path for keyword in ("cart", "checkout", "account", "my-account")):
        return "commerce / account"
    if first in {"category", "tag", "brand", "product-category", "product-tag"}:
        return "taxonomy / listing"
    if any(keyword in title_lower for keyword in ("career", "job")):
        return "careers"
    return "other public pages"


def extract_endpoint_candidates(
    html: str,
    page_url: str,
    embeds: list[str],
) -> list[EndpointRecord]:
    candidates: list[EndpointRecord] = []
    url_pattern = re.compile(r"""(?P<quote>["'])?(?P<url>https?://[^"'\\\s<>]+|/[^"'\\\s<>]+(?:wp-json|admin-ajax\.php|api|ajax)[^"'\\\s<>]*)""")
    google_map_pattern = re.compile(r"https://www\.google\.com/maps/d/[^\"' ]+")
    elfsight_pattern = re.compile(r"elfsight-app-([a-f0-9-]{36})", re.I)

    for match in url_pattern.finditer(html):
        raw = match.group("url")
        endpoint = normalize_url(urljoin(page_url, raw))
        lowered = endpoint.lower()
        if is_noise_endpoint(endpoint):
            continue
        if "wp-json" in lowered:
            candidates.append(
                EndpointRecord(page_url, endpoint, "WordPress REST", False, False, "Detected from page markup.", [])
            )
        elif "admin-ajax.php" in lowered:
            candidates.append(
                EndpointRecord(page_url, endpoint, "AJAX endpoint", False, False, "Detected from inline script or markup.", [])
            )
        elif any(token in lowered for token in ("/api/", "ajax", ".php?")):
            candidates.append(
                EndpointRecord(page_url, endpoint, "Custom endpoint", False, False, "Detected from page markup.", [])
            )

    for embed in embeds:
        if is_noise_endpoint(embed):
            continue
        if "google.com/maps/d/" in embed:
            notes = "Google My Maps embed detected."
            query = parse_qs(urlparse(embed).query)
            mid = query.get("mid", [None])[0]
            sample_fields = ["name", "address", "coordinates"]
            endpoint = f"https://www.google.com/maps/d/kml?mid={mid}&forcekml=1" if mid else embed
            if mid:
                notes += " KML export pattern is likely available."
            candidates.append(EndpointRecord(page_url, endpoint, "Google My Maps KML", False, False, notes, sample_fields))
        if "elfsight" in embed:
            app_match = elfsight_pattern.search(html)
            notes = "Elfsight widget detected."
            endpoint = "https://core.service.elfsight.com/p/boot/"
            if app_match:
                notes += f" App id: {app_match.group(1)}."
            candidates.append(EndpointRecord(page_url, endpoint, "Elfsight boot", False, False, notes, []))

    for match in google_map_pattern.finditer(html):
        embed = normalize_url(match.group(0))
        if is_noise_endpoint(embed):
            continue
        candidates.append(
            EndpointRecord(page_url, embed, "Google My Maps embed", False, False, "Map embed found in page HTML.", [])
        )

    return dedupe_endpoint_records(candidates)


def dedupe_endpoint_records(records: Iterable[EndpointRecord]) -> list[EndpointRecord]:
    seen: set[tuple[str, str]] = set()
    deduped: list[EndpointRecord] = []
    for record in records:
        key = (record.endpoint, record.kind)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(record)
    return deduped


def collect_downloads(soup: BeautifulSoup, current_url: str) -> list[str]:
    downloads: list[str] = []
    for tag in soup.select("a[href]"):
        absolute = normalize_url(urljoin(current_url, tag.get("href", "").strip()))
        if absolute and is_asset_url(absolute):
            downloads.append(absolute)
    return dedupe_preserve_order(downloads)


def extract_object_hints(page: PageRecord, downloads: list[str]) -> list[ObjectRecord]:
    objects: list[ObjectRecord] = []
    family = classify_url_family(page.final_url, title=page.title)
    page_text = " ".join(page.headings).lower() + " " + page.text_sample.lower()

    if family == "branches / store locator":
        fields = ["branch_name", "address", "city_or_region", "phone", "coordinates", "directions_url"]
        objects.append(ObjectRecord("branch / store records", page.final_url, fields, "Store locator and contact pages usually expose these fields.", "High"))
    if family == "products / vehicles":
        fields = ["name", "brand", "model", "price", "specifications", "images", "product_url"]
        if "financing" in page_text or "monthly" in page_text:
            fields.append("financing_offer")
        objects.append(ObjectRecord("product / vehicle records", page.final_url, fields, "Product and model pages expose catalog data and merchandising fields.", "High"))
    if family == "promos / campaigns":
        fields = ["title", "promo_copy", "start_date", "end_date", "image", "cta_url"]
        objects.append(ObjectRecord("promo / campaign records", page.final_url, fields, "Promo pages usually expose campaign text and timing.", "Medium"))
    if family == "blog / news":
        fields = ["title", "published_at", "author", "summary", "body", "image", "article_url"]
        objects.append(ObjectRecord("article / news records", page.final_url, fields, "News and blog pages expose editorial fields.", "High"))
    if family == "careers":
        fields = ["job_title", "location", "description", "requirements", "apply_url"]
        objects.append(ObjectRecord("career / job records", page.final_url, fields, "Career pages expose openings or application calls to action.", "Medium"))
    if family == "forms / applications / calculators" or page.forms:
        field_names = sorted({field for form in page.forms for field in form.visible_fields})
        fields = field_names[:8] or ["name", "mobile", "email", "message"]
        objects.append(ObjectRecord("form / application schema", page.final_url, fields, "Forms can be scraped structurally without submitting.", "High"))
    if page.json_ld_types:
        fields = ["@type", "name", "url", "image", "description"]
        objects.append(ObjectRecord("structured-data entities", page.final_url, fields, f"JSON-LD types: {', '.join(page.json_ld_types[:5])}", "Medium"))
    if downloads:
        fields = ["asset_url", "asset_type"]
        objects.append(ObjectRecord("download / media assets", page.final_url, fields, "Publicly linked downloadable assets are discoverable.", "Low"))
    return objects


def parse_page(response: requests.Response, base_url: str, requested_url: str) -> tuple[PageRecord, list[EndpointRecord], list[ObjectRecord]]:
    content_type = response.headers.get("content-type", "")
    html = response.text if is_html_response(content_type, response.url) else ""
    soup = BeautifulSoup(html, "lxml") if html else BeautifulSoup("", "lxml")
    scripts = [normalize_url(urljoin(response.url, tag.get("src", "").strip())) for tag in soup.find_all("script", src=True)]
    headings = [tag.get_text(" ", strip=True) for tag in soup.find_all(["h1", "h2", "h3"]) if tag.get_text(" ", strip=True)]
    page = PageRecord(
        url=requested_url,
        final_url=normalize_url(response.url),
        status=response.status_code,
        content_type=content_type,
        title=(soup.title.string.strip() if soup.title and soup.title.string else ""),
        meta_description=meta_content(soup, "description"),
        canonical_url=meta_content(soup, "canonical", attr="property") or (soup.find("link", rel=lambda value: value and "canonical" in value).get("href", "").strip() if soup.find("link", rel=lambda value: value and "canonical" in value) else ""),
        headings=headings[:20],
        internal_links=extract_internal_links(soup, base_url, response.url),
        scripts=dedupe_preserve_order(scripts),
        forms=extract_forms(soup, response.url),
        json_ld_types=extract_json_ld_types(soup),
        tech_flags=detect_tech_flags(html, response.url, response.headers, scripts),
        notes=[],
        embeds=extract_embeds(soup, response.url),
        text_sample=summarize_text(soup),
    )

    if response.status_code in {401, 403, 429}:
        page.notes.append(f"HTTP {response.status_code} indicates restricted or rate-limited access.")
    if "cloudflare" in page.tech_flags and response.status_code in {403, 503}:
        page.notes.append("Cloudflare protection is affecting normal HTTP access.")

    endpoints = extract_endpoint_candidates(html, response.url, page.embeds)
    objects = extract_object_hints(page, collect_downloads(soup, response.url))
    return page, endpoints, objects


def should_use_playwright(page: PageRecord) -> bool:
    family = classify_url_family(page.final_url, title=page.title)
    if page.status in {401, 403, 429}:
        return True
    if family in {"branches / store locator", "forms / applications / calculators"}:
        return True
    if "nextjs" in page.tech_flags:
        return True
    if any(flag in page.tech_flags for flag in ("elfsight", "agile-store-locator", "google-my-maps")):
        return True
    if len(page.text_sample) < 250 and len(page.scripts) > 5:
        return True
    return False


def analyze_site(
    dealer: DealerRow,
    session: requests.Session,
    *,
    headed: bool = False,
) -> ReportBundle:
    homepage_response = fetch(session, dealer.website)
    homepage_status = homepage_response.status_code if homepage_response is not None else None
    final_homepage_url = normalize_url(homepage_response.url if homepage_response is not None else dealer.website)

    robots_status, disallows, sitemap_urls = discover_sitemaps(session, final_homepage_url)
    discovered_urls = gather_sitemap_urls(session, sitemap_urls, final_homepage_url)
    if not discovered_urls:
        discovered_urls = bfs_discover_urls(session, final_homepage_url, homepage_response)
    if normalize_url(final_homepage_url) not in discovered_urls:
        discovered_urls.insert(0, normalize_url(final_homepage_url))
    fetch_urls = select_urls_for_fetch(discovered_urls, final_homepage_url)

    pages: list[PageRecord] = []
    endpoints: list[EndpointRecord] = []
    objects: list[ObjectRecord] = []
    forms: list[tuple[str, FormRecord]] = []
    constraints: list[str] = []
    overview_notes: list[str] = []
    open_questions: list[str] = []
    blocked_pages = 0
    errors = 0

    browser_inspector: BrowserInspector | None = None
    browser_started = False
    playwright_pages_used = 0
    playwright_family_counts: Counter[str] = Counter()

    for url in fetch_urls:
        response = fetch(session, url)
        if response is None:
            errors += 1
            open_questions.append(f"Failed to fetch {url} over normal HTTP.")
            continue

        if not is_html_response(response.headers.get("content-type", ""), response.url):
            continue

        page, page_endpoints, page_objects = parse_page(response, final_homepage_url, url)
        pages.append(page)
        endpoints.extend(page_endpoints)
        objects.extend(page_objects)
        forms.extend((page.final_url, form) for form in page.forms)

        if page.status in {401, 403, 429}:
            blocked_pages += 1
            constraints.append(f"{page.final_url} returned HTTP {page.status}.")

        if should_use_playwright(page):
            family = classify_url_family(page.final_url, title=page.title)
            should_inspect = False
            if page.status in {401, 403, 429}:
                should_inspect = True
            elif playwright_pages_used < MAX_PLAYWRIGHT_PAGES_PER_SITE and playwright_family_counts[family] == 0:
                should_inspect = True
            elif playwright_pages_used < MAX_PLAYWRIGHT_PAGES_PER_SITE and family in {"branches / store locator", "forms / applications / calculators"} and playwright_family_counts[family] < 2:
                should_inspect = True

            if not should_inspect:
                continue
            if sync_playwright is None:
                page.notes.append("Playwright was not available for dynamic follow-up.")
                continue
            if page.status in {401, 403, 429} and not headed:
                page.notes.append("Headed browser mode may be required for deeper inspection on this protected page.")
                continue
            if not browser_started:
                browser_inspector = BrowserInspector(headed=headed).__enter__()
                browser_started = True
            browser_requests, browser_notes, text_sample, challenge = browser_inspector.inspect(page.final_url)
            playwright_pages_used += 1
            playwright_family_counts[family] += 1
            page.notes.extend(browser_notes)
            if text_sample and len(page.text_sample) < 400:
                page.text_sample = text_sample
            if challenge and page.status in {401, 403, 429}:
                constraints.append(f"Manual verification is likely required to inspect {page.final_url} fully.")
            for request_url in browser_requests:
                normalized_request = normalize_url(request_url)
                if is_noise_endpoint(normalized_request):
                    continue
                kind = "Browser network"
                lowered = normalized_request.lower()
                if "wp-json" in lowered:
                    kind = "WordPress REST"
                elif "admin-ajax.php" in lowered:
                    kind = "AJAX endpoint"
                elif "_next" in lowered:
                    kind = "Next.js asset/data"
                elif "maps" in lowered:
                    kind = "Map or geodata request"
                endpoints.append(
                    EndpointRecord(
                        source_url=page.final_url,
                        endpoint=normalized_request,
                        kind=kind,
                        auth_required=False,
                        blocked=False,
                        notes="Observed during Playwright inspection.",
                        sample_fields=[],
                    )
                )

    if browser_inspector is not None:
        browser_inspector.__exit__(None, None, None)

    tech_flags = sorted({flag for page in pages for flag in page.tech_flags})
    families = summarize_url_families_from_discovery(discovered_urls, pages)
    objects = summarize_objects(objects)
    endpoints = dedupe_endpoint_records(endpoints)
    endpoints = enrich_endpoint_samples(session, endpoints)

    if not pages:
        overview_notes.append("No HTML pages were successfully crawled.")
    else:
        overview_notes.append(f"Detected technology markers: {', '.join(tech_flags) if tech_flags else 'no strong CMS markers'}")
        if sitemap_urls:
            overview_notes.append(f"Sitemaps discovered: {len(sitemap_urls)}")
        else:
            overview_notes.append("No usable sitemap was discovered; report relied on same-site link discovery.")
        if len(fetch_urls) < len(discovered_urls):
            overview_notes.append(
                f"Representative page bodies were fetched for {len(fetch_urls)} URLs while coverage counts include {len(discovered_urls)} discovered public URLs."
            )
        if disallows:
            overview_notes.append(f"robots.txt declared {len(disallows)} disallow rules.")

    recommendation = build_scrape_recommendation(tech_flags, endpoints, constraints, pages, headed)
    crawl_summary = CrawlSummary(
        homepage_status=homepage_status,
        final_homepage_url=final_homepage_url,
        robots_status=robots_status,
        sitemap_urls=sitemap_urls,
        discovered_urls=len(discovered_urls),
        fetched_pages=len(fetch_urls),
        html_pages=len(pages),
        blocked_pages=blocked_pages,
        errors=errors,
    )

    for page in pages:
        for note in page.notes:
            if "Cloudflare" in note or "verification" in note.lower():
                constraints.append(note)

    if not endpoints:
        open_questions.append("No explicit public API endpoint was confirmed; scraping may rely on rendered HTML.")
    if not any(obj.kind == "branch / store records" for obj in objects):
        open_questions.append("Branch/store objects were not strongly confirmed from the crawled pages.")

    return ReportBundle(
        dealer=dealer.dealer,
        website=dealer.website,
        crawl_summary=crawl_summary,
        url_families=families,
        objects=objects,
        endpoints=endpoints,
        forms=forms,
        tech_flags=tech_flags,
        constraints=dedupe_preserve_order(constraints),
        overview_notes=dedupe_preserve_order(overview_notes),
        open_questions=dedupe_preserve_order(open_questions),
        scrape_recommendation=recommendation,
    )


def summarize_url_families(pages: list[PageRecord]) -> list[tuple[str, int, list[str], str]]:
    grouped: dict[str, list[PageRecord]] = {}
    for page in pages:
        family = classify_url_family(page.final_url, title=page.title)
        grouped.setdefault(family, []).append(page)

    rows: list[tuple[str, int, list[str], str]] = []
    for family, family_pages in sorted(grouped.items(), key=lambda item: (-len(item[1]), item[0])):
        notes: list[str] = []
        techs = sorted({flag for page in family_pages for flag in page.tech_flags})
        if techs:
            notes.append("Tech markers: " + ", ".join(techs[:4]))
        if any(page.forms for page in family_pages):
            notes.append("Contains forms or interactive inputs.")
        if any(page.json_ld_types for page in family_pages):
            notes.append("Structured data present.")
        rows.append(
            (
                family,
                len(family_pages),
                [page.final_url for page in family_pages[:5]],
                " ".join(notes) if notes else "Public pages discovered in this family.",
            )
        )
    return rows


def select_urls_for_fetch(discovered_urls: list[str], homepage_url: str) -> list[str]:
    grouped: dict[str, list[str]] = {}
    for url in discovered_urls:
        family = classify_url_family(url)
        grouped.setdefault(family, []).append(url)

    selected: list[str] = [normalize_url(homepage_url)]
    family_limits = {
        "homepage": 1,
        "branches / store locator": 8,
        "forms / applications / calculators": 8,
        "products / vehicles": 5,
        "promos / campaigns": 5,
        "blog / news": 4,
        "careers": 3,
        "taxonomy / listing": 4,
        "other public pages": 3,
    }

    for family, urls in grouped.items():
        ordered = sorted(dedupe_preserve_order(urls), key=score_link)
        limit = family_limits.get(family, DEFAULT_FAMILY_FETCH_LIMIT)
        selected.extend(ordered[:limit])

    return dedupe_preserve_order(selected)


def summarize_url_families_from_discovery(
    discovered_urls: list[str],
    fetched_pages: list[PageRecord],
) -> list[tuple[str, int, list[str], str]]:
    grouped: dict[str, list[str]] = {}
    for url in discovered_urls:
        family = classify_url_family(url)
        grouped.setdefault(family, []).append(url)

    fetched_by_family: dict[str, list[PageRecord]] = {}
    for page in fetched_pages:
        family = classify_url_family(page.final_url, title=page.title)
        fetched_by_family.setdefault(family, []).append(page)

    rows: list[tuple[str, int, list[str], str]] = []
    for family, urls in sorted(grouped.items(), key=lambda item: (-len(item[1]), item[0])):
        notes: list[str] = []
        fetched = fetched_by_family.get(family, [])
        techs = sorted({flag for page in fetched for flag in page.tech_flags})
        if techs:
            notes.append("Tech markers: " + ", ".join(techs[:4]))
        if any(page.forms for page in fetched):
            notes.append("Representative pages contain forms or interactive inputs.")
        if any(page.json_ld_types for page in fetched):
            notes.append("Representative pages expose structured data.")
        if not fetched:
            notes.append("Family counted from sitemap/link discovery only.")
        rows.append(
            (
                family,
                len(urls),
                urls[:5],
                " ".join(notes) if notes else "Public URLs discovered in this family.",
            )
        )
    return rows


def summarize_objects(objects: list[ObjectRecord]) -> list[ObjectRecord]:
    grouped: dict[str, ObjectRecord] = {}
    for obj in objects:
        current = grouped.get(obj.kind)
        if current is None:
            grouped[obj.kind] = ObjectRecord(obj.kind, obj.source_url, list(obj.fields), obj.extraction_hint, obj.confidence)
            continue
        current.fields = dedupe_preserve_order([*current.fields, *obj.fields])
        if confidence_rank(obj.confidence) > confidence_rank(current.confidence):
            current.confidence = obj.confidence
        if len(current.extraction_hint) < len(obj.extraction_hint):
            current.extraction_hint = obj.extraction_hint
    return sorted(grouped.values(), key=lambda item: item.kind)


def confidence_rank(value: str) -> int:
    return {"Low": 1, "Medium": 2, "High": 3}.get(value, 0)


def enrich_endpoint_samples(session: requests.Session, endpoints: list[EndpointRecord]) -> list[EndpointRecord]:
    enriched: list[EndpointRecord] = []
    for index, endpoint in enumerate(endpoints):
        if index >= MAX_ENDPOINT_SAMPLES_PER_SITE:
            endpoint.notes = endpoint.notes + " Endpoint sampling skipped to keep the report run bounded."
            enriched.append(endpoint)
            continue
        if endpoint.sample_fields:
            enriched.append(endpoint)
            continue
        lowered = endpoint.endpoint.lower()
        if not any(token in lowered for token in ("wp-json", "api", "ajax", "kml", ".php")):
            enriched.append(endpoint)
            continue
        response = fetch(session, endpoint.endpoint)
        if response is None:
            endpoint.blocked = True
            endpoint.notes = endpoint.notes + " Endpoint fetch failed during sampling."
            enriched.append(endpoint)
            continue
        if response.status_code in {401, 403, 429}:
            endpoint.blocked = True
            endpoint.notes = endpoint.notes + f" Endpoint returned HTTP {response.status_code}."
            enriched.append(endpoint)
            continue
        content_type = response.headers.get("content-type", "").lower()
        sample_fields: list[str] = []
        try:
            if "json" in content_type:
                payload = response.json()
                sample_fields = extract_json_fields(payload)
            elif "xml" in content_type or lowered.endswith("kml"):
                sample_fields = extract_xml_field_hints(response.text)
        except Exception:
            pass
        endpoint.sample_fields = sample_fields[:8]
        enriched.append(endpoint)
    return enriched


def extract_json_fields(payload: object) -> list[str]:
    if isinstance(payload, list) and payload:
        return extract_json_fields(payload[0])
    if isinstance(payload, dict):
        return sorted(str(key) for key in payload.keys())[:12]
    return []


def extract_xml_field_hints(text: str) -> list[str]:
    hints = []
    for token in ("Placemark", "name", "description", "coordinates", "address"):
        if token.lower() in text.lower():
            hints.append(token.lower())
    return hints


def build_scrape_recommendation(
    tech_flags: list[str],
    endpoints: list[EndpointRecord],
    constraints: list[str],
    pages: list[PageRecord],
    headed: bool,
) -> ScrapeRecommendation:
    endpoint_kinds = {endpoint.kind for endpoint in endpoints if not endpoint.blocked}
    if any(kind in endpoint_kinds for kind in {"WordPress REST", "Google My Maps KML"}):
        primary = "Direct public endpoint extraction"
        fallback = "Server-rendered HTML parsing for pages not covered by the endpoint."
        tooling = "`requests` + endpoint pagination/JSON-or-KML parsing"
        blockers = "Low, unless endpoint contracts or plugin configuration change."
        confidence = "High"
    elif any(kind in endpoint_kinds for kind in {"AJAX endpoint", "Custom endpoint"}):
        primary = "Direct AJAX/custom endpoint extraction"
        fallback = "Rendered DOM parsing with Playwright network confirmation."
        tooling = "`requests` for endpoint calls, Playwright for validation"
        blockers = "Medium if CSRF tokens or nonce values rotate."
        confidence = "Medium"
    elif any("nextjs" == flag for flag in tech_flags):
        primary = "Playwright-assisted rendered-page and network scraping"
        fallback = "Static HTML plus internal-link crawl where routes expose content server-side."
        tooling = "Playwright + `requests`"
        blockers = "Medium due to client rendering and missing sitemap coverage."
        confidence = "Medium"
    else:
        primary = "Server-rendered HTML parsing"
        fallback = "Playwright for dynamic templates or locators."
        tooling = "`requests` + BeautifulSoup"
        blockers = "Low for static pages; medium for widgets and locators."
        confidence = "Medium"

    if any("Cloudflare" in constraint or "HTTP 403" in constraint for constraint in constraints):
        primary = "Manual-assisted Playwright for protected surfaces, with direct endpoints where available"
        fallback = primary if headed else "Partial static crawl with blockers documented"
        tooling = "Headed Playwright + `requests`"
        blockers = "High due to anti-bot controls and possible manual verification."
        confidence = "Low" if not headed else "Medium"

    if not pages:
        primary = "Initial connectivity triage before scraping"
        fallback = "Manual browser review to locate stable public sources."
        tooling = "`requests` + Playwright"
        blockers = "High because crawl coverage is currently minimal."
        confidence = "Low"

    return ScrapeRecommendation(primary, fallback, tooling, blockers, confidence)


def markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    if not rows:
        rows = [["No data", "", "", ""][: len(headers)]]
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        padded = row + [""] * (len(headers) - len(row))
        lines.append("| " + " | ".join(cell.replace("\n", "<br>") for cell in padded[: len(headers)]) + " |")
    return "\n".join(lines)


def render_report(bundle: ReportBundle) -> str:
    overview_lines = [
        f"- Website provided in workbook: `{bundle.website}`",
        f"- Final homepage URL: `{bundle.crawl_summary.final_homepage_url}`",
        f"- Homepage status: `{bundle.crawl_summary.homepage_status}`",
    ]
    overview_lines.extend(f"- {note}" for note in bundle.overview_notes)

    coverage_lines = [
        f"- Discovered URLs: `{bundle.crawl_summary.discovered_urls}`",
        f"- Representative pages fetched: `{bundle.crawl_summary.fetched_pages}`",
        f"- HTML pages parsed: `{bundle.crawl_summary.html_pages}`",
        f"- Blocked pages: `{bundle.crawl_summary.blocked_pages}`",
        f"- Errors: `{bundle.crawl_summary.errors}`",
        f"- `robots.txt` status: `{bundle.crawl_summary.robots_status}`",
        f"- Usable sitemaps: `{len(bundle.crawl_summary.sitemap_urls)}`",
    ]

    url_family_rows = [
        [family, str(count), "<br>".join(samples), notes]
        for family, count, samples, notes in bundle.url_families
    ]
    object_rows = [
        [obj.kind, obj.source_url, ", ".join(obj.fields), f"{obj.extraction_hint} Confidence: {obj.confidence}."]
        for obj in bundle.objects
    ]
    endpoint_rows = [
        [
            endpoint.endpoint,
            endpoint.kind,
            "blocked" if endpoint.blocked else "public",
            ", ".join(endpoint.sample_fields) if endpoint.sample_fields else "n/a",
            f"Source: {endpoint.source_url}. {endpoint.notes}",
        ]
        for endpoint in bundle.endpoints
    ]

    form_lines = []
    if bundle.forms:
        for page_url, form in bundle.forms:
            fields = ", ".join(form.visible_fields[:10]) if form.visible_fields else "no named visible fields detected"
            hidden = ", ".join(form.hidden_fields[:6]) if form.hidden_fields else "none detected"
            submits = ", ".join(form.submit_labels[:4]) if form.submit_labels else "not labeled"
            form_lines.append(
                f"- `{page_url}` -> `{form.method}` `{form.action}`; visible fields: {fields}; hidden fields: {hidden}; submit labels: {submits}. Inspect-only, no submission."
            )
    else:
        form_lines.append("- No public forms with named fields were detected during this crawl.")

    constraint_lines = [f"- {item}" for item in bundle.constraints] or ["- No severe access constraints were confirmed on the crawled public pages."]
    open_question_lines = [f"- {item}" for item in bundle.open_questions] or ["- No major open questions remain from this crawl."]

    recommendation_lines = [
        f"- Primary method: {bundle.scrape_recommendation.primary_method}",
        f"- Fallback: {bundle.scrape_recommendation.fallback_method}",
        f"- Tooling: {bundle.scrape_recommendation.tooling}",
        f"- Expected blockers: {bundle.scrape_recommendation.blockers}",
        f"- Confidence: {bundle.scrape_recommendation.confidence}",
    ]

    architecture = ", ".join(bundle.tech_flags) if bundle.tech_flags else "No strong framework marker was confirmed."

    sections = [
        f"# {bundle.dealer}",
        "",
        "## Overview",
        "\n".join(overview_lines),
        "",
        "## Crawl Coverage",
        "\n".join(coverage_lines),
        "",
        "## Site Architecture and Rendering Model",
        f"- Detected markers: {architecture}",
        "",
        "## Discoverable URL Families",
        markdown_table(["Family / Template", "Count", "Sample URLs", "Notes"], url_family_rows),
        "",
        "## Scrapable Objects and Fields",
        markdown_table(["Object Type", "Where Found", "Likely Fields", "Extraction Notes"], object_rows),
        "",
        "## APIs, Feeds, and Embedded Data Sources",
        markdown_table(["Endpoint / Source", "Type", "Auth / Challenge", "Sample Fields", "Notes"], endpoint_rows),
        "",
        "## Forms and Interactive Flows",
        "\n".join(form_lines),
        "",
        "## Anti-Bot, Access, and Constraints",
        "\n".join(constraint_lines),
        "",
        "## Recommended Scrape Method",
        "\n".join(recommendation_lines),
        "",
        "## Open Questions and Risks",
        "\n".join(open_question_lines),
        "",
    ]
    return "\n".join(sections)


def write_report(output_dir: Path, bundle: ReportBundle) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{slugify(bundle.dealer)}.md"
    path.write_text(render_report(bundle), encoding="utf-8")
    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate structured reconnaissance reports for dealer websites.")
    parser.add_argument("--workbook", default="Dealers.xlsx", help="Workbook path containing Dealer and Website columns.")
    parser.add_argument(
        "--output-dir",
        default="working/outputs/reports",
        help="Folder where Markdown reports will be written.",
    )
    parser.add_argument("--only", help="Run the report for one exact dealer name.")
    parser.add_argument("--headed", action="store_true", help="Use a visible browser for Playwright follow-up when needed.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        print(f"Workbook not found: {workbook_path}", file=sys.stderr)
        return 1

    dealers = load_dealers(workbook_path)
    if args.only:
        dealers = [dealer for dealer in dealers if dealer.dealer == args.only]

    if not dealers:
        print("No dealers matched the requested scope.", file=sys.stderr)
        return 1

    session = build_session()
    output_dir = Path(args.output_dir)

    written: list[Path] = []
    for dealer in dealers:
        print(f"[INFO] Analyzing {dealer.dealer} ({dealer.website})", flush=True)
        bundle = analyze_site(dealer, session, headed=args.headed)
        report_path = write_report(output_dir, bundle)
        written.append(report_path)
        print(f"[OK] Wrote {report_path}", flush=True)

    print()
    print(f"Generated {len(written)} report(s) in {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
