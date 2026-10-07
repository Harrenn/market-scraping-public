#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from full_scrape.motortrade.scrape_motortrade_analytics import (
    DATABASE_SCHEMAS,
    DATABASE_SHEETS,
    TIMEZONE,
    USER_AGENT,
    TableStore,
    extract_form_fields,
    get_rendered,
    html_to_text,
    normalize_url,
    parse_currency_amount,
    sanitize_excel_value,
    split_numeric_unit,
)


DEALER_ID = "guanzon_motorcycle_group"
DEALER_NAME = "Guanzon Motorcycle Group"
PRIMARY_DOMAIN = "https://www.guanzongroup.com.ph"
FORMS_DOMAIN = PRIMARY_DOMAIN
REST_BASE = f"{PRIMARY_DOMAIN}/wp-json/wp/v2"
REQUEST_DELAY_SECONDS = 0.05

REST_COLLECTIONS = [
    ("pages", "page", "pages"),
    ("posts", "post", "posts"),
    ("product", "product", "product"),
    ("wpsl_stores", "wpsl_stores", "store_locator"),
    ("blocks", "blocks", "ux_blocks"),
]

TERM_COLLECTIONS = [
    ("categories", "category"),
    ("tags", "post_tag"),
    ("product_cat", "product_cat"),
    ("product_tag", "product_tag"),
    ("product_brand", "product_brand"),
    ("wpsl_store_category", "wpsl_store_category"),
]

BRANCH_PAGES = [
    ("motorcycle_branch", f"{PRIMARY_DOMAIN}/branches/motorcycle-branches/"),
    ("mobile_phone_store", f"{PRIMARY_DOMAIN}/branches/mobile-phone-stores/"),
]

TRACKED_FORM_URLS = [
    f"{PRIMARY_DOMAIN}/career-application/",
    f"{PRIMARY_DOMAIN}/career/",
    f"{PRIMARY_DOMAIN}/branches/branch-locator/",
    f"{PRIMARY_DOMAIN}/contact-us/",
    f"{PRIMARY_DOMAIN}/mc-credit-application/",
    f"{PRIMARY_DOMAIN}/mp-credit-application/",
    f"{PRIMARY_DOMAIN}/mobile-phones/apply-for-a-loan/",
]

SERVICE_URLS = [
    ("customer_care", f"{PRIMARY_DOMAIN}/contact-us/"),
    ("branch_locator", f"{PRIMARY_DOMAIN}/branches/branch-locator/"),
    ("motorcycle_credit_application", f"{PRIMARY_DOMAIN}/mc-credit-application/"),
    ("mobile_phone_credit_application", f"{PRIMARY_DOMAIN}/mp-credit-application/"),
    ("mobile_phone_loan", f"{PRIMARY_DOMAIN}/mobile-phones/apply-for-a-loan/"),
    ("careers", f"{PRIMARY_DOMAIN}/career/"),
    ("career_application", f"{PRIMARY_DOMAIN}/career-application/"),
]

MOTORCYCLE_BRANDS = {
    "honda": "Honda",
    "kawasaki": "Kawasaki",
    "suzuki": "Suzuki",
    "yamaha": "Yamaha",
    "hatasu": "Hatasu",
}

MOBILE_PHONE_BRANDS = {
    "apple",
    "honor",
    "huawei",
    "oppo",
    "realme",
    "samsung",
    "tecno",
    "vivo",
    "xiaomi",
}

PROMO_KEYWORDS = {
    "promo",
    "promos",
    "promotion",
    "discount",
    "bundle",
    "sale",
    "free",
    "raffle",
    "cashback",
    "installment",
    "deal",
}


class GuanzonDatabaseScraper:
    def __init__(self, output_path: Path):
        self.output_path = output_path
        self.session = requests.Session()
        self.session.headers.update(
            {
                "User-Agent": USER_AGENT,
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/json;q=0.8,*/*;q=0.7",
                "Accept-Language": "en-US,en;q=0.9",
            }
        )
        self.tables = TableStore(DATABASE_SCHEMAS)
        self.rest_objects: dict[str, list[dict[str, object]]] = defaultdict(list)
        self.terms_by_id: dict[int, dict[str, object]] = {}
        self.terms_by_taxonomy_id: dict[tuple[str, int], dict[str, object]] = {}
        self.product_category_tree: dict[int, set[str]] = {}
        self.source_page_by_url: dict[str, str] = {}
        self.product_by_url: dict[str, dict[str, object]] = {}
        self.motorcycle_by_url: dict[str, str] = {}
        self.motorcycle_summary: dict[str, dict[str, str]] = {}
        self.branch_seen: set[str] = set()
        self.coverage_seen: set[tuple[str, str, str]] = set()
        self.external_cta_seen: set[tuple[str, str]] = set()
        self.form_seen: set[tuple[str, str]] = set()
        self.promo_seen: set[str] = set()
        self.promo_link_seen: set[tuple[str, str, str]] = set()
        self.pricing_seen: set[str] = set()
        self.pending_related_cards: list[dict[str, object]] = []
        self.source_counter = 0

        now = datetime.now(TIMEZONE)
        self.crawl_date = now.date().isoformat()
        self.run_started_at = now.isoformat()
        self.run_id = f"{DEALER_ID}-run-{now.strftime('%Y%m%d%H%M%S')}"

    def run(self) -> None:
        self.log("Starting Guanzon database scrape")
        self.add_dealer()
        self.log("Collecting taxonomy/filter references")
        self.collect_terms()
        self.log("Collecting WordPress REST objects")
        self.collect_rest_objects()
        self.log("Collecting products and motorcycle detail pages")
        self.collect_products()
        self.log("Collecting branch pages")
        self.collect_branches()
        self.log("Collecting promo and campaign content")
        self.collect_posts_and_promos()
        self.log("Collecting forms")
        self.collect_forms()
        self.log("Collecting service and job surfaces")
        self.collect_services_and_jobs()
        self.collect_relationship_probe()
        self.collect_data_quality_notes()
        self.finalize_crawl_run()
        self.log("Writing workbook")
        self.write_workbook()

    def log(self, message: str) -> None:
        print(f"[Guanzon] {message}", flush=True)

    def add_dealer(self) -> None:
        self.tables.append(
            "dealer",
            {
                "dealer_id": DEALER_ID,
                "dealer_name": DEALER_NAME,
                "primary_domain": PRIMARY_DOMAIN,
                "forms_domain": FORMS_DOMAIN,
                "crawl_timezone": str(TIMEZONE),
                "crawl_date": self.crawl_date,
                "notes": "Public Guanzon WordPress/WooCommerce snapshot using the shared dealer database schema.",
            },
        )

    def fetch(self, url: str, *, params: dict[str, object] | None = None, timeout: int = 45) -> requests.Response:
        response = self.session.get(url, params=params, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
        return response

    def fetch_html(self, url: str) -> tuple[str, BeautifulSoup, int, str]:
        time.sleep(REQUEST_DELAY_SECONDS)
        response = self.fetch(url)
        soup = BeautifulSoup(response.text, "lxml")
        return response.text, soup, response.status_code, response.url

    def fetch_rest_collection(self, rest_base: str, *, per_page: int = 100) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        page = 1
        total_pages = 1
        while page <= total_pages:
            url = f"{REST_BASE}/{rest_base}"
            try:
                response = self.fetch(url, params={"per_page": per_page, "page": page, "_embed": 1})
            except requests.HTTPError as exc:
                self.add_data_quality_issue(
                    entity_type="rest_collection",
                    entity_id=rest_base,
                    source_url=f"{url}?per_page={per_page}&page={page}",
                    issue_type="rest_page_error",
                    issue_detail=f"Skipped REST page after HTTP error: {exc}",
                    severity="warning",
                )
                page += 1
                continue
            data = response.json()
            if isinstance(data, list):
                rows.extend(data)
            total_pages = int(response.headers.get("X-WP-TotalPages") or total_pages)
            self.log(f"REST {rest_base}: page {page}/{total_pages}, rows so far {len(rows)}")
            page += 1
        return rows

    def collect_terms(self) -> None:
        for rest_base, taxonomy in TERM_COLLECTIONS:
            for term in self.fetch_rest_collection(rest_base, per_page=100):
                term_id = int(term.get("id") or 0)
                self.terms_by_id[term_id] = term
                self.terms_by_taxonomy_id[(taxonomy, term_id)] = term
                self.tables.append(
                    "website_filter_reference",
                    {
                        "dealer_id": DEALER_ID,
                        "reference_type": taxonomy,
                        "filter_name": taxonomy,
                        "option_value": term.get("slug", ""),
                        "option_label": clean_text(term.get("name", "")),
                        "source_url": term.get("link", ""),
                        "notes": "Source category/tag/reference value only. Use direct business columns for analytics unless an exact bridge exists.",
                    },
                )
        self.product_category_tree = {term_id: self.category_slug_closure(term_id) for term_id in self.terms_by_id}

    def category_slug_closure(self, term_id: int) -> set[str]:
        slugs: set[str] = set()
        seen: set[int] = set()
        current_id = term_id
        while current_id and current_id not in seen:
            seen.add(current_id)
            term = self.terms_by_taxonomy_id.get(("product_cat", current_id), {})
            if not term:
                break
            slug = str(term.get("slug", "")).strip()
            if slug:
                slugs.add(slug)
            current_id = int(term.get("parent") or 0)
        return slugs

    def collect_rest_objects(self) -> None:
        for rest_base, wp_post_type, source_type in REST_COLLECTIONS:
            rows = self.fetch_rest_collection(rest_base, per_page=100)
            self.rest_objects[wp_post_type].extend(rows)
            for row in rows:
                self.add_source_page(row, source_type=source_type, wp_post_type=wp_post_type)

    def add_source_page(self, row: dict[str, object], *, source_type: str, wp_post_type: str) -> str:
        wp_id = str(row.get("id", ""))
        link = normalize_url(str(row.get("link") or get_rendered(row.get("guid", {})) or ""))
        source_page_id = f"source-{wp_post_type}-{wp_id}"
        if link:
            self.source_page_by_url[link] = source_page_id
        yoast = row.get("yoast_head_json") if isinstance(row.get("yoast_head_json"), dict) else {}
        content_html = get_rendered(row.get("content", {}))
        excerpt_html = get_rendered(row.get("excerpt", {}))
        category_ids = row.get("categories") or row.get("product_cat") or row.get("wpsl_store_category") or []
        tag_ids = row.get("tags") or row.get("product_tag") or []
        self.tables.append(
            "source_page",
            {
                "source_page_id": source_page_id,
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "source_type": source_type,
                "wp_post_type": wp_post_type,
                "wp_id": wp_id,
                "url": link,
                "slug": row.get("slug", ""),
                "title": clean_text(get_rendered(row.get("title", {}))),
                "status": row.get("status", ""),
                "published_at": row.get("date", ""),
                "modified_at": row.get("modified", ""),
                "author_id": row.get("author", ""),
                "excerpt_text": html_to_text(excerpt_html),
                "content_html": content_html,
                "content_text": html_to_text(content_html),
                "featured_media_id": row.get("featured_media", ""),
                "featured_image_url": embedded_featured_image(row),
                "category_ids": join_values(category_ids),
                "tag_ids": join_values(tag_ids),
                "job_type_ids": "",
                "seo_title": clean_text(yoast.get("title", "")),
                "seo_description": clean_text(yoast.get("description", "")),
                "canonical_url": yoast.get("canonical", link),
            },
        )
        self.source_counter += 1
        return source_page_id

    def collect_products(self) -> None:
        products = self.rest_objects.get("product", [])
        for index, product in enumerate(products, start=1):
            url = normalize_url(str(product.get("link", "")))
            if not url:
                continue
            self.product_by_url[url] = product
            category_slugs = self.product_category_slugs(product)
            brand_terms = self.product_brand_terms(product)
            tag_terms = self.product_tag_terms(product)
            is_vehicle = is_vehicle_product(category_slugs)
            image_url = embedded_featured_image(product) or image_from_yoast(product)
            if not is_vehicle:
                self.add_non_vehicle_product(product, image_url)
                continue

            self.log(f"Motorcycle detail {index}/{len(products)}: {url}")
            html = ""
            soup = BeautifulSoup("", "lxml")
            final_url = url
            try:
                html, soup, _, final_url = self.fetch_html(url)
            except Exception as exc:  # noqa: BLE001
                self.add_data_quality_issue(
                    entity_type="product",
                    entity_id=str(product.get("id", "")),
                    source_url=url,
                    issue_type="detail_page_fetch_failed",
                    issue_detail=str(exc),
                    severity="warning",
                )
            image_url = image_url or product_image_from_page(soup)
            self.add_vehicle_product(product, category_slugs, brand_terms, tag_terms, soup, html, image_url, final_url)
            if soup.find():
                self.collect_external_links_from_soup(soup, final_url, self.source_page_by_url.get(url, ""))

        self.collect_related_card_pricing()

    def add_vehicle_product(
        self,
        product: dict[str, object],
        category_slugs: set[str],
        brand_terms: list[str],
        tag_terms: list[str],
        soup: BeautifulSoup,
        html: str,
        image_url: str,
        final_url: str,
    ) -> None:
        url = normalize_url(str(product.get("link", "")))
        product_id = str(product.get("id", ""))
        motorcycle_id = f"motorcycle-{product_id}"
        title = clean_text(get_rendered(product.get("title", {})))
        brand_name = infer_brand_name(category_slugs, brand_terms, title)
        specs = extract_guanzon_specs(soup, get_rendered(product.get("content", {})))
        flattened = flatten_specs(specs)
        pricing = extract_guanzon_pricing(soup)
        excerpt_text = html_to_text(get_rendered(product.get("excerpt", {}))) or main_product_text(soup)
        vehicle_type, vehicle_type_source = infer_vehicle_type(category_slugs, tag_terms, specs, title)
        apply_url = find_internal_link(soup, "mc-credit-application") or f"{PRIMARY_DOMAIN}/mc-credit-application/"
        inquiry_url = find_internal_link(soup, "contact-us") or f"{PRIMARY_DOMAIN}/contact-us/"
        source_page_id = self.source_page_by_url.get(url, "")
        model_name = infer_model_name(title, brand_name)
        raw_taxonomies = sorted(set(category_slugs) | set(normalize_token(value) for value in brand_terms + tag_terms if value))
        primary_monthly = pricing.get("monthly_36_amount") or pricing.get("first_monthly_amount") or ""

        self.tables.append(
            "motorcycle",
            {
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "display_name": title,
                "brand_name": brand_name,
                "model_name": model_name,
                "vehicle_type": vehicle_type,
                "detail_url": final_url,
                "listing_url": category_listing_url(category_slugs),
                "srp_amount": pricing.get("srp_amount", ""),
                "downpayment_amount": pricing.get("downpayment_amount", ""),
                "monthly_amount": primary_monthly,
                "price_contexts": "product_detail_price_card" if pricing else "",
                "promo_titles": "",
                "is_blue_tag": "",
                "is_featured": "TRUE" if "featured" in " ".join(tag_terms).lower() else "",
                "is_top_pick": "",
                "has_vehicle_promo_flag": "TRUE" if any("promo" in value.lower() for value in tag_terms) else "",
                "spec_make": brand_name,
                "spec_model": flattened.get("spec_model") or model_name,
                "spec_transmission": flattened.get("spec_transmission", ""),
                "spec_engine_type": flattened.get("spec_engine_type", ""),
                "spec_seat_height": flattened.get("spec_seat_height", ""),
                "spec_brake_system": flattened.get("spec_brake_system", ""),
                "spec_displacement_cc": flattened.get("spec_displacement_cc", ""),
                "spec_front_tire": flattened.get("spec_front_tire", ""),
                "spec_rear_tire": flattened.get("spec_rear_tire", ""),
                "spec_wheels_type": flattened.get("spec_wheels_type", ""),
                "spec_starting_system": flattened.get("spec_starting_system", ""),
                "spec_overall_dimensions": flattened.get("spec_overall_dimensions", ""),
                "spec_ground_clearance": flattened.get("spec_ground_clearance", ""),
                "spec_fuel_system": flattened.get("spec_fuel_system", ""),
                "primary_image_url": image_url,
                "excerpt_text": excerpt_text,
                "apply_loan_url": apply_url,
                "inquiry_url": inquiry_url,
                "service_url": "",
                "compare_enabled": "",
                "published_at": product.get("date", ""),
                "modified_at": product.get("modified", ""),
                "vehicle_type_source": vehicle_type_source,
                "classification_notes": "Classified from Guanzon product category/spec evidence; uncertain exact big-bike membership is not guessed.",
                "raw_taxonomies": join_values(raw_taxonomies),
                "source_page_id": source_page_id,
                "source_page_url": url,
            },
        )
        self.motorcycle_by_url[url] = motorcycle_id
        self.motorcycle_summary[motorcycle_id] = {
            "display_name": title,
            "brand_name": brand_name,
            "model_name": model_name,
        }
        self.add_pricing_rows(motorcycle_id, title, brand_name, model_name, final_url, pricing)
        self.add_financing_rows(motorcycle_id, title, brand_name, model_name, final_url, pricing, apply_url)

        for index, spec in enumerate(specs, start=1):
            numeric, unit = split_numeric_unit(spec["value"])
            self.tables.append(
                "motorcycle_spec",
                {
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": title,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "spec_group": spec.get("group", "") or "Specifications",
                    "spec_label": spec["label"],
                    "spec_value_raw": spec["value"],
                    "spec_value_numeric": numeric,
                    "spec_unit": unit,
                    "display_order": index,
                    "source_url": final_url,
                },
            )

        for card in extract_related_cards(soup, final_url):
            self.pending_related_cards.append(card)
        if not specs:
            self.add_data_quality_issue(
                entity_type="motorcycle",
                entity_id=motorcycle_id,
                source_url=final_url,
                issue_type="missing_specs",
                issue_detail="No Guanzon specification block was parsed from this product detail page.",
                severity="info",
            )
        if not pricing:
            self.add_data_quality_issue(
                entity_type="motorcycle",
                entity_id=motorcycle_id,
                source_url=final_url,
                issue_type="missing_price_card",
                issue_detail="No Guanzon product price card was parsed from this product detail page.",
                severity="info",
            )

    def add_pricing_rows(
        self,
        motorcycle_id: str,
        display_name: str,
        brand_name: str,
        model_name: str,
        source_url: str,
        pricing: dict[str, object],
    ) -> None:
        if not pricing:
            return
        rows = []
        if pricing.get("srp_amount") or pricing.get("downpayment_amount") or pricing.get("monthly_36_amount"):
            rows.append(
                {
                    "pricing_id": f"pricing-{motorcycle_id}-detail-36",
                    "source_context": "product_detail_price_card_36_months",
                    "monthly_text": pricing.get("monthly_36_text", ""),
                    "monthly_amount": pricing.get("monthly_36_amount", ""),
                }
            )
        for term in (24, 18):
            amount = pricing.get(f"monthly_{term}_amount")
            text = pricing.get(f"monthly_{term}_text")
            if amount or text:
                rows.append(
                    {
                        "pricing_id": f"pricing-{motorcycle_id}-detail-{term}",
                        "source_context": f"product_detail_price_card_{term}_months",
                        "monthly_text": text,
                        "monthly_amount": amount,
                    }
                )
        if pricing.get("daily_text"):
            rows.append(
                {
                    "pricing_id": f"pricing-{motorcycle_id}-detail-daily",
                    "source_context": "product_detail_daily_installment",
                    "monthly_text": pricing.get("daily_text", ""),
                    "monthly_amount": pricing.get("daily_amount", ""),
                }
            )

        for row in rows:
            pricing_id = str(row["pricing_id"])
            if pricing_id in self.pricing_seen:
                continue
            self.pricing_seen.add(pricing_id)
            self.tables.append(
                "pricing",
                {
                    "pricing_id": pricing_id,
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": display_name,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "source_url": source_url,
                    "source_context": row.get("source_context", ""),
                    "currency": "PHP",
                    "srp_text": pricing.get("srp_text", ""),
                    "srp_amount": pricing.get("srp_amount", ""),
                    "downpayment_text": pricing.get("downpayment_text", ""),
                    "downpayment_amount": pricing.get("downpayment_amount", ""),
                    "monthly_text": row.get("monthly_text", ""),
                    "monthly_amount": row.get("monthly_amount", ""),
                    "price_note": pricing.get("price_note", ""),
                    "is_indicative": "TRUE",
                    "captured_at": datetime.now(TIMEZONE).isoformat(),
                },
            )

    def add_financing_rows(
        self,
        motorcycle_id: str,
        display_name: str,
        brand_name: str,
        model_name: str,
        source_url: str,
        pricing: dict[str, object],
        apply_url: str,
    ) -> None:
        for term in (36, 24, 18):
            amount = pricing.get(f"monthly_{term}_amount")
            if not amount:
                continue
            self.tables.append(
                "financing",
                {
                    "financing_id": f"financing-{motorcycle_id}-{term}",
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": display_name,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "source_url": source_url,
                    "financing_type": "installment_price_card",
                    "program_name": f"Guanzon {term}-month installment",
                    "monthly_amount": amount,
                    "downpayment_amount": pricing.get("downpayment_amount", ""),
                    "loan_amount": "",
                    "loan_term_months": term,
                    "estimated_income_amount": "",
                    "form_url": apply_url,
                    "cta_url": apply_url,
                    "note_text": pricing.get("price_note", ""),
                    "captured_at": datetime.now(TIMEZONE).isoformat(),
                },
            )
        if pricing.get("daily_amount"):
            self.tables.append(
                "financing",
                {
                    "financing_id": f"financing-{motorcycle_id}-daily",
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": display_name,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "source_url": source_url,
                    "financing_type": "daily_advertised_installment",
                    "program_name": "Guanzon daily advertised installment",
                    "monthly_amount": "",
                    "downpayment_amount": pricing.get("downpayment_amount", ""),
                    "loan_amount": "",
                    "loan_term_months": "",
                    "estimated_income_amount": "",
                    "form_url": apply_url,
                    "cta_url": apply_url,
                    "note_text": pricing.get("daily_text", ""),
                    "captured_at": datetime.now(TIMEZONE).isoformat(),
                },
            )

    def collect_related_card_pricing(self) -> None:
        for card in self.pending_related_cards:
            target_url = normalize_url(str(card.get("target_url", "")))
            motorcycle_id = self.motorcycle_by_url.get(target_url)
            if not motorcycle_id:
                continue
            summary = self.motorcycle_summary.get(motorcycle_id, {})
            srp_text = str(card.get("srp_text", ""))
            daily_text = str(card.get("daily_text", ""))
            pricing_id = f"pricing-{motorcycle_id}-related-{len(self.tables.rows['pricing']) + 1}"
            self.tables.append(
                "pricing",
                {
                    "pricing_id": pricing_id,
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": summary.get("display_name", card.get("display_name", "")),
                    "brand_name": summary.get("brand_name", card.get("brand_name", "")),
                    "model_name": summary.get("model_name", ""),
                    "source_url": card.get("source_url", ""),
                    "source_context": "related_product_card",
                    "currency": "PHP",
                    "srp_text": srp_text,
                    "srp_amount": parse_currency_amount(srp_text),
                    "downpayment_text": "",
                    "downpayment_amount": "",
                    "monthly_text": daily_text,
                    "monthly_amount": parse_currency_amount(daily_text),
                    "price_note": "Related-card price shown on another motorcycle detail page; daily amount is not a monthly payment.",
                    "is_indicative": "TRUE",
                    "captured_at": datetime.now(TIMEZONE).isoformat(),
                },
            )

    def add_non_vehicle_product(self, product: dict[str, object], image_url: str) -> None:
        url = normalize_url(str(product.get("link", "")))
        title = clean_text(get_rendered(product.get("title", {})))
        self.tables.append(
            "parts_catalog",
            {
                "part_id": f"product-{product.get('id')}",
                "dealer_id": DEALER_ID,
                "source_page_id": self.source_page_by_url.get(url, ""),
                "title": title,
                "slug": product.get("slug", ""),
                "detail_url": url,
                "image_url": image_url,
                "description_text": html_to_text(get_rendered(product.get("content", {}))) or html_to_text(get_rendered(product.get("excerpt", {}))),
                "published_at": product.get("date", ""),
                "modified_at": product.get("modified", ""),
            },
        )

    def product_category_slugs(self, product: dict[str, object]) -> set[str]:
        slugs: set[str] = set()
        for term_id in product.get("product_cat", []) or []:
            slugs.update(self.product_category_tree.get(int(term_id), set()))
        return slugs

    def product_brand_terms(self, product: dict[str, object]) -> list[str]:
        names: list[str] = []
        for term_id in product.get("product_brand", []) or []:
            term = self.terms_by_taxonomy_id.get(("product_brand", int(term_id)), {})
            name = clean_text(term.get("name", ""))
            if name:
                names.append(name)
        return names

    def product_tag_terms(self, product: dict[str, object]) -> list[str]:
        names: list[str] = []
        for term_id in product.get("product_tag", []) or []:
            term = self.terms_by_taxonomy_id.get(("product_tag", int(term_id)), {})
            name = clean_text(term.get("name", ""))
            if name:
                names.append(name)
        return names

    def collect_branches(self) -> None:
        for source_type, url in BRANCH_PAGES:
            try:
                _, soup, _, final_url = self.fetch_html(url)
            except Exception as exc:  # noqa: BLE001
                self.add_data_quality_issue(
                    entity_type="branch_page",
                    entity_id=source_type,
                    source_url=url,
                    issue_type="branch_page_fetch_failed",
                    issue_detail=str(exc),
                    severity="warning",
                )
                continue
            source_page_id = self.source_page_by_url.get(normalize_url(final_url), "")
            branches = parse_guanzon_branch_page(soup, final_url, source_type)
            self.log(f"Parsed {len(branches)} branches from {final_url}")
            for branch in branches:
                store_id = branch.get("locator_id", "") or slugify_text(f"{branch.get('branch_name', '')}-{branch.get('full_address', '')}")
                branch_id = f"branch-{store_id}"
                if branch_id in self.branch_seen:
                    self.collect_branch_coverage(branch_id, final_url, branch)
                    continue
                self.branch_seen.add(branch_id)
                self.tables.append(
                    "branch",
                    {
                        "branch_id": branch_id,
                        "dealer_id": DEALER_ID,
                        "source_url": final_url,
                        "locator_id": branch.get("locator_id", ""),
                        "branch_name": branch.get("branch_name", ""),
                        "slug": slugify_text(str(branch.get("branch_name", ""))),
                        "street": branch.get("street", ""),
                        "city": branch.get("city", ""),
                        "state_or_province": branch.get("state_or_province", ""),
                        "postal_code": branch.get("postal_code", ""),
                        "full_address": branch.get("full_address", ""),
                        "latitude": branch.get("latitude", ""),
                        "longitude": branch.get("longitude", ""),
                        "phone": branch.get("phone", ""),
                        "fax": "",
                        "email": branch.get("email", ""),
                        "website": branch.get("facebook_url", ""),
                        "open_hours_raw": "",
                        "days_str": "",
                        "brand_raw": branch.get("brand_raw", ""),
                        "product_raw": branch.get("product_raw", ""),
                        "categories_raw": branch.get("categories_raw", ""),
                        "inquire_url": branch.get("facebook_url", "") or f"{PRIMARY_DOMAIN}/contact-us/",
                    },
                )
                self.collect_branch_coverage(branch_id, final_url, branch)
            self.collect_external_links_from_soup(soup, final_url, source_page_id)

    def collect_branch_coverage(self, branch_id: str, source_url: str, branch: dict[str, str]) -> None:
        candidates = [
            ("source_branch_page", branch.get("source_type", "")),
            ("branch_category", branch.get("categories_raw", "")),
            ("company", branch.get("company_name", "")),
            ("product_family", branch.get("product_raw", "")),
        ]
        if branch.get("brand_raw"):
            for brand in split_pipe(branch.get("brand_raw", "")):
                candidates.append(("brand", brand))
        for coverage_type, raw_value in candidates:
            normalized = normalize_token(raw_value)
            if not normalized:
                continue
            key = (branch_id, coverage_type, normalized)
            if key in self.coverage_seen:
                continue
            self.coverage_seen.add(key)
            self.tables.append(
                "branch_product_coverage",
                {
                    "coverage_id": f"coverage-{branch_id}-{coverage_type}-{normalized}",
                    "dealer_id": DEALER_ID,
                    "branch_id": branch_id,
                    "coverage_type": coverage_type,
                    "coverage_value_raw": raw_value,
                    "coverage_value_normalized": normalized,
                    "source_url": source_url,
                    "source_field": "static branch page",
                    "confidence": "coverage_only",
                    "observed_at": datetime.now(TIMEZONE).isoformat(),
                    "notes": "Broad Guanzon branch category/source coverage only; not exact model availability.",
                },
            )

    def collect_posts_and_promos(self) -> None:
        products_by_url = {url: row for url, row in self.product_by_url.items()}
        category_terms = self.terms_by_taxonomy_id
        for post in self.rest_objects.get("post", []):
            url = normalize_url(str(post.get("link", "")))
            title = clean_text(get_rendered(post.get("title", {})))
            content_html = get_rendered(post.get("content", {}))
            content_text = html_to_text(content_html)
            category_labels = [
                clean_text(category_terms.get(("category", int(term_id)), {}).get("name", ""))
                for term_id in post.get("categories", []) or []
            ]
            if not looks_like_promo(title, content_text, category_labels):
                self.collect_external_links_from_html(content_html, url, self.source_page_by_url.get(url, ""))
                continue
            promo_id = f"promo-{post.get('id')}"
            if promo_id in self.promo_seen:
                continue
            self.promo_seen.add(promo_id)
            image_url = embedded_featured_image(post) or image_from_yoast(post)
            promo_text = f"{title} {content_text} {' '.join(category_labels)}"
            self.tables.append(
                "promo",
                {
                    "promo_id": promo_id,
                    "dealer_id": DEALER_ID,
                    "source_page_id": self.source_page_by_url.get(url, ""),
                    "source_page_url": url,
                    "source_url": url,
                    "title": title,
                    "promo_type": infer_promo_type(promo_text),
                    "description_text": content_text,
                    "image_url": image_url,
                    "cta_url": first_link_from_html(content_html),
                    "target_listing_url": f"{PRIMARY_DOMAIN}/promos/",
                    "applies_to_brand": infer_brand_from_text(promo_text),
                    "applies_to_vehicle_type": infer_vehicle_type_from_text(promo_text),
                    "applies_to_tag": join_values(category_labels),
                    "has_vehicle_promo_flag": "TRUE" if "motorcycle" in promo_text.lower() or infer_brand_from_text(promo_text) else "",
                    "linked_motorcycle_count": "",
                    "linked_motorcycles": "",
                    "link_types": "",
                    "link_source_urls": "",
                    "published_at": post.get("date", ""),
                    "modified_at": post.get("modified", ""),
                },
            )
            normalized_html = content_html.lower()
            haystack = f"{title} {content_text}".lower()
            for product_url, product in products_by_url.items():
                motorcycle_id = self.motorcycle_by_url.get(product_url)
                if not motorcycle_id:
                    continue
                product_title = clean_text(get_rendered(product.get("title", {})))
                if product_url.lower() in normalized_html or exact_product_title_mentioned(product_title, haystack):
                    self.add_promo_link(motorcycle_id, promo_id, "post_product_link_or_title_match", url)
            self.collect_external_links_from_html(content_html, url, self.source_page_by_url.get(url, ""))

    def add_promo_link(self, motorcycle_id: str, promo_id: str, link_type: str, source_url: str) -> None:
        key = (motorcycle_id, promo_id, link_type)
        if key in self.promo_link_seen:
            return
        self.promo_link_seen.add(key)
        motorcycle = self.motorcycle_summary.get(motorcycle_id, {})
        promo_title = ""
        for row in self.tables.rows["promo"]:
            if row.get("promo_id") == promo_id:
                promo_title = str(row.get("title", ""))
                break
        self.tables.append(
            "motorcycle_promo_link",
            {
                "motorcycle_id": motorcycle_id,
                "promo_id": promo_id,
                "dealer_id": DEALER_ID,
                "display_name": motorcycle.get("display_name", ""),
                "promo_title": promo_title,
                "link_type": link_type,
                "source_url": source_url,
                "notes": "Linked only when a promo exposed the product URL or an exact product-title text match.",
            },
        )

    def collect_forms(self) -> None:
        for url in TRACKED_FORM_URLS:
            try:
                _, soup, _, final_url = self.fetch_html(url)
            except Exception as exc:  # noqa: BLE001
                self.add_data_quality_issue(
                    entity_type="form",
                    entity_id=url,
                    source_url=url,
                    issue_type="form_page_fetch_failed",
                    issue_detail=str(exc),
                    severity="warning",
                )
                continue
            page_title = soup.title.get_text(" ", strip=True) if soup.title else ""
            for index, form in enumerate(soup.find_all("form"), start=1):
                if not is_meaningful_form(form):
                    continue
                action = urljoin(final_url, form.get("action", ""))
                form_key = (final_url, action)
                if form_key in self.form_seen:
                    continue
                self.form_seen.add(form_key)
                form_category = infer_form_category(final_url, page_title, form)
                form_id = f"form-{slug_from_url_path(final_url)}-{index}"
                fields = extract_form_fields(form, final_url)
                self.tables.append(
                    "form_definition",
                    {
                        "form_id": form_id,
                        "dealer_id": DEALER_ID,
                        "form_name": page_title.replace(" | Guanzon Merchandising Corporation", ""),
                        "form_category": form_category,
                        "source_url": url,
                        "final_url": final_url,
                        "page_title": page_title,
                        "form_action": action,
                        "form_method": (form.get("method") or "get").upper(),
                        "requires_file_upload": "TRUE" if form.get("enctype") == "multipart/form-data" or form.select('input[type="file"]') else "",
                        "requires_otp": "",
                        "multi_step": "",
                        "has_privacy_checkbox": "TRUE" if "privacy" in form.get_text(" ", strip=True).lower() or "consent" in form.get_text(" ", strip=True).lower() else "",
                        "notes": "Public form/filter schema only. No live submission performed.",
                    },
                )
                field_order = 1
                for field in fields:
                    options = field.get("options") or [{}]
                    for option in options:
                        self.tables.append(
                            "form_field",
                            {
                                "form_id": form_id,
                                "dealer_id": DEALER_ID,
                                "form_name": page_title.replace(" | Guanzon Merchandising Corporation", ""),
                                "form_category": form_category,
                                "field_order": field_order,
                                "field_name": field.get("field_name", ""),
                                "field_label": field.get("field_label", ""),
                                "field_type": field.get("field_type", ""),
                                "is_required": field.get("is_required", ""),
                                "default_value": field.get("default_value", ""),
                                "placeholder": field.get("placeholder", ""),
                                "option_value": option.get("option_value", ""),
                                "option_label": option.get("option_label", ""),
                                "section_name": field.get("section_name", ""),
                                "source_url": final_url,
                            },
                        )
                    field_order += 1
            self.collect_external_links_from_soup(soup, final_url, self.source_page_by_url.get(normalize_url(final_url), ""))

    def collect_services_and_jobs(self) -> None:
        career_soup = BeautifulSoup("", "lxml")
        for category, url in SERVICE_URLS:
            try:
                _, soup, _, final_url = self.fetch_html(url)
            except Exception:
                soup = BeautifulSoup("", "lxml")
                final_url = url
            title = soup.title.get_text(" ", strip=True).replace(" | Guanzon Merchandising Corporation", "") if soup.title else category.replace("_", " ").title()
            text = main_text(soup)
            if category == "careers":
                career_soup = soup
            self.tables.append(
                "aftersales_service",
                {
                    "service_id": f"service-{category}",
                    "dealer_id": DEALER_ID,
                    "source_page_id": self.source_page_by_url.get(normalize_url(final_url), ""),
                    "service_category": category,
                    "title": title,
                    "description_text": text,
                    "partner_name": "",
                    "cta_url": first_link_from_soup(soup),
                    "external_form_url": final_url if "application" in category or "credit" in category else "",
                    "notes": "Guanzon public service/support surface. Some credit/application pages are currently under construction.",
                },
            )
        self.collect_job_listings(career_soup)
        self.add_generic_financing_rows()

    def collect_job_listings(self, career_soup: BeautifulSoup) -> None:
        application_url = f"{PRIMARY_DOMAIN}/career-application/"
        positions = career_positions(career_soup)
        if not positions:
            try:
                _, app_soup, _, _ = self.fetch_html(application_url)
                positions = career_positions(app_soup)
            except Exception:
                positions = []
        if not positions:
            positions = ["Guanzon Careers"]
        for index, position in enumerate(positions, start=1):
            self.tables.append(
                "job_listing",
                {
                    "job_id": f"job-guanzon-{index}",
                    "dealer_id": DEALER_ID,
                    "source_page_id": self.source_page_by_url.get(normalize_url(f"{PRIMARY_DOMAIN}/career/"), ""),
                    "title": position,
                    "slug": slugify_text(position),
                    "job_location": "Philippines",
                    "application_target": application_url,
                    "company_name": DEALER_NAME,
                    "company_website": PRIMARY_DOMAIN,
                    "company_tagline": "",
                    "job_salary": "",
                    "salary_currency": "",
                    "salary_unit": "",
                    "is_featured": "",
                    "is_filled": "",
                    "is_remote": "",
                    "job_type_ids": "",
                    "description_text": "Position option exposed on the public Guanzon careers/application pages.",
                    "published_at": "",
                    "modified_at": "",
                },
            )

    def add_generic_financing_rows(self) -> None:
        for slug, label in [
            ("mc-credit-application", "Motorcycle Credit Application"),
            ("mp-credit-application", "Mobile Phone Credit Application"),
            ("mobile-phones/apply-for-a-loan", "Mobile Phone Apply For A Loan"),
        ]:
            source_url = f"{PRIMARY_DOMAIN}/{slug}/"
            self.tables.append(
                "financing",
                {
                    "financing_id": f"financing-guanzon-{slugify_text(slug)}",
                    "motorcycle_id": "",
                    "dealer_id": DEALER_ID,
                    "display_name": "",
                    "brand_name": "",
                    "model_name": "",
                    "source_url": source_url,
                    "financing_type": "generic_public_application_surface",
                    "program_name": label,
                    "monthly_amount": "",
                    "downpayment_amount": "",
                    "loan_amount": "",
                    "loan_term_months": "",
                    "estimated_income_amount": "",
                    "form_url": source_url,
                    "cta_url": source_url,
                    "note_text": "Public application page observed, but page content is currently under construction and no application form fields were exposed.",
                    "captured_at": datetime.now(TIMEZONE).isoformat(),
                },
            )

    def collect_relationship_probe(self) -> None:
        self.tables.append(
            "relationship_probe",
            {
                "probe_id": "probe-guanzon-branch-model-availability",
                "dealer_id": DEALER_ID,
                "relationship_name": "branch_to_motorcycle_model_availability",
                "source_url": f"{PRIMARY_DOMAIN}/branches/motorcycle-branches/",
                "source_type": "static branch pages, product detail pages, WordPress REST",
                "status": "coverage_only",
                "evidence_found": "FALSE",
                "rows_created": 0,
                "notes": "Guanzon branch pages expose branch categories, contacts, and coordinates. Public pages did not expose exact branch plus exact motorcycle model inventory evidence.",
                "checked_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def collect_data_quality_notes(self) -> None:
        self.add_data_quality_issue(
            entity_type="branch_model_availability",
            entity_id=DEALER_ID,
            source_url=f"{PRIMARY_DOMAIN}/branches/motorcycle-branches/",
            issue_type="no_exact_public_relationship",
            issue_detail="No exact public branch-to-motorcycle model bridge was found. branch_product_coverage is broad coverage only.",
            severity="warning",
        )
        self.add_data_quality_issue(
            entity_type="application_forms",
            entity_id=DEALER_ID,
            source_url=f"{PRIMARY_DOMAIN}/mc-credit-application/",
            issue_type="under_construction",
            issue_detail="Motorcycle/mobile-phone credit application pages are public but currently under construction; only page surfaces and CTAs were captured.",
            severity="info",
        )

    def add_data_quality_issue(
        self,
        *,
        entity_type: str,
        entity_id: str,
        source_url: str,
        issue_type: str,
        issue_detail: str,
        severity: str,
    ) -> None:
        self.tables.append(
            "data_quality_issue",
            {
                "issue_id": f"issue-{len(self.tables.rows['data_quality_issue']) + 1}",
                "dealer_id": DEALER_ID,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "source_url": source_url,
                "issue_type": issue_type,
                "issue_detail": issue_detail,
                "severity": severity,
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def collect_external_links_from_soup(self, soup: BeautifulSoup, source_url: str, source_page_id: str) -> None:
        for link in soup.find_all("a", href=True):
            href = urljoin(source_url, link.get("href", ""))
            self.add_external_cta(source_page_id, source_url, href, link.get_text(" ", strip=True))

    def collect_external_links_from_html(self, html: str, source_url: str, source_page_id: str) -> None:
        self.collect_external_links_from_soup(BeautifulSoup(html, "lxml"), source_url, source_page_id)

    def add_external_cta(self, source_page_id: str, source_url: str, target_url: str, anchor_text: str) -> None:
        parsed = urlparse(target_url)
        if not parsed.scheme.startswith("http"):
            return
        target_domain = parsed.netloc.lower()
        path = parsed.path.lower()
        is_guanzon = target_domain in {"guanzongroup.com.ph", "www.guanzongroup.com.ph"}
        if is_guanzon and not any(
            fragment in path
            for fragment in (
                "/mc-credit-application",
                "/mp-credit-application",
                "/mobile-phones/apply-for-a-loan",
                "/contact-us",
                "/career",
                "/branches/branch-locator",
                "/product/",
            )
        ):
            return
        key = (source_url, target_url)
        if key in self.external_cta_seen:
            return
        self.external_cta_seen.add(key)
        self.tables.append(
            "external_cta",
            {
                "cta_id": f"cta-{len(self.tables.rows['external_cta']) + 1}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "source_url": source_url,
                "target_url": target_url,
                "target_domain": target_domain,
                "cta_type": infer_cta_type(target_url, anchor_text),
                "anchor_text": anchor_text,
                "notes": "Public link/CTA evidence.",
            },
        )

    def finalize_crawl_run(self) -> None:
        self.tables.append(
            "crawl_run",
            {
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "run_started_at": self.run_started_at,
                "run_finished_at": datetime.now(TIMEZONE).isoformat(),
                "source_count": self.source_counter,
                "notes": "Current-state public Guanzon snapshot using WordPress REST, WooCommerce product HTML, static branch pages, public forms, and promo/career pages.",
            },
        )

    def write_workbook(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in DATABASE_SHEETS:
            schema = DATABASE_SCHEMAS[sheet_name]
            sheet = workbook.create_sheet(title=sheet_name[:31])
            sheet.append(schema)
            for row in self.tables.rows.get(sheet_name, []):
                sheet.append([safe_excel_value(row.get(column)) for column in schema])
        workbook.save(self.output_path)


def clean_text(value: object) -> str:
    if isinstance(value, dict):
        value = get_rendered(value)
    return BeautifulSoup(str(value or ""), "lxml").get_text(" ", strip=True).replace("\xa0", " ")


def safe_excel_value(value: object) -> object:
    value = sanitize_excel_value(value)
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", " ", value)
    return value


def join_values(values: object) -> str:
    if isinstance(values, (list, tuple, set)):
        return " | ".join(str(value) for value in values if str(value).strip())
    return str(values or "")


def split_pipe(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"\s*\|\s*|,\s*", value or "") if part.strip()]


def embedded_featured_image(row: dict[str, object]) -> str:
    embedded = row.get("_embedded")
    if isinstance(embedded, dict):
        media_rows = embedded.get("wp:featuredmedia")
        if isinstance(media_rows, list) and media_rows and isinstance(media_rows[0], dict):
            return str(media_rows[0].get("source_url") or media_rows[0].get("link") or "")
    return image_from_yoast(row)


def image_from_yoast(row: dict[str, object]) -> str:
    yoast = row.get("yoast_head_json") if isinstance(row.get("yoast_head_json"), dict) else {}
    images = yoast.get("og_image")
    if isinstance(images, list) and images and isinstance(images[0], dict):
        return str(images[0].get("url", ""))
    return ""


def product_image_from_page(soup: BeautifulSoup) -> str:
    image = soup.select_one(".woocommerce-product-gallery img, .product img")
    if image and image.get("src"):
        return str(image.get("src"))
    return ""


def is_vehicle_product(category_slugs: set[str]) -> bool:
    return "motorcycle" in category_slugs and "motorcycle-parts" not in category_slugs and "mobile-phone" not in category_slugs


def extract_guanzon_specs(soup: BeautifulSoup, fallback_html: str = "") -> list[dict[str, str]]:
    specs: list[dict[str, str]] = []
    for container in soup.select(".sp-spicifications-container, .sp-specifications-container"):
        label_tag = None
        value_tag = container.select_one(".sp-specs-value")
        for tag in container.find_all("p"):
            if tag is value_tag:
                continue
            text = clean_text(tag.get_text(" ", strip=True)).strip(" :")
            if text:
                label_tag = text
                break
        value = clean_text(value_tag.get_text(" ", strip=True)) if value_tag else ""
        if not label_tag:
            text = container.get_text(" ", strip=True)
            if ":" in text:
                label_tag, value = [part.strip() for part in text.split(":", 1)]
        if label_tag and value:
            specs.append({"group": "Specifications", "label": label_tag, "value": value})
    if specs:
        return specs

    fallback_soup = BeautifulSoup(fallback_html or "", "lxml")
    for row in fallback_soup.select("tr"):
        cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
        if len(cells) >= 2 and cells[0] and cells[1]:
            specs.append({"group": "Specifications", "label": cells[0].strip(" :"), "value": cells[1]})
    return specs


def flatten_specs(specs: list[dict[str, str]]) -> dict[str, str]:
    flattened: dict[str, str] = {}
    brake_values: list[str] = []
    for spec in specs:
        label = normalize_label(spec["label"])
        value = spec["value"]
        if label == "model":
            flattened["spec_model"] = value
        elif "transmission" in label:
            flattened["spec_transmission"] = value
        elif "engine type" in label:
            flattened["spec_engine_type"] = value
        elif "displacement" in label:
            flattened["spec_displacement_cc"] = value
        elif "front" in label and ("tire" in label or "tyre" in label):
            flattened["spec_front_tire"] = value
        elif "rear" in label and ("tire" in label or "tyre" in label):
            flattened["spec_rear_tire"] = value
        elif "wheel" in label and "type" in label:
            flattened["spec_wheels_type"] = value
        elif "starting system" in label:
            flattened["spec_starting_system"] = value
        elif "overall dimensions" in label or "length x width x height" in label or label in {"dimension", "dimensions", "l x w x h"}:
            flattened["spec_overall_dimensions"] = value
        elif "ground clearance" in label:
            flattened["spec_ground_clearance"] = value
        elif "fuel system" in label:
            flattened["spec_fuel_system"] = value
        elif "seat height" in label:
            flattened["spec_seat_height"] = value
        elif "brake" in label:
            brake_values.append(f"{spec['label']}: {value}")
    if brake_values:
        flattened["spec_brake_system"] = " | ".join(brake_values)
    return flattened


def extract_guanzon_pricing(soup: BeautifulSoup) -> dict[str, object]:
    card = soup.select_one(".sp-product-price")
    if not card:
        return {}
    raw = card.get_text(" ", strip=True)
    normalized = re.sub(r"\s+", " ", raw)
    pricing: dict[str, object] = {"raw_text": normalized}
    srp_match = re.search(r"SRP:\s*₱?\s*([\d,]+)", normalized, re.I)
    dp_match = re.search(r"\bDP:\s*₱?\s*([\d,]+)", normalized, re.I)
    daily_match = re.search(r"₱\s*([\d,]+)\s*/\s*day", normalized, re.I)
    if srp_match:
        amount_text = f"₱{srp_match.group(1)}"
        pricing["srp_text"] = amount_text
        pricing["srp_amount"] = parse_currency_amount(amount_text)
    if dp_match:
        amount_text = f"₱{dp_match.group(1)}"
        pricing["downpayment_text"] = amount_text
        pricing["downpayment_amount"] = parse_currency_amount(amount_text)
    if daily_match:
        amount_text = f"₱{daily_match.group(1)}/day"
        pricing["daily_text"] = amount_text
        pricing["daily_amount"] = parse_currency_amount(amount_text)
    first_monthly = ""
    for term, value in re.findall(r"(\d{2})\s*Months:\s*₱?\s*([\d,]+)", normalized, re.I):
        amount_text = f"₱{value}"
        pricing[f"monthly_{term}_text"] = amount_text
        pricing[f"monthly_{term}_amount"] = parse_currency_amount(amount_text)
        if not first_monthly:
            first_monthly = term
    if first_monthly:
        pricing["first_monthly_amount"] = pricing.get(f"monthly_{first_monthly}_amount", "")
    note = ""
    for note_tag in card.find_all(string=re.compile(r"prices shown|subject to change|vary by branch", re.I)):
        note = str(note_tag).strip()
        break
    pricing["price_note"] = note or "*Prices shown are for reference only and may vary by branch. All prices are subject to change without prior notice."
    return pricing


def extract_related_cards(soup: BeautifulSoup, source_url: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for card in soup.select(".pa-product"):
        link = card.find("a", href=True)
        if not link:
            continue
        target_url = normalize_url(urljoin(source_url, link.get("href", "")))
        text = card.get_text(" ", strip=True)
        srp_match = re.search(r"SRP:\s*₱?\s*([\d,]+)", text, re.I)
        daily_match = re.search(r"As low as\s*₱?\s*([\d,]+)\s*/\s*day", text, re.I)
        title_bits = [part.strip() for part in re.split(r"\s{2,}|\s+\|\s+", text) if part.strip()]
        rows.append(
            {
                "target_url": target_url,
                "source_url": source_url,
                "display_name": title_bits[0] if title_bits else "",
                "brand_name": infer_brand_from_text(text),
                "srp_text": f"₱{srp_match.group(1)}" if srp_match else "",
                "daily_text": f"₱{daily_match.group(1)}/day" if daily_match else "",
            }
        )
    return rows


def parse_guanzon_branch_page(soup: BeautifulSoup, source_url: str, source_type: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    content = soup.select_one(".br-content")
    if not content:
        return rows
    current_category = ""
    for child in content.find_all(recursive=False):
        category_tag = child.select_one(".br-slug-text")
        if category_tag:
            current_category = clean_text(category_tag.get_text(" ", strip=True))
            continue
        if "br-sub-category-container" not in child.get("class", []):
            continue
        number = select_text(child, ".br-item-total")
        company = select_text(child, ".br-post-meta-text")
        branch_short = select_text(child, ".br-post-text")
        address = select_text(child, ".br-post-address-text")
        phone = select_text(child, ".br-post-cn-text")
        email = select_text(child, ".br-post-email-text")
        facebook_tag = child.select_one(".br-post-fb-text a[href]")
        map_tag = child.select_one(".br-map-text a")
        lat, lng = split_lat_lng(map_tag.get("data-store", "") if map_tag else "")
        store_id = map_tag.get("data-store-id", "") if map_tag else ""
        province = infer_province(address)
        postal = infer_postal_code(address)
        product_family = "motorcycle" if source_type == "motorcycle_branch" else "mobile-phone"
        brands = infer_branch_brands(current_category, company, branch_short)
        branch_name = " - ".join(part for part in [company, branch_short] if part)
        rows.append(
            {
                "locator_id": store_id or number,
                "branch_name": branch_name,
                "company_name": company,
                "street": address,
                "city": branch_short,
                "state_or_province": province,
                "postal_code": postal,
                "full_address": address,
                "latitude": lat,
                "longitude": lng,
                "phone": phone,
                "email": email,
                "facebook_url": facebook_tag.get("href", "") if facebook_tag else "",
                "categories_raw": current_category,
                "brand_raw": " | ".join(brands),
                "product_raw": product_family,
                "source_type": source_type,
                "source_url": source_url,
            }
        )
    return rows


def select_text(parent: Tag, selector: str) -> str:
    tag = parent.select_one(selector)
    return clean_text(tag.get_text(" ", strip=True)) if tag else ""


def split_lat_lng(value: str) -> tuple[str, str]:
    parts = [part.strip() for part in str(value or "").split(",", 1)]
    if len(parts) == 2:
        return parts[0], parts[1]
    return "", ""


def infer_province(address: str) -> str:
    parts = [part.strip() for part in address.split(",") if part.strip()]
    return parts[-1] if len(parts) >= 2 else ""


def infer_postal_code(address: str) -> str:
    match = re.search(r"\b(\d{4})\b", address or "")
    return match.group(1) if match else ""


def infer_branch_brands(category: str, company: str, branch_short: str) -> list[str]:
    haystack = f"{category} {company} {branch_short}".lower()
    brands = []
    for slug, brand in MOTORCYCLE_BRANDS.items():
        if slug in haystack:
            brands.append(brand)
    for brand in MOBILE_PHONE_BRANDS:
        if brand in haystack:
            brands.append(brand.title() if brand != "tecno" else "Tecno Mobile")
    if "multibrand" in haystack:
        brands.append("Multibrand")
    return sorted(set(brands))


def infer_brand_name(category_slugs: set[str], brand_terms: list[str], title: str) -> str:
    if brand_terms:
        return brand_terms[0]
    lower = title.lower()
    for slug, brand in MOTORCYCLE_BRANDS.items():
        if slug in category_slugs or re.search(rf"\b{re.escape(slug)}\b", lower):
            return brand
    return ""


def infer_brand_from_text(text: str) -> str:
    lower = text.lower()
    for slug, brand in MOTORCYCLE_BRANDS.items():
        if re.search(rf"\b{re.escape(slug)}\b", lower):
            return brand
    for brand in MOBILE_PHONE_BRANDS:
        if re.search(rf"\b{re.escape(brand)}\b", lower):
            return brand.title() if brand != "tecno" else "Tecno Mobile"
    return ""


def infer_model_name(title: str, brand_name: str) -> str:
    model = title
    if brand_name and model.lower().startswith(brand_name.lower()):
        model = model[len(brand_name) :].strip()
    return model


def infer_vehicle_type(
    category_slugs: set[str],
    tag_terms: list[str],
    specs: list[dict[str, str]],
    title: str,
) -> tuple[str, str]:
    text = " ".join([title, " ".join(category_slugs), " ".join(tag_terms)]).lower()
    if "hatasu" in text or "e-bike" in text or "electric" in text:
        return "electric", "product_cat_or_tag"
    displacement = displacement_cc(specs)
    if displacement and displacement >= 400:
        return "big-bike", "spec_displacement_cc"
    if "big bike" in text or "big-bike" in text or "bigbike" in text:
        return "big-bike", "product_text"
    if "motorcycle" in category_slugs:
        return "regular-bike", "product_cat"
    return "", ""


def infer_vehicle_type_from_text(text: str) -> str:
    lower = text.lower()
    if "e-bike" in lower or "electric" in lower:
        return "electric"
    displacement_match = re.search(r"\b([4-9]\d{2}|1\d{3})\s*cc\b", lower)
    if "big bike" in lower or "bigbike" in lower or displacement_match:
        return "big-bike"
    if "motorcycle" in lower or infer_brand_from_text(lower):
        return "regular-bike"
    return ""


def displacement_cc(specs: list[dict[str, str]]) -> float:
    for spec in specs:
        if "displacement" in normalize_label(spec.get("label", "")):
            match = re.search(r"([\d,.]+)", spec.get("value", ""))
            if match:
                try:
                    return float(match.group(1).replace(",", ""))
                except ValueError:
                    return 0.0
    return 0.0


def category_listing_url(category_slugs: set[str]) -> str:
    if "honda" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/motorcycle/honda/"
    if "kawasaki" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/motorcycle/kawasaki/"
    if "suzuki" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/motorcycle/suzuki/"
    if "yamaha" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/motorcycle/yamaha/"
    if "hatasu" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/motorcycle/hatasu/"
    return f"{PRIMARY_DOMAIN}/product-category/motorcycle/"


def find_internal_link(soup: BeautifulSoup, path_fragment: str) -> str:
    for link in soup.find_all("a", href=True):
        href = normalize_url(urljoin(PRIMARY_DOMAIN, link.get("href", "")))
        if path_fragment in urlparse(href).path:
            return href
    return ""


def first_link_from_html(html: str) -> str:
    return first_link_from_soup(BeautifulSoup(html or "", "lxml"))


def first_link_from_soup(soup: BeautifulSoup) -> str:
    for link in soup.find_all("a", href=True):
        href = normalize_url(urljoin(PRIMARY_DOMAIN, link.get("href", "")))
        if href.startswith("http"):
            return href
    return ""


def main_product_text(soup: BeautifulSoup) -> str:
    product = soup.select_one(".product") or soup.select_one("main") or soup
    return html_to_text(str(product))


def main_text(soup: BeautifulSoup) -> str:
    main = soup.select_one("main") or soup.select_one("#main") or soup
    return html_to_text(str(main))


def looks_like_promo(title: str, content_text: str, category_labels: list[str]) -> bool:
    haystack = f"{title} {content_text} {' '.join(category_labels)}".lower()
    if any("promo" in label.lower() for label in category_labels):
        return True
    return any(keyword in haystack for keyword in PROMO_KEYWORDS)


def infer_promo_type(text: str) -> str:
    lower = text.lower()
    if "raffle" in lower:
        return "raffle"
    if "installment" in lower or "downpayment" in lower:
        return "financing"
    if "discount" in lower or "sale" in lower or "deal" in lower:
        return "discount"
    if "bundle" in lower or "free" in lower:
        return "bundle"
    return "campaign"


def exact_product_title_mentioned(product_title: str, haystack: str) -> bool:
    normalized_title = re.sub(r"\s+", " ", product_title.lower()).strip()
    if not normalized_title or len(normalized_title) < 4:
        return False
    if normalized_title in haystack:
        return True
    compact_title = re.sub(r"[^a-z0-9]+", "", normalized_title)
    compact_haystack = re.sub(r"[^a-z0-9]+", "", haystack)
    return len(compact_title) >= 5 and compact_title in compact_haystack


def is_meaningful_form(form: Tag) -> bool:
    field_names = [str(field.get("name") or "") for field in form.select("input, select, textarea")]
    meaningful = [name for name in field_names if name and name not in {"username", "password", "rememberme", "woocommerce-login-nonce", "_wp_http_referer"}]
    if any(name.startswith("wpsl-") for name in meaningful):
        return True
    if {"carPositn", "fname", "address", "contact_no", "email_add"} & set(meaningful):
        return True
    if "submit_contact_form" in meaningful:
        return True
    return len(meaningful) >= 3


def infer_form_category(url: str, page_title: str, form: Tag) -> str:
    lower = f"{url} {page_title}".lower()
    names = " ".join(str(field.get("name") or "") for field in form.select("input, select, textarea")).lower()
    if "wpsl-" in names or "branch-locator" in lower:
        return "branch_locator_filter"
    if "career" in lower or "carpositn" in names:
        return "career_application"
    if "credit" in lower or "loan" in lower:
        return "financing_application"
    if "contact" in lower:
        return "customer_care"
    return "public_form"


def career_positions(soup: BeautifulSoup) -> list[str]:
    positions: list[str] = []
    for option in soup.select('select[name="carPositn"] option'):
        text = clean_text(option.get_text(" ", strip=True))
        value = str(option.get("value") or "").strip()
        if not text or text.startswith("--") or not value:
            continue
        positions.append(text)
    if positions:
        return positions
    text = main_text(soup)
    for match in re.finditer(r"\b(?:ASA|CCA|DSA|FS)\s*-\s*[A-Za-z][A-Za-z /-]+|Technician", text):
        positions.append(match.group(0).strip())
    return sorted(set(positions))


def infer_cta_type(target_url: str, anchor_text: str) -> str:
    lower = f"{target_url} {anchor_text}".lower()
    if "facebook" in lower:
        return "social"
    if "maps" in lower or "branch-locator" in lower:
        return "map_or_locator"
    if "credit" in lower or "loan" in lower:
        return "financing_application"
    if "career" in lower:
        return "career"
    if "/product/" in lower:
        return "product_detail"
    if "contact" in lower:
        return "contact"
    return "link"


def normalize_label(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def normalize_token(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(value or "").lower()).strip("-")


def slugify_text(value: str) -> str:
    return normalize_token(value) or "item"


def slug_from_url_path(url: str) -> str:
    path = urlparse(url).path.strip("/")
    return slugify_text(path) or "home"


def main() -> int:
    output_path = Path(__file__).resolve().parents[2] / "working/outputs/dealer-databases/guanzon_database.xlsx"
    scraper = GuanzonDatabaseScraper(output_path)
    scraper.run()
    print(f"[OK] Wrote {output_path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
