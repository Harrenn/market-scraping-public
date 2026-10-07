#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from full_scrape.motortrade.scrape_motortrade_analytics import (
    DATABASE_SCHEMAS,
    DATABASE_SHEETS,
    TIMEZONE,
    USER_AGENT,
    TableStore,
    extract_form_fields,
    get_rendered,
    html_to_text,
    normalize_url,
    parse_currency_amount,
    sanitize_excel_value,
    split_numeric_unit,
)


DEALER_ID = "kservico"
DEALER_NAME = "KServico"
PRIMARY_DOMAIN = "https://kservico.com.ph"
FORMS_DOMAIN = "https://kservico.com.ph"
REST_BASE = f"{PRIMARY_DOMAIN}/wp-json/wp/v2"
REQUEST_DELAY_SECONDS = 0.05

REST_COLLECTIONS = [
    ("pages", "page", "pages"),
    ("posts", "post", "posts"),
    ("project", "project", "project"),
    ("store", "store", "store"),
    ("product", "product", "product"),
    ("um_form", "um_form", "um_form"),
]

TERM_COLLECTIONS = [
    ("categories", "category"),
    ("tags", "post_tag"),
    ("project_category", "project_category"),
    ("project_tag", "project_tag"),
    ("product_cat", "product_cat"),
    ("product_tag", "product_tag"),
]

TRACKED_FORM_URLS = [
    "https://kservico.com.ph/application-form/",
    "https://kservico.com.ph/request-for-quote/",
    "https://kservico.com.ph/contact/",
    "https://kservico.com.ph/join-our-team/",
    "https://kservico.com.ph/faq/",
]

SERVICE_URLS = [
    ("financing", "https://kservico.com.ph/application-form/"),
    ("quote_request", "https://kservico.com.ph/request-for-quote/"),
    ("customer_care", "https://kservico.com.ph/contact/"),
    ("careers", "https://kservico.com.ph/join-our-team/"),
    ("faq", "https://kservico.com.ph/faq/"),
    ("bike_repair", "https://kservico.com.ph/project/bike-repair/"),
]

VEHICLE_ROOT_SLUGS = {
    "motorcycles",
    "bajaj-re",
    "e-bike",
}

MOTORCYCLE_BRANDS = {
    "honda": "HONDA",
    "kawasaki": "KAWASAKI",
    "kymco": "KYMCO",
    "suzuki": "SUZUKI",
    "yamaha": "YAMAHA",
    "hatasu": "HATASU",
    "bajaj": "BAJAJ",
}

SPEC_SECTION_HEADINGS = {
    "engine",
    "dimensions",
    "suspension & brakes",
    "suspension and brakes",
    "tyres",
    "tires",
    "electricals",
    "features",
}

KNOWN_SPEC_LABELS = sorted(
    {
        "Engine Type",
        "Displacement",
        "Bore & Stroke",
        "Bore x Stroke",
        "Compression Ratio",
        "Maximum Power",
        "Max Power",
        "Maximum Torque",
        "Max Torque",
        "Starting System",
        "Engine Oil Capacity",
        "Fuel System",
        "Fuel Consumption",
        "Transmission Type",
        "Transmission",
        "Suspension Type (Front)",
        "Suspension Type (Rear)",
        "Front Suspension",
        "Rear Suspension",
        "Brake type (Front)",
        "Brake Type (Rear)",
        "Front Brake",
        "Rear Brake",
        "Tire Size & Type (Front)",
        "Tire Size & Type (Rear)",
        "Front & Rear Tyre Size",
        "Front & Rear Tyre Pressure",
        "Wheel Type",
        "Overall Dimensions (L x W x H)",
        "Length x Width x Height",
        "Dimensions",
        "Dimension",
        "Curb Weight",
        "Kerb Weight",
        "Gross Vehicle Weight",
        "Seat Height",
        "Seat Weight",
        "Wheelbase",
        "Ground Clearance",
        "Minimum Ground Clearance",
        "Full Tank Capacity",
        "Fuel Tank Capacity",
        "Ignition Type",
        "Battery Type",
        "Battery",
        "Motor",
        "Range",
        "Speed",
        "Load WT",
        "Max Speed",
        "Cooling System",
        "Clutch",
        "Turning Radius",
        "Seating Capacity",
        "System Voltage",
        "Headlamp",
        "Tail/Stop Lamp",
        "Reverse Lamp",
        "Horn",
        "Dashboard",
        "Seating",
        "Cargo Capacity",
        "Color Options",
    },
    key=len,
    reverse=True,
)


class KservicoDatabaseScraper:
    def __init__(self, output_path: Path):
        self.output_path = output_path
        self.session = requests.Session()
        self.session.headers.update(
            {
                "User-Agent": USER_AGENT,
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/json;q=0.8,*/*;q=0.7",
                "Accept-Language": "en-US,en;q=0.9",
            }
        )
        self.tables = TableStore(DATABASE_SCHEMAS)
        self.source_counter = 0
        self.rest_objects: dict[str, list[dict[str, object]]] = defaultdict(list)
        self.terms_by_id: dict[int, dict[str, object]] = {}
        self.product_category_tree: dict[int, set[str]] = {}
        self.source_page_by_url: dict[str, str] = {}
        self.motorcycle_by_url: dict[str, str] = {}
        self.product_by_url: dict[str, dict[str, object]] = {}
        self.branch_pages: dict[str, dict[str, object]] = {}
        self.external_cta_seen: set[tuple[str, str]] = set()
        self.form_seen: set[tuple[str, str]] = set()
        self.coverage_seen: set[tuple[str, str, str]] = set()
        self.promo_seen: set[str] = set()
        self.promo_link_seen: set[tuple[str, str, str]] = set()

        now = datetime.now(TIMEZONE)
        self.crawl_date = now.date().isoformat()
        self.run_started_at = now.isoformat()
        self.run_id = f"{DEALER_ID}-run-{now.strftime('%Y%m%d%H%M%S')}"

    def run(self) -> None:
        self.log("Starting KServico database scrape")
        self.add_dealer()
        self.log("Collecting taxonomy/filter references")
        self.collect_terms()
        self.log("Collecting WordPress REST objects")
        self.collect_rest_objects()
        self.log("Collecting products and vehicle detail pages")
        self.collect_products()
        self.log("Collecting branch/store pages")
        self.collect_branches()
        self.log("Collecting promos and blog campaign surfaces")
        self.collect_posts_and_promos()
        self.log("Collecting public forms")
        self.collect_forms()
        self.log("Collecting services and jobs")
        self.collect_services_and_jobs()
        self.collect_relationship_probe()
        self.collect_data_quality_notes()
        self.finalize_crawl_run()
        self.log("Writing workbook")
        self.write_workbook()

    def log(self, message: str) -> None:
        print(f"[KServico] {message}", flush=True)

    def add_dealer(self) -> None:
        self.tables.append(
            "dealer",
            {
                "dealer_id": DEALER_ID,
                "dealer_name": DEALER_NAME,
                "primary_domain": PRIMARY_DOMAIN,
                "forms_domain": FORMS_DOMAIN,
                "crawl_timezone": str(TIMEZONE),
                "crawl_date": self.crawl_date,
                "notes": "Public KServico WordPress/WooCommerce snapshot using the shared dealer database schema.",
            },
        )

    def fetch(self, url: str, *, params: dict[str, object] | None = None, timeout: int = 45) -> requests.Response:
        response = self.session.get(url, params=params, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
        return response

    def fetch_html(self, url: str) -> tuple[str, BeautifulSoup, int, str]:
        time.sleep(REQUEST_DELAY_SECONDS)
        response = self.fetch(url)
        soup = BeautifulSoup(response.text, "lxml")
        return response.text, soup, response.status_code, response.url

    def fetch_json(self, url: str, *, params: dict[str, object] | None = None) -> object:
        time.sleep(REQUEST_DELAY_SECONDS)
        return self.fetch(url, params=params).json()

    def fetch_rest_collection(self, rest_base: str, *, per_page: int = 20) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        page = 1
        total_pages = 1
        while page <= total_pages:
            url = f"{REST_BASE}/{rest_base}"
            try:
                response = self.fetch(url, params={"per_page": per_page, "page": page, "_embed": 1})
            except requests.HTTPError as exc:
                self.add_data_quality_issue(
                    entity_type="rest_collection",
                    entity_id=rest_base,
                    source_url=f"{url}?per_page={per_page}&page={page}",
                    issue_type="rest_page_error",
                    issue_detail=f"Skipped REST page after HTTP error: {exc}",
                    severity="warning",
                )
                page += 1
                continue
            data = response.json()
            if isinstance(data, list):
                rows.extend(data)
            total_pages = int(response.headers.get("X-WP-TotalPages") or total_pages)
            self.log(f"REST {rest_base}: page {page}/{total_pages}, rows so far {len(rows)}")
            page += 1
        return rows

    def collect_terms(self) -> None:
        for rest_base, taxonomy in TERM_COLLECTIONS:
            per_page = 100 if rest_base != "tags" else 50
            for term in self.fetch_rest_collection(rest_base, per_page=per_page):
                term_id = int(term.get("id") or 0)
                self.terms_by_id[term_id] = term
                self.tables.append(
                    "website_filter_reference",
                    {
                        "dealer_id": DEALER_ID,
                        "reference_type": taxonomy,
                        "filter_name": taxonomy,
                        "option_value": term.get("slug", ""),
                        "option_label": clean_text(term.get("name", "")),
                        "source_url": term.get("link", ""),
                        "notes": "Source category/tag reference only. Do not join as exact product membership without product category evidence.",
                    },
                )
        self.product_category_tree = {term_id: self.category_slug_closure(term_id) for term_id in self.terms_by_id}

    def category_slug_closure(self, term_id: int) -> set[str]:
        slugs: set[str] = set()
        seen: set[int] = set()
        current_id = term_id
        while current_id and current_id not in seen:
            seen.add(current_id)
            term = self.terms_by_id.get(current_id, {})
            if str(term.get("taxonomy", "")) == "product_cat":
                slug = str(term.get("slug", "")).strip()
                if slug:
                    slugs.add(slug)
            current_id = int(term.get("parent") or 0)
        return slugs

    def collect_rest_objects(self) -> None:
        for rest_base, wp_post_type, source_type in REST_COLLECTIONS:
            per_page = 20
            if rest_base == "pages":
                per_page = 5
            rows = self.fetch_rest_collection(rest_base, per_page=per_page)
            self.rest_objects[wp_post_type].extend(rows)
            for row in rows:
                self.add_source_page(row, source_type=source_type, wp_post_type=wp_post_type)

    def add_source_page(self, row: dict[str, object], *, source_type: str, wp_post_type: str) -> str:
        wp_id = str(row.get("id", ""))
        link = normalize_url(str(row.get("link") or get_rendered(row.get("guid", {})) or ""))
        source_page_id = f"source-{wp_post_type}-{wp_id}"
        if link:
            self.source_page_by_url[link] = source_page_id
        yoast = row.get("yoast_head_json") if isinstance(row.get("yoast_head_json"), dict) else {}
        content_html = get_rendered(row.get("content", {}))
        excerpt_html = get_rendered(row.get("excerpt", {}))
        featured_image_url = embedded_featured_image(row)
        category_ids = row.get("categories") or row.get("product_cat") or []
        tag_ids = row.get("tags") or row.get("product_tag") or []
        self.tables.append(
            "source_page",
            {
                "source_page_id": source_page_id,
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "source_type": source_type,
                "wp_post_type": wp_post_type,
                "wp_id": wp_id,
                "url": link,
                "slug": row.get("slug", ""),
                "title": clean_text(get_rendered(row.get("title", {}))),
                "status": row.get("status", ""),
                "published_at": row.get("date", ""),
                "modified_at": row.get("modified", ""),
                "author_id": row.get("author", ""),
                "excerpt_text": html_to_text(excerpt_html),
                "content_html": content_html,
                "content_text": html_to_text(content_html),
                "featured_media_id": row.get("featured_media", ""),
                "featured_image_url": featured_image_url,
                "category_ids": join_values(category_ids),
                "tag_ids": join_values(tag_ids),
                "job_type_ids": "",
                "seo_title": clean_text(yoast.get("title", "")),
                "seo_description": clean_text(yoast.get("description", "")),
                "canonical_url": yoast.get("canonical", link),
            },
        )
        self.source_counter += 1
        return source_page_id

    def collect_products(self) -> None:
        products = self.rest_objects.get("product", [])
        for index, product in enumerate(products, start=1):
            url = normalize_url(str(product.get("link", "")))
            if not url:
                continue
            self.product_by_url[url] = product
            category_slugs = self.product_category_slugs(product)
            is_vehicle = bool(category_slugs & VEHICLE_ROOT_SLUGS)
            html = ""
            soup = BeautifulSoup("", "lxml")
            final_url = url
            image_url = embedded_featured_image(product) or product_image_from_schema(soup)
            if is_vehicle:
                self.log(f"Vehicle detail {index}/{len(products)}: {url}")
                try:
                    html, soup, _, final_url = self.fetch_html(url)
                except Exception as exc:  # noqa: BLE001
                    self.add_data_quality_issue(
                        entity_type="product",
                        entity_id=str(product.get("id", "")),
                        source_url=url,
                        issue_type="detail_page_fetch_failed",
                        issue_detail=str(exc),
                        severity="warning",
                    )
                price = extract_product_price(soup)
                image_url = embedded_featured_image(product) or product_image_from_schema(soup)
                self.add_vehicle_product(product, category_slugs, soup, price, image_url, final_url)
            else:
                self.add_non_vehicle_product(product, image_url)
            if soup.find():
                self.collect_external_links_from_soup(soup, final_url, self.source_page_by_url.get(url, ""))

    def add_vehicle_product(
        self,
        product: dict[str, object],
        category_slugs: set[str],
        soup: BeautifulSoup,
        price: dict[str, object],
        image_url: str,
        final_url: str,
    ) -> None:
        url = normalize_url(str(product.get("link", "")))
        product_id = str(product.get("id", ""))
        motorcycle_id = f"motorcycle-{product_id}"
        title = clean_text(get_rendered(product.get("title", {})))
        brand_name = infer_brand_name(category_slugs, title)
        vehicle_type = infer_vehicle_type(category_slugs)
        specs = parse_specs_from_html(get_rendered(product.get("content", {})))
        flattened = flatten_specs(specs)
        excerpt_text = html_to_text(get_rendered(product.get("excerpt", {})))
        apply_url = find_internal_link(soup, "application-form") or f"{PRIMARY_DOMAIN}/application-form/"
        quote_url = find_internal_link(soup, "request-for-quote") or f"{PRIMARY_DOMAIN}/request-for-quote/"
        source_page_id = self.source_page_by_url.get(url, "")

        self.tables.append(
            "motorcycle",
            {
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "display_name": title,
                "brand_name": brand_name,
                "model_name": infer_model_name(title, brand_name),
                "vehicle_type": vehicle_type,
                "detail_url": final_url,
                "listing_url": category_listing_url(category_slugs),
                "srp_amount": price.get("srp_amount", ""),
                "downpayment_amount": "",
                "monthly_amount": "",
                "price_contexts": "product_detail_schema" if price.get("srp_amount") not in ("", None) else "",
                "promo_titles": "",
                "is_blue_tag": "",
                "is_featured": "",
                "is_top_pick": "",
                "has_vehicle_promo_flag": "",
                "spec_make": brand_name,
                "spec_model": flattened.get("spec_model") or infer_model_name(title, brand_name),
                "spec_transmission": flattened.get("spec_transmission", ""),
                "spec_engine_type": flattened.get("spec_engine_type", ""),
                "spec_seat_height": flattened.get("spec_seat_height", ""),
                "spec_brake_system": flattened.get("spec_brake_system", ""),
                "spec_displacement_cc": flattened.get("spec_displacement_cc", ""),
                "spec_front_tire": flattened.get("spec_front_tire", ""),
                "spec_rear_tire": flattened.get("spec_rear_tire", ""),
                "spec_wheels_type": flattened.get("spec_wheels_type", ""),
                "spec_starting_system": flattened.get("spec_starting_system", ""),
                "spec_overall_dimensions": flattened.get("spec_overall_dimensions", ""),
                "spec_ground_clearance": flattened.get("spec_ground_clearance", ""),
                "spec_fuel_system": flattened.get("spec_fuel_system", ""),
                "primary_image_url": image_url,
                "excerpt_text": excerpt_text,
                "apply_loan_url": apply_url,
                "inquiry_url": quote_url,
                "service_url": "",
                "compare_enabled": "",
                "published_at": product.get("date", ""),
                "modified_at": product.get("modified", ""),
                "vehicle_type_source": "product_cat",
                "classification_notes": "Classified from KServico product categories.",
                "raw_taxonomies": join_values(sorted(category_slugs)),
                "source_page_id": source_page_id,
                "source_page_url": url,
            },
        )
        self.motorcycle_by_url[url] = motorcycle_id
        self.add_pricing_row(
            pricing_id=f"pricing-{motorcycle_id}-detail",
            motorcycle_id=motorcycle_id,
            source_url=final_url,
            source_context="product_detail_schema",
            price=price,
            note_text=excerpt_text,
        )
        self.add_financing_row(
            financing_id=f"financing-{motorcycle_id}-installment",
            motorcycle_id=motorcycle_id,
            source_url=final_url,
            financing_type="installment_cta",
            program_name="KServico Installment",
            form_url=apply_url,
            cta_url=apply_url,
            note_text=excerpt_text,
        )
        for index, spec in enumerate(specs, start=1):
            numeric, unit = split_numeric_unit(spec["value"])
            self.tables.append(
                "motorcycle_spec",
                {
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": title,
                    "brand_name": brand_name,
                    "model_name": infer_model_name(title, brand_name),
                    "spec_group": spec.get("group", "") or "Specifications",
                    "spec_label": spec["label"],
                    "spec_value_raw": spec["value"],
                    "spec_value_numeric": numeric,
                    "spec_unit": unit,
                    "display_order": index,
                    "source_url": final_url,
                },
            )

    def add_non_vehicle_product(self, product: dict[str, object], image_url: str) -> None:
        url = normalize_url(str(product.get("link", "")))
        self.tables.append(
            "parts_catalog",
            {
                "part_id": f"product-{product.get('id')}",
                "dealer_id": DEALER_ID,
                "source_page_id": self.source_page_by_url.get(url, ""),
                "title": clean_text(get_rendered(product.get("title", {}))),
                "slug": product.get("slug", ""),
                "detail_url": url,
                "image_url": image_url,
                "description_text": html_to_text(get_rendered(product.get("content", {}))) or html_to_text(get_rendered(product.get("excerpt", {}))),
                "published_at": product.get("date", ""),
                "modified_at": product.get("modified", ""),
            },
        )

    def add_pricing_row(
        self,
        *,
        pricing_id: str,
        motorcycle_id: str,
        source_url: str,
        source_context: str,
        price: dict[str, object],
        note_text: str,
    ) -> None:
        if price.get("srp_amount") in ("", None) and not price.get("srp_text"):
            return
        self.tables.append(
            "pricing",
            {
                "pricing_id": pricing_id,
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "display_name": "",
                "brand_name": "",
                "model_name": "",
                "source_url": source_url,
                "source_context": source_context,
                "currency": "PHP",
                "srp_text": price.get("srp_text", ""),
                "srp_amount": price.get("srp_amount", ""),
                "downpayment_text": "",
                "downpayment_amount": "",
                "monthly_text": "",
                "monthly_amount": "",
                "price_note": note_text,
                "is_indicative": "TRUE",
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def add_financing_row(
        self,
        *,
        financing_id: str,
        motorcycle_id: str,
        source_url: str,
        financing_type: str,
        program_name: str,
        form_url: str,
        cta_url: str,
        note_text: str,
    ) -> None:
        self.tables.append(
            "financing",
            {
                "financing_id": financing_id,
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "display_name": "",
                "brand_name": "",
                "model_name": "",
                "source_url": source_url,
                "financing_type": financing_type,
                "program_name": program_name,
                "monthly_amount": "",
                "downpayment_amount": "",
                "loan_amount": "",
                "loan_term_months": "",
                "estimated_income_amount": "",
                "form_url": form_url,
                "cta_url": cta_url,
                "note_text": note_text,
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def product_category_slugs(self, product: dict[str, object]) -> set[str]:
        slugs: set[str] = set()
        for term_id in product.get("product_cat", []) or []:
            slugs.update(self.product_category_tree.get(int(term_id), set()))
        return slugs

    def collect_branches(self) -> None:
        stores = self.rest_objects.get("store", [])
        for index, store in enumerate(stores, start=1):
            url = normalize_url(str(store.get("link", "")))
            if not url:
                continue
            self.log(f"Branch detail {index}/{len(stores)}: {url}")
            title = clean_text(get_rendered(store.get("title", {})))
            source_page_id = self.source_page_by_url.get(url, "")
            html = ""
            soup = BeautifulSoup("", "lxml")
            final_url = url
            try:
                html, soup, _, final_url = self.fetch_html(url)
            except Exception as exc:  # noqa: BLE001
                self.add_data_quality_issue(
                    entity_type="branch",
                    entity_id=str(store.get("id", "")),
                    source_url=url,
                    issue_type="branch_page_fetch_failed",
                    issue_detail=str(exc),
                    severity="warning",
                )
            branch = parse_store_page(soup, title)
            branch_id = f"branch-{store.get('id')}"
            self.tables.append(
                "branch",
                {
                    "branch_id": branch_id,
                    "dealer_id": DEALER_ID,
                    "source_url": final_url,
                    "locator_id": store.get("id", ""),
                    "branch_name": title,
                    "slug": store.get("slug", ""),
                    "street": "",
                    "city": "",
                    "state_or_province": "",
                    "postal_code": branch.get("postal_code", ""),
                    "full_address": branch.get("address", ""),
                    "latitude": "",
                    "longitude": "",
                    "phone": branch.get("phone", ""),
                    "fax": "",
                    "email": branch.get("email", ""),
                    "website": final_url,
                    "open_hours_raw": branch.get("open_hours", ""),
                    "days_str": branch.get("days_str", ""),
                    "brand_raw": branch.get("brand_raw", ""),
                    "product_raw": branch.get("product_raw", ""),
                    "categories_raw": branch.get("categories_raw", ""),
                    "inquire_url": f"{PRIMARY_DOMAIN}/request-for-quote/",
                },
            )
            self.branch_pages[branch_id] = {"url": final_url, "title": title, "text": html_to_text(html)}
            self.collect_branch_coverage(branch_id, final_url, title, branch)
            self.collect_external_links_from_soup(soup, final_url, source_page_id)

    def collect_branch_coverage(self, branch_id: str, source_url: str, title: str, branch: dict[str, str]) -> None:
        haystack = f"{title} {branch.get('categories_raw', '')} {branch.get('product_raw', '')} {branch.get('brand_raw', '')}".lower()
        coverage_values = {
            "motorcycle": "motorcycle" in haystack,
            "three-wheeler": any(term in haystack for term in ("bajaj", "three", "tricycle")),
            "e-bike": "e-bike" in haystack or "ebike" in haystack,
            "appliance": "appliance" in haystack,
            "gadget": "gadget" in haystack or "cellphone" in haystack or "laptop" in haystack,
            "furniture": "furniture" in haystack,
            "parts": "parts" in haystack,
            "honda": "honda" in haystack,
            "kawasaki": "kawasaki" in haystack,
            "yamaha": "yamaha" in haystack,
            "bajaj": "bajaj" in haystack,
        }
        for value, present in coverage_values.items():
            if not present:
                continue
            key = (branch_id, "public_branch_text", value)
            if key in self.coverage_seen:
                continue
            self.coverage_seen.add(key)
            self.tables.append(
                "branch_product_coverage",
                {
                    "coverage_id": f"coverage-{branch_id}-{value}",
                    "dealer_id": DEALER_ID,
                    "branch_id": branch_id,
                    "coverage_type": "public_branch_text",
                    "coverage_value_raw": value,
                    "coverage_value_normalized": value,
                    "source_url": source_url,
                    "source_field": "store detail text",
                    "confidence": "coverage_only",
                    "observed_at": datetime.now(TIMEZONE).isoformat(),
                    "notes": "Broad branch/store page coverage hint only; not exact model availability.",
                },
            )

    def collect_posts_and_promos(self) -> None:
        products_by_url = {url: row for url, row in self.product_by_url.items()}
        for post in self.rest_objects.get("post", []):
            url = normalize_url(str(post.get("link", "")))
            title = clean_text(get_rendered(post.get("title", {})))
            content_text = html_to_text(get_rendered(post.get("content", {})))
            if not looks_like_promo(title, content_text):
                continue
            promo_id = f"promo-{post.get('id')}"
            if promo_id in self.promo_seen:
                continue
            self.promo_seen.add(promo_id)
            source_page_id = self.source_page_by_url.get(url, "")
            image_url = embedded_featured_image(post)
            self.tables.append(
                "promo",
                {
                    "promo_id": promo_id,
                    "dealer_id": DEALER_ID,
                    "source_page_id": source_page_id,
                    "source_page_url": url,
                    "source_url": url,
                    "title": title,
                    "promo_type": infer_promo_type(title, content_text),
                    "description_text": content_text,
                    "image_url": image_url,
                    "cta_url": first_link_from_html(get_rendered(post.get("content", {}))),
                    "target_listing_url": "",
                    "applies_to_brand": infer_brand_name(set(), title + " " + content_text),
                    "applies_to_vehicle_type": infer_vehicle_type_from_text(title + " " + content_text),
                    "applies_to_tag": "",
                    "has_vehicle_promo_flag": "TRUE" if infer_vehicle_type_from_text(title + " " + content_text) else "",
                    "linked_motorcycle_count": "",
                    "linked_motorcycles": "",
                    "link_types": "",
                    "link_source_urls": "",
                    "published_at": post.get("date", ""),
                    "modified_at": post.get("modified", ""),
                },
            )
            content_html = get_rendered(post.get("content", {}))
            normalized_html = content_html.lower()
            for product_url in products_by_url:
                motorcycle_id = self.motorcycle_by_url.get(product_url)
                if not motorcycle_id:
                    continue
                product = products_by_url[product_url]
                product_title = clean_text(get_rendered(product.get("title", {})))
                if product_url.lower() in normalized_html or product_title.lower() in (title + " " + content_text).lower():
                    self.add_promo_link(motorcycle_id, promo_id, "post_product_link_or_title_match", url)
            self.collect_external_links_from_html(content_html, url, source_page_id)

    def add_promo_link(self, motorcycle_id: str, promo_id: str, link_type: str, source_url: str) -> None:
        key = (motorcycle_id, promo_id, link_type)
        if key in self.promo_link_seen:
            return
        self.promo_link_seen.add(key)
        motorcycle_title = ""
        promo_title = ""
        for row in self.tables.rows["motorcycle"]:
            if row.get("motorcycle_id") == motorcycle_id:
                motorcycle_title = str(row.get("display_name", ""))
                break
        for row in self.tables.rows["promo"]:
            if row.get("promo_id") == promo_id:
                promo_title = str(row.get("title", ""))
                break
        self.tables.append(
            "motorcycle_promo_link",
            {
                "motorcycle_id": motorcycle_id,
                "promo_id": promo_id,
                "dealer_id": DEALER_ID,
                "display_name": motorcycle_title,
                "promo_title": promo_title,
                "link_type": link_type,
                "source_url": source_url,
                "notes": "Linked only when a promo post exposed the product URL or exact product title text.",
            },
        )

    def collect_forms(self) -> None:
        for url in TRACKED_FORM_URLS:
            try:
                _, soup, _, final_url = self.fetch_html(url)
            except Exception as exc:  # noqa: BLE001
                self.add_data_quality_issue(
                    entity_type="form",
                    entity_id=url,
                    source_url=url,
                    issue_type="form_page_fetch_failed",
                    issue_detail=str(exc),
                    severity="warning",
                )
                continue
            page_title = soup.title.get_text(" ", strip=True) if soup.title else ""
            for index, form in enumerate(soup.find_all("form"), start=1):
                if not is_meaningful_form(form):
                    continue
                action = urljoin(final_url, form.get("action", ""))
                form_key = (final_url, action)
                if form_key in self.form_seen:
                    continue
                self.form_seen.add(form_key)
                form_category = infer_form_category(final_url, page_title)
                form_id = f"form-{slug_from_url_path(final_url)}-{index}"
                fields = extract_form_fields(form, final_url)
                self.tables.append(
                    "form_definition",
                    {
                        "form_id": form_id,
                        "dealer_id": DEALER_ID,
                        "form_name": page_title.replace(" - KServico", ""),
                        "form_category": form_category,
                        "source_url": url,
                        "final_url": final_url,
                        "page_title": page_title,
                        "form_action": action,
                        "form_method": (form.get("method") or "get").upper(),
                        "requires_file_upload": "TRUE" if form.get("enctype") == "multipart/form-data" or form.select('input[type="file"]') else "",
                        "requires_otp": "",
                        "multi_step": "TRUE" if form.select(".step, .wpcf7cf-step") else "",
                        "has_privacy_checkbox": "TRUE" if "privacy" in form.get_text(" ", strip=True).lower() or "consent" in form.get_text(" ", strip=True).lower() else "",
                        "notes": "Contact Form 7/public form schema only. No live submission performed.",
                    },
                )
                field_order = 1
                for field in fields:
                    options = field.get("options") or [{}]
                    for option in options:
                        self.tables.append(
                            "form_field",
                            {
                                "form_id": form_id,
                                "dealer_id": DEALER_ID,
                                "form_name": page_title.replace(" - KServico", ""),
                                "form_category": form_category,
                                "field_order": field_order,
                                "field_name": field.get("field_name", ""),
                                "field_label": field.get("field_label", ""),
                                "field_type": field.get("field_type", ""),
                                "is_required": field.get("is_required", ""),
                                "default_value": field.get("default_value", ""),
                                "placeholder": field.get("placeholder", ""),
                                "option_value": option.get("option_value", ""),
                                "option_label": option.get("option_label", ""),
                                "section_name": field.get("section_name", ""),
                                "source_url": final_url,
                            },
                        )
                    field_order += 1
            self.collect_external_links_from_soup(soup, final_url, self.source_page_by_url.get(normalize_url(final_url), ""))

    def collect_services_and_jobs(self) -> None:
        for category, url in SERVICE_URLS:
            try:
                _, soup, _, final_url = self.fetch_html(url)
            except Exception:
                final_url = url
                soup = BeautifulSoup("", "lxml")
            title = soup.title.get_text(" ", strip=True).replace(" - KServico", "") if soup.title else category.replace("_", " ").title()
            text = main_text(soup)
            self.tables.append(
                "aftersales_service",
                {
                    "service_id": f"service-{category}",
                    "dealer_id": DEALER_ID,
                    "source_page_id": self.source_page_by_url.get(normalize_url(final_url), ""),
                    "service_category": category,
                    "title": title,
                    "description_text": text,
                    "partner_name": "",
                    "cta_url": first_link_from_soup(soup),
                    "external_form_url": final_url if category in {"financing", "quote_request", "customer_care"} else "",
                    "notes": "KServico public page/service surface.",
                },
            )
            if category == "careers":
                self.tables.append(
                    "job_listing",
                    {
                        "job_id": "job-kservico-careers-general",
                        "dealer_id": DEALER_ID,
                        "source_page_id": self.source_page_by_url.get(normalize_url(final_url), ""),
                        "title": "KServico Careers",
                        "slug": "join-our-team",
                        "job_location": "Philippines",
                        "application_target": final_url,
                        "company_name": DEALER_NAME,
                        "company_website": PRIMARY_DOMAIN,
                        "company_tagline": "",
                        "job_salary": "",
                        "salary_currency": "",
                        "salary_unit": "",
                        "is_featured": "",
                        "is_filled": "",
                        "is_remote": "",
                        "job_type_ids": "",
                        "description_text": text,
                        "published_at": "",
                        "modified_at": "",
                    },
                )

        self.add_financing_row(
            financing_id="financing-kservico-generic-application",
            motorcycle_id="",
            source_url=f"{PRIMARY_DOMAIN}/application-form/",
            financing_type="generic_application_form",
            program_name="KServico Installment Application",
            form_url=f"{PRIMARY_DOMAIN}/application-form/",
            cta_url=f"{PRIMARY_DOMAIN}/application-form/",
            note_text="Generic public financing/application flow for products and vehicles.",
        )

    def collect_relationship_probe(self) -> None:
        self.tables.append(
            "relationship_probe",
            {
                "probe_id": "probe-kservico-branch-model-availability",
                "dealer_id": DEALER_ID,
                "relationship_name": "branch_to_motorcycle_model_availability",
                "source_url": f"{PRIMARY_DOMAIN}/store/",
                "source_type": "store detail pages and product detail pages",
                "status": "coverage_only",
                "evidence_found": "FALSE",
                "rows_created": 0,
                "notes": "KServico store pages expose branch details and broad product/brand hints. Public pages did not expose exact branch plus exact model inventory evidence.",
                "checked_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def collect_data_quality_notes(self) -> None:
        self.add_data_quality_issue(
            entity_type="api_endpoint",
            entity_id="wc-store-products",
            source_url=f"{PRIMARY_DOMAIN}/wp-json/wc/store/v1/products",
            issue_type="cloudflare_challenge",
            issue_detail="WooCommerce Store API returned a Cloudflare challenge during exploration, so product REST plus public HTML/schema were used instead.",
            severity="info",
        )
        self.add_data_quality_issue(
            entity_type="branch_model_availability",
            entity_id=DEALER_ID,
            source_url=f"{PRIMARY_DOMAIN}/store/",
            issue_type="no_exact_public_relationship",
            issue_detail="No exact public branch-to-motorcycle model bridge was found. branch_product_coverage is broad coverage only.",
            severity="warning",
        )

    def add_data_quality_issue(
        self,
        *,
        entity_type: str,
        entity_id: str,
        source_url: str,
        issue_type: str,
        issue_detail: str,
        severity: str,
    ) -> None:
        self.tables.append(
            "data_quality_issue",
            {
                "issue_id": f"issue-{len(self.tables.rows['data_quality_issue']) + 1}",
                "dealer_id": DEALER_ID,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "source_url": source_url,
                "issue_type": issue_type,
                "issue_detail": issue_detail,
                "severity": severity,
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def collect_external_links_from_soup(self, soup: BeautifulSoup, source_url: str, source_page_id: str) -> None:
        for link in soup.find_all("a", href=True):
            href = urljoin(source_url, link.get("href", ""))
            self.add_external_cta(source_page_id, source_url, href, link.get_text(" ", strip=True))

    def collect_external_links_from_html(self, html: str, source_url: str, source_page_id: str) -> None:
        self.collect_external_links_from_soup(BeautifulSoup(html, "lxml"), source_url, source_page_id)

    def add_external_cta(self, source_page_id: str, source_url: str, target_url: str, anchor_text: str) -> None:
        parsed = urlparse(target_url)
        if not parsed.scheme.startswith("http"):
            return
        target_domain = parsed.netloc.lower()
        if target_domain in {"kservico.com.ph", "www.kservico.com.ph"}:
            if not any(path in parsed.path for path in ("/application-form", "/request-for-quote", "/contact", "/join-our-team")):
                return
        key = (source_url, target_url)
        if key in self.external_cta_seen:
            return
        self.external_cta_seen.add(key)
        self.tables.append(
            "external_cta",
            {
                "cta_id": f"cta-{len(self.tables.rows['external_cta']) + 1}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "source_url": source_url,
                "target_url": target_url,
                "target_domain": target_domain,
                "cta_type": infer_cta_type(target_url, anchor_text),
                "anchor_text": anchor_text,
                "notes": "Public link/CTA evidence.",
            },
        )

    def finalize_crawl_run(self) -> None:
        self.tables.append(
            "crawl_run",
            {
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "run_started_at": self.run_started_at,
                "run_finished_at": datetime.now(TIMEZONE).isoformat(),
                "source_count": self.source_counter,
                "notes": "Current-state public KServico snapshot using WordPress REST, product HTML/schema, store pages, and public form pages.",
            },
        )

    def write_workbook(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in DATABASE_SHEETS:
            schema = DATABASE_SCHEMAS[sheet_name]
            sheet = workbook.create_sheet(title=sheet_name[:31])
            sheet.append(schema)
            for row in self.tables.rows.get(sheet_name, []):
                sheet.append([safe_excel_value(row.get(column)) for column in schema])
        workbook.save(self.output_path)


def clean_text(value: object) -> str:
    if isinstance(value, dict):
        value = get_rendered(value)
    return BeautifulSoup(str(value or ""), "lxml").get_text(" ", strip=True).replace("\xa0", " ")


def safe_excel_value(value: object) -> object:
    value = sanitize_excel_value(value)
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", " ", value)
    return value


def join_values(values: object) -> str:
    if isinstance(values, (list, tuple, set)):
        return " | ".join(str(value) for value in values if str(value).strip())
    return str(values or "")


def embedded_featured_image(row: dict[str, object]) -> str:
    embedded = row.get("_embedded")
    if not isinstance(embedded, dict):
        return ""
    media_rows = embedded.get("wp:featuredmedia")
    if not isinstance(media_rows, list) or not media_rows:
        return ""
    media = media_rows[0]
    if not isinstance(media, dict):
        return ""
    return str(media.get("source_url") or media.get("link") or "")


def iter_jsonld_objects(soup: BeautifulSoup) -> list[dict[str, object]]:
    objects: list[dict[str, object]] = []
    for script in soup.find_all("script", type=lambda value: value and "ld+json" in value):
        try:
            parsed = json.loads(script.get_text())
        except json.JSONDecodeError:
            continue
        stack = [parsed]
        while stack:
            item = stack.pop()
            if isinstance(item, dict):
                objects.append(item)
                graph = item.get("@graph")
                if isinstance(graph, list):
                    stack.extend(graph)
            elif isinstance(item, list):
                stack.extend(item)
    return objects


def extract_product_price(soup: BeautifulSoup) -> dict[str, object]:
    for item in iter_jsonld_objects(soup):
        item_type = item.get("@type")
        if item_type == "Product" or (isinstance(item_type, list) and "Product" in item_type):
            offers = item.get("offers")
            offer = offers[0] if isinstance(offers, list) and offers else offers
            if isinstance(offer, dict) and offer.get("price"):
                amount = parse_json_price(offer.get("price"))
                return {
                    "srp_text": f"PHP {amount:,.2f}" if isinstance(amount, float) else str(offer.get("price", "")),
                    "srp_amount": amount,
                }
    price_tag = soup.select_one(".single-product .et_pb_wc_price .price, .single-product p.price, .summary .price")
    if price_tag:
        text = price_tag.get_text(" ", strip=True)
        return {"srp_text": text, "srp_amount": parse_currency_amount(text)}
    return {"srp_text": "", "srp_amount": ""}


def parse_json_price(value: object) -> float | str:
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return parse_currency_amount(value)


def product_image_from_schema(soup: BeautifulSoup) -> str:
    for item in iter_jsonld_objects(soup):
        item_type = item.get("@type")
        if item_type == "Product" or (isinstance(item_type, list) and "Product" in item_type):
            image = item.get("image", "")
            if isinstance(image, list):
                return str(image[0]) if image else ""
            return str(image or "")
    return ""


def parse_specs_from_html(html: str) -> list[dict[str, str]]:
    soup = BeautifulSoup(html or "", "lxml")
    specs: list[dict[str, str]] = []
    for row in soup.select("tr"):
        cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
        if len(cells) >= 2 and cells[0] and cells[1]:
            specs.append({"group": "Specifications", "label": cells[0], "value": cells[1]})
    if specs:
        return specs

    lines = [line.strip() for line in soup.get_text("\n", strip=True).splitlines() if line.strip()]
    current_group = "Specifications"
    pending_label = ""
    for index, line in enumerate(lines):
        normalized = normalize_label(line)
        if normalized in SPEC_SECTION_HEADINGS:
            current_group = line
            pending_label = ""
            continue
        matched = split_known_label(line)
        if matched:
            label, value = matched
            specs.append({"group": current_group, "label": label, "value": value})
            pending_label = ""
            continue
        if pending_label:
            specs.append({"group": current_group, "label": pending_label, "value": line})
            pending_label = ""
            continue
        if looks_like_label(line):
            pending_label = line
        elif index == 0 and len(line) <= 80:
            specs.append({"group": current_group, "label": "Description", "value": line})
    return specs


def split_known_label(line: str) -> tuple[str, str] | None:
    for label in KNOWN_SPEC_LABELS:
        if line.lower().startswith(label.lower() + " "):
            return label, line[len(label) :].strip(" :-")
        if line.lower() == label.lower():
            return None
    return None


def looks_like_label(line: str) -> bool:
    if len(line) > 60:
        return False
    if re.search(r"\d", line):
        return False
    words = line.split()
    return 1 <= len(words) <= 6


def normalize_label(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def flatten_specs(specs: list[dict[str, str]]) -> dict[str, str]:
    flattened: dict[str, str] = {}
    brake_values: list[str] = []
    for spec in specs:
        label = normalize_label(spec["label"])
        value = spec["value"]
        if label == "model":
            flattened["spec_model"] = value
        elif "transmission" in label:
            flattened["spec_transmission"] = value
        elif "engine type" in label:
            flattened["spec_engine_type"] = value
        elif label == "displacement":
            flattened["spec_displacement_cc"] = value
        elif "front" in label and ("tire" in label or "tyre" in label):
            flattened["spec_front_tire"] = value
        elif "rear" in label and ("tire" in label or "tyre" in label):
            flattened["spec_rear_tire"] = value
        elif "front rear tyre size" in label:
            flattened["spec_front_tire"] = value
            flattened["spec_rear_tire"] = value
        elif "wheel type" in label:
            flattened["spec_wheels_type"] = value
        elif "starting system" in label:
            flattened["spec_starting_system"] = value
        elif "overall dimensions" in label or "length x width x height" in label or label == "dimension" or label == "dimensions":
            flattened["spec_overall_dimensions"] = value
        elif "ground clearance" in label:
            flattened["spec_ground_clearance"] = value
        elif "fuel system" in label:
            flattened["spec_fuel_system"] = value
        elif "seat height" in label or "seat weight" in label:
            flattened["spec_seat_height"] = value
        elif "brake" in label:
            brake_values.append(f"{spec['label']}: {value}")
    if brake_values:
        flattened["spec_brake_system"] = " | ".join(brake_values)
    return flattened


def infer_brand_name(category_slugs: set[str], text: str) -> str:
    lower = text.lower()
    for slug, brand in MOTORCYCLE_BRANDS.items():
        if slug in category_slugs or re.search(rf"\b{re.escape(slug)}\b", lower):
            return brand
    return ""


def infer_model_name(title: str, brand_name: str) -> str:
    model = title
    if brand_name and model.upper().startswith(brand_name.upper()):
        model = model[len(brand_name) :].strip()
    return model


def infer_vehicle_type(category_slugs: set[str]) -> str:
    if "e-bike" in category_slugs:
        return "electric"
    if "bajaj-re" in category_slugs or "bajaj" in category_slugs or "tricycle-model" in category_slugs:
        return "tri-wheeler"
    if "big-bikes" in category_slugs:
        return "big-bike"
    if "motorcycles" in category_slugs or category_slugs & {"honda", "kawasaki", "kymco", "suzuki", "yamaha", "scooter", "solo-model"}:
        return "regular-bike"
    return ""


def infer_vehicle_type_from_text(text: str) -> str:
    lower = text.lower()
    if "e-bike" in lower or "ebike" in lower:
        return "electric"
    if "bajaj" in lower or "three-wheeler" in lower or "tricycle" in lower:
        return "tri-wheeler"
    if "big bike" in lower or "big-bike" in lower:
        return "big-bike"
    if "motorcycle" in lower or any(brand in lower for brand in MOTORCYCLE_BRANDS):
        return "regular-bike"
    return ""


def category_listing_url(category_slugs: set[str]) -> str:
    if "e-bike" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/e-bike/"
    if "bajaj-re" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/bajaj-re/"
    if "big-bikes" in category_slugs:
        return f"{PRIMARY_DOMAIN}/product-category/motorcycles/big-bikes/"
    return f"{PRIMARY_DOMAIN}/product-category/motorcycles/"


def find_internal_link(soup: BeautifulSoup, path_fragment: str) -> str:
    for link in soup.find_all("a", href=True):
        href = urljoin(PRIMARY_DOMAIN, link.get("href", ""))
        if path_fragment in urlparse(href).path:
            return href
    return ""


def first_link_from_html(html: str) -> str:
    return first_link_from_soup(BeautifulSoup(html or "", "lxml"))


def first_link_from_soup(soup: BeautifulSoup) -> str:
    for link in soup.find_all("a", href=True):
        href = urljoin(PRIMARY_DOMAIN, link.get("href", ""))
        if href.startswith("http"):
            return href
    return ""


def parse_store_page(soup: BeautifulSoup, fallback_title: str) -> dict[str, str]:
    lines = [line.strip() for line in soup.get_text("\n", strip=True).splitlines() if line.strip()]
    address = text_between(lines, "Address", {"Get Directions", "Contact"})
    phone = value_after_label(lines, "Mobile:")
    email = first_mailto(soup) or value_after_label(lines, "E-mail:")
    hours = parse_open_hours(lines)
    about = text_between(lines, "About Us", {"360 Street View", "Our Commitment", "Get update"})
    postal_match = re.search(r"\b(\d{4})\b", address)
    branch_text = f"{fallback_title} {about}"
    return {
        "address": address,
        "postal_code": postal_match.group(1) if postal_match else "",
        "phone": phone,
        "email": email,
        "open_hours": " | ".join(hours),
        "days_str": " | ".join(hour.split(":", 1)[0] for hour in hours),
        "brand_raw": ", ".join(brand for slug, brand in MOTORCYCLE_BRANDS.items() if slug in branch_text.lower()),
        "product_raw": about,
        "categories_raw": about,
    }


def text_between(lines: list[str], start_label: str, stop_labels: set[str]) -> str:
    collecting = False
    values: list[str] = []
    for line in lines:
        if line == start_label:
            collecting = True
            continue
        if collecting and (line in stop_labels or any(line.startswith(stop) for stop in stop_labels)):
            break
        if collecting:
            values.append(line)
    return " ".join(values).strip()


def value_after_label(lines: list[str], label: str) -> str:
    for index, line in enumerate(lines):
        if line == label and index + 1 < len(lines):
            return lines[index + 1].strip()
        if line.startswith(label):
            return line[len(label) :].strip()
    return ""


def first_mailto(soup: BeautifulSoup) -> str:
    tag = soup.find("a", href=lambda href: href and href.startswith("mailto:"))
    if not tag:
        return ""
    return tag.get("href", "").replace("mailto:", "").strip()


def parse_open_hours(lines: list[str]) -> list[str]:
    days = {"Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"}
    hours: list[str] = []
    for index, line in enumerate(lines):
        day = line.rstrip(":")
        if day in days and index + 1 < len(lines):
            hours.append(f"{day}: {lines[index + 1]}")
    return hours


def looks_like_promo(title: str, content_text: str) -> bool:
    haystack = f"{title} {content_text}".lower()
    keywords = [
        "promo",
        "promotion",
        "sale",
        "discount",
        "downpayment",
        "raffle",
        "hotdeal",
        "summer",
        "christmas",
        "cashback",
        "bundle",
        "installment",
    ]
    return any(keyword in haystack for keyword in keywords)


def infer_promo_type(title: str, content_text: str) -> str:
    haystack = f"{title} {content_text}".lower()
    if "raffle" in haystack:
        return "raffle"
    if "downpayment" in haystack or "installment" in haystack:
        return "financing"
    if "sale" in haystack or "discount" in haystack or "hotdeal" in haystack:
        return "discount"
    return "campaign"


def is_meaningful_form(form: Tag) -> bool:
    action = form.get("action", "")
    method = (form.get("method") or "").lower()
    if method == "get" and form.find("input", attrs={"name": "s"}):
        return False
    if "wpcf7" in action:
        return True
    field_names = [field.get("name", "") for field in form.select("input, select, textarea")]
    return len([name for name in field_names if name and name not in {"s", "post_type", "dgwt_wcas"}]) >= 3


def infer_form_category(url: str, page_title: str) -> str:
    lower = f"{url} {page_title}".lower()
    if "application" in lower:
        return "financing_application"
    if "quote" in lower:
        return "quote_request"
    if "contact" in lower:
        return "customer_care"
    if "career" in lower or "join" in lower:
        return "career_application"
    if "faq" in lower:
        return "newsletter"
    return "public_form"


def infer_cta_type(target_url: str, anchor_text: str) -> str:
    lower = f"{target_url} {anchor_text}".lower()
    if "facebook" in lower:
        return "social"
    if "maps" in lower or "google" in lower:
        return "map"
    if "application" in lower:
        return "application_form"
    if "quote" in lower:
        return "quote_form"
    if "career" in lower or "join" in lower:
        return "career"
    return "link"


def main_text(soup: BeautifulSoup) -> str:
    main = soup.select_one("#main-content") or soup.select_one("main") or soup
    return html_to_text(str(main))


def slug_from_url_path(url: str) -> str:
    path = urlparse(url).path.strip("/")
    return re.sub(r"[^a-z0-9]+", "-", path.lower()).strip("-") or "home"


def main() -> int:
    output_path = Path(__file__).resolve().parents[2] / "working/outputs/dealer-databases/kservico_database.xlsx"
    scraper = KservicoDatabaseScraper(output_path)
    scraper.run()
    print(f"[OK] Wrote {output_path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
