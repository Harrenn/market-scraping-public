#!/usr/bin/env python3
from __future__ import annotations

import argparse
import ast
import json
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from full_scrape.motortrade.scrape_motortrade_analytics import (  # noqa: E402
    DATABASE_SCHEMAS,
    DATABASE_SHEETS,
    TIMEZONE,
    USER_AGENT,
    TableStore,
    html_to_text,
    normalize_url,
    parse_currency_amount,
    sanitize_excel_value,
    split_numeric_unit,
)


DEALER_ID = "ropali"
DEALER_NAME = "Ropali"
BASE_URL = "https://ropali.ph"
OUTPUT_PATH = Path(__file__).resolve().parents[2] / "working/outputs/dealer-databases/ropali_database.xlsx"
REQUEST_DELAY_SECONDS = 0.12
REQUEST_TIMEOUT_SECONDS = 35

MOTORCYCLE_CATEGORY_SLUGS = {"business", "scooter", "underbone", "lifestyle", "off-road"}
PART_CATEGORY_SLUGS = {"helmet", "raincoat", "top-box", "tunnel-bag"}
BRANDS = [
    "Honda",
    "Yamaha",
    "Suzuki",
    "Kawasaki",
    "Lambretta",
    "Bajaj",
    "Kymco",
    "TVS",
    "KTM",
    "Vespa",
]
SERVICE_SLUGS = {
    "about-us": "company_profile",
    "contact-us": "customer_care",
    "our-stores": "branch_locator",
    "locations": "branch_locator",
    "faqs": "faq",
    "data-privacy": "privacy",
    "careers-and-rewards": "careers",
}


@dataclass
class FetchResult:
    url: str
    status_code: int
    text: str
    content_type: str
    final_url: str


def normalize_spaces(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def slugify(value: object) -> str:
    text = normalize_spaces(value).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text or "item"


def text_from_html(value: object) -> str:
    return html_to_text(str(value or ""))


def json_text(value: object) -> str:
    try:
        return json.dumps(value, ensure_ascii=True, sort_keys=True)
    except TypeError:
        return normalize_spaces(value)


def record_title(record: dict[str, object]) -> str:
    title = record.get("title")
    if isinstance(title, dict):
        return text_from_html(title.get("rendered", ""))
    return normalize_spaces(title)


def record_content_html(record: dict[str, object]) -> str:
    content = record.get("content")
    if isinstance(content, dict):
        return str(content.get("rendered") or "")
    return str(content or "")


def record_excerpt(record: dict[str, object]) -> str:
    excerpt = record.get("excerpt")
    if isinstance(excerpt, dict):
        return text_from_html(excerpt.get("rendered", ""))
    return normalize_spaces(excerpt)


def parse_amount(value: object) -> float | str:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_spaces(value)
    amount = parse_currency_amount(text)
    if amount != "":
        return amount
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", text)
    if not match:
        return ""
    try:
        return float(match.group(0).replace(",", ""))
    except ValueError:
        return ""


def store_api_price_amount(prices: dict[str, object]) -> float | str:
    raw = prices.get("sale_price") or prices.get("price") or prices.get("regular_price")
    if raw in (None, "", "0", 0):
        return ""
    try:
        minor_unit = int(prices.get("currency_minor_unit") or 2)
        return int(raw) / (10**minor_unit)
    except (TypeError, ValueError):
        return parse_amount(raw)


def store_api_price_text(prices: dict[str, object]) -> str:
    amount = store_api_price_amount(prices)
    if amount == "":
        return ""
    currency = prices.get("currency_code") or "PHP"
    return f"{currency} {float(amount):,.2f}"


def guess_brand(text: str) -> str:
    lower = text.lower()
    for brand in BRANDS:
        if brand.lower() in lower:
            return brand
    return ""


def infer_vehicle_type(name: str, categories: list[dict[str, object]], specs: list[dict[str, str]]) -> tuple[str, str]:
    slugs = {str(category.get("slug") or "").lower() for category in categories}
    joined = " ".join([name, " ".join(s.get("value", "") for s in specs)]).lower()
    cc = 0.0
    for match in re.finditer(r"(\d{2,4}(?:\.\d+)?)\s*(?:cc|cm3|cm³)\b", joined):
        cc = max(cc, float(match.group(1)))
    if "off-road" in slugs:
        return "regular-bike", "product category"
    if cc >= 400:
        return "big-bike", "displacement"
    if "lifestyle" in slugs:
        return "regular-bike", "product category"
    if slugs & {"business", "scooter", "underbone"}:
        return "regular-bike", "product category"
    return "", ""


def is_motorcycle_product(product: dict[str, object]) -> bool:
    slugs = {str(category.get("slug") or "").lower() for category in product.get("categories", [])}
    return bool(slugs & MOTORCYCLE_CATEGORY_SLUGS)


def product_image(product: dict[str, object]) -> str:
    images = product.get("images")
    if isinstance(images, list) and images and isinstance(images[0], dict):
        return str(images[0].get("src") or "")
    return ""


def product_categories_text(product: dict[str, object]) -> str:
    return ", ".join(
        normalize_spaces(category.get("name") or category.get("slug"))
        for category in product.get("categories", [])
        if isinstance(category, dict)
    )


def extract_specs_from_soup(soup: BeautifulSoup) -> list[dict[str, str]]:
    specs: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for table in soup.find_all("table"):
        for row in table.find_all("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
            if len(cells) < 2:
                continue
            label = normalize_spaces(cells[0]).rstrip(":")
            value = normalize_spaces(" ".join(cells[1:]))
            if not label or not value:
                continue
            lowered = f"{label} {value}".lower()
            if "review" in lowered or "comment" in lowered:
                continue
            key = (label.lower(), value.lower())
            if key in seen:
                continue
            seen.add(key)
            specs.append({"group": "Additional information", "label": label, "value": value})
    return specs


def spec_lookup(specs: list[dict[str, str]], *labels: str) -> str:
    wanted = [label.lower() for label in labels]
    for spec in specs:
        label = spec.get("label", "").lower()
        if any(w in label for w in wanted):
            return spec.get("value", "")
    return ""


def extract_forms_from_soup(soup: BeautifulSoup, source_url: str, page_title: str) -> list[dict[str, object]]:
    forms: list[dict[str, object]] = []
    for index, form in enumerate(soup.find_all("form"), start=1):
        action = normalize_url(urljoin(source_url, form.get("action") or source_url))
        method = normalize_spaces(form.get("method") or "get").lower()
        fields: list[dict[str, object]] = []
        for order, field in enumerate(form.find_all(["input", "select", "textarea", "button"]), start=1):
            field_name = normalize_spaces(field.get("name") or field.get("id") or field.name)
            field_type = normalize_spaces(field.get("type") or field.name)
            label = ""
            if field.get("id"):
                label_tag = soup.find("label", attrs={"for": field.get("id")})
                if label_tag:
                    label = label_tag.get_text(" ", strip=True)
            if not label:
                label = normalize_spaces(field.get("aria-label") or field.get("placeholder") or field_name)
            options = []
            if field.name == "select":
                options = [
                    (normalize_spaces(option.get("value")), option.get_text(" ", strip=True))
                    for option in field.find_all("option")
                ]
            if not options:
                options = [("", "")]
            for option_value, option_label in options:
                fields.append(
                    {
                        "field_order": order,
                        "field_name": field_name,
                        "field_label": normalize_spaces(label),
                        "field_type": field_type,
                        "is_required": bool(field.get("required")),
                        "default_value": normalize_spaces(field.get("value")),
                        "placeholder": normalize_spaces(field.get("placeholder")),
                        "option_value": option_value,
                        "option_label": option_label,
                    }
                )
        forms.append(
            {
                "index": index,
                "form_name": f"{page_title or 'Ropali page'} form {index}",
                "category": classify_form(action, fields),
                "source_url": source_url,
                "action": action,
                "method": method,
                "requires_file_upload": any(f.get("field_type") == "file" for f in fields),
                "has_privacy_checkbox": any("privacy" in (f.get("field_label") or "").lower() for f in fields),
                "fields": fields,
            }
        )
    return forms


def classify_form(action: str, fields: list[dict[str, object]]) -> str:
    evidence = " ".join([action, " ".join(str(f.get("field_name") or "") for f in fields)]).lower()
    if "comment" in evidence or "review" in evidence:
        return "review"
    if "add-to-cart" in evidence or "quantity" in evidence:
        return "cart"
    if "login" in evidence or "register" in evidence or "password" in evidence:
        return "account"
    if "search" in evidence or "product_cat" in evidence:
        return "search"
    return "website"


def extract_external_links(soup: BeautifulSoup, source_url: str) -> list[dict[str, str]]:
    links: list[dict[str, str]] = []
    seen: set[str] = set()
    base_domain = urlparse(BASE_URL).netloc
    for anchor in soup.find_all("a", href=True):
        href = anchor.get("href") or ""
        if href.startswith(("mailto:", "tel:")):
            target = href
        else:
            target = normalize_url(urljoin(source_url, href))
        parsed = urlparse(target)
        domain = parsed.netloc
        if not target or target in seen:
            continue
        if domain and domain != base_domain or target.startswith(("mailto:", "tel:")):
            seen.add(target)
            links.append(
                {
                    "target_url": target,
                    "target_domain": domain or target.split(":", 1)[0],
                    "anchor_text": anchor.get_text(" ", strip=True),
                    "cta_type": "contact" if target.startswith(("mailto:", "tel:")) else "external",
                }
            )
    return links


def extract_mapbox_store_features(js_text: str) -> list[dict[str, object]]:
    start = js_text.find("const stores=")
    if start < 0:
        return []
    start += len("const stores=")
    end_match = re.search(r"}\s*;\s*function\s+isValidCoordinate", js_text[start:])
    end = start + end_match.start() if end_match else -1
    if end < 0:
        stores_match = re.search(r";\s*stores\.features", js_text[start:])
        end = start + stores_match.start() if stores_match else -1
    if end < 0:
        return []
    chunk = js_text[start : end + 1]
    try:
        stores = ast.literal_eval(chunk)
    except (SyntaxError, ValueError):
        return []
    features = stores.get("features") if isinstance(stores, dict) else []
    return [feature for feature in features if isinstance(feature, dict)]


class RopaliDatabaseScraper:
    def __init__(self, output_path: Path = OUTPUT_PATH):
        self.output_path = output_path
        self.session = requests.Session()
        self.session.headers.update(
            {
                "User-Agent": USER_AGENT,
                "Accept": "text/html,application/json,*/*",
                "Referer": BASE_URL + "/",
            }
        )
        self.tables = TableStore(DATABASE_SCHEMAS)
        self.now = datetime.now(TIMEZONE)
        self.run_id = f"{DEALER_ID}-run-{self.now.strftime('%Y%m%d%H%M%S')}"
        self.run_started_at = self.now.isoformat()
        self.source_counter = 0
        self.source_pages_by_url: dict[str, str] = {}
        self.wp_products_by_id: dict[int, dict[str, object]] = {}
        self.store_products: list[dict[str, object]] = []
        self.product_categories: list[dict[str, object]] = []
        self.product_tags: list[dict[str, object]] = []
        self.category_by_id: dict[int, dict[str, object]] = {}
        self.tag_by_id: dict[int, dict[str, object]] = {}
        self.motorcycle_rows: dict[int, dict[str, object]] = {}
        self.motorcycle_by_slug: dict[str, str] = {}
        self.promo_rows_by_slug: dict[str, dict[str, object]] = {}
        self.form_seen: set[tuple[str, str, str]] = set()
        self.external_seen: set[tuple[str, str]] = set()
        self.issues_seen: set[tuple[str, str, str]] = set()
        self.motorcycles_without_specs: list[str] = []
        self.motorcycles_without_pricing: list[str] = []

    def fetch(self, url: str, *, expect_json: bool = False, retries: int = 4) -> FetchResult:
        last: FetchResult | None = None
        for attempt in range(1, retries + 1):
            try:
                response = self.session.get(url, timeout=REQUEST_TIMEOUT_SECONDS, allow_redirects=True)
                result = FetchResult(
                    url=url,
                    status_code=response.status_code,
                    text=response.text,
                    content_type=response.headers.get("content-type", ""),
                    final_url=normalize_url(response.url),
                )
                last = result
                if response.status_code == 200 and (not expect_json or "json" in result.content_type.lower()):
                    time.sleep(REQUEST_DELAY_SECONDS)
                    return result
            except requests.RequestException as exc:
                last = FetchResult(url=url, status_code=0, text=str(exc), content_type="", final_url=url)
            time.sleep(0.8 * attempt)
        assert last is not None
        self.add_issue("source", "", url, "fetch_warning", f"Fetch did not return expected response. Last status: {last.status_code}", "medium")
        return last

    def fetch_json(self, url: str) -> object:
        result = self.fetch(url, expect_json=True)
        if result.status_code != 200:
            return [] if "?" in url else {}
        try:
            return json.loads(result.text)
        except json.JSONDecodeError:
            self.add_issue("source", "", url, "json_parse_warning", "Response was not parseable JSON.", "medium")
            return []

    def add_issue(self, entity_type: str, entity_id: str, source_url: str, issue_type: str, detail: str, severity: str) -> None:
        key = (entity_type, entity_id, issue_type)
        if key in self.issues_seen:
            return
        self.issues_seen.add(key)
        self.tables.append(
            "data_quality_issue",
            {
                "issue_id": f"issue-{DEALER_ID}-{slugify(entity_type)}-{slugify(entity_id or issue_type)}-{len(self.issues_seen)}",
                "dealer_id": DEALER_ID,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "source_url": source_url,
                "issue_type": issue_type,
                "issue_detail": detail,
                "severity": severity,
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def add_source_page(self, source_type: str, record: dict[str, object], url: str, title: str = "") -> str:
        clean_url = normalize_url(url)
        if clean_url in self.source_pages_by_url:
            return self.source_pages_by_url[clean_url]
        self.source_counter += 1
        source_page_id = f"source-{DEALER_ID}-{self.source_counter:04d}"
        content_html = record_content_html(record)
        self.tables.append(
            "source_page",
            {
                "source_page_id": source_page_id,
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "source_type": source_type,
                "wp_post_type": record.get("type") or source_type,
                "wp_id": record.get("id", ""),
                "url": clean_url,
                "slug": record.get("slug") or slugify(title or clean_url),
                "title": title or record_title(record),
                "status": record.get("status", "publish"),
                "published_at": record.get("date", ""),
                "modified_at": record.get("modified", ""),
                "author_id": record.get("author", ""),
                "excerpt_text": record_excerpt(record),
                "content_html": content_html,
                "content_text": text_from_html(content_html),
                "featured_media_id": record.get("featured_media", ""),
                "featured_image_url": self.featured_image_url(record),
                "category_ids": ",".join(str(v) for v in record.get("categories", []) if v),
                "tag_ids": ",".join(str(v) for v in record.get("tags", []) if v),
                "job_type_ids": "",
                "seo_title": self.yoast_value(record, "title"),
                "seo_description": self.yoast_value(record, "description"),
                "canonical_url": self.yoast_value(record, "canonical") or clean_url,
            },
        )
        self.source_pages_by_url[clean_url] = source_page_id
        return source_page_id

    def featured_image_url(self, record: dict[str, object]) -> str:
        yoast = record.get("yoast_head_json")
        if isinstance(yoast, dict):
            images = yoast.get("og_image")
            if isinstance(images, list) and images and isinstance(images[0], dict):
                return str(images[0].get("url") or "")
        return ""

    def yoast_value(self, record: dict[str, object], key: str) -> str:
        yoast = record.get("yoast_head_json")
        return str(yoast.get(key) or "") if isinstance(yoast, dict) else ""

    def load_rest_collections(self) -> None:
        types = self.fetch_json(f"{BASE_URL}/wp-json/wp/v2/types")
        if isinstance(types, dict):
            for type_name, type_info in types.items():
                self.tables.append(
                    "website_filter_reference",
                    {
                        "dealer_id": DEALER_ID,
                        "reference_type": "wp_post_type",
                        "filter_name": "wp_post_type",
                        "option_value": type_name,
                        "option_label": type_info.get("name", type_name) if isinstance(type_info, dict) else type_name,
                        "source_url": f"{BASE_URL}/wp-json/wp/v2/types",
                        "notes": "Public WordPress REST post type.",
                    },
                )
        self.product_categories = self.safe_list(f"{BASE_URL}/wp-json/wc/store/v1/products/categories?per_page=100")
        self.product_tags = self.safe_list(f"{BASE_URL}/wp-json/wp/v2/product_tag?per_page=100")
        self.category_by_id = {int(c["id"]): c for c in self.product_categories if isinstance(c, dict) and c.get("id")}
        self.tag_by_id = {int(t["id"]): t for t in self.product_tags if isinstance(t, dict) and t.get("id")}
        for category in self.product_categories:
            self.add_filter_reference("product_category", "product_cat", category.get("slug"), category.get("name"), category.get("permalink") or category.get("link"), f"count={category.get('count', '')}")
        for tag in self.product_tags:
            self.add_filter_reference("product_tag", "product_tag", tag.get("slug"), tag.get("name"), tag.get("link"), f"count={tag.get('count', '')}")
        self.store_products = self.safe_list(f"{BASE_URL}/wp-json/wc/store/v1/products?per_page=100")
        wp_products = self.safe_list(f"{BASE_URL}/wp-json/wp/v2/product?per_page=100")
        self.wp_products_by_id = {int(record["id"]): record for record in wp_products if isinstance(record, dict) and record.get("id")}
        for endpoint in ["pages", "posts"]:
            records = self.safe_list(f"{BASE_URL}/wp-json/wp/v2/{endpoint}?per_page=100")
            for record in records:
                if not isinstance(record, dict):
                    continue
                title = record_title(record)
                source_id = self.add_source_page(endpoint[:-1], record, str(record.get("link") or ""), title)
                self.extract_page_objects(endpoint[:-1], record, source_id, title)

    def safe_list(self, url: str) -> list[dict[str, object]]:
        payload = self.fetch_json(url)
        if isinstance(payload, list):
            return [row for row in payload if isinstance(row, dict)]
        return []

    def add_filter_reference(self, reference_type: str, filter_name: str, value: object, label: object, source_url: object, notes: str) -> None:
        self.tables.append(
            "website_filter_reference",
            {
                "dealer_id": DEALER_ID,
                "reference_type": reference_type,
                "filter_name": filter_name,
                "option_value": normalize_spaces(value),
                "option_label": normalize_spaces(label),
                "source_url": normalize_url(str(source_url or BASE_URL)),
                "notes": notes,
            },
        )

    def extract_page_objects(self, source_type: str, record: dict[str, object], source_page_id: str, title: str) -> None:
        url = normalize_url(str(record.get("link") or ""))
        slug = str(record.get("slug") or "")
        html = record_content_html(record)
        soup = BeautifulSoup(html, "lxml")
        text = text_from_html(html)
        if source_type == "post":
            categories = [int(v) for v in record.get("categories", []) if str(v).isdigit()]
            category_names = [str(self.category_by_id.get(v, {}).get("name", "")) for v in categories]
            if "promo" in f"{slug} {title} {' '.join(category_names)}".lower() or "promo" in text.lower():
                self.add_promo(source_page_id, record, title, url, text, "post")
        if slug in SERVICE_SLUGS or any(term in f"{slug} {title}".lower() for term in ["contact", "store", "faq", "career", "privacy"]):
            self.add_service(source_page_id, SERVICE_SLUGS.get(slug, "website_page"), title, text, url)
        self.add_forms_from_soup(soup, url, title, source_page_id)
        self.add_external_links_from_soup(soup, url, source_page_id)

    def add_promo(self, source_page_id: str, record: dict[str, object], title: str, url: str, text: str, promo_type: str) -> str:
        promo_id = f"promo-{DEALER_ID}-{slugify(url or title)}"
        if promo_id in self.promo_rows_by_slug:
            return promo_id
        row = {
            "promo_id": promo_id,
            "dealer_id": DEALER_ID,
            "source_page_id": source_page_id,
            "source_page_url": url,
            "source_url": url,
            "title": title,
            "promo_type": promo_type,
            "description_text": text,
            "image_url": self.featured_image_url(record),
            "cta_url": url,
            "target_listing_url": "",
            "applies_to_brand": guess_brand(text),
            "applies_to_vehicle_type": "",
            "applies_to_tag": "",
            "has_vehicle_promo_flag": True,
            "linked_motorcycle_count": 0,
            "linked_motorcycles": "",
            "link_types": "",
            "link_source_urls": "",
            "published_at": record.get("date", ""),
            "modified_at": record.get("modified", ""),
        }
        self.tables.append("promo", row)
        self.promo_rows_by_slug[promo_id] = row
        return promo_id

    def add_service(self, source_page_id: str, category: str, title: str, text: str, url: str) -> None:
        self.tables.append(
            "aftersales_service",
            {
                "service_id": f"service-{DEALER_ID}-{slugify(url or title)}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "service_category": category,
                "title": title,
                "description_text": text,
                "partner_name": "",
                "cta_url": url,
                "external_form_url": "",
                "notes": "Public page/service surface from Ropali WordPress content.",
            },
        )

    def add_forms_from_soup(self, soup: BeautifulSoup, source_url: str, title: str, source_page_id: str = "") -> None:
        for form in extract_forms_from_soup(soup, source_url, title):
            key = (form["source_url"], form["action"], form["form_name"])
            if key in self.form_seen:
                continue
            self.form_seen.add(key)
            form_id = f"form-{DEALER_ID}-{slugify(form['source_url'])}-{form['index']}"
            self.tables.append(
                "form_definition",
                {
                    "form_id": form_id,
                    "dealer_id": DEALER_ID,
                    "form_name": form["form_name"],
                    "form_category": form["category"],
                    "source_url": form["source_url"],
                    "final_url": form["source_url"],
                    "page_title": title,
                    "form_action": form["action"],
                    "form_method": form["method"],
                    "requires_file_upload": form["requires_file_upload"],
                    "requires_otp": False,
                    "multi_step": False,
                    "has_privacy_checkbox": form["has_privacy_checkbox"],
                    "notes": f"Structural capture only; source_page_id={source_page_id}. No submission performed.",
                },
            )
            for field in form["fields"]:
                self.tables.append(
                    "form_field",
                    {
                        "form_id": form_id,
                        "dealer_id": DEALER_ID,
                        "form_name": form["form_name"],
                        "form_category": form["category"],
                        "field_order": field["field_order"],
                        "field_name": field["field_name"],
                        "field_label": field["field_label"],
                        "field_type": field["field_type"],
                        "is_required": field["is_required"],
                        "default_value": field["default_value"],
                        "placeholder": field["placeholder"],
                        "option_value": field["option_value"],
                        "option_label": field["option_label"],
                        "section_name": "",
                        "source_url": form["source_url"],
                    },
                )

    def add_external_links_from_soup(self, soup: BeautifulSoup, source_url: str, source_page_id: str = "") -> None:
        for link in extract_external_links(soup, source_url):
            key = (source_url, link["target_url"])
            if key in self.external_seen:
                continue
            self.external_seen.add(key)
            self.tables.append(
                "external_cta",
                {
                    "cta_id": f"cta-{DEALER_ID}-{len(self.external_seen):04d}",
                    "dealer_id": DEALER_ID,
                    "source_page_id": source_page_id,
                    "source_url": source_url,
                    "target_url": link["target_url"],
                    "target_domain": link["target_domain"],
                    "cta_type": link["cta_type"],
                    "anchor_text": link["anchor_text"],
                    "notes": "Public outbound/contact CTA; target not scraped as a private data source.",
                },
            )

    def scrape_products(self) -> None:
        for product in self.store_products:
            if not isinstance(product, dict) or not product.get("id"):
                continue
            product_id = int(product["id"])
            wp_record = self.wp_products_by_id.get(product_id, {})
            url = normalize_url(str(product.get("permalink") or wp_record.get("link") or ""))
            title = normalize_spaces(product.get("name") or record_title(wp_record))
            source_page_id = self.add_source_page("product", wp_record or {"id": product_id, "slug": slugify(title)}, url, title)
            if is_motorcycle_product(product):
                self.add_motorcycle(product, wp_record, source_page_id)
            else:
                self.add_part(product, wp_record, source_page_id)

    def add_motorcycle(self, product: dict[str, object], wp_record: dict[str, object], source_page_id: str) -> str:
        product_id = int(product["id"])
        title = normalize_spaces(product.get("name") or record_title(wp_record))
        url = normalize_url(str(product.get("permalink") or wp_record.get("link") or ""))
        detail = self.fetch(url)
        soup = BeautifulSoup(detail.text, "lxml") if detail.status_code == 200 else BeautifulSoup("", "lxml")
        if detail.status_code != 200:
            self.add_issue("motorcycle", str(product_id), url, "detail_fetch_warning", "Product API row exists but detail page could not be fetched for specs/forms.", "medium")
        specs = extract_specs_from_soup(soup)
        if not specs:
            content_specs = extract_specs_from_soup(BeautifulSoup(record_content_html(wp_record), "lxml"))
            specs = content_specs
        if not specs:
            self.motorcycles_without_specs.append(title)
        brand = guess_brand(f"{title} {text_from_html(record_content_html(wp_record))}")
        model = title
        if brand and model.lower().startswith(brand.lower()):
            model = normalize_spaces(model[len(brand) :])
        vehicle_type, vehicle_type_source = infer_vehicle_type(title, product.get("categories", []), specs)
        categories = product_categories_text(product)
        tags = [self.tag_by_id.get(int(tag_id), {}).get("slug", "") for tag_id in wp_record.get("product_tag", []) if str(tag_id).isdigit()]
        price_amount = store_api_price_amount(product.get("prices", {}) if isinstance(product.get("prices"), dict) else {})
        price_text = store_api_price_text(product.get("prices", {}) if isinstance(product.get("prices"), dict) else {})
        motorcycle_id = f"motorcycle-{DEALER_ID}-{product_id}"
        spec_displacement = spec_lookup(specs, "engine displacement", "displacement")
        spec_displacement_numeric, _ = split_numeric_unit(spec_displacement)
        row = {
            "motorcycle_id": motorcycle_id,
            "dealer_id": DEALER_ID,
            "display_name": title,
            "brand_name": brand,
            "model_name": model,
            "vehicle_type": vehicle_type,
            "detail_url": url,
            "listing_url": f"{BASE_URL}/shop/",
            "srp_amount": "",
            "downpayment_amount": price_amount,
            "monthly_amount": "",
            "price_contexts": "WooCommerce Store API price; Ropali detail pages state displayed product price is downpayment.",
            "promo_titles": "",
            "is_blue_tag": False,
            "is_featured": False,
            "is_top_pick": False,
            "has_vehicle_promo_flag": bool(tags),
            "spec_make": brand or spec_lookup(specs, "make"),
            "spec_model": spec_lookup(specs, "model") or model,
            "spec_transmission": spec_lookup(specs, "transmission"),
            "spec_engine_type": spec_lookup(specs, "engine type"),
            "spec_seat_height": spec_lookup(specs, "seat height"),
            "spec_brake_system": spec_lookup(specs, "brake system", "brake type"),
            "spec_displacement_cc": spec_displacement_numeric,
            "spec_front_tire": spec_lookup(specs, "tire size (front)", "front tire"),
            "spec_rear_tire": spec_lookup(specs, "tire size (rear)", "rear tire"),
            "spec_wheels_type": spec_lookup(specs, "wheels type"),
            "spec_starting_system": spec_lookup(specs, "starting system"),
            "spec_overall_dimensions": spec_lookup(specs, "overall dimensions"),
            "spec_ground_clearance": spec_lookup(specs, "ground clearance"),
            "spec_fuel_system": spec_lookup(specs, "fuel system"),
            "primary_image_url": product_image(product),
            "excerpt_text": text_from_html(product.get("short_description") or record_excerpt(wp_record)),
            "apply_loan_url": "",
            "inquiry_url": "",
            "service_url": "",
            "compare_enabled": "compare" in detail.text.lower(),
            "published_at": wp_record.get("date", ""),
            "modified_at": wp_record.get("modified", ""),
            "vehicle_type_source": vehicle_type_source,
            "classification_notes": f"Category evidence: {categories}. No branch-level inventory evidence found.",
            "raw_taxonomies": json_text({"categories": product.get("categories", []), "tags": tags}),
            "source_page_id": source_page_id,
            "source_page_url": url,
        }
        self.tables.append("motorcycle", row)
        self.motorcycle_rows[product_id] = row
        self.motorcycle_by_slug[slugify(title)] = motorcycle_id
        for order, spec in enumerate(specs, start=1):
            value_numeric, unit = split_numeric_unit(spec["value"])
            self.tables.append(
                "motorcycle_spec",
                {
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": title,
                    "brand_name": brand,
                    "model_name": model,
                    "spec_group": spec["group"],
                    "spec_label": spec["label"],
                    "spec_value_raw": spec["value"],
                    "spec_value_numeric": value_numeric,
                    "spec_unit": unit,
                    "display_order": order,
                    "source_url": url,
                },
            )
        if price_amount != "":
            self.tables.append(
                "pricing",
                {
                    "pricing_id": f"pricing-{DEALER_ID}-{product_id}-downpayment",
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": title,
                    "brand_name": brand,
                    "model_name": model,
                    "source_url": url,
                    "source_context": "WooCommerce product price",
                    "currency": "PHP",
                    "srp_text": "",
                    "srp_amount": "",
                    "downpayment_text": price_text,
                    "downpayment_amount": price_amount,
                    "monthly_text": "",
                    "monthly_amount": "",
                    "price_note": "Ropali product detail pages show disclaimer: the displayed product price is only a downpayment.",
                    "is_indicative": True,
                    "captured_at": datetime.now(TIMEZONE).isoformat(),
                },
            )
            self.tables.append(
                "financing",
                {
                    "financing_id": f"financing-{DEALER_ID}-{product_id}-downpayment",
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": title,
                    "brand_name": brand,
                    "model_name": model,
                    "source_url": url,
                    "financing_type": "downpayment",
                    "program_name": "Ropali displayed downpayment",
                    "monthly_amount": "",
                    "downpayment_amount": price_amount,
                    "loan_amount": "",
                    "loan_term_months": "",
                    "estimated_income_amount": "",
                    "form_url": "",
                    "cta_url": url,
                    "note_text": "Installment/payment availability is stated generally on Ropali pages; term/monthly details are not exposed in the scraped public product record.",
                    "captured_at": datetime.now(TIMEZONE).isoformat(),
                },
            )
        else:
            self.motorcycles_without_pricing.append(title)
        if soup:
            self.add_forms_from_soup(soup, url, title, source_page_id)
            self.add_external_links_from_soup(soup, url, source_page_id)
        return motorcycle_id

    def add_part(self, product: dict[str, object], wp_record: dict[str, object], source_page_id: str) -> None:
        title = normalize_spaces(product.get("name") or record_title(wp_record))
        url = normalize_url(str(product.get("permalink") or wp_record.get("link") or ""))
        categories = product_categories_text(product)
        self.tables.append(
            "parts_catalog",
            {
                "part_id": f"part-{DEALER_ID}-{product.get('id')}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "title": title,
                "slug": product.get("slug") or slugify(title),
                "detail_url": url,
                "image_url": product_image(product),
                "description_text": f"{text_from_html(product.get('short_description') or record_excerpt(wp_record))} Categories: {categories}".strip(),
                "published_at": wp_record.get("date", ""),
                "modified_at": wp_record.get("modified", ""),
            },
        )

    def scrape_branches(self) -> None:
        result = self.fetch(f"{BASE_URL}/our-stores/")
        if result.status_code != 200:
            self.add_issue("branch", "", f"{BASE_URL}/our-stores/", "branch_page_fetch_warning", "Branch map page could not be fetched.", "high")
            return
        soup = BeautifulSoup(result.text, "lxml")
        source_page_id = self.add_source_page("html_page", {"id": "our-stores", "slug": "our-stores"}, result.final_url, "Our Stores")
        self.add_forms_from_soup(soup, result.final_url, "Our Stores", source_page_id)
        self.add_external_links_from_soup(soup, result.final_url, source_page_id)
        js_urls = [urljoin(result.final_url, script.get("src")) for script in soup.find_all("script", src=True)]
        features: list[dict[str, object]] = []
        for js_url in js_urls:
            if "autoptimize" not in js_url and "map" not in js_url.lower():
                continue
            js_result = self.fetch(js_url)
            features = extract_mapbox_store_features(js_result.text)
            if features:
                break
        if not features:
            self.add_issue("branch", "", result.final_url, "branch_map_parse_warning", "Mapbox branch feature collection was not found in page JavaScript.", "high")
            return
        for index, feature in enumerate(features, start=1):
            props = feature.get("properties", {}) if isinstance(feature.get("properties"), dict) else {}
            geometry = feature.get("geometry", {}) if isinstance(feature.get("geometry"), dict) else {}
            coords = geometry.get("coordinates") if isinstance(geometry.get("coordinates"), list) else ["", ""]
            longitude = coords[0] if len(coords) > 0 else ""
            latitude = coords[1] if len(coords) > 1 else ""
            name = normalize_spaces(props.get("address") or f"Ropali location {index}")
            city_address = normalize_spaces(props.get("city"))
            branch_id = f"branch-{DEALER_ID}-{index:03d}-{slugify(name)}"
            self.tables.append(
                "branch",
                {
                    "branch_id": branch_id,
                    "dealer_id": DEALER_ID,
                    "source_url": result.final_url,
                    "locator_id": index,
                    "branch_name": name,
                    "slug": slugify(name),
                    "street": normalize_spaces(props.get("crossStreet")),
                    "city": city_address,
                    "state_or_province": normalize_spaces(props.get("state")),
                    "postal_code": normalize_spaces(props.get("postalCode")),
                    "full_address": city_address,
                    "latitude": latitude,
                    "longitude": longitude,
                    "phone": normalize_spaces(props.get("phone") or props.get("phoneFormatted")),
                    "fax": "",
                    "email": "",
                    "website": "",
                    "open_hours_raw": "",
                    "days_str": "",
                    "brand_raw": "",
                    "product_raw": "Motorcycle dealer location",
                    "categories_raw": "Mapbox branch/location feature",
                    "inquire_url": "",
                },
            )
            self.tables.append(
                "branch_product_coverage",
                {
                    "coverage_id": f"coverage-{DEALER_ID}-{index:03d}",
                    "dealer_id": DEALER_ID,
                    "branch_id": branch_id,
                    "coverage_type": "public_branch_listing",
                    "coverage_value_raw": "Motorcycle sales/services branch location",
                    "coverage_value_normalized": "motorcycle dealer location",
                    "source_url": result.final_url,
                    "source_field": "Mapbox feature collection",
                    "confidence": "coverage only",
                    "observed_at": datetime.now(TIMEZONE).isoformat(),
                    "notes": "Ropali map exposes branch/location coordinates and addresses, but not exact model availability by branch.",
                },
            )

    def add_tag_promos_and_links(self) -> None:
        for tag in self.product_tags:
            name = normalize_spaces(tag.get("name"))
            slug = normalize_spaces(tag.get("slug"))
            if not name:
                continue
            promo_id = f"promo-{DEALER_ID}-tag-{slugify(slug)}"
            self.tables.append(
                "promo",
                {
                    "promo_id": promo_id,
                    "dealer_id": DEALER_ID,
                    "source_page_id": "",
                    "source_page_url": normalize_url(str(tag.get("link") or "")),
                    "source_url": normalize_url(str(tag.get("link") or "")),
                    "title": name,
                    "promo_type": "product_tag",
                    "description_text": f"Ropali WooCommerce product tag with {tag.get('count', '')} tagged products.",
                    "image_url": "",
                    "cta_url": normalize_url(str(tag.get("link") or "")),
                    "target_listing_url": normalize_url(str(tag.get("link") or "")),
                    "applies_to_brand": "",
                    "applies_to_vehicle_type": "",
                    "applies_to_tag": slug,
                    "has_vehicle_promo_flag": True,
                    "linked_motorcycle_count": 0,
                    "linked_motorcycles": "",
                    "link_types": "",
                    "link_source_urls": "",
                    "published_at": "",
                    "modified_at": "",
                },
            )
            tag_id = tag.get("id")
            linked_titles: list[str] = []
            linked_urls: list[str] = []
            for product_id, row in self.motorcycle_rows.items():
                wp_record = self.wp_products_by_id.get(product_id, {})
                tag_ids = {str(v) for v in wp_record.get("product_tag", [])}
                if str(tag_id) not in tag_ids:
                    continue
                linked_titles.append(str(row["display_name"]))
                linked_urls.append(str(row["detail_url"]))
                self.tables.append(
                    "motorcycle_promo_link",
                    {
                        "motorcycle_id": row["motorcycle_id"],
                        "promo_id": promo_id,
                        "dealer_id": DEALER_ID,
                        "display_name": row["display_name"],
                        "promo_title": name,
                        "link_type": "woocommerce_product_tag",
                        "source_url": row["detail_url"],
                        "notes": "Exact WooCommerce product tag assigned to this product.",
                    },
                )
            if linked_titles:
                # Update is intentionally append-safe in memory before workbook write.
                for promo_row in reversed(self.tables.rows["promo"]):
                    if promo_row["promo_id"] == promo_id:
                        promo_row["linked_motorcycle_count"] = len(linked_titles)
                        promo_row["linked_motorcycles"] = "; ".join(linked_titles)
                        promo_row["link_types"] = "woocommerce_product_tag"
                        promo_row["link_source_urls"] = "; ".join(linked_urls)
                        break

    def add_generic_financing_and_relationship_probes(self) -> None:
        self.tables.append(
            "financing",
            {
                "financing_id": f"financing-{DEALER_ID}-generic-installment",
                "motorcycle_id": "",
                "dealer_id": DEALER_ID,
                "display_name": "",
                "brand_name": "",
                "model_name": "",
                "source_url": BASE_URL,
                "financing_type": "generic_installment",
                "program_name": "Payment Options Available",
                "monthly_amount": "",
                "downpayment_amount": "",
                "loan_amount": "",
                "loan_term_months": "",
                "estimated_income_amount": "",
                "form_url": "",
                "cta_url": BASE_URL,
                "note_text": "Site footer/page copy says customers may pay by cash or installment; no public calculator terms were found.",
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )
        self.tables.append(
            "relationship_probe",
            {
                "probe_id": f"probe-{DEALER_ID}-branch-model",
                "dealer_id": DEALER_ID,
                "relationship_name": "branch_to_motorcycle_availability",
                "source_url": f"{BASE_URL}/our-stores/",
                "source_type": "Mapbox branch list + WooCommerce catalog",
                "status": "no evidence",
                "evidence_found": False,
                "rows_created": 0,
                "notes": "Ropali exposes branch locations and dealer-wide products separately. Do not join branch to motorcycle by dealer_id for availability.",
                "checked_at": datetime.now(TIMEZONE).isoformat(),
            },
        )
        self.tables.append(
            "relationship_probe",
            {
                "probe_id": f"probe-{DEALER_ID}-promo-model",
                "dealer_id": DEALER_ID,
                "relationship_name": "promo_to_motorcycle",
                "source_url": f"{BASE_URL}/wp-json/wp/v2/product_tag?per_page=100",
                "source_type": "WooCommerce product tags",
                "status": "bridge when product tags are populated",
                "evidence_found": self.tables.row_count("motorcycle_promo_link") > 0,
                "rows_created": self.tables.row_count("motorcycle_promo_link"),
                "notes": "Exact links are created only from WooCommerce product tags assigned to product records.",
                "checked_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def add_dealer_and_run_rows(self) -> None:
        finished = datetime.now(TIMEZONE).isoformat()
        self.tables.append(
            "dealer",
            {
                "dealer_id": DEALER_ID,
                "dealer_name": DEALER_NAME,
                "primary_domain": BASE_URL,
                "forms_domain": "",
                "crawl_timezone": "Asia/Manila",
                "crawl_date": self.now.date().isoformat(),
                "notes": "Ropali public WordPress/WooCommerce scrape. Homepage/source endpoints can intermittently return 522; scraper uses retries and records warnings.",
            },
        )
        self.tables.append(
            "crawl_run",
            {
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "run_started_at": self.run_started_at,
                "run_finished_at": finished,
                "source_count": self.source_counter,
                "notes": "Current-state public snapshot. Branch/model availability is intentionally empty without exact source evidence.",
            },
        )

    def write_workbook(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in DATABASE_SHEETS:
            schema = DATABASE_SCHEMAS[sheet_name]
            sheet = workbook.create_sheet(title=sheet_name[:31])
            sheet.append(schema)
            for row in self.tables.rows[sheet_name]:
                sheet.append([sanitize_excel_value(row.get(column)) for column in schema])
        workbook.save(self.output_path)

    def run(self) -> None:
        self.load_rest_collections()
        self.scrape_products()
        self.scrape_branches()
        self.add_tag_promos_and_links()
        self.add_generic_financing_and_relationship_probes()
        if self.motorcycles_without_specs:
            self.add_issue(
                "motorcycle_spec",
                "partial",
                BASE_URL,
                "missing_public_specs",
                f"{len(self.motorcycles_without_specs)} exact motorcycle products did not expose parseable public spec tables: "
                + "; ".join(self.motorcycles_without_specs[:15]),
                "medium",
            )
        if self.motorcycles_without_pricing:
            self.add_issue(
                "pricing",
                "partial",
                BASE_URL,
                "missing_public_pricing",
                f"{len(self.motorcycles_without_pricing)} exact motorcycle products had blank/zero WooCommerce price/downpayment values: "
                + "; ".join(self.motorcycles_without_pricing[:15]),
                "medium",
            )
        if not self.tables.row_count("motorcycle_spec"):
            self.add_issue("motorcycle_spec", "", BASE_URL, "empty_specs", "No motorcycle specs were extracted; detail page parser should be rechecked.", "high")
        if not self.tables.row_count("branch"):
            self.add_issue("branch", "", f"{BASE_URL}/our-stores/", "empty_branches", "No branches were extracted from the public branch map.", "high")
        self.add_dealer_and_run_rows()
        self.write_workbook()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Scrape Ropali public dealer data into the shared database workbook format.")
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH, help="Output workbook path.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    scraper = RopaliDatabaseScraper(args.output)
    scraper.run()
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
