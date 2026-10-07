#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
import time
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from full_scrape.motortrade.scrape_motortrade_analytics import (
    DATABASE_SCHEMAS,
    DATABASE_SHEETS,
    TIMEZONE,
    USER_AGENT,
    TableStore,
    html_to_text,
    normalize_url,
    parse_currency_amount,
    sanitize_excel_value,
    split_numeric_unit,
)


BRANDS = [
    "Honda",
    "Yamaha",
    "Suzuki",
    "Kawasaki",
    "Bajaj",
    "Kymco",
    "KYMCO",
    "TVS",
    "KTM",
    "Royal Enfield",
    "BMW",
    "CF Moto",
    "CFMOTO",
    "Rieju",
    "Sherco",
    "Monarch",
    "Hatasu",
]

VEHICLE_HINTS = [
    "motorcycle",
    "motorcycles",
    "scooter",
    "underbone",
    "big bike",
    "big-bike",
    "tri-wheeler",
    "tricycle",
    "bajaj",
    "e-bike",
]

PART_HINTS = ["part", "parts", "accessory", "accessories", "helmet", "oil", "yamalube", "brake", "v-belt", "cleaner"]
PROMO_HINTS = ["promo", "promotion", "discount", "sale", "deal", "news", "event"]
SERVICE_HINTS = ["service", "contact", "loan", "credit", "financing", "application", "about", "faq", "kando"]
JOB_HINTS = ["career", "job", "hiring", "mechanic"]
BRANCH_HINTS = ["branch", "branches", "store", "stores", "locator", "location"]


@dataclass
class DealerConfig:
    dealer_id: str
    dealer_name: str
    primary_domain: str
    folder: str
    output_filename: str
    script_filename: str
    forms_domain: str = ""
    known_pages: list[str] = field(default_factory=list)
    notes: str = ""
    max_sitemap_pages: int = 450
    request_delay_seconds: float = 0.04
    request_timeout_seconds: float = 10.0


def normalize_spaces(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def slugify(value: object) -> str:
    text = normalize_spaces(value).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text or "item"


def parse_amount(value: object) -> float | str:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_spaces(value)
    amount = parse_currency_amount(text)
    if amount != "":
        return amount
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", text)
    if not match:
        return ""
    try:
        return float(match.group(0).replace(",", ""))
    except ValueError:
        return ""


def parse_motorcentral_labelled_amount(value: object) -> tuple[float | str, str]:
    text = normalize_spaces(value)
    text = re.sub(r"(?<=\d),\s+(?=\d)", ",", text)
    match = re.search(r"(?<![\d,])(?:PHP|PhP|P|\u20b1)?\s*((?:0|\d{1,3}(?:,\d{3})+|\d{4,})(?:\.\d+)?)(?![\d,])", text)
    if not match:
        return "", ""
    raw_text = match.group(0).strip()
    number_text = match.group(1).replace(",", "")
    try:
        amount = float(number_text)
    except ValueError:
        return "", ""
    if amount == 6619640522 or amount > 1_000_000:
        return "", ""
    if amount not in {0.0} and "," not in match.group(1) and amount < 1000:
        return "", ""
    return amount, raw_text


def extract_motorcentral_pricing_from_soup(soup: BeautifulSoup, display_name: str = "") -> dict[str, object]:
    pricing_root = soup
    normalized_target = re.sub(r"[\u2013\u2014-]+", "-", normalize_spaces(display_name)).lower()
    headings = soup.select("h1")
    heading = next(
        (
            candidate
            for candidate in headings
            if re.sub(r"[\u2013\u2014-]+", "-", normalize_spaces(candidate.get_text(" ", strip=True))).lower()
            == normalized_target
        ),
        headings[0] if headings else None,
    )
    if heading:
        for parent in heading.parents:
            scoped_text = normalize_spaces(parent.get_text(" ", strip=True))
            scoped_lower = scoped_text.lower()
            if all(label in scoped_lower for label in ("price", "down payment", "monthly")) and len(scoped_text) <= 1200:
                pricing_root = parent
                break

    lines = [normalize_spaces(line) for line in pricing_root.get_text("\n", strip=True).splitlines() if normalize_spaces(line)]
    if not lines:
        return {}
    start_index = 0
    if display_name:
        target = normalize_spaces(display_name).lower()
        for index, line in enumerate(lines):
            if line.lower() == target:
                start_index = index
                break
    window = lines[start_index : start_index + 35]
    stop_words = {"specifications", "description", "reviews", "related products"}
    clipped: list[str] = []
    for line in window:
        if line.lower() in stop_words:
            break
        clipped.append(line)

    labels = {
        "srp": {"price"},
        "downpayment": {"estimated down payment", "down payment"},
        "monthly": {"estimated monthly", "monthly"},
    }
    all_labels = {label for values in labels.values() for label in values}
    result: dict[str, object] = {}
    for index, line in enumerate(clipped):
        lower = line.lower().strip(":")
        for key, label_values in labels.items():
            if f"{key}_amount" in result:
                continue
            if lower not in label_values and not any(lower.startswith(label + " ") for label in label_values):
                continue
            candidates = [line]
            for candidate in clipped[index + 1 : index + 4]:
                if candidate.lower().strip(":") in all_labels:
                    break
                candidates.append(candidate)
            for candidate in candidates:
                candidate_text = candidate
                for label in label_values:
                    candidate_text = re.sub(rf"^{re.escape(label)}\s*:?", "", candidate_text, flags=re.I).strip()
                amount, raw_text = parse_motorcentral_labelled_amount(candidate_text)
                if amount != "":
                    result[f"{key}_amount"] = amount
                    result[f"{key}_text"] = raw_text
                    break
    return result


def amount_text(value: object) -> str:
    amount = parse_amount(value)
    if amount == "":
        return normalize_spaces(value)
    return f"PHP {float(amount):,.0f}"


def text_from_html(value: object) -> str:
    return html_to_text(str(value or ""))


def json_text(value: object) -> str:
    try:
        return json.dumps(value, ensure_ascii=True, sort_keys=True)
    except TypeError:
        return normalize_spaces(value)


def image_from_wp(record: dict[str, object]) -> str:
    yoast = record.get("yoast_head_json")
    if isinstance(yoast, dict):
        og_image = yoast.get("og_image")
        if isinstance(og_image, list) and og_image and isinstance(og_image[0], dict):
            return str(og_image[0].get("url") or "")
    better = record.get("better_featured_image")
    if isinstance(better, dict):
        return str(better.get("source_url") or "")
    return ""


def record_title(record: dict[str, object]) -> str:
    title = record.get("title")
    if isinstance(title, dict):
        return text_from_html(title.get("rendered", ""))
    return normalize_spaces(title)


def record_excerpt(record: dict[str, object]) -> str:
    excerpt = record.get("excerpt")
    if isinstance(excerpt, dict):
        return text_from_html(excerpt.get("rendered", ""))
    return normalize_spaces(excerpt)


def record_content_html(record: dict[str, object]) -> str:
    content = record.get("content")
    if isinstance(content, dict):
        return str(content.get("rendered") or "")
    return str(content or "")


def guess_brand(text: str) -> str:
    normalized = text.lower()
    for brand in BRANDS:
        if brand.lower() in normalized:
            return "CFMOTO" if brand.lower() in {"cf moto", "cfmoto"} else brand.upper() if brand == "KYMCO" else brand
    return ""


def displacement_cc(text: str) -> float:
    match = re.search(r"(\d{2,4}(?:\.\d+)?)\s*cc\b", text.lower())
    if not match:
        return 0.0
    return float(match.group(1))


def infer_vehicle_type(text: str) -> tuple[str, str]:
    lower = text.lower()
    if any(term in lower for term in ["tri-wheeler", "tricycle", "bajaj re", "3-wheeler", "three wheeler"]):
        return "tri-wheeler", "keyword"
    if any(term in lower for term in ["electric motorcycle", "electric scooter", "electric vehicle", "e-bike", " ebike", "em1 e", "cuv e", "iqube"]):
        return "electric", "keyword"
    cc = displacement_cc(text)
    if cc >= 400:
        return "big-bike", "displacement"
    if any(term in lower for term in ["big bike", "big-bike", "superbike"]):
        return "big-bike", "keyword"
    if any(term in lower for term in VEHICLE_HINTS) or guess_brand(text):
        return "regular-bike", "keyword"
    return "", ""


def looks_like_vehicle(text: str, url: str = "") -> bool:
    parsed = urlparse(url)
    url_evidence = " ".join([parsed.path, parsed.query])
    lower = f"{text} {url_evidence}".lower()
    if "/product-page/" in lower and any(term in lower for term in PART_HINTS):
        return False
    return any(term in lower for term in VEHICLE_HINTS) or bool(guess_brand(lower))


def looks_like_part(text: str, url: str = "") -> bool:
    parsed = urlparse(url)
    url_evidence = " ".join([parsed.path, parsed.query])
    lower = f"{text} {url_evidence}".lower()
    return any(term in lower for term in PART_HINTS) or "/product-page/" in lower


def classify_service_category(text: str) -> str:
    lower = text.lower()
    if any(term in lower for term in ["loan", "credit", "financing", "application"]):
        return "financing"
    if "career" in lower or "job" in lower:
        return "career"
    if "contact" in lower:
        return "contact"
    if "service" in lower:
        return "service"
    if "branch" in lower or "store" in lower or "locator" in lower:
        return "branch_locator"
    return "public_page"


SPEC_LABELS = [
    "Overall Dimensions: L x W x H",
    "Fuel Consumption (WMTC Test Method)",
    "Tire Size & Type (Front)",
    "Tire Size & Type (Rear)",
    "Brake Type (Front)",
    "Brake Type (Rear)",
    "Tire Size (Front)",
    "Tire Size (Rear)",
    "Suspension Type (Front)",
    "Suspension Type (Rear)",
    "Suspension (front)",
    "Suspension (rear)",
    "Cylinder Arrangement",
    "Overall Dimensions",
    "Maximum Horse Power",
    "Maximum Power",
    "Maximum Torque",
    "Starting System",
    "Ignition System",
    "Lubrication System",
    "Engine Oil Capacity",
    "Transmission Type",
    "Fuel Tank Capacity",
    "Ground Clearance",
    "Compression Ratio",
    "Bore & Stroke",
    "Engine Type",
    "Displacement",
    "Ignition Type",
    "Wheel Type",
    "Seat Height",
    "Fuel System",
    "Battery Type",
    "Curb Weight",
    "Gear Shift Pattern",
    "Front Tire",
    "Rear Tire",
    "Front Brake",
    "Rear Brake",
    "Frame Type",
    "Front Suspension",
    "Rear Suspension",
    "Engine",
    "Cylinder",
]


def split_compact_spec(text: str) -> tuple[str, str] | None:
    cleaned = normalize_spaces(text).rstrip(":")
    if not cleaned:
        return None
    for label in sorted(SPEC_LABELS, key=len, reverse=True):
        lower = cleaned.lower()
        label_lower = label.lower()
        if lower == label_lower:
            return None
        if lower.startswith(label_lower + " "):
            return label, cleaned[len(label):].strip(" :-")
        if lower.startswith(label_lower + ":"):
            return label, cleaned[len(label):].strip(" :-")
    return None


def extract_motorcentral_sequential_specs(soup: BeautifulSoup) -> list[dict[str, str]]:
    lines = [normalize_spaces(line) for line in soup.get_text("\n", strip=True).splitlines() if normalize_spaces(line)]
    try:
        start = next(i for i, line in enumerate(lines) if line.lower() == "specifications") + 1
    except StopIteration:
        return []
    stop = len(lines)
    for i in range(start, len(lines)):
        if lines[i].lower() in {"related motorcycle", "motorcentral group of companies"}:
            stop = i
            break
    section = lines[start:stop]
    label_lookup = {label.lower(): label for label in SPEC_LABELS}
    specs: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    i = 0
    while i < len(section):
        label = label_lookup.get(section[i].lower())
        if not label:
            i += 1
            continue
        value_parts: list[str] = []
        j = i + 1
        while j < len(section) and section[j].lower() not in label_lookup:
            value_parts.append(section[j])
            j += 1
        value = normalize_spaces(" ".join(value_parts))
        if value:
            key = (label, value)
            if key not in seen:
                specs.append({"group": "Specifications", "label": label, "value": value})
                seen.add(key)
        i = j
    return specs


def extract_specs_from_soup(soup: BeautifulSoup) -> list[dict[str, str]]:
    sequential_specs = extract_motorcentral_sequential_specs(soup)
    if sequential_specs:
        return sequential_specs
    specs: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()

    def add_spec(label: str, value: str, group: str = "Specifications") -> None:
        label = normalize_spaces(label).rstrip(":")
        value = normalize_spaces(value)
        if not label or not value:
            return
        key = (label, value)
        if key not in seen:
            specs.append({"group": group, "label": label, "value": value})
            seen.add(key)

    for row in soup.select("tr"):
        cells = [normalize_spaces(cell.get_text(" ", strip=True)) for cell in row.find_all(["th", "td"]) if normalize_spaces(cell.get_text(" ", strip=True))]
        if len(cells) == 2 and 1 <= len(cells[0]) <= 80 and cells[1]:
            add_spec(cells[0], cells[1])
        for cell in cells:
            compact = split_compact_spec(cell)
            if compact:
                add_spec(compact[0], compact[1])

    for block in soup.select("li, p, div"):
        text = normalize_spaces(block.get_text(" ", strip=True))
        if not text or len(text) > 220:
            continue
        compact = split_compact_spec(text)
        if compact:
            add_spec(compact[0], compact[1])
            continue
        match = re.match(r"^([A-Za-z][A-Za-z0-9 /()&.,-]{2,70})\s*[:\-]\s*(.{2,120})$", text)
        if match and any(term in match.group(1).lower() for term in ["engine", "displacement", "transmission", "brake", "tire", "seat", "fuel", "dimension", "wheel", "starter", "ground", "suspension"]):
            add_spec(match.group(1), match.group(2))
    return specs


def flatten_specs(specs: list[dict[str, str]], brand_name: str, model_name: str) -> dict[str, object]:
    flattened: dict[str, object] = {"spec_make": brand_name, "spec_model": model_name}
    brake_values: list[str] = []
    for spec in specs:
        label = spec["label"].lower()
        value = spec["value"]
        if "transmission" in label:
            flattened["spec_transmission"] = value
        elif "engine" in label and "type" in label:
            flattened["spec_engine_type"] = value
        elif "seat" in label and "height" in label:
            flattened["spec_seat_height"] = value
        elif "displacement" in label or label in {"cc", "engine displacement"}:
            flattened["spec_displacement_cc"] = value
        elif "front tire" in label:
            flattened["spec_front_tire"] = value
        elif "rear tire" in label:
            flattened["spec_rear_tire"] = value
        elif "tire" in label and "front" not in label and "rear" not in label:
            flattened["spec_front_tire"] = value
            flattened["spec_rear_tire"] = value
        elif "wheel" in label:
            flattened["spec_wheels_type"] = value
        elif "start" in label:
            flattened["spec_starting_system"] = value
        elif "dimension" in label or "l x w x h" in label:
            flattened["spec_overall_dimensions"] = value
        elif "ground" in label:
            flattened["spec_ground_clearance"] = value
        elif "fuel" in label and "system" in label:
            flattened["spec_fuel_system"] = value
        elif "brake" in label:
            brake_values.append(f"{spec['label']}: {value}")
    if brake_values:
        flattened["spec_brake_system"] = " | ".join(brake_values)
    return flattened


def parse_sitemap_locs(xml_text: str) -> tuple[list[str], bool]:
    root = ET.fromstring(xml_text.encode("utf-8"))
    is_index = root.tag.endswith("sitemapindex")
    ns = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    return [loc.text or "" for loc in root.findall(".//sm:loc", ns) if loc.text], is_index


class MotorcentralGroupDatabaseScraper:
    def __init__(self, config: DealerConfig, output_path: Path):
        self.config = config
        self.output_path = output_path
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": USER_AGENT, "Accept": "text/html,application/json,application/xml;q=0.9,*/*;q=0.8"})
        self.tables = TableStore(DATABASE_SCHEMAS)
        self.source_counter = 0
        self.motorcycle_seen: set[str] = set()
        self.part_seen: set[str] = set()
        self.promo_seen: set[str] = set()
        self.service_seen: set[str] = set()
        self.branch_seen: set[str] = set()
        self.form_seen: set[str] = set()
        self.cta_seen: set[tuple[str, str]] = set()
        self.coverage_seen: set[tuple[str, str, str]] = set()
        self.sitemap_urls: list[str] = []
        self.wp_available = False
        self.wp_types: set[str] = set()
        self.wp_taxonomy_labels: dict[str, dict[int, str]] = {}
        self.store_products_by_id: dict[int, dict[str, object]] = {}
        now = datetime.now(TIMEZONE)
        self.crawl_date = now.date().isoformat()
        self.run_started_at = now.isoformat()
        self.run_id = f"{config.dealer_id}-run-{now.strftime('%Y%m%d%H%M%S')}"

    def log(self, message: str) -> None:
        print(f"[{self.config.dealer_name}] {message}", flush=True)

    def get(self, url: str, **kwargs: object) -> requests.Response:
        timeout = kwargs.pop("timeout", self.config.request_timeout_seconds)
        response = self.session.get(url, timeout=timeout, allow_redirects=True, **kwargs)
        time.sleep(self.config.request_delay_seconds)
        return response

    def post(self, url: str, **kwargs: object) -> requests.Response:
        response = self.session.post(url, timeout=self.config.request_timeout_seconds, **kwargs)
        time.sleep(self.config.request_delay_seconds)
        return response

    def run(self) -> None:
        self.log("Starting scrape")
        self.add_dealer()
        self.collect_discovery_pages()
        self.collect_wordpress()
        if self.config.dealer_id == "dueksam":
            self.collect_dueksam_ajax()
        self.collect_sitemap_pages()
        self.collect_relationship_probe()
        self.collect_data_quality_notes()
        self.finalize_crawl_run()
        if self.write_workbook():
            self.log(f"Wrote {self.output_path}")
        else:
            self.log(f"Preserved existing {self.output_path}")

    @property
    def base(self) -> str:
        return self.config.primary_domain.rstrip("/")

    def abs_url(self, url: str) -> str:
        return urljoin(self.base + "/", url)

    def add_dealer(self) -> None:
        self.tables.append(
            "dealer",
            {
                "dealer_id": self.config.dealer_id,
                "dealer_name": self.config.dealer_name,
                "primary_domain": self.config.primary_domain,
                "forms_domain": self.config.forms_domain or self.config.primary_domain,
                "crawl_timezone": str(TIMEZONE),
                "crawl_date": self.crawl_date,
                "notes": self.config.notes or "Public dealer website snapshot using the shared dealer database schema.",
            },
        )

    def add_source_page(self, url: str, source_type: str, title: str = "", html: str = "", record: dict[str, object] | None = None) -> str:
        final_url = normalize_url(url)
        source_id = f"source-{self.config.dealer_id}-{self.source_counter + 1}"
        self.source_counter += 1
        soup = BeautifulSoup(html or "", "lxml") if html else None
        page_title = title or (soup.title.get_text(" ", strip=True) if soup and soup.title else "")
        canonical = ""
        if soup:
            canonical_tag = soup.select_one("link[rel='canonical']")
            canonical = canonical_tag.get("href", "") if canonical_tag else ""
        row = {
            "source_page_id": source_id,
            "run_id": self.run_id,
            "dealer_id": self.config.dealer_id,
            "source_type": source_type,
            "url": final_url,
            "slug": slugify(urlparse(final_url).path or page_title),
            "title": page_title,
            "status": "publish",
            "content_html": html[:50000] if html else "",
            "content_text": text_from_html(html)[:20000] if html else "",
            "canonical_url": canonical,
        }
        if record:
            row.update(
                {
                    "wp_post_type": record.get("type", ""),
                    "wp_id": record.get("id", ""),
                    "slug": record.get("slug", row["slug"]),
                    "status": record.get("status", ""),
                    "published_at": record.get("date", ""),
                    "modified_at": record.get("modified", ""),
                    "author_id": record.get("author", ""),
                    "excerpt_text": record_excerpt(record),
                    "content_html": record_content_html(record)[:50000],
                    "content_text": text_from_html(record_content_html(record))[:20000],
                    "featured_media_id": record.get("featured_media", ""),
                    "featured_image_url": image_from_wp(record),
                    "category_ids": ",".join(map(str, record.get("categories", []) or [])),
                    "tag_ids": ",".join(map(str, record.get("tags", []) or [])),
                    "seo_title": (record.get("yoast_head_json") or {}).get("title", "") if isinstance(record.get("yoast_head_json"), dict) else "",
                    "seo_description": (record.get("yoast_head_json") or {}).get("description", "") if isinstance(record.get("yoast_head_json"), dict) else "",
                }
            )
        self.tables.append("source_page", row)
        return source_id

    def collect_discovery_pages(self) -> None:
        for path in ["/", "/robots.txt", "/sitemap.xml", "/wp-sitemap.xml", *self.config.known_pages]:
            url = self.abs_url(path)
            try:
                response = self.get(url)
            except requests.RequestException as exc:
                self.add_issue("source_page", path, url, "request_failed", str(exc), "warning")
                continue
            if response.status_code >= 400:
                continue
            content_type = response.headers.get("content-type", "")
            if "xml" in content_type or path.endswith(".xml"):
                self.collect_sitemap(response.url, response.text)
            elif "html" in content_type:
                source_type = self.classify_url(response.url)
                source_id = self.add_source_page(response.url, source_type, html=response.text)
                self.inspect_html_page(response.url, response.text, source_type, source_id=source_id)

    def collect_sitemap(self, url: str, xml_text: str, depth: int = 0) -> None:
        try:
            locs, is_index = parse_sitemap_locs(xml_text)
        except Exception as exc:
            self.add_issue("sitemap", slugify(url), url, "sitemap_parse_failed", str(exc), "warning")
            return
        if is_index and depth < 2:
            for loc in locs[:60]:
                try:
                    response = self.get(loc)
                    if response.status_code == 200:
                        self.collect_sitemap(response.url, response.text, depth + 1)
                except requests.RequestException:
                    continue
        else:
            for loc in locs:
                cleaned = normalize_url(loc)
                if cleaned not in self.sitemap_urls:
                    self.sitemap_urls.append(cleaned)
                    self.add_filter_reference("sitemap_url", "sitemap", cleaned, cleaned, url, "Public sitemap URL.")
                    parsed = urlparse(cleaned)
                    for key, values in parse_qs(parsed.query).items():
                        for value in values:
                            self.add_filter_reference("query_filter", key, value, value, cleaned, "Sitemap query filter option.")

    def classify_url(self, url: str) -> str:
        parsed = urlparse(url)
        lower = " ".join([parsed.path, parsed.query]).lower()
        if any(term in lower for term in BRANCH_HINTS):
            return "branch_page"
        if any(term in lower for term in JOB_HINTS):
            return "job_page"
        if any(term in lower for term in PROMO_HINTS):
            return "promo_page"
        if any(term in lower for term in SERVICE_HINTS):
            return "service_page"
        if looks_like_part("", lower):
            return "parts_page"
        if looks_like_vehicle("", lower):
            return "motorcycle_page"
        return "public_page"

    def collect_wordpress(self) -> None:
        types_url = self.abs_url("/wp-json/wp/v2/types")
        try:
            response = self.get(types_url)
            self.wp_available = response.status_code == 200 and "json" in response.headers.get("content-type", "")
            if self.wp_available:
                payload = response.json()
                if isinstance(payload, dict):
                    self.wp_types = set(payload)
        except requests.RequestException:
            self.wp_available = False
        if not self.wp_available:
            return
        self.log("Collecting WordPress/WooCommerce sources")
        self.collect_taxonomies()
        self.collect_store_products()
        candidates = ["pages", "posts", "product", "store", "wpsl_stores", "job_listing", "job-listings", "parts", "preowned"]
        for collection in candidates:
            type_key = "page" if collection == "pages" else "post" if collection == "posts" else collection
            if type_key not in self.wp_types and collection not in {"pages", "posts"}:
                continue
            records = self.fetch_wp_collection(collection)
            if not records:
                continue
            for record in records:
                self.process_wp_record(collection, record)

    def fetch_wp_collection(self, collection: str) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        page = 1
        while page <= 30:
            url = self.abs_url(f"/wp-json/wp/v2/{collection}")
            try:
                response = self.get(url, params={"per_page": 100, "page": page})
            except requests.RequestException:
                break
            if response.status_code != 200 or "json" not in response.headers.get("content-type", ""):
                break
            batch = response.json()
            if not isinstance(batch, list) or not batch:
                break
            rows.extend(batch)
            total_pages = int(response.headers.get("X-WP-TotalPages") or page)
            if page >= total_pages:
                break
            page += 1
        if rows:
            self.log(f"Collected {len(rows)} WP {collection} rows")
        return rows

    def collect_taxonomies(self) -> None:
        for taxonomy in ["categories", "tags", "product_cat", "product_tag", "product_brand", "store-category", "wpsl_store_category"]:
            terms = self.fetch_wp_collection(taxonomy)
            for term in terms:
                term_id = int(term.get("id") or 0)
                name = text_from_html(term.get("name", ""))
                if term_id:
                    self.wp_taxonomy_labels.setdefault(taxonomy, {})[term_id] = name
                self.add_filter_reference(taxonomy, taxonomy, str(term.get("slug") or term_id), name, str(term.get("link") or self.base), "WordPress taxonomy/filter reference.")

    def collect_store_products(self) -> None:
        page = 1
        while page <= 30:
            try:
                response = self.get(self.abs_url("/wp-json/wc/store/v1/products"), params={"per_page": 100, "page": page})
            except requests.RequestException:
                break
            if response.status_code != 200 or "json" not in response.headers.get("content-type", ""):
                break
            batch = response.json()
            if not isinstance(batch, list) or not batch:
                break
            for item in batch:
                if isinstance(item, dict):
                    self.store_products_by_id[int(item.get("id") or 0)] = item
                    for category in item.get("categories", []) or []:
                        if isinstance(category, dict):
                            self.add_filter_reference("store_product_category", "store_product_category", str(category.get("slug") or category.get("id") or ""), str(category.get("name") or ""), str(item.get("permalink") or self.base), "WooCommerce Store API product category.")
            total_pages = int(response.headers.get("X-WP-TotalPages") or page)
            if page >= total_pages:
                break
            page += 1

    def process_wp_record(self, collection: str, record: dict[str, object]) -> None:
        link = str(record.get("link") or self.abs_url(f"/{record.get('slug', '')}/"))
        title = record_title(record)
        html = record_content_html(record)
        text = " ".join([title, link, record_excerpt(record), text_from_html(html), self.taxonomy_text(record)])
        source_type = self.classify_url(link)
        if collection in {"product", "parts", "preowned"}:
            source_type = "product"
        source_id = self.add_source_page(link, source_type, title=title, record=record)
        if collection in {"store", "wpsl_stores"}:
            self.add_branch_from_wp(record, source_id)
        elif collection in {"job_listing", "job-listings"} or any(term in text.lower() for term in JOB_HINTS):
            self.add_job(record, source_id)
        elif collection == "product":
            self.add_product_record(record, source_id, text)
        elif collection == "parts":
            self.add_part(record, source_id, title, link, image_from_wp(record), record_excerpt(record))
        elif collection == "preowned":
            self.add_preowned(record, source_id, title, link, image_from_wp(record), record_excerpt(record))
        else:
            self.inspect_html_page(link, html, source_type, source_id=source_id, title=title)

    def taxonomy_text(self, record: dict[str, object]) -> str:
        values: list[str] = []
        for key, labels in self.wp_taxonomy_labels.items():
            ids = record.get(key) or []
            if isinstance(ids, list):
                values.extend(labels.get(int(term_id), "") for term_id in ids if str(term_id).isdigit())
        return " ".join(values)

    def add_product_record(self, record: dict[str, object], source_id: str, evidence_text: str) -> None:
        title = record_title(record)
        link = str(record.get("link") or "")
        store_item = self.store_products_by_id.get(int(record.get("id") or 0), {})
        category_text = " ".join(str(cat.get("name", "")) for cat in store_item.get("categories", []) or [] if isinstance(cat, dict))
        detail_html = ""
        detail_text = ""
        detail_specs: list[dict[str, str]] = []
        detail_image = ""
        pricing: dict[str, object] = {}
        if link:
            try:
                detail_response = self.get(link, timeout=min(6.0, self.config.request_timeout_seconds))
                if detail_response.status_code == 200 and "html" in detail_response.headers.get("content-type", ""):
                    detail_html = detail_response.text
                    detail_soup = BeautifulSoup(detail_html, "lxml")
                    detail_text = detail_soup.get_text(" ", strip=True)
                    detail_specs = extract_specs_from_soup(detail_soup)
                    pricing = extract_motorcentral_pricing_from_soup(detail_soup, title)
                    image_tag = detail_soup.select_one("meta[property='og:image']") or detail_soup.select_one(".woocommerce-product-gallery img") or detail_soup.select_one("img")
                    if image_tag:
                        detail_image = image_tag.get("content") or image_tag.get("src") or ""
            except requests.RequestException as exc:
                self.add_issue("product", str(record.get("id") or title), link, "detail_page_fetch_failed", str(exc), "warning")
        spec_text = " ".join(f"{spec.get('label', '')} {spec.get('value', '')}" for spec in detail_specs)
        classification_text = " ".join([title, link, category_text, self.taxonomy_text(record), spec_text])
        price_text = " ".join([classification_text, detail_text])
        if self.is_product_motorcycle(record, store_item, classification_text, link):
            self.add_motorcycle(
                source_id=source_id,
                detail_url=link,
                listing_url=self.abs_url("/shop/"),
                display_name=title,
                brand_name=guess_brand(classification_text),
                raw_text=classification_text,
                image_url=urljoin(link, detail_image) if detail_image else self.product_image(record, store_item),
                excerpt=record_excerpt(record) or detail_text[:500],
                published_at=str(record.get("date") or ""),
                modified_at=str(record.get("modified") or ""),
                price=pricing.get("srp_amount", ""),
                price_text=str(pricing.get("srp_text", "")),
                downpayment=pricing.get("downpayment_amount", ""),
                downpayment_text=str(pricing.get("downpayment_text", "")),
                monthly=pricing.get("monthly_amount", ""),
                monthly_text=str(pricing.get("monthly_text", "")),
                specs=detail_specs,
            )
        else:
            self.add_part(record, source_id, title, link, urljoin(link, detail_image) if detail_image else self.product_image(record, store_item), record_excerpt(record) or detail_text[:1000])

    def is_product_motorcycle(self, record: dict[str, object], store_item: dict[str, object], combined_text: str, link: str) -> bool:
        category_text = " ".join(str(cat.get("name", "")) for cat in store_item.get("categories", []) or [] if isinstance(cat, dict)).lower()
        product_type_text = " ".join([category_text, self.taxonomy_text(record).lower(), link.lower()])
        if any(term in product_type_text for term in ["parts", "part", "accessories", "accessory", "helmet", "oil", "yamalube", "apparel"]):
            return False
        if any(term in product_type_text for term in ["motorcycle", "motorcycles", "scooter", "underbone", "big bike", "big-bike", "tri-wheeler", "tricycle", "automatic", "cub", "moped", "street", "sport"]):
            return True
        return looks_like_vehicle(combined_text, link)

    def product_image(self, record: dict[str, object], store_item: dict[str, object]) -> str:
        images = store_item.get("images")
        if isinstance(images, list) and images and isinstance(images[0], dict):
            return str(images[0].get("src") or "")
        return image_from_wp(record)

    def product_price(self, store_item: dict[str, object], text: str) -> float | str:
        return ""

    def add_motorcycle(
        self,
        source_id: str,
        detail_url: str,
        listing_url: str,
        display_name: str,
        brand_name: str,
        raw_text: str,
        image_url: str = "",
        excerpt: str = "",
        published_at: str = "",
        modified_at: str = "",
        price: float | str = "",
        price_text: str = "",
        downpayment: float | str = "",
        downpayment_text: str = "",
        monthly: float | str = "",
        monthly_text: str = "",
        specs: list[dict[str, str]] | None = None,
    ) -> str:
        key = normalize_url(detail_url) or f"{self.config.dealer_id}-{slugify(display_name)}"
        if key in self.motorcycle_seen:
            return f"motorcycle-{self.config.dealer_id}-{slugify(display_name)}"
        self.motorcycle_seen.add(key)
        model_name = display_name
        brand_name = brand_name or guess_brand(display_name)
        vehicle_type, vehicle_type_source = infer_vehicle_type(raw_text)
        specs = specs or []
        flat = flatten_specs(specs, brand_name, model_name)
        parsed_detail = urlparse(detail_url)
        url_slug = slugify(" ".join([parsed_detail.path, parsed_detail.query]) or display_name)
        motorcycle_id = f"motorcycle-{self.config.dealer_id}-{url_slug}"
        srp_amount = price
        has_pricing = any(value != "" for value in [srp_amount, downpayment, monthly])
        row = {
            "motorcycle_id": motorcycle_id,
            "dealer_id": self.config.dealer_id,
            "display_name": display_name,
            "brand_name": brand_name,
            "model_name": model_name,
            "vehicle_type": vehicle_type,
            "detail_url": detail_url,
            "listing_url": listing_url,
            "srp_amount": srp_amount,
            "downpayment_amount": downpayment,
            "monthly_amount": monthly,
            "price_contexts": "labelled_detail_page_price_block" if has_pricing else "",
            "primary_image_url": image_url,
            "excerpt_text": excerpt,
            "apply_loan_url": self.find_financing_url(),
            "inquiry_url": self.find_contact_url(),
            "published_at": published_at,
            "modified_at": modified_at,
            "vehicle_type_source": vehicle_type_source,
            "classification_notes": "Generic classifier based on public URL, taxonomy, title, text, and displacement.",
            "raw_taxonomies": raw_text[:5000],
            "source_page_id": source_id,
            "source_page_url": detail_url,
        }
        row.update(flat)
        self.tables.append("motorcycle", row)
        if has_pricing:
            self.tables.append(
                "pricing",
                {
                    "pricing_id": f"pricing-{motorcycle_id}-public",
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": self.config.dealer_id,
                    "display_name": display_name,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "source_url": detail_url,
                    "source_context": "public_product_or_model_page",
                    "currency": "PHP",
                    "srp_text": price_text or amount_text(srp_amount),
                    "srp_amount": srp_amount,
                    "downpayment_text": downpayment_text or amount_text(downpayment),
                    "downpayment_amount": downpayment,
                    "monthly_text": monthly_text or amount_text(monthly),
                    "monthly_amount": monthly,
                    "price_note": "Parsed from labelled Motorcentral detail-page price block; WooCommerce Store API price is ignored because it exposes a repeated fake value.",
                    "is_indicative": "yes",
                    "captured_at": self.run_started_at,
                },
            )
        for order, spec in enumerate(specs, start=1):
            number, unit = split_numeric_unit(spec["value"])
            self.tables.append(
                "motorcycle_spec",
                {
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": self.config.dealer_id,
                    "display_name": display_name,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "spec_group": spec["group"],
                    "spec_label": spec["label"],
                    "spec_value_raw": spec["value"],
                    "spec_value_numeric": number,
                    "spec_unit": unit,
                    "display_order": order,
                    "source_url": detail_url,
                },
            )
        if self.find_financing_url():
            self.tables.append(
                "financing",
                {
                    "financing_id": f"financing-{motorcycle_id}-cta",
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": self.config.dealer_id,
                    "display_name": display_name,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "source_url": detail_url,
                    "financing_type": "application_cta",
                    "form_url": self.find_financing_url(),
                    "cta_url": self.find_financing_url(),
                    "note_text": "Model row has a dealer-level public financing/application CTA; not exact approval terms.",
                    "captured_at": self.run_started_at,
                },
            )
        return motorcycle_id

    def inspect_html_page(self, url: str, html: str, source_type: str, source_id: str = "", title: str = "") -> None:
        soup = BeautifulSoup(html or "", "lxml")
        page_title = title or (soup.title.get_text(" ", strip=True) if soup.title else normalize_spaces(urlparse(url).path))
        text = soup.get_text(" ", strip=True)
        if not source_id:
            source_id = f"source-{self.config.dealer_id}-inline"
        if source_type == "motorcycle_page" and self.is_exact_motorcycle_page(url, page_title, text):
            specs = extract_specs_from_soup(soup)
            image = ""
            image_tag = soup.select_one("meta[property='og:image']") or soup.select_one("img")
            if image_tag:
                image = image_tag.get("content") or image_tag.get("src") or ""
            clean_title = re.sub(r"\s*\|\s*CDO\s*2\s*Cyc(?:les|els)\s*$", "", page_title, flags=re.I).strip() or page_title
            focused_text = " ".join([clean_title, url, " ".join(f"{spec['label']} {spec['value']}" for spec in specs)])
            pricing = extract_motorcentral_pricing_from_soup(soup, clean_title)
            self.add_motorcycle(
                source_id,
                url,
                self.abs_url("/"),
                clean_title,
                self.brand_from_url(url) or guess_brand(focused_text),
                focused_text,
                image_url=urljoin(url, image),
                excerpt=text[:500],
                price=pricing.get("srp_amount", ""),
                price_text=str(pricing.get("srp_text", "")),
                downpayment=pricing.get("downpayment_amount", ""),
                downpayment_text=str(pricing.get("downpayment_text", "")),
                monthly=pricing.get("monthly_amount", ""),
                monthly_text=str(pricing.get("monthly_text", "")),
                specs=specs,
            )
        elif source_type == "parts_page" and page_title:
            self.add_part({}, source_id, page_title, url, "", text[:1000])
        elif source_type == "promo_page":
            self.add_promo(source_id, url, page_title, text)
        elif source_type == "job_page":
            self.add_job({"slug": slugify(page_title), "date": "", "modified": ""}, source_id, title_override=page_title, description_override=text)
        elif source_type in {"service_page", "branch_page", "public_page"}:
            if any(term in f"{url} {page_title}".lower() for term in SERVICE_HINTS + BRANCH_HINTS):
                self.add_service(source_id, page_title, url, text)
            if source_type == "branch_page":
                self.parse_branch_html(url, html)
        self.extract_forms(source_id, url, soup, page_title)
        self.extract_ctas(source_id, url, soup)

    def is_exact_motorcycle_page(self, url: str, title: str, text: str) -> bool:
        path_parts = [part for part in urlparse(url).path.lower().split("/") if part]
        return len(path_parts) == 2 and path_parts[0] == "motorcycle"

    def brand_from_url(self, url: str) -> str:
        path = urlparse(url).path.lower()
        if "yamaha" in path:
            return "Yamaha"
        if "kawasaki" in path:
            return "Kawasaki"
        if "suzuki" in path:
            return "Suzuki"
        if "bajaj" in path:
            return "Bajaj"
        if "honda" in path:
            return "Honda"
        return ""

    def add_part(self, record: dict[str, object], source_id: str, title: str, detail_url: str, image_url: str, description: str) -> None:
        key = normalize_url(detail_url) or title
        if key in self.part_seen:
            return
        self.part_seen.add(key)
        self.tables.append(
            "parts_catalog",
            {
                "part_id": f"part-{self.config.dealer_id}-{slugify(title)}",
                "dealer_id": self.config.dealer_id,
                "source_page_id": source_id,
                "title": title,
                "slug": slugify(title),
                "detail_url": detail_url,
                "image_url": image_url,
                "description_text": description,
                "published_at": record.get("date", "") if record else "",
                "modified_at": record.get("modified", "") if record else "",
            },
        )

    def add_preowned(self, record: dict[str, object], source_id: str, title: str, detail_url: str, image_url: str, description: str) -> None:
        self.tables.append(
            "preowned_motorcycle",
            {
                "preowned_id": f"preowned-{self.config.dealer_id}-{slugify(title)}",
                "dealer_id": self.config.dealer_id,
                "source_page_id": source_id,
                "title": title,
                "slug": slugify(title),
                "detail_url": detail_url,
                "image_url": image_url,
                "description_text": description,
                "published_at": record.get("date", "") if record else "",
                "modified_at": record.get("modified", "") if record else "",
            },
        )

    def add_promo(self, source_id: str, url: str, title: str, text: str) -> None:
        key = normalize_url(url)
        if key in self.promo_seen:
            return
        self.promo_seen.add(key)
        self.tables.append(
            "promo",
            {
                "promo_id": f"promo-{self.config.dealer_id}-{slugify(title or url)}",
                "dealer_id": self.config.dealer_id,
                "source_page_id": source_id,
                "source_page_url": url,
                "source_url": url,
                "title": title,
                "promo_type": "promo_or_news",
                "description_text": text[:5000],
                "applies_to_brand": guess_brand(text),
                "linked_motorcycle_count": 0,
            },
        )

    def add_service(self, source_id: str, title: str, url: str, text: str) -> None:
        key = normalize_url(url)
        if key in self.service_seen:
            return
        self.service_seen.add(key)
        category = classify_service_category(f"{title} {url} {text}")
        self.tables.append(
            "aftersales_service",
            {
                "service_id": f"service-{self.config.dealer_id}-{slugify(title or url)}",
                "dealer_id": self.config.dealer_id,
                "source_page_id": source_id,
                "service_category": category,
                "title": title,
                "description_text": text[:5000],
                "cta_url": url,
                "external_form_url": url if category == "financing" else "",
                "notes": "Generic public page/service surface.",
            },
        )
        if category == "financing":
            self.tables.append(
                "financing",
                {
                    "financing_id": f"financing-{self.config.dealer_id}-{slugify(title or url)}",
                    "dealer_id": self.config.dealer_id,
                    "source_url": url,
                    "financing_type": "generic_application_surface",
                    "form_url": url,
                    "cta_url": url,
                    "note_text": "Generic financing/application page; not attached to a specific motorcycle.",
                    "captured_at": self.run_started_at,
                },
            )

    def add_job(self, record: dict[str, object], source_id: str, title_override: str = "", description_override: str = "") -> None:
        title = title_override or record_title(record)
        if not title:
            return
        self.tables.append(
            "job_listing",
            {
                "job_id": f"job-{self.config.dealer_id}-{slugify(title)}",
                "dealer_id": self.config.dealer_id,
                "source_page_id": source_id,
                "title": title,
                "slug": record.get("slug", slugify(title)),
                "application_target": str(record.get("link") or ""),
                "description_text": description_override or text_from_html(record_content_html(record)),
                "published_at": record.get("date", ""),
                "modified_at": record.get("modified", ""),
            },
        )

    def add_branch_from_wp(self, record: dict[str, object], source_id: str) -> None:
        title = record_title(record)
        link = str(record.get("link") or "")
        meta = record.get("meta") if isinstance(record.get("meta"), dict) else {}
        address = " ".join(str(meta.get(key) or "") for key in ["address", "wpsl_address", "wpsl_city", "wpsl_state", "wpsl_zip"])
        if not address:
            address = record_excerpt(record) or text_from_html(record_content_html(record))[:500]
        self.add_branch(title, link, address, phone=str(meta.get("phone") or meta.get("wpsl_phone") or ""), email=str(meta.get("email") or meta.get("wpsl_email") or ""), source_id=source_id)

    def parse_branch_html(self, url: str, html: str) -> None:
        soup = BeautifulSoup(html, "lxml")
        for table in soup.select("table"):
            headers = [normalize_spaces(th.get_text(" ", strip=True)).lower() for th in table.select("th")]
            if not headers or not any("branch" in h or "store" in h for h in headers):
                continue
            for row in table.select("tr"):
                cells = [normalize_spaces(td.get_text(" ", strip=True)) for td in row.select("td")]
                if len(cells) < 2:
                    continue
                data = dict(zip(headers, cells))
                branch_name = data.get("branch") or data.get("store") or cells[0]
                store_type = data.get("store", "")
                address = data.get("address", "")
                phone = " / ".join(v for k, v in data.items() if k in {"landline", "smart", "globe", "phone"} and v and v.lower() != "n/a")
                self.add_branch(branch_name, url, address, phone=phone, product_raw=store_type, source_id="")
        text_lines = [normalize_spaces(line) for line in soup.get_text("\n", strip=True).splitlines() if normalize_spaces(line)]
        for idx, line in enumerate(text_lines):
            if line.lower() == "contact details:" and idx > 0:
                name = text_lines[idx - 1]
                details = text_lines[idx + 1 : idx + 5]
                address = details[0] if details else ""
                email = next((item for item in details if "@" in item), "")
                phone = next((item for item in details if re.search(r"\d{7,}", item)), "")
                self.add_branch(name, url, address, phone=phone, email=email, source_id="")

    def add_branch(
        self,
        branch_name: str,
        source_url: str,
        address: str,
        phone: str = "",
        email: str = "",
        product_raw: str = "",
        categories_raw: str = "",
        source_id: str = "",
    ) -> str:
        branch_name = normalize_spaces(branch_name)
        if not branch_name:
            return ""
        branch_id = f"branch-{self.config.dealer_id}-{slugify(branch_name)}"
        if branch_id in self.branch_seen:
            return branch_id
        self.branch_seen.add(branch_id)
        city, province = infer_city_province(address)
        self.tables.append(
            "branch",
            {
                "branch_id": branch_id,
                "dealer_id": self.config.dealer_id,
                "source_url": source_url,
                "locator_id": slugify(branch_name),
                "branch_name": branch_name,
                "slug": slugify(branch_name),
                "city": city,
                "state_or_province": province,
                "full_address": address,
                "phone": phone,
                "email": email,
                "website": self.base,
                "brand_raw": guess_brand(f"{branch_name} {product_raw}"),
                "product_raw": product_raw or "motorcycle" if looks_like_vehicle(branch_name + " " + product_raw) else product_raw,
                "categories_raw": categories_raw,
            },
        )
        for coverage_type, value in [("product_family", product_raw or "motorcycle"), ("brand", guess_brand(branch_name + " " + product_raw)), ("category", categories_raw)]:
            if value:
                self.add_branch_coverage(branch_id, coverage_type, value, source_url)
        return branch_id

    def add_branch_coverage(self, branch_id: str, coverage_type: str, value: str, source_url: str) -> None:
        normalized = slugify(value)
        key = (branch_id, coverage_type, normalized)
        if key in self.coverage_seen:
            return
        self.coverage_seen.add(key)
        self.tables.append(
            "branch_product_coverage",
            {
                "coverage_id": f"coverage-{self.config.dealer_id}-{slugify(branch_id)}-{coverage_type}-{normalized}",
                "dealer_id": self.config.dealer_id,
                "branch_id": branch_id,
                "coverage_type": coverage_type,
                "coverage_value_raw": value,
                "coverage_value_normalized": normalized,
                "source_url": source_url,
                "source_field": "public_branch_text",
                "confidence": "coverage_only",
                "observed_at": self.run_started_at,
                "notes": "Broad branch coverage hint; not exact model availability.",
            },
        )

    def extract_forms(self, source_id: str, source_url: str, soup: BeautifulSoup, page_title: str) -> None:
        for form_index, form in enumerate(soup.find_all("form"), start=1):
            form_id = f"form-{self.config.dealer_id}-{slugify(page_title or source_url)}-{form_index}"
            if form_id in self.form_seen:
                continue
            self.form_seen.add(form_id)
            action = urljoin(source_url, form.get("action", ""))
            method = str(form.get("method") or "get").upper()
            inputs = form.find_all(["input", "select", "textarea", "button"])
            self.tables.append(
                "form_definition",
                {
                    "form_id": form_id,
                    "dealer_id": self.config.dealer_id,
                    "form_name": page_title or f"Form {form_index}",
                    "form_category": classify_service_category(f"{page_title} {source_url}"),
                    "source_url": source_url,
                    "final_url": source_url,
                    "page_title": page_title,
                    "form_action": action,
                    "form_method": method,
                    "requires_file_upload": "yes" if form.find("input", {"type": "file"}) else "no",
                    "has_privacy_checkbox": "yes" if "privacy" in form.get_text(" ", strip=True).lower() else "unknown",
                    "notes": "Form inspected structurally only; no live submission.",
                },
            )
            order = 1
            for field in inputs:
                name = field.get("name") or field.get("id") or field.get("type") or field.name
                label = self.find_label(form, field) or field.get("aria-label") or field.get("placeholder") or name
                options = field.find_all("option") if field.name == "select" else []
                if options:
                    for option in options:
                        self.add_form_field(form_id, page_title, source_url, order, name, label, "select", field, option)
                        order += 1
                else:
                    self.add_form_field(form_id, page_title, source_url, order, name, label, field.get("type") or field.name, field, None)
                    order += 1

    def add_form_field(self, form_id: str, form_name: str, source_url: str, order: int, name: str, label: str, field_type: str, field: object, option: object | None) -> None:
        self.tables.append(
            "form_field",
            {
                "form_id": form_id,
                "dealer_id": self.config.dealer_id,
                "form_name": form_name,
                "form_category": classify_service_category(f"{form_name} {source_url}"),
                "field_order": order,
                "field_name": normalize_spaces(name),
                "field_label": normalize_spaces(label),
                "field_type": field_type,
                "is_required": "yes" if getattr(field, "has_attr", lambda _: False)("required") else "no",
                "default_value": getattr(field, "get", lambda *_: "")("value", ""),
                "placeholder": getattr(field, "get", lambda *_: "")("placeholder", ""),
                "option_value": getattr(option, "get", lambda *_: "")("value", "") if option else "",
                "option_label": normalize_spaces(option.get_text(" ", strip=True)) if option else "",
                "source_url": source_url,
            },
        )

    def find_label(self, form: BeautifulSoup, field: object) -> str:
        field_id = field.get("id") if hasattr(field, "get") else ""
        if field_id:
            label = form.find("label", {"for": field_id})
            if label:
                return normalize_spaces(label.get_text(" ", strip=True))
        parent = getattr(field, "parent", None)
        if parent:
            label = parent.find("label")
            if label:
                return normalize_spaces(label.get_text(" ", strip=True))
        return ""

    def extract_ctas(self, source_id: str, source_url: str, soup: BeautifulSoup) -> None:
        for anchor in soup.find_all("a", href=True):
            href = urljoin(source_url, anchor["href"])
            text = normalize_spaces(anchor.get_text(" ", strip=True))
            if not href.startswith("http"):
                continue
            key = (source_url, href)
            if key in self.cta_seen:
                continue
            self.cta_seen.add(key)
            self.tables.append(
                "external_cta",
                {
                    "cta_id": f"cta-{self.config.dealer_id}-{len(self.cta_seen)}",
                    "dealer_id": self.config.dealer_id,
                    "source_page_id": source_id,
                    "source_url": source_url,
                    "target_url": href,
                    "target_domain": urlparse(href).netloc,
                    "cta_type": infer_cta_type(href, text),
                    "anchor_text": text,
                    "notes": "Public link/CTA evidence.",
                },
            )

    def collect_sitemap_pages(self) -> None:
        if self.wp_available:
            return
        selected = []
        for url in self.sitemap_urls:
            lower = url.lower()
            if any(term in lower for term in VEHICLE_HINTS + PART_HINTS + PROMO_HINTS + SERVICE_HINTS + JOB_HINTS + BRANCH_HINTS):
                selected.append(url)
        for url in selected[: self.config.max_sitemap_pages]:
            try:
                response = self.get(url)
            except requests.RequestException:
                continue
            if response.status_code != 200 or "html" not in response.headers.get("content-type", ""):
                continue
            source_type = self.classify_url(response.url)
            source_id = self.add_source_page(response.url, source_type, html=response.text)
            self.inspect_html_page(response.url, response.text, source_type, source_id=source_id)

    def collect_dueksam_ajax(self) -> None:
        if self.config.dealer_id != "dueksam":
            return
        self.log("Collecting Dueksam AJAX motorcycle and branch feeds")
        seen_entities: set[str] = set()
        for brand_filter in ["Honda", "Kawasaki", "Suzuki", "Yamaha", ""]:
            empty_pages = 0
            for page_id in [0, *range(1, 80)]:
                try:
                    response = self.post(self.abs_url("/includes/motorcycles.php"), data={"type": "fetchDT", "brand": brand_filter, "pageId": page_id})
                    data = response.json()
                except Exception as exc:
                    if not seen_entities:
                        self.add_issue("motorcycle", "dueksam_ajax", self.abs_url("/includes/motorcycles.php"), "ajax_failed", str(exc), "warning")
                    break
                rows = data.get("data", []) if isinstance(data, dict) else []
                if not rows:
                    empty_pages += 1
                    if empty_pages >= 3:
                        break
                    continue
                empty_pages = 0
                for item in rows:
                    if not isinstance(item, dict):
                        continue
                    entity_id = str(item.get("entity_id") or "")
                    if entity_id in seen_entities:
                        continue
                    seen_entities.add(entity_id)
                    model = normalize_spaces(item.get("model"))
                    if not model:
                        continue
                    brand = brand_filter or guess_brand(" ".join([str(item.get("sku") or ""), model]))
                    image = self.abs_url("/media/catalog/product" + str(item.get("img") or ""))
                    source_id = self.add_source_page(self.abs_url(f"/motorcycles?entity_id={entity_id}"), "motorcycle_ajax_record", title=model)
                    self.add_motorcycle(
                        source_id,
                        self.abs_url(f"/motorcycles?entity_id={entity_id}"),
                        self.abs_url("/motorcycles"),
                        model,
                        brand,
                        "motorcycle " + brand + " " + json_text(item),
                        image_url=image,
                        price=parse_amount(item.get("sp_price") or item.get("price")),
                    )
        for page in range(0, 60):
            try:
                response = self.post(self.abs_url("/includes/branches.php"), data={"type": "fetch", "pageId": page, "mySearch": ""})
                payload = response.json()
            except Exception:
                break
            html = payload.get("data", "") if isinstance(payload, dict) else ""
            if not html:
                break
            self.parse_branch_html(self.abs_url("/branches"), html)
            if "changePagination" not in str(payload.get("page", "")):
                break

    def collect_relationship_probe(self) -> None:
        self.tables.append(
            "relationship_probe",
            {
                "probe_id": f"probe-{self.config.dealer_id}-branch-model",
                "dealer_id": self.config.dealer_id,
                "relationship_name": "branch_to_motorcycle_model_availability",
                "source_url": self.base,
                "source_type": "public_site_review",
                "status": "no_evidence",
                "evidence_found": "false",
                "rows_created": self.tables.row_count("branch_model_availability"),
                "notes": "No exact public branch-to-model inventory evidence was found. Do not join branch to motorcycle by dealer_id.",
                "checked_at": self.run_started_at,
            },
        )

    def collect_data_quality_notes(self) -> None:
        if not self.wp_available:
            self.add_issue("site", self.config.dealer_id, self.base, "wordpress_rest_unavailable", "WordPress REST/WooCommerce collections were not available or not applicable; scraper used sitemap/static/AJAX sources.", "info")
        if not self.tables.row_count("branch_model_availability"):
            self.add_issue("relationship", "branch_model_availability", self.base, "no_exact_branch_model_relationship", "Public site did not expose exact branch-level model availability.", "info")
    def add_issue(self, entity_type: str, entity_id: str, source_url: str, issue_type: str, detail: str, severity: str) -> None:
        self.tables.append(
            "data_quality_issue",
            {
                "issue_id": f"issue-{self.config.dealer_id}-{self.tables.row_count('data_quality_issue') + 1}",
                "dealer_id": self.config.dealer_id,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "source_url": source_url,
                "issue_type": issue_type,
                "issue_detail": detail,
                "severity": severity,
                "captured_at": self.run_started_at,
            },
        )

    def add_filter_reference(self, reference_type: str, filter_name: str, option_value: str, option_label: str, source_url: str, notes: str) -> None:
        self.tables.append(
            "website_filter_reference",
            {
                "dealer_id": self.config.dealer_id,
                "reference_type": reference_type,
                "filter_name": filter_name,
                "option_value": option_value,
                "option_label": option_label,
                "source_url": source_url,
                "notes": notes,
            },
        )

    def find_financing_url(self) -> str:
        for url in self.sitemap_urls + [self.abs_url(path) for path in self.config.known_pages]:
            if any(term in url.lower() for term in ["loan", "credit", "financing", "application", "inquire"]):
                return url
        return ""

    def find_contact_url(self) -> str:
        for url in self.sitemap_urls + [self.abs_url(path) for path in self.config.known_pages]:
            if "contact" in url.lower() or "inquire" in url.lower():
                return url
        return ""

    def finalize_crawl_run(self) -> None:
        source_count = sum(self.tables.row_count(sheet) for sheet in DATABASE_SHEETS if sheet not in {"dealer", "crawl_run"})
        self.tables.append(
            "crawl_run",
            {
                "run_id": self.run_id,
                "dealer_id": self.config.dealer_id,
                "run_started_at": self.run_started_at,
                "run_finished_at": datetime.now(TIMEZONE).isoformat(),
                "source_count": source_count,
                "notes": f"Generated from public website sources. sitemap_urls={len(self.sitemap_urls)} wp_available={self.wp_available}",
            },
        )

    def write_workbook(self) -> bool:
        if self.should_preserve_existing_workbook():
            self.log(f"Skipped overwrite of {self.output_path}; current run has no scrapeable records and an existing workbook has data")
            return False
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in DATABASE_SHEETS:
            sheet = workbook.create_sheet(sheet_name)
            schema = DATABASE_SCHEMAS[sheet_name]
            sheet.append(schema)
            for row in self.tables.rows.get(sheet_name, []):
                sheet.append([sanitize_excel_value(row.get(column)) for column in schema])
        workbook.save(self.output_path)
        return True

    def should_preserve_existing_workbook(self) -> bool:
        if not self.output_path.exists():
            return False
        evidence_sheets = ("source_page", "motorcycle", "motorcycle_spec", "branch", "pricing", "promo", "aftersales_service", "form_definition")
        current_records = sum(len(self.tables.rows.get(sheet, [])) for sheet in evidence_sheets)
        if current_records > 0:
            return False
        try:
            workbook = load_workbook(self.output_path, read_only=True, data_only=True)
        except Exception:
            return False
        existing_records = 0
        for sheet_name in evidence_sheets:
            if sheet_name in workbook.sheetnames:
                existing_records += max(workbook[sheet_name].max_row - 1, 0)
        return existing_records > 0


def infer_city_province(address: str) -> tuple[str, str]:
    parts = [normalize_spaces(part) for part in str(address or "").split(",") if normalize_spaces(part)]
    if len(parts) >= 2:
        return parts[-2], parts[-1]
    if parts:
        return parts[0], ""
    return "", ""


def infer_cta_type(target_url: str, anchor_text: str) -> str:
    lower = f"{target_url} {anchor_text}".lower()
    if "facebook" in lower or "instagram" in lower or "linkedin" in lower:
        return "social"
    if "map" in lower or "google.com/maps" in lower:
        return "map"
    if "loan" in lower or "credit" in lower or "application" in lower:
        return "application"
    if "mailto:" in lower or "tel:" in lower:
        return "contact"
    return "link"


CONFIG = DealerConfig(
    dealer_id='motorcentralgroup',
    dealer_name='Motorcentral Group',
    primary_domain='https://motorcentralgroup.com',
    folder='motorcentralgroup',
    output_filename='motorcentralgroup_database.xlsx',
    script_filename='scrape_motorcentralgroup_database.py',
    known_pages=['/motorcycles/', '/branches/', '/contact-us/', '/job-career/', '/news-and-promos/', '/inquire-here/'],
    notes='Public Motorcentral Group WordPress/WooCommerce scrape using shared dealer database schema.',
)


def main() -> None:
    output_path = Path(__file__).resolve().parents[2] / "working/outputs/dealer-databases" / CONFIG.output_filename
    MotorcentralGroupDatabaseScraper(CONFIG, output_path).run()


if __name__ == "__main__":
    main()
