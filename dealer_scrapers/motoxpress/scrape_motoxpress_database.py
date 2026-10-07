#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, quote, urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from full_scrape.motortrade.scrape_motortrade_analytics import (
    DATABASE_SCHEMAS,
    DATABASE_SHEETS,
    TIMEZONE,
    USER_AGENT,
    TableStore,
    html_to_text,
    normalize_url,
    parse_currency_amount,
    sanitize_excel_value,
    split_numeric_unit,
)


DEALER_ID = "motoxpress"
DEALER_NAME = "MotoXpress"
PRIMARY_DOMAIN = "https://motoxpress.ph"
WWW_DOMAIN = "https://www.motoxpress.ph"
FORMS_DOMAIN = PRIMARY_DOMAIN
API_BASE = "https://api.motoxpress.ph/v1"
REQUEST_DELAY_SECONDS = 0.05

PAGE_URLS = {
    "home": f"{PRIMARY_DOMAIN}/",
    "shop": f"{PRIMARY_DOMAIN}/shop",
    "location": f"{PRIMARY_DOMAIN}/location",
    "news_promos": f"{PRIMARY_DOMAIN}/news-promos",
    "about": f"{PRIMARY_DOMAIN}/about",
    "financing": f"{PRIMARY_DOMAIN}/financing?form=true&make=Honda&model=ADV%20160",
}

SPEC_FIELDS = [
    ("fuel_type", "Fuel Type"),
    ("transmission", "Transmission"),
    ("price_range", "Price Range"),
    ("exterior_dimensions", "Exterior Dimensions"),
    ("wheel_type", "Wheel Type"),
    ("ground_clearance", "Ground Clearance"),
    ("seat_height", "Seat Height"),
    ("full_tank_capacity", "Full Tank Capacity"),
    ("engine_type", "Engine Type"),
    ("displacement", "Displacement"),
    ("starting_system", "Starting System"),
    ("battery_type", "Battery Type"),
    ("fuel_system", "Fuel System"),
    ("front_brakes_types", "Front Brakes Type"),
    ("rear_brakes_types", "Rear Brakes Type"),
    ("tires", "Tires"),
    ("meter_panel", "Meter Panel"),
    ("additional_features", "Additional Features"),
    ("others", "Others"),
]

FINANCING_TERMS = ["12", "18", "24", "36", "48", "60"]
DOWNPAYMENT_OPTIONS = ["15", "20", "30", "40", "50", "60", "70"]

FINANCING_FORM_FIELDS = [
    ("make", "Make", "select", "Motorcycle and Loan Details", ["Honda", "Yamaha", "Suzuki", "Kawasaki", "KYMCO", "TVS"]),
    ("model", "Model", "select", "Motorcycle and Loan Details", []),
    ("preferred_monthly_term", "Preferred Monthly Term", "select", "Motorcycle and Loan Details", FINANCING_TERMS),
    ("downpayment", "Downpayment (%)", "select", "Motorcycle and Loan Details", DOWNPAYMENT_OPTIONS),
    ("firstname", "First name", "text", "Borrower's Information", []),
    ("middlename", "Middle name", "text", "Borrower's Information", []),
    ("lastname", "Last name", "text", "Borrower's Information", []),
    ("email", "Email", "email", "Borrower's Information", []),
    ("civil_status", "Civil Status", "select", "Borrower's Information", ["Single", "Married", "Widowed", "Separated", "Annulled"]),
    ("mobile_number", "Mobile Number", "number", "Borrower's Information", []),
    ("birthdate", "Birthdate", "date", "Borrower's Information", []),
    ("gender", "Gender", "text", "Borrower's Information", []),
    ("birth_place", "Birth Place", "text", "Borrower's Information", []),
    ("nationality", "Nationality", "text", "Borrower's Information", []),
    ("current_address", "House Number & Street", "text", "Borrower's Information", []),
    ("region", "Region", "select", "Borrower's Information", []),
    ("province", "Province", "select", "Borrower's Information", []),
    ("city", "City", "select", "Borrower's Information", []),
    ("barangay", "Barangay", "select", "Borrower's Information", []),
    ("type_of_residency", "Type of Residency", "select", "Borrower's Information", ["Owned", "Living w/Parents", "Living w/ Relatives", "Rented", "Stay in"]),
    ("permanent_or_province_address", "Complete permanent/province address", "text", "Borrower's Information", []),
    ("monthly_rental", "Monthly Rental", "text", "Borrower's Information", []),
    ("relationship_to_owner_permanent_address", "Relationship to owner of permanent address", "text", "Borrower's Information", []),
    ("spouse_firstname", "Spouse first name", "text", "Spouse Information", []),
    ("spouse_middlename", "Spouse middle name", "text", "Spouse Information", []),
    ("spouse_lastname", "Spouse last name", "text", "Spouse Information", []),
    ("spouse_email", "Spouse email", "email", "Spouse Information", []),
    ("spouse_contact_number", "Spouse contact number", "number", "Spouse Information", []),
    ("spouse_birthdate", "Spouse birthdate", "date", "Spouse Information", []),
    ("spouse_source_of_income", "Spouse source of income", "text", "Spouse Information", []),
    ("spouse_company_name", "Spouse company name", "text", "Spouse Information", []),
    ("spouse_position", "Spouse position", "text", "Spouse Information", []),
    ("spouse_gross_monthly_income", "Spouse gross monthly income", "text", "Spouse Information", []),
    ("references[0].fullname", "Character reference 1 name", "text", "Character References", []),
    ("references[0].contact_number", "Character reference 1 contact number", "number", "Character References", []),
    ("references[0].relationship", "Character reference 1 relationship", "text", "Character References", []),
    ("references[1].fullname", "Character reference 2 name", "text", "Character References", []),
    ("references[1].contact_number", "Character reference 2 contact number", "number", "Character References", []),
    ("references[1].relationship", "Character reference 2 relationship", "text", "Character References", []),
    ("references[2].fullname", "Character reference 3 name", "text", "Character References", []),
    ("references[2].contact_number", "Character reference 3 contact number", "number", "Character References", []),
    ("references[2].relationship", "Character reference 3 relationship", "text", "Character References", []),
    ("childrens[0].fullname", "Child/dependent 1 name", "text", "Children", []),
    ("childrens[0].contact_number", "Child/dependent 1 contact number", "number", "Children", []),
    ("childrens[0].relationship", "Child/dependent 1 relationship", "text", "Children", []),
    ("employer_name", "Employer name", "text", "Employment", []),
    ("work_address", "Work address", "text", "Employment", []),
    ("position", "Position", "text", "Employment", []),
    ("tenure", "Tenure", "select", "Employment", ["below 1 year", "1 year", "2 years", "3 years", "4 years", "5 years", "6 years up"]),
    ("net_income", "Net income", "text", "Employment", []),
    ("employment_status", "Employment status", "text", "Employment", []),
    ("work_contact_person", "Work contact person", "text", "Employment", []),
    ("work_contact_number", "Work contact number", "number", "Employment", []),
    ("sender_name", "Remittance sender name", "text", "Remittance", []),
    ("sender_relationship", "Remittance sender relationship", "text", "Remittance", []),
    ("monthly_remittance", "Monthly remittance", "text", "Remittance", []),
    ("country", "Country", "text", "Remittance", []),
    ("bank_and_branch", "Bank and branch", "text", "Financial Information", []),
    ("account_number", "Account number", "text", "Financial Information", []),
    ("bank_telephone_no", "Bank telephone number", "text", "Financial Information", []),
    ("loans_mortgages_cc", "Loans / mortgages / credit cards", "text", "Financial Information", []),
    ("type_of_loan", "Type of loan", "text", "Financial Information", []),
    ("loan_card_no", "Loan/card number", "text", "Financial Information", []),
    ("name_address_of_creditor", "Name/address of creditor", "text", "Financial Information", []),
    ("contact_person_and_number", "Creditor contact person and number", "text", "Financial Information", []),
    ("loan_monthly_amortization", "Loan monthly amortization", "text", "Financial Information", []),
    ("contact_method", "Preferred contact method", "select", "Contact Preference", ["Phone Call", "Text Message", "E-Mail", "Viber", "FB Messenger"]),
    ("best_time_to_call", "Best time to call", "text", "Contact Preference", []),
    ("other_contact", "Other contact", "text", "Contact Preference", []),
    ("files", "Requirement uploads", "file", "Attachments", ["image/jpeg", "image/png", "application/pdf"]),
    ("terms", "Terms and privacy agreement", "checkbox", "Consent", []),
]


class MotoXpressDatabaseScraper:
    def __init__(self, output_path: Path):
        self.output_path = output_path
        self.session = requests.Session()
        self.session.headers.update(
            {
                "User-Agent": USER_AGENT,
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/json;q=0.8,*/*;q=0.7",
                "Accept-Language": "en-US,en;q=0.9",
            }
        )
        self.tables = TableStore(DATABASE_SCHEMAS)
        self.source_page_by_url: dict[str, str] = {}
        self.motorcycle_by_slug: dict[str, str] = {}
        self.motorcycle_by_model: dict[str, str] = {}
        self.motorcycle_summary: dict[str, dict[str, str]] = {}
        self.external_cta_seen: set[tuple[str, str]] = set()
        self.coverage_seen: set[tuple[str, str, str]] = set()
        self.promo_seen: set[str] = set()
        self.promo_link_seen: set[tuple[str, str, str]] = set()
        self.source_counter = 0
        self.build_id = ""
        self.brands: list[dict[str, object]] = []
        self.motorcycles: list[dict[str, object]] = []
        self.locations: list[dict[str, object]] = []
        self.articles: list[dict[str, object]] = []
        self.article_categories: list[dict[str, object]] = []
        self.about_sections: list[dict[str, object]] = []
        self.faqs: list[dict[str, object]] = []
        self.financing_page: dict[str, object] = {}

        now = datetime.now(TIMEZONE)
        self.crawl_date = now.date().isoformat()
        self.run_started_at = now.isoformat()
        self.run_id = f"{DEALER_ID}-run-{now.strftime('%Y%m%d%H%M%S')}"

    def run(self) -> None:
        self.log("Starting MotoXpress database scrape")
        self.add_dealer()
        self.discover_build_id()
        self.collect_static_pages()
        self.collect_brands()
        self.collect_motorcycles()
        self.collect_locations()
        self.collect_articles()
        self.collect_about_and_services()
        self.collect_forms()
        self.collect_relationship_probe()
        self.collect_data_quality_notes()
        self.finalize_crawl_run()
        self.log("Writing workbook")
        self.write_workbook()

    def log(self, message: str) -> None:
        print(f"[MotoXpress] {message}", flush=True)

    def add_dealer(self) -> None:
        self.tables.append(
            "dealer",
            {
                "dealer_id": DEALER_ID,
                "dealer_name": DEALER_NAME,
                "primary_domain": PRIMARY_DOMAIN,
                "forms_domain": FORMS_DOMAIN,
                "crawl_timezone": str(TIMEZONE),
                "crawl_date": self.crawl_date,
                "notes": "Public MotoXpress Next.js/API snapshot using the shared dealer database schema.",
            },
        )

    def fetch(self, url: str, *, params: dict[str, object] | None = None, timeout: int = 45) -> requests.Response:
        time.sleep(REQUEST_DELAY_SECONDS)
        response = self.session.get(url, params=params, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
        return response

    def fetch_json(self, url: str, *, params: dict[str, object] | None = None) -> object:
        return self.fetch(url, params=params).json()

    def fetch_html(self, url: str) -> tuple[str, BeautifulSoup, str]:
        response = self.fetch(url)
        return response.text, BeautifulSoup(response.text, "lxml"), response.url

    def discover_build_id(self) -> None:
        html, soup, final_url = self.fetch_html(PAGE_URLS["home"])
        next_data = parse_next_data(soup)
        self.build_id = str(next_data.get("buildId") or "")
        if not self.build_id:
            raise RuntimeError("MotoXpress Next.js build id was not found.")
        self.log(f"Discovered Next build id {self.build_id}")
        self.add_source_page(
            source_page_id="source-next-home",
            source_type="next_page",
            wp_post_type="next_page",
            wp_id="home",
            url=final_url,
            slug="home",
            title=soup.title.get_text(" ", strip=True) if soup.title else "MotoXpress PH",
            content_text=soup.get_text(" ", strip=True),
        )

    def next_json(self, page: str, *, params: dict[str, object] | None = None) -> dict[str, object]:
        data = self.fetch_json(f"{PRIMARY_DOMAIN}/_next/data/{self.build_id}/{page}.json", params=params)
        return data.get("pageProps", {}) if isinstance(data, dict) else {}

    def collect_static_pages(self) -> None:
        for slug, url in PAGE_URLS.items():
            if slug == "home":
                continue
            html, soup, final_url = self.fetch_html(url)
            self.add_source_page(
                source_page_id=f"source-next-{slug}",
                source_type="next_page",
                wp_post_type="next_page",
                wp_id=slug,
                url=final_url,
                slug=slug,
                title=soup.title.get_text(" ", strip=True) if soup.title else slug,
                content_text=soup.get_text(" ", strip=True),
            )
            self.collect_external_links_from_soup(soup, final_url, f"source-next-{slug}")

    def collect_brands(self) -> None:
        payload = self.fetch_json(f"{API_BASE}/brands")
        self.brands = payload.get("brands", []) if isinstance(payload, dict) else []
        for brand in self.brands:
            self.tables.append(
                "website_filter_reference",
                {
                    "dealer_id": DEALER_ID,
                    "reference_type": "brand",
                    "filter_name": "brand",
                    "option_value": brand.get("name", ""),
                    "option_label": brand.get("name", ""),
                    "source_url": f"{API_BASE}/brands",
                    "notes": "Public MotoXpress brand/filter reference. Use motorcycle.brand_name for analytics filtering.",
                },
            )

    def collect_motorcycles(self) -> None:
        payload = self.fetch_json(f"{API_BASE}/motorcycles", params={"page": 1, "per_page": 100})
        self.motorcycles = payload.get("data", []) if isinstance(payload, dict) else []
        self.log(f"Collected {len(self.motorcycles)} motorcycle API records")
        for motorcycle in self.motorcycles:
            self.add_motorcycle(motorcycle)

    def add_motorcycle(self, motorcycle: dict[str, object]) -> None:
        source_url = f"{API_BASE}/motorcycles/{motorcycle.get('slug', '')}"
        detail_url = f"{PRIMARY_DOMAIN}/shop/{motorcycle.get('slug', '')}"
        source_page_id = f"source-api-motorcycle-{motorcycle.get('id')}"
        brand_name = brand_name_from_motorcycle(motorcycle)
        model_name = str(motorcycle.get("model_name") or "")
        display_name = f"{brand_name} {model_name}".strip()
        motorcycle_id = f"motorcycle-{motorcycle.get('id')}"
        specs = motorcycle_specs(motorcycle)
        flattened = flatten_specs(specs, brand_name, model_name, motorcycle)
        srp_amount = parse_published_price(motorcycle.get("suggested_retail_price"))
        downpayment_amount = parse_published_price(motorcycle.get("all_in_downpayment"))
        monthly_amount = parse_published_price(motorcycle.get("monthly_amortization"))
        vehicle_type = infer_vehicle_type(motorcycle)
        image_url = primary_image_url(motorcycle)

        self.add_source_page(
            source_page_id=source_page_id,
            source_type="api_motorcycle",
            wp_post_type="api_motorcycle",
            wp_id=str(motorcycle.get("id", "")),
            url=source_url,
            slug=str(motorcycle.get("slug", "")),
            title=display_name,
            content_text=json.dumps(motorcycle, ensure_ascii=False),
            published_at=str(motorcycle.get("createdAt") or ""),
            modified_at=str(motorcycle.get("updatedAt") or ""),
            featured_image_url=image_url,
        )
        self.tables.append(
            "motorcycle",
            {
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "display_name": display_name,
                "brand_name": brand_name,
                "model_name": model_name,
                "vehicle_type": vehicle_type,
                "detail_url": detail_url,
                "listing_url": f"{PRIMARY_DOMAIN}/shop",
                "srp_amount": srp_amount,
                "downpayment_amount": downpayment_amount,
                "monthly_amount": monthly_amount,
                "price_contexts": "public_motorcycles_api",
                "promo_titles": "On Sale" if motorcycle.get("is_sale") else "",
                "is_blue_tag": "",
                "is_featured": "",
                "is_top_pick": "",
                "has_vehicle_promo_flag": "TRUE" if motorcycle.get("is_sale") else "",
                "spec_make": brand_name,
                "spec_model": model_name,
                "spec_transmission": flattened.get("spec_transmission", ""),
                "spec_engine_type": flattened.get("spec_engine_type", ""),
                "spec_seat_height": flattened.get("spec_seat_height", ""),
                "spec_brake_system": flattened.get("spec_brake_system", ""),
                "spec_displacement_cc": flattened.get("spec_displacement_cc", ""),
                "spec_front_tire": flattened.get("spec_front_tire", ""),
                "spec_rear_tire": flattened.get("spec_rear_tire", ""),
                "spec_wheels_type": flattened.get("spec_wheels_type", ""),
                "spec_starting_system": flattened.get("spec_starting_system", ""),
                "spec_overall_dimensions": flattened.get("spec_overall_dimensions", ""),
                "spec_ground_clearance": flattened.get("spec_ground_clearance", ""),
                "spec_fuel_system": flattened.get("spec_fuel_system", ""),
                "primary_image_url": image_url,
                "excerpt_text": html_to_text(str(motorcycle.get("additional_features") or "")),
                "apply_loan_url": financing_url(brand_name, model_name),
                "inquiry_url": f"{PRIMARY_DOMAIN}/financing?form=true&make={quote(brand_name)}&model={quote(model_name)}",
                "service_url": "",
                "compare_enabled": "",
                "published_at": motorcycle.get("createdAt", ""),
                "modified_at": motorcycle.get("updatedAt", ""),
                "vehicle_type_source": "api_displacement_and_fuel_type",
                "classification_notes": f"API is_active={motorcycle.get('is_active')}; stock={motorcycle.get('stock')}; is_sale={motorcycle.get('is_sale')}.",
                "raw_taxonomies": join_values([brand_name, motorcycle.get("fuel_type", ""), motorcycle.get("transmission", ""), "active" if motorcycle.get("is_active") else "inactive"]),
                "source_page_id": source_page_id,
                "source_page_url": source_url,
            },
        )
        self.motorcycle_by_slug[str(motorcycle.get("slug", ""))] = motorcycle_id
        self.motorcycle_by_model[normalize_token(model_name)] = motorcycle_id
        self.motorcycle_summary[motorcycle_id] = {"display_name": display_name, "brand_name": brand_name, "model_name": model_name}

        self.add_specs(motorcycle_id, display_name, brand_name, model_name, detail_url, specs)
        self.add_pricing(motorcycle_id, display_name, brand_name, model_name, source_url, srp_amount, downpayment_amount, monthly_amount, motorcycle)
        self.add_financing(motorcycle_id, display_name, brand_name, model_name, source_url, downpayment_amount, monthly_amount)
        if motorcycle.get("brochure_file_url"):
            self.add_external_cta(source_page_id, source_url, str(motorcycle.get("brochure_file_url")), "Brochure")

    def add_specs(
        self,
        motorcycle_id: str,
        display_name: str,
        brand_name: str,
        model_name: str,
        source_url: str,
        specs: list[dict[str, str]],
    ) -> None:
        for index, spec in enumerate(specs, start=1):
            numeric, unit = split_numeric_unit(spec["value"])
            self.tables.append(
                "motorcycle_spec",
                {
                    "motorcycle_id": motorcycle_id,
                    "dealer_id": DEALER_ID,
                    "display_name": display_name,
                    "brand_name": brand_name,
                    "model_name": model_name,
                    "spec_group": spec.get("group", "Specifications"),
                    "spec_label": spec["label"],
                    "spec_value_raw": spec["value"],
                    "spec_value_numeric": numeric,
                    "spec_unit": unit,
                    "display_order": index,
                    "source_url": source_url,
                },
            )

    def add_pricing(
        self,
        motorcycle_id: str,
        display_name: str,
        brand_name: str,
        model_name: str,
        source_url: str,
        srp_amount: float | str,
        downpayment_amount: float | str,
        monthly_amount: float | str,
        motorcycle: dict[str, object],
    ) -> None:
        self.tables.append(
            "pricing",
            {
                "pricing_id": f"pricing-{motorcycle_id}-api",
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "display_name": display_name,
                "brand_name": brand_name,
                "model_name": model_name,
                "source_url": source_url,
                "source_context": "public_motorcycles_api",
                "currency": "PHP",
                "srp_text": amount_text(srp_amount),
                "srp_amount": srp_amount,
                "downpayment_text": amount_text(downpayment_amount),
                "downpayment_amount": downpayment_amount,
                "monthly_text": amount_text(monthly_amount),
                "monthly_amount": monthly_amount,
                "price_note": f"API price fields. price_range={motorcycle.get('price_range')}; is_sale={motorcycle.get('is_sale')}; stock={motorcycle.get('stock')}.",
                "is_indicative": "TRUE",
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def add_financing(
        self,
        motorcycle_id: str,
        display_name: str,
        brand_name: str,
        model_name: str,
        source_url: str,
        downpayment_amount: float | str,
        monthly_amount: float | str,
    ) -> None:
        self.tables.append(
            "financing",
            {
                "financing_id": f"financing-{motorcycle_id}-api",
                "motorcycle_id": motorcycle_id,
                "dealer_id": DEALER_ID,
                "display_name": display_name,
                "brand_name": brand_name,
                "model_name": model_name,
                "source_url": source_url,
                "financing_type": "api_model_financing_fields",
                "program_name": "MotoXpress motorcycle financing",
                "monthly_amount": monthly_amount,
                "downpayment_amount": downpayment_amount,
                "loan_amount": "",
                "loan_term_months": "",
                "estimated_income_amount": "",
                "form_url": financing_url(brand_name, model_name),
                "cta_url": financing_url(brand_name, model_name),
                "note_text": "Model-specific public API downpayment/monthly fields and prefilled financing URL.",
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def collect_locations(self) -> None:
        page = 1
        total_pages = 1
        while page <= total_pages:
            props = self.next_json("location", params={"page": page})
            pagination = props.get("pagination") if isinstance(props.get("pagination"), dict) else {}
            total_pages = int(pagination.get("pageCount") or total_pages)
            page_locations = props.get("locations", []) if isinstance(props.get("locations"), list) else []
            self.locations.extend(page_locations)
            if page == 1:
                for city in props.get("allCities", []) or []:
                    self.tables.append(
                        "website_filter_reference",
                        {
                            "dealer_id": DEALER_ID,
                            "reference_type": "branch_city_filter",
                            "filter_name": "city",
                            "option_value": normalize_spaces(city.get("city", "")),
                            "option_label": normalize_spaces(city.get("city", "")),
                            "source_url": f"{PRIMARY_DOMAIN}/location",
                            "notes": f"Public location city filter/count. Count shown: {city.get('count')}. Coverage metadata only.",
                        },
                    )
            page += 1
        self.log(f"Collected {len(self.locations)} location rows")
        for location in self.locations:
            self.add_branch(location)

    def add_branch(self, location: dict[str, object]) -> None:
        branch_id = f"branch-{location.get('id')}"
        source_url = f"{PRIMARY_DOMAIN}/location"
        name = normalize_spaces(location.get("branch", ""))
        city = normalize_spaces(location.get("city", ""))
        address = normalize_spaces(location.get("address", ""))
        lat, lng = parse_google_maps_coordinates(str(location.get("map_link") or ""))
        brand = infer_brand_from_text(name)
        self.tables.append(
            "branch",
            {
                "branch_id": branch_id,
                "dealer_id": DEALER_ID,
                "source_url": source_url,
                "locator_id": location.get("id", ""),
                "branch_name": name,
                "slug": slugify_text(name),
                "street": address,
                "city": city,
                "state_or_province": infer_province(address, city),
                "postal_code": infer_postal_code(address),
                "full_address": address,
                "latitude": lat,
                "longitude": lng,
                "phone": location.get("contact", ""),
                "fax": "",
                "email": "",
                "website": location.get("map_link", ""),
                "open_hours_raw": location.get("schedule", ""),
                "days_str": parse_days(location.get("schedule", "")),
                "brand_raw": brand,
                "product_raw": "motorcycle",
                "categories_raw": "MotoXpress Store",
                "inquire_url": location.get("map_link", ""),
            },
        )
        self.add_branch_coverage(branch_id, source_url, "product_family", "motorcycle")
        self.add_branch_coverage(branch_id, source_url, "city", city)
        if brand:
            self.add_branch_coverage(branch_id, source_url, "brand", brand)
        if location.get("map_link"):
            self.add_external_cta("source-next-location", source_url, str(location.get("map_link")), "Map")

    def add_branch_coverage(self, branch_id: str, source_url: str, coverage_type: str, raw_value: object) -> None:
        normalized = normalize_token(raw_value)
        if not normalized:
            return
        key = (branch_id, coverage_type, normalized)
        if key in self.coverage_seen:
            return
        self.coverage_seen.add(key)
        self.tables.append(
            "branch_product_coverage",
            {
                "coverage_id": f"coverage-{branch_id}-{coverage_type}-{normalized}",
                "dealer_id": DEALER_ID,
                "branch_id": branch_id,
                "coverage_type": coverage_type,
                "coverage_value_raw": raw_value,
                "coverage_value_normalized": normalized,
                "source_url": source_url,
                "source_field": "Next.js location page data",
                "confidence": "coverage_only",
                "observed_at": datetime.now(TIMEZONE).isoformat(),
                "notes": "Broad branch coverage hint only; not exact model availability.",
            },
        )

    def collect_articles(self) -> None:
        page = 1
        total_pages = 1
        categories_added = False
        while page <= total_pages:
            props = self.next_json("news-promos", params={"page": page})
            pagination = props.get("pagination") if isinstance(props.get("pagination"), dict) else {}
            total_pages = int(pagination.get("pageCount") or total_pages)
            if not categories_added:
                self.article_categories = props.get("categories", []) if isinstance(props.get("categories"), list) else []
                for category in self.article_categories:
                    self.tables.append(
                        "website_filter_reference",
                        {
                            "dealer_id": DEALER_ID,
                            "reference_type": "article_category",
                            "filter_name": "article_category",
                            "option_value": category.get("slug", ""),
                            "option_label": category.get("name", ""),
                            "source_url": f"{PRIMARY_DOMAIN}/news-promos",
                            "notes": "Public news/promos category filter. Coverage metadata only.",
                        },
                    )
                categories_added = True
            for article in props.get("articles", []) or []:
                if str(article.get("id")) not in {str(item.get("id")) for item in self.articles}:
                    self.articles.append(article)
            page += 1
        self.log(f"Collected {len(self.articles)} article/promo rows")
        for article in self.articles:
            self.add_article(article)

    def add_article(self, article: dict[str, object]) -> None:
        source_url = f"{PRIMARY_DOMAIN}/news-promos/{article.get('slug', '')}"
        source_page_id = f"source-article-{article.get('id')}"
        title = normalize_spaces(article.get("title", ""))
        description = normalize_spaces(article.get("short_description", "")) or rich_text_to_text(article.get("content"))
        self.add_source_page(
            source_page_id=source_page_id,
            source_type="next_article",
            wp_post_type="next_article",
            wp_id=str(article.get("id", "")),
            url=source_url,
            slug=str(article.get("slug", "")),
            title=title,
            content_text=description,
            published_at=str(article.get("published_date") or article.get("publishedAt") or article.get("createdAt") or ""),
            modified_at=str(article.get("updatedAt") or ""),
            featured_image_url=image_url_from_object(article.get("image")),
        )
        if not looks_like_promo(article):
            return
        promo_id = f"promo-{article.get('id')}"
        if promo_id in self.promo_seen:
            return
        self.promo_seen.add(promo_id)
        text = f"{title} {description}"
        self.tables.append(
            "promo",
            {
                "promo_id": promo_id,
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "source_page_url": source_url,
                "source_url": source_url,
                "title": title,
                "promo_type": infer_promo_type(text),
                "description_text": description,
                "image_url": image_url_from_object(article.get("image")),
                "cta_url": source_url,
                "target_listing_url": f"{PRIMARY_DOMAIN}/news-promos",
                "applies_to_brand": infer_brand_from_text(text),
                "applies_to_vehicle_type": infer_vehicle_type_from_text(text),
                "applies_to_tag": article_category_label(article),
                "has_vehicle_promo_flag": "TRUE" if "motorcycle" in text.lower() or infer_brand_from_text(text) else "",
                "linked_motorcycle_count": "",
                "linked_motorcycles": "",
                "link_types": "",
                "link_source_urls": "",
                "published_at": article.get("published_date", ""),
                "modified_at": article.get("updatedAt", ""),
            },
        )
        for model_key, motorcycle_id in self.motorcycle_by_model.items():
            if model_key and model_key in normalize_token(text):
                self.add_promo_link(motorcycle_id, promo_id, "article_exact_model_text_match", source_url)

    def add_promo_link(self, motorcycle_id: str, promo_id: str, link_type: str, source_url: str) -> None:
        key = (motorcycle_id, promo_id, link_type)
        if key in self.promo_link_seen:
            return
        self.promo_link_seen.add(key)
        motorcycle = self.motorcycle_summary.get(motorcycle_id, {})
        promo_title = ""
        for row in self.tables.rows["promo"]:
            if row.get("promo_id") == promo_id:
                promo_title = str(row.get("title", ""))
                break
        self.tables.append(
            "motorcycle_promo_link",
            {
                "motorcycle_id": motorcycle_id,
                "promo_id": promo_id,
                "dealer_id": DEALER_ID,
                "display_name": motorcycle.get("display_name", ""),
                "promo_title": promo_title,
                "link_type": link_type,
                "source_url": source_url,
                "notes": "Linked only when article/promo text contained an exact normalized model name.",
            },
        )

    def collect_about_and_services(self) -> None:
        about = self.next_json("about")
        self.about_sections = about.get("aboutSections", []) if isinstance(about.get("aboutSections"), list) else []
        self.faqs = about.get("faqs", []) if isinstance(about.get("faqs"), list) else []
        financing = self.next_json("financing", params={"form": "true", "make": "Honda", "model": "ADV 160"})
        self.financing_page = financing.get("financing", {}) if isinstance(financing.get("financing"), dict) else {}
        service_rows = [
            ("financing", PAGE_URLS["financing"], self.financing_page.get("title", "Finance a Motorcycle"), self.financing_page.get("subtitle", "")),
            ("store_locations", PAGE_URLS["location"], "Store Locations", f"{len(self.locations)} public location rows."),
            ("about", PAGE_URLS["about"], "About MotoXpress", " ".join(rich_text_to_text(section.get("description")) for section in self.about_sections)),
            ("faq", PAGE_URLS["about"], "Frequently Asked Questions", " | ".join(f"{faq.get('title')}: {rich_text_to_text(faq.get('description'))}" for faq in self.faqs)),
        ]
        for category, url, title, description in service_rows:
            self.tables.append(
                "aftersales_service",
                {
                    "service_id": f"service-{category}",
                    "dealer_id": DEALER_ID,
                    "source_page_id": self.source_page_by_url.get(normalize_url(url), ""),
                    "service_category": category,
                    "title": normalize_spaces(title),
                    "description_text": normalize_spaces(description),
                    "partner_name": "",
                    "cta_url": url,
                    "external_form_url": PAGE_URLS["financing"] if category == "financing" else "",
                    "notes": "MotoXpress public Next.js service/content surface.",
                },
            )
        self.add_generic_financing_row()

    def add_generic_financing_row(self) -> None:
        self.tables.append(
            "financing",
            {
                "financing_id": "financing-motoxpress-generic-application",
                "motorcycle_id": "",
                "dealer_id": DEALER_ID,
                "display_name": "",
                "brand_name": "",
                "model_name": "",
                "source_url": PAGE_URLS["financing"],
                "financing_type": "generic_application_form",
                "program_name": "MotoXpress online loan application",
                "monthly_amount": "",
                "downpayment_amount": "",
                "loan_amount": "",
                "loan_term_months": "",
                "estimated_income_amount": "",
                "form_url": PAGE_URLS["financing"],
                "cta_url": PAGE_URLS["financing"],
                "note_text": "Generic rendered financing form. Model-specific rows populate motorcycle_id separately.",
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def collect_forms(self) -> None:
        self.add_financing_form()
        self.add_search_form()
        self.add_shop_filter_form()

    def add_financing_form(self) -> None:
        form_id = "form-motoxpress-financing-application"
        self.tables.append(
            "form_definition",
            {
                "form_id": form_id,
                "dealer_id": DEALER_ID,
                "form_name": "MotoXpress Loan Application",
                "form_category": "financing_application",
                "source_url": PAGE_URLS["financing"],
                "final_url": PAGE_URLS["financing"],
                "page_title": "Financing | MotoXpress PH",
                "form_action": f"{API_BASE}/financings",
                "form_method": "POST",
                "requires_file_upload": "TRUE",
                "requires_otp": "",
                "multi_step": "",
                "has_privacy_checkbox": "TRUE",
                "notes": "Client-rendered form inspected with Playwright and Next.js chunk. No live submission performed. Files POST to /v1/financings/{id} after JSON submission.",
            },
        )
        self.add_form_fields(form_id, "MotoXpress Loan Application", "financing_application", FINANCING_FORM_FIELDS, PAGE_URLS["financing"])

    def add_search_form(self) -> None:
        form_id = "form-motoxpress-site-search"
        self.tables.append(
            "form_definition",
            {
                "form_id": form_id,
                "dealer_id": DEALER_ID,
                "form_name": "Site Search",
                "form_category": "site_search",
                "source_url": PAGE_URLS["home"],
                "final_url": PAGE_URLS["home"],
                "page_title": "MotoXpress PH",
                "form_action": PAGE_URLS["home"],
                "form_method": "GET",
                "requires_file_upload": "",
                "requires_otp": "",
                "multi_step": "",
                "has_privacy_checkbox": "",
                "notes": "Public search input schema only.",
            },
        )
        self.add_form_fields(form_id, "Site Search", "site_search", [("search", "Search", "search", "Search", [])], PAGE_URLS["home"])

    def add_shop_filter_form(self) -> None:
        form_id = "form-motoxpress-shop-filters"
        self.tables.append(
            "form_definition",
            {
                "form_id": form_id,
                "dealer_id": DEALER_ID,
                "form_name": "Shop Filters",
                "form_category": "product_filter",
                "source_url": PAGE_URLS["shop"],
                "final_url": PAGE_URLS["shop"],
                "page_title": "Shop online | MotoXpress Philippines | MotoXpress PH",
                "form_action": PAGE_URLS["shop"],
                "form_method": "GET",
                "requires_file_upload": "",
                "requires_otp": "",
                "multi_step": "",
                "has_privacy_checkbox": "",
                "notes": "Public listing filter schema only.",
            },
        )
        fields = [
            ("brand", "Brand", "select", "Shop Filters", [str(brand.get("name")) for brand in self.brands]),
            ("price_min", "Price min", "range", "Shop Filters", ["40000"]),
            ("price_max", "Price max", "range", "Shop Filters", ["300000"]),
            ("transmission", "Transmission", "select", "Shop Filters", ["Automatic", "Manual", "Semi Manual"]),
            ("deals", "Deals", "checkbox", "Shop Filters", ["On Sale"]),
            ("sort", "Sort By", "select", "Shop Filters", ["Default", "Price low to high", "Price high to low"]),
        ]
        self.add_form_fields(form_id, "Shop Filters", "product_filter", fields, PAGE_URLS["shop"])
        for name, label, _, _, options in fields:
            for option in options:
                self.tables.append(
                    "website_filter_reference",
                    {
                        "dealer_id": DEALER_ID,
                        "reference_type": "shop_filter",
                        "filter_name": name,
                        "option_value": option,
                        "option_label": option,
                        "source_url": PAGE_URLS["shop"],
                        "notes": f"Public shop filter option for {label}. Reference metadata only.",
                    },
                )

    def add_form_fields(
        self,
        form_id: str,
        form_name: str,
        form_category: str,
        fields: list[tuple[str, str, str, str, list[str]]],
        source_url: str,
    ) -> None:
        order = 1
        for field_name, field_label, field_type, section_name, options in fields:
            option_rows = options or [""]
            for option in option_rows:
                self.tables.append(
                    "form_field",
                    {
                        "form_id": form_id,
                        "dealer_id": DEALER_ID,
                        "form_name": form_name,
                        "form_category": form_category,
                        "field_order": order,
                        "field_name": field_name,
                        "field_label": field_label,
                        "field_type": field_type,
                        "is_required": "TRUE" if form_category == "financing_application" and field_name in {"make", "model", "preferred_monthly_term", "downpayment", "firstname", "lastname", "email", "mobile_number"} else "",
                        "default_value": "",
                        "placeholder": "",
                        "option_value": option,
                        "option_label": option,
                        "section_name": section_name,
                        "source_url": source_url,
                    },
                )
            order += 1

    def collect_relationship_probe(self) -> None:
        self.tables.append(
            "relationship_probe",
            {
                "probe_id": "probe-motoxpress-branch-model-availability",
                "dealer_id": DEALER_ID,
                "relationship_name": "branch_to_motorcycle_model_availability",
                "source_url": PAGE_URLS["location"],
                "source_type": "Next.js location rows and public motorcycle API",
                "status": "coverage_only",
                "evidence_found": "FALSE",
                "rows_created": 0,
                "notes": "MotoXpress exposes branch locations and dealer-wide motorcycle catalog/stock fields, but no exact public branch plus exact model inventory bridge.",
                "checked_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def collect_data_quality_notes(self) -> None:
        self.add_data_quality_issue(
            entity_type="branch_model_availability",
            entity_id=DEALER_ID,
            source_url=PAGE_URLS["location"],
            issue_type="no_exact_public_relationship",
            issue_detail="No exact branch-to-motorcycle model availability was found. Dealer-wide stock exists on motorcycle API rows, but it is not branch-level inventory.",
            severity="warning",
        )
        self.add_data_quality_issue(
            entity_type="form_definition",
            entity_id="form-motoxpress-financing-application",
            source_url=PAGE_URLS["financing"],
            issue_type="client_rendered_form",
            issue_detail="Financing form fields are client-rendered. Schema was verified by Playwright and the Next.js financing chunk; no form submission was performed.",
            severity="info",
        )
        self.add_data_quality_issue(
            entity_type="sitemap",
            entity_id="sitemap.xml",
            source_url=f"{PRIMARY_DOMAIN}/sitemap.xml",
            issue_type="not_found",
            issue_detail="MotoXpress sitemap.xml and robots.txt returned 404 during exploration. Next data/API sources were used instead.",
            severity="info",
        )

    def add_data_quality_issue(
        self,
        *,
        entity_type: str,
        entity_id: str,
        source_url: str,
        issue_type: str,
        issue_detail: str,
        severity: str,
    ) -> None:
        self.tables.append(
            "data_quality_issue",
            {
                "issue_id": f"issue-{len(self.tables.rows['data_quality_issue']) + 1}",
                "dealer_id": DEALER_ID,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "source_url": source_url,
                "issue_type": issue_type,
                "issue_detail": issue_detail,
                "severity": severity,
                "captured_at": datetime.now(TIMEZONE).isoformat(),
            },
        )

    def add_source_page(
        self,
        *,
        source_page_id: str,
        source_type: str,
        wp_post_type: str,
        wp_id: str,
        url: str,
        slug: str,
        title: str,
        content_text: str,
        published_at: str = "",
        modified_at: str = "",
        featured_image_url: str = "",
    ) -> None:
        normalized_url = normalize_url(url)
        if normalized_url in self.source_page_by_url:
            return
        self.source_page_by_url[normalized_url] = source_page_id
        self.tables.append(
            "source_page",
            {
                "source_page_id": source_page_id,
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "source_type": source_type,
                "wp_post_type": wp_post_type,
                "wp_id": wp_id,
                "url": normalized_url,
                "slug": slug,
                "title": normalize_spaces(title),
                "status": "public",
                "published_at": published_at,
                "modified_at": modified_at,
                "author_id": "",
                "excerpt_text": normalize_spaces(content_text)[:1000],
                "content_html": "",
                "content_text": normalize_spaces(content_text),
                "featured_media_id": "",
                "featured_image_url": featured_image_url,
                "category_ids": "",
                "tag_ids": "",
                "job_type_ids": "",
                "seo_title": normalize_spaces(title),
                "seo_description": normalize_spaces(content_text)[:300],
                "canonical_url": normalized_url,
            },
        )
        self.source_counter += 1

    def collect_external_links_from_soup(self, soup: BeautifulSoup, source_url: str, source_page_id: str) -> None:
        for link in soup.find_all("a", href=True):
            href = urljoin(source_url, link.get("href", ""))
            self.add_external_cta(source_page_id, source_url, href, link.get_text(" ", strip=True))

    def add_external_cta(self, source_page_id: str, source_url: str, target_url: str, anchor_text: str) -> None:
        parsed = urlparse(target_url)
        if not parsed.scheme.startswith("http"):
            return
        target_domain = parsed.netloc.lower()
        is_internal = target_domain in {"motoxpress.ph", "www.motoxpress.ph"}
        if is_internal and not any(part in parsed.path for part in ("/shop", "/financing", "/location", "/news-promos", "/about")):
            return
        key = (source_url, target_url)
        if key in self.external_cta_seen:
            return
        self.external_cta_seen.add(key)
        self.tables.append(
            "external_cta",
            {
                "cta_id": f"cta-{len(self.tables.rows['external_cta']) + 1}",
                "dealer_id": DEALER_ID,
                "source_page_id": source_page_id,
                "source_url": source_url,
                "target_url": target_url,
                "target_domain": target_domain,
                "cta_type": infer_cta_type(target_url, anchor_text),
                "anchor_text": normalize_spaces(anchor_text),
                "notes": "Public link/CTA evidence.",
            },
        )

    def finalize_crawl_run(self) -> None:
        self.tables.append(
            "crawl_run",
            {
                "run_id": self.run_id,
                "dealer_id": DEALER_ID,
                "run_started_at": self.run_started_at,
                "run_finished_at": datetime.now(TIMEZONE).isoformat(),
                "source_count": self.source_counter,
                "notes": "Current-state public MotoXpress snapshot using Next.js data, public API endpoints, rendered-form inspection, and public page HTML.",
            },
        )

    def write_workbook(self) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in DATABASE_SHEETS:
            schema = DATABASE_SCHEMAS[sheet_name]
            sheet = workbook.create_sheet(title=sheet_name[:31])
            sheet.append(schema)
            for row in self.tables.rows.get(sheet_name, []):
                sheet.append([safe_excel_value(row.get(column)) for column in schema])
        workbook.save(self.output_path)


def parse_next_data(soup: BeautifulSoup) -> dict[str, object]:
    tag = soup.select_one("#__NEXT_DATA__")
    if not tag:
        return {}
    try:
        return json.loads(tag.get_text())
    except json.JSONDecodeError:
        return {}


def motorcycle_specs(motorcycle: dict[str, object]) -> list[dict[str, str]]:
    specs: list[dict[str, str]] = []
    for key, label in SPEC_FIELDS:
        value = motorcycle.get(key)
        if value in (None, ""):
            continue
        if key == "additional_features":
            value = html_to_text(str(value))
        specs.append({"group": "Specifications", "label": label, "value": normalize_spaces(value)})
    for image in motorcycle.get("images") or []:
        color = normalize_spaces(image.get("color", ""))
        if color:
            specs.append({"group": "Color Options", "label": "Color", "value": color})
    return specs


def flatten_specs(
    specs: list[dict[str, str]],
    brand_name: str,
    model_name: str,
    motorcycle: dict[str, object],
) -> dict[str, str]:
    flattened = {"spec_make": brand_name, "spec_model": model_name}
    brake_values: list[str] = []
    for spec in specs:
        label = normalize_label(spec["label"])
        value = spec["value"]
        if "transmission" in label:
            flattened["spec_transmission"] = value
        elif "engine type" in label:
            flattened["spec_engine_type"] = value
        elif "seat height" in label:
            flattened["spec_seat_height"] = value
        elif "displacement" in label:
            flattened["spec_displacement_cc"] = value
        elif "wheel type" in label:
            flattened["spec_wheels_type"] = value
        elif "starting system" in label:
            flattened["spec_starting_system"] = value
        elif "exterior dimensions" in label:
            flattened["spec_overall_dimensions"] = value
        elif "ground clearance" in label:
            flattened["spec_ground_clearance"] = value
        elif "fuel system" in label:
            flattened["spec_fuel_system"] = value
        elif "brake" in label:
            brake_values.append(f"{spec['label']}: {value}")
        elif "tires" in label:
            flattened["spec_front_tire"] = value
            flattened["spec_rear_tire"] = value
    if brake_values:
        flattened["spec_brake_system"] = " | ".join(brake_values)
    return flattened


def brand_name_from_motorcycle(motorcycle: dict[str, object]) -> str:
    brand = motorcycle.get("brand")
    if isinstance(brand, dict):
        return str(brand.get("name") or "")
    return ""


def primary_image_url(motorcycle: dict[str, object]) -> str:
    images = motorcycle.get("images")
    if isinstance(images, list) and images:
        return str(images[0].get("url") or "")
    return ""


def image_url_from_object(value: object) -> str:
    if not isinstance(value, dict):
        return ""
    url = str(value.get("url") or "")
    if url.startswith("http"):
        return url
    if url.startswith("sgp1.digitaloceanspaces.com"):
        return f"https://{url}"
    return url


def parse_amount(value: object) -> float | str:
    if value in (None, "", 0, "0"):
        return "" if value in (None, "") else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    amount = parse_currency_amount(text)
    if amount != "":
        return amount
    match = re.search(r"^-?\d[\d,]*(?:\.\d+)?$", text)
    if match:
        return float(text.replace(",", ""))
    return ""


def parse_published_price(value: object) -> float | str:
    amount = parse_amount(value)
    return amount if isinstance(amount, (int, float)) and amount > 0 else ""


def amount_text(value: object) -> str:
    if value in ("", None):
        return ""
    try:
        return f"₱{float(value):,.0f}"
    except (TypeError, ValueError):
        return str(value)


def financing_url(brand_name: str, model_name: str) -> str:
    return f"{PRIMARY_DOMAIN}/financing?form=true&make={quote(brand_name)}&model={quote(model_name)}"


def infer_vehicle_type(motorcycle: dict[str, object]) -> str:
    fuel_type = str(motorcycle.get("fuel_type") or "").lower()
    if "electric" in fuel_type:
        return "electric"
    displacement = displacement_cc(str(motorcycle.get("displacement") or ""))
    if displacement >= 400:
        return "big-bike"
    return "regular-bike"


def infer_vehicle_type_from_text(text: str) -> str:
    lower = text.lower()
    if "electric" in lower or "e-bike" in lower:
        return "electric"
    if re.search(r"\b([4-9]\d{2}|1\d{3})\s*cc\b", lower) or "big bike" in lower:
        return "big-bike"
    if "motorcycle" in lower or infer_brand_from_text(text):
        return "regular-bike"
    return ""


def displacement_cc(value: str) -> float:
    match = re.search(r"([\d,.]+)", value or "")
    if not match:
        return 0.0
    try:
        return float(match.group(1).replace(",", ""))
    except ValueError:
        return 0.0


def parse_google_maps_coordinates(url: str) -> tuple[str, str]:
    for pattern in (r"@(-?\d+\.\d+),(-?\d+\.\d+)", r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)", r"!8m2!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)"):
        match = re.search(pattern, url)
        if match:
            return match.group(1), match.group(2)
    return "", ""


def infer_province(address: str, city: str) -> str:
    parts = [normalize_spaces(part) for part in str(address or "").split(",") if normalize_spaces(part)]
    if len(parts) >= 2:
        last = parts[-1]
        if not re.search(r"\d", last) and normalize_token(last) != normalize_token(city):
            return last
        if len(parts) >= 3:
            return parts[-2]
    return ""


def infer_postal_code(address: str) -> str:
    match = re.search(r"\b(\d{4})\b", address or "")
    return match.group(1) if match else ""


def parse_days(schedule: object) -> str:
    text = normalize_spaces(schedule)
    match = re.match(r"([A-Za-z]+(?:\s+to\s+[A-Za-z]+)?)", text)
    return match.group(1) if match else ""


def infer_brand_from_text(text: object) -> str:
    lower = str(text or "").lower()
    for brand in ["Honda", "Yamaha", "Suzuki", "Kawasaki", "KYMCO", "TVS"]:
        if re.search(rf"\b{re.escape(brand.lower())}\b", lower):
            return brand
    return ""


def looks_like_promo(article: dict[str, object]) -> bool:
    text = f"{article.get('title', '')} {article.get('short_description', '')} {rich_text_to_text(article.get('content'))} {article_category_label(article)}".lower()
    return any(word in text for word in ["promo", "promotion", "cashback", "deal", "launch", "christmas", "masaya", "sale"])


def infer_promo_type(text: str) -> str:
    lower = text.lower()
    if "cashback" in lower:
        return "cashback"
    if "launch" in lower:
        return "launch"
    if "deal" in lower or "sale" in lower:
        return "discount"
    return "campaign"


def article_category_label(article: dict[str, object]) -> str:
    category = article.get("category")
    if isinstance(category, dict):
        return str(category.get("name") or category.get("slug") or "")
    return ""


def rich_text_to_text(value: object) -> str:
    text: list[str] = []

    def walk(node: object) -> None:
        if isinstance(node, dict):
            if "text" in node:
                text.append(str(node.get("text") or ""))
            for child in node.get("children", []) or []:
                walk(child)
        elif isinstance(node, list):
            for child in node:
                walk(child)

    walk(value)
    return normalize_spaces(" ".join(text))


def infer_cta_type(target_url: str, anchor_text: str) -> str:
    lower = f"{target_url} {anchor_text}".lower()
    if "google.com/maps" in lower:
        return "map"
    if "facebook" in lower:
        return "social"
    if "financing" in lower:
        return "financing_application"
    if "/shop" in lower:
        return "product_or_shop"
    if "brochure" in lower or lower.endswith(".pdf"):
        return "download"
    return "link"


def join_values(values: object) -> str:
    if isinstance(values, (list, tuple, set)):
        return " | ".join(str(value) for value in values if str(value).strip())
    return str(values or "")


def normalize_label(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def normalize_token(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(value or "").lower()).strip("-")


def slugify_text(value: object) -> str:
    return normalize_token(value) or "item"


def normalize_spaces(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def safe_excel_value(value: object) -> object:
    value = sanitize_excel_value(value)
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", " ", value)
    return value


def main() -> int:
    output_path = Path(__file__).resolve().parents[2] / "working/outputs/dealer-databases/motoxpress_database.xlsx"
    scraper = MotoXpressDatabaseScraper(output_path)
    scraper.run()
    print(f"[OK] Wrote {output_path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
